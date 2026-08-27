"""The own-check period restriction, pinned in both directions.

Registered round: docs/pierce/own-check-periods.md. Proof 1A failed
because a covenant threshold parked in a scalar column was quoted as
a failing period. These tests pin that it is not, that a genuine
non-zero in a real period still is, and that a sheet whose axis
cannot be read keeps the old behaviour rather than falling silent.
"""

import tempfile
from pathlib import Path

from polar.tieout.analytics import run_analytics
from polar.tieout.structure import read_structure
from polar.tieout.workbook import read_workbook


def _analytics(build):
    from openpyxl import Workbook as Book

    book = Book()
    build(book.active)
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "built.xlsx"
        book.save(path)
        loaded = read_workbook(str(path))
    return loaded, run_analytics(loaded, read_structure(loaded))


def _periods(sheet, first_col: int = 8) -> None:
    """A year-labelled period grid starting at `first_col`, with a
    scalar column at E — Kelso's shape."""
    from openpyxl.utils import get_column_letter

    for offset in range(6):
        letter = get_column_letter(first_col + offset)
        sheet[f"{letter}2"] = f"FY{2015 + offset}"
        for row in range(10, 16):
            sheet[f"{letter}{row}"] = 0.0
    for row in range(10, 16):
        sheet[f"A{row}"] = f"Line {row}"


def _own(result):
    return [f for f in result.findings if f.rule == "model-own-check"]


def test_a_threshold_in_a_scalar_column_is_not_a_failing_period() -> None:
    """Proof 1A's eight false alarms, distilled: 1.15 parked in the
    parameter column of a check row whose real series is zero."""

    def build(sheet) -> None:
        _periods(sheet)
        sheet["A12"] = "Check: Minimum ADSCR > breach level"
        sheet["E12"] = 1.15

    _, result = _analytics(build)
    assert not [f for f in _own(result) if f.ref.endswith("E12")]


def test_a_genuine_break_in_a_period_column_still_fires() -> None:
    """The other direction — the restriction must not silence the
    check where it is actually about a period."""
    from openpyxl.utils import get_column_letter

    def build(sheet) -> None:
        _periods(sheet)
        sheet["A12"] = "Check"
        sheet[f"{get_column_letter(10)}12"] = -2.32

    _, result = _analytics(build)
    found = _own(result)
    assert len(found) == 1
    assert found[0].ref.endswith("J12")


def test_a_sheet_without_a_readable_axis_keeps_the_old_behaviour() -> None:
    """The registered fallback: « we cannot see the axis » must never
    silently become « we stop checking »."""

    def build(sheet) -> None:
        for row in range(10, 16):
            sheet[f"A{row}"] = f"Line {row}"
            for col in "BCDEFG":
                sheet[f"{col}{row}"] = 0.0
        sheet["A12"] = "Check"
        sheet["E12"] = -2.32

    book, result = _analytics(build)
    assert not read_structure(book).axes.get("Sheet")
    assert [f for f in _own(result) if f.ref.endswith("E12")]
