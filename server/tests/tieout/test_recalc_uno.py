"""B1 on the real machine: the UNO adapter, end to end.

These tests run only where a LibreOffice ≥ 25.8 install exists (this
container after `dev/setup-libreoffice`; skipped honestly anywhere
else — a skip is not a pass). Every workbook here is written by
openpyxl with **no stored answers**, so any value that comes back was
genuinely computed by LibreOffice through the UNO socket — the same
proof the lead ran by hand on 26 August, now standing in the suite.

The round trip test is B2's whole mechanism in miniature: recalculate
and store a copy (now it has stored values), gate the copy against a
fresh recalculation (pass), plant a one-cell lie in the stored values
(the gate must name that cell), and B1's DONE sentence — a changed
input produces changed downstream values, unattended — is the pool
test at the end.
"""

import shutil
import zipfile
from collections.abc import Generator
from pathlib import Path

import pytest
from openpyxl import Workbook

from polar.tieout.recalc.gate import gate_file
from polar.tieout.recalc.pool import WorkerPool
from polar.tieout.recalc.uno_calc import UnoCalculator, find_install
from polar.tieout.workbook import read_workbook

pytestmark = pytest.mark.skipif(
    find_install() is None,
    reason="no LibreOffice >= 25.8 here — run dev/setup-libreoffice",
)


@pytest.fixture(scope="module")
def calculator() -> Generator[UnoCalculator]:
    calc = UnoCalculator()
    calc.start()
    yield calc
    calc.stop()


def _write_model(path: Path, unit_price: float = 2.0) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "M"
    sheet["A1"] = unit_price
    sheet["A2"] = 3
    sheet["A3"] = "=SUM(A1:A2)*10"
    sheet["A4"] = "=1/0"
    sheet["A5"] = '=IF(A1>1,"yes","no")'
    book.save(path)


def test_uncached_formulas_are_actually_computed(
    calculator: UnoCalculator, tmp_path: Path
) -> None:
    path = tmp_path / "model.xlsx"
    _write_model(path)
    result = calculator.recalculate(str(path))
    assert result.values["M!A3"] == 50.0
    assert str(result.values["M!A4"]).startswith("#ERR:")
    assert result.values["M!A5"] == "yes"
    assert result.engine.startswith("LibreOffice 25.")


def test_the_files_own_iteration_settings_are_pushed(
    calculator: UnoCalculator, tmp_path: Path
) -> None:
    # A circular pair converging to A1=4/3, B1=2/3 — computable only
    # if the file's calcPr (iteration on) was pushed into the engine;
    # with LibreOffice's default (iteration off) both cells error.
    book = Workbook()
    sheet = book.active
    sheet.title = "M"
    sheet["A1"] = "=B1/2+1"
    sheet["B1"] = "=A1/2"
    book.calculation.iterate = True
    book.calculation.iterateCount = 200
    book.calculation.iterateDelta = 1e-9
    path = tmp_path / "circular.xlsx"
    book.save(path)
    result = calculator.recalculate(str(path))
    a1, b1 = result.values["M!A1"], result.values["M!B1"]
    assert isinstance(a1, float)
    assert abs(a1 - 4 / 3) < 1e-6
    assert isinstance(b1, float)
    assert abs(b1 - 2 / 3) < 1e-6


def _write_labelled_model(path: Path) -> None:
    # Laid out the way the reader reads a model: labels down column A,
    # a period header across row 1, figures in the grid.
    book = Workbook()
    sheet = book.active
    sheet.title = "M"
    sheet["B1"] = "FY2025A"
    sheet["C1"] = "FY2026E"
    sheet["A2"] = "Price"
    sheet["B2"] = 2
    sheet["C2"] = 4
    sheet["A3"] = "Volume"
    sheet["B3"] = 3
    sheet["C3"] = 6
    sheet["A4"] = "Revenue"
    sheet["B4"] = "=SUM(B2:B3)*10"
    sheet["C4"] = "=SUM(C2:C3)*10"
    book.save(path)


def test_gate_round_trip_and_a_planted_lie(
    calculator: UnoCalculator, tmp_path: Path
) -> None:
    fresh = tmp_path / "fresh.xlsx"
    stored = tmp_path / "stored.xlsx"
    _write_labelled_model(fresh)
    # Recalculate the uncached file and store a copy: the copy now
    # carries stored values, like any file Excel once saved.
    calculator.recalculate(str(fresh), store_to=str(stored))

    cells = read_workbook(str(stored)).cells
    report = gate_file(cells, calculator.recalculate(str(stored)).values)
    assert report.verdict == "pass"
    assert report.compared >= 2
    assert report.match_rate == 1.0

    # Plant a lie in the stored values only — the formula stays right,
    # the cached answer goes wrong, which is precisely the fidelity
    # defect class the gate exists to catch.
    lied = tmp_path / "lied.xlsx"
    _tamper_stored_value(stored, lied, old=">50<", new=">51<")
    lied_cells = read_workbook(str(lied)).cells
    lied_report = gate_file(lied_cells, calculator.recalculate(str(lied)).values)
    assert lied_report.verdict == "fail"
    assert [diff.ref for diff in lied_report.mismatches] == ["M!B4"]


def test_changed_input_changes_downstream_unattended(tmp_path: Path) -> None:
    # B1's DONE sentence, verbatim — through the pool, no hand-holding.
    two = tmp_path / "price2.xlsx"
    five = tmp_path / "price5.xlsx"
    _write_model(two, unit_price=2.0)
    _write_model(five, unit_price=5.0)
    pool = WorkerPool(UnoCalculator, size=1, recycle_after=10)
    report = pool.recalculate_many([str(two), str(five)])
    assert report.failures == {}
    assert report.results[str(two)].values["M!A3"] == 50.0
    assert report.results[str(five)].values["M!A3"] == 80.0


def _tamper_stored_value(source: Path, target: Path, *, old: str, new: str) -> None:
    """Rewrite one stored value in the sheet XML, byte-surgically."""
    shutil.copyfile(source, target)
    with zipfile.ZipFile(source) as archive:
        names = archive.namelist()
        contents = {name: archive.read(name) for name in names}
    sheet_names = [n for n in names if n.startswith("xl/worksheets/sheet")]
    hit = False
    for name in sheet_names:
        text = contents[name].decode("utf-8")
        if old in text:
            contents[name] = text.replace(old, new, 1).encode("utf-8")
            hit = True
            break
    assert hit, f"marker {old!r} not found in {sheet_names}"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in names:
            archive.writestr(name, contents[name])
