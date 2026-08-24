"""Track A1: the write path's first promises, tested.

The writer's contract is surgical: everything an edit does not touch
keeps its exact bytes, and everything it does touch reads back
correctly in an independent library.
"""

import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook as Book
from openpyxl import load_workbook

from polar.tieout.workbook import read_workbook
from polar.tieout.writer import WorkbookWriter, WriteRefused


def _host(folder: str) -> Path:
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Rate"
    sheet["B1"] = 0.08
    sheet["A2"] = "Base"
    sheet["B2"] = 100
    sheet["A3"] = "Grown"
    sheet["B3"] = "=B2*(1+B1)"
    other = book.create_sheet("Notes")
    other["A1"] = "untouched"
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return path


def _member_bytes(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {i.filename: archive.read(i.filename) for i in archive.infolist()}


def test_a_no_op_save_keeps_every_members_exact_bytes() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        out = WorkbookWriter(host).save(Path(folder) / "out.xlsx")
        assert _member_bytes(host) == _member_bytes(out)


def test_a_value_edit_changes_only_its_own_sheet() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        writer = WorkbookWriter(host)
        edit = writer.set_cell("Model", "B2", formula=None, value="250")
        out = writer.save(Path(folder) / "out.xlsx")

        assert edit.before_value == "100"
        before, after = _member_bytes(host), _member_bytes(out)
        changed = {name for name in before if before[name] != after.get(name)}
        assert changed == {"xl/worksheets/sheet1.xml"}

        seen = load_workbook(out)
        assert seen["Model"]["B2"].value == 250
        assert seen["Notes"]["A1"].value == "untouched"


def test_a_formula_edit_reads_back_and_drops_the_calc_chain() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "B3", formula="=B2*(1+B1*2)", value="116")
        out = writer.save(Path(folder) / "out.xlsx")

        seen = load_workbook(out)
        assert seen["Model"]["B3"].value == "=B2*(1+B1*2)"
        with zipfile.ZipFile(out) as archive:
            names = archive.namelist()
            assert "xl/calcChain.xml" not in names
            kinds = archive.read("[Content_Types].xml").decode("utf-8")
            assert "calcChain" not in kinds

        book = read_workbook(str(out))
        cell = book.cells["Model!B3"]
        assert cell.formula == "=B2*(1+B1*2)"
        assert str(cell.value) == "116"


def test_a_the_writer_refuses_what_it_cannot_do_safely() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        writer = WorkbookWriter(host)
        with pytest.raises(WriteRefused, match="create=True"):
            writer.set_cell("Model", "Z99", formula=None, value="1")
        with pytest.raises(WriteRefused, match="no sheet"):
            writer.set_cell("Missing", "A1", formula=None, value="1")
        with pytest.raises(WriteRefused, match="not a cell reference"):
            writer.set_cell("Model", "1A", formula=None, value="1", create=True)


def test_a_edits_record_before_and_after_for_the_changeset() -> None:
    with tempfile.TemporaryDirectory() as folder:
        writer = WorkbookWriter(_host(folder))
        edit = writer.set_cell("Model", "B3", formula="=B2*2", value="200")
        assert edit.before_formula == "=B2*(1+B1)"
        assert edit.formula == "=B2*2"
        assert [e.ref for e in writer.edits] == ["B3"]


def _share_row(path: Path) -> Path:
    """Rewrite B3:D3 of a generated workbook as one shared group —
    the layout Excel itself produces when a formula is filled right."""
    import re

    with zipfile.ZipFile(path) as archive:
        members = {i.filename: archive.read(i.filename) for i in archive.infolist()}
        order = archive.infolist()
    xml = members["xl/worksheets/sheet1.xml"].decode("utf-8")
    xml = re.sub(
        r'(<c r="B3"[^>]*>)<f>[^<]*</f>',
        r'\1<f t="shared" ref="B3:D3" si="0">B2*(1+B1)</f>',
        xml,
    )
    for ref in ("C3", "D3"):
        xml = re.sub(
            rf'(<c r="{ref}"[^>]*>)<f>[^<]*</f>',
            r'\1<f t="shared" si="0"/>',
            xml,
        )
    members["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
    out = path.with_name("shared.xlsx")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in order:
            archive.writestr(item, members[item.filename])
    return out


def _shared_host(folder: str) -> Path:
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["B1"], sheet["C1"], sheet["D1"] = 0.05, 0.06, 0.07
    sheet["B2"], sheet["C2"], sheet["D2"] = 100, 200, 300
    sheet["B3"] = "=B2*(1+B1)"
    sheet["C3"] = "=C2*(1+C1)"
    sheet["D3"] = "=D2*(1+D1)"
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return _share_row(path)


def test_a3_editing_a_shared_member_unshares_the_whole_group() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _shared_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "C3", formula="=C2*2", value="400")
        out = writer.save(Path(folder) / "out.xlsx")

        seen = load_workbook(out)
        assert seen["Model"]["C3"].value == "=C2*2"
        #: The siblings keep computing what the master declared,
        #: translated to their own positions.
        assert seen["Model"]["B3"].value == "=B2*(1+B1)"
        assert seen["Model"]["D3"].value == "=D2*(1+D1)"


def test_a3_the_master_itself_can_be_edited() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _shared_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "B3", formula="=B2", value="100")
        out = writer.save(Path(folder) / "out.xlsx")
        seen = load_workbook(out)
        assert seen["Model"]["B3"].value == "=B2"
        assert seen["Model"]["C3"].value == "=C2*(1+C1)"
        assert seen["Model"]["D3"].value == "=D2*(1+D1)"


# --- F1: creation — a cell that does not exist yet, materialized ---


def _gapped_host(folder: str) -> Path:
    """Rows 1–3 and 5 exist; row 4 does not. Row 5 has A5 and C5 with
    B5 absent — the two creation positions the XML makes awkward."""
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Rate"
    sheet["B1"] = 0.08
    sheet["A2"] = "Base"
    sheet["B2"] = 100
    sheet["A3"] = "Grown"
    sheet["B3"] = "=B2*(1+B1)"
    sheet["A5"] = "Tail"
    sheet["C5"] = 7
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return path


def test_f1_creating_a_cell_appends_to_its_existing_row() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        edit = writer.set_cell("Model", "C2", formula=None, value="55", create=True)
        out = writer.save(Path(folder) / "out.xlsx")

        assert edit.created is True
        assert edit.before_formula is None and edit.before_value is None
        before, after = _member_bytes(host), _member_bytes(out)
        changed = {name for name in before if before[name] != after.get(name)}
        assert changed == {"xl/worksheets/sheet1.xml"}

        seen = load_workbook(out)
        assert seen["Model"]["C2"].value == 55
        assert seen["Model"]["A2"].value == "Base"
        assert seen["Model"]["B2"].value == 100


def test_f1_creating_a_cell_between_its_row_neighbours() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "B5", formula=None, value="3.5", create=True)
        out = writer.save(Path(folder) / "out.xlsx")

        #: Column order inside the row must hold — B5 lands between
        #: A5 and C5, not appended after them.
        with zipfile.ZipFile(out) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        row = xml[xml.index('<row r="5"') :]
        assert row.index('r="A5"') < row.index('r="B5"') < row.index('r="C5"')

        seen = load_workbook(out)
        assert seen["Model"]["B5"].value == 3.5
        assert seen["Model"]["A5"].value == "Tail"
        assert seen["Model"]["C5"].value == 7


def test_f1_creating_a_cell_materializes_its_missing_row_in_order() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "B4", formula=None, value="41", create=True)
        writer.set_cell("Model", "B9", formula=None, value="91", create=True)
        out = writer.save(Path(folder) / "out.xlsx")

        with zipfile.ZipFile(out) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        #: Row 4 between 3 and 5; row 9 after 5 — ascending, as Excel
        #: writes them.
        assert (
            xml.index('<row r="3"')
            < xml.index('<row r="4"')
            < xml.index('<row r="5"')
            < xml.index('<row r="9"')
        )

        seen = load_workbook(out)
        assert seen["Model"]["B4"].value == 41
        assert seen["Model"]["B9"].value == 91
        assert seen["Model"]["A5"].value == "Tail"


def test_f1_created_text_reads_back_as_text() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "D1", formula=None, value="see note", create=True)
        out = writer.save(Path(folder) / "out.xlsx")
        seen = load_workbook(out)
        assert seen["Model"]["D1"].value == "see note"


def test_f1_created_formula_reads_back_in_both_readers() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell(
            "Model", "B6", formula="=B2*2", value="200", create=True
        )
        out = writer.save(Path(folder) / "out.xlsx")

        seen = load_workbook(out)
        assert seen["Model"]["B6"].value == "=B2*2"
        with zipfile.ZipFile(out) as archive:
            assert "xl/calcChain.xml" not in archive.namelist()

        book = read_workbook(str(out))
        cell = book.cells["Model!B6"]
        assert cell.formula == "=B2*2"
        assert str(cell.value) == "200"


def test_f1_creation_expands_the_dimension_and_row_spans() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        writer.set_cell("Model", "E1", formula=None, value="9", create=True)
        out = writer.save(Path(folder) / "out.xlsx")

        with zipfile.ZipFile(out) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        import re as _re

        dimension = _re.search(r'<dimension ref="([^"]+)"', xml)
        assert dimension is not None and dimension.group(1) == "A1:E5"
        seen = load_workbook(out)
        assert seen["Model"]["E1"].value == 9


def test_f1_creation_updates_the_rows_spans_when_excel_wrote_them() -> None:
    """openpyxl writes no spans; Excel does. Injected here, the way
    the shared-formula tests inject Excel's own layout."""
    import re

    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        with zipfile.ZipFile(host) as archive:
            members = {i.filename: archive.read(i.filename) for i in archive.infolist()}
            order = archive.infolist()
        xml = members["xl/worksheets/sheet1.xml"].decode("utf-8")
        xml = re.sub(r'<row r="1"', '<row r="1" spans="1:2"', xml)
        members["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
        spanned = Path(folder) / "spanned.xlsx"
        with zipfile.ZipFile(spanned, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in order:
                archive.writestr(item, members[item.filename])

        writer = WorkbookWriter(spanned)
        writer.set_cell("Model", "E1", formula=None, value="9", create=True)
        out = writer.save(Path(folder) / "out.xlsx")

        with zipfile.ZipFile(out) as archive:
            written = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        row_one = re.search(r'<row r="1"[^>]*spans="([^"]+)"', written)
        assert row_one is not None and row_one.group(1) == "1:5"
        seen = load_workbook(out)
        assert seen["Model"]["E1"].value == 9
        assert seen["Model"]["A1"].value == "Rate"


def test_f1_without_create_the_writer_still_refuses() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _gapped_host(folder)
        writer = WorkbookWriter(host)
        with pytest.raises(WriteRefused, match="create=True"):
            writer.set_cell("Model", "B4", formula=None, value="41")
