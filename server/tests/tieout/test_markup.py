"""The marked-up model, tested against the founder's spec (§4 of
swens-product-plan.md), sentence by sentence.

« Your own model handed back to you, with every problem cell coloured
in and a note stuck on it saying what's wrong. Plus a first sheet
listing all findings so you can sort them and tick them off. ...
nothing in the model is altered. Not one formula, not one number.
Only colour and notes. It's a copy, with a different filename, so the
original is never at risk. »
"""

import io
import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook as Book
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

from polar.tieout.markup import (
    MarkupFinding,
    MarkupRefused,
    marked_up_copy,
    marked_up_name,
)


def _host(folder: str) -> Path:
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Rate"
    sheet["B1"] = 0.08
    sheet["A2"] = "Base"
    sheet["B2"] = 100
    sheet["A3"] = "Grown"
    sheet["B3"] = "=B2*(1+B1)"
    sheet["A4"] = "Total"
    sheet["B4"] = "=SUM(B2:B3)"
    other = book.create_sheet("Notes")
    other["A1"] = "untouched"
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return path


FINDINGS = [
    MarkupFinding(
        severity="error",
        sheet="Model",
        ref="B3",
        text="Typed over: the row's own formula belongs here.",
    ),
    MarkupFinding(
        severity="smell",
        sheet="Model",
        ref="B4",
        text="The sum skips a row the structure includes.",
    ),
]


def test_nothing_in_the_model_is_altered() -> None:
    """Not one formula, not one number — checked cell by cell."""
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(host.read_bytes(), FINDINGS)

        original = load_workbook(host)
        marked = load_workbook(io.BytesIO(copy))
        assert marked.sheetnames == ["Findings", *original.sheetnames]
        for name in original.sheetnames:
            held = {
                c.coordinate: c.value
                for row in original[name].iter_rows()
                for c in row
                if c.value is not None
            }
            now = {
                c.coordinate: c.value
                for row in marked[name].iter_rows()
                for c in row
                if c.value is not None
            }
            assert held == now


def test_untouched_members_keep_their_exact_bytes() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(host.read_bytes(), FINDINGS)

        with zipfile.ZipFile(host) as archive:
            before = {i.filename: archive.read(i.filename) for i in archive.infolist()}
        with zipfile.ZipFile(io.BytesIO(copy)) as archive:
            after = {i.filename: archive.read(i.filename) for i in archive.infolist()}
        changed = {name for name in before if before[name] != after.get(name)}
        #: The parts the markup must touch, and no other: the touched
        #: sheet (colour + note anchors), the styles (new fills), the
        #: workbook + rels + content types (the new first sheet), and
        #: the Notes sheet is NOT among them.
        assert "xl/worksheets/sheet2.xml" not in changed
        assert "xl/sharedStrings.xml" not in changed
        assert changed <= {
            "xl/worksheets/sheet1.xml",
            "xl/styles.xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "[Content_Types].xml",
        }


def test_problem_cells_are_coloured_by_severity_with_a_note() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(host.read_bytes(), FINDINGS)
        marked = load_workbook(io.BytesIO(copy))

        error_cell = marked["Model"]["B3"]
        assert error_cell.fill.fgColor.rgb == "FFFFC7CE"
        assert error_cell.comment is not None
        assert "Typed over" in error_cell.comment.text
        assert error_cell.comment.author == "Swens"

        smell_cell = marked["Model"]["B4"]
        assert smell_cell.fill.fgColor.rgb == "FFFFEB9C"
        assert smell_cell.comment is not None
        assert "skips a row" in smell_cell.comment.text


def test_sheet_one_is_the_sortable_list() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(host.read_bytes(), FINDINGS)
        marked = load_workbook(io.BytesIO(copy))

        listing = marked["Findings"]
        assert [c.value for c in listing[1]] == [
            "Severity",
            "Sheet",
            "Cell",
            "What's wrong",
            "Notes",
            "Done",
        ]
        assert [c.value for c in listing[2]][:4] == [
            "Error",
            "Model",
            "B3",
            "Typed over: the row's own formula belongs here.",
        ]
        assert listing[2][4].value is None  # their own blank columns
        assert listing[2][5].value is None
        assert listing.auto_filter.ref == "A1:F3"


def test_a_finding_on_an_absent_cell_still_gets_its_colour() -> None:
    """Skipped-cell findings point at cells the file never wrote."""
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(
            host.read_bytes(),
            [
                MarkupFinding(
                    severity="error",
                    sheet="Model",
                    ref="B5",
                    text="Skipped: the column's run breaks here.",
                )
            ],
        )
        marked = load_workbook(io.BytesIO(copy))
        cell = marked["Model"]["B5"]
        assert cell.value is None  # still not one number added
        assert cell.fill.fgColor.rgb == "FFFFC7CE"
        assert cell.comment is not None


def test_two_findings_on_one_cell_share_one_note_worst_colour() -> None:
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        copy = marked_up_copy(
            host.read_bytes(),
            [
                MarkupFinding(
                    severity="smell", sheet="Model", ref="B3", text="First thing."
                ),
                MarkupFinding(
                    severity="error", sheet="Model", ref="B3", text="Second thing."
                ),
            ],
        )
        marked = load_workbook(io.BytesIO(copy))
        cell = marked["Model"]["B3"]
        assert cell.fill.fgColor.rgb == "FFFFC7CE"  # error wins
        assert cell.comment is not None
        assert "First thing." in cell.comment.text
        assert "Second thing." in cell.comment.text


def test_scoped_defined_names_survive_the_new_first_sheet() -> None:
    """localSheetId is an index into sheet order; prepending a sheet
    must not silently re-point every scoped name."""
    with tempfile.TemporaryDirectory() as folder:
        book = Book()
        sheet = book.active
        sheet.title = "Model"
        sheet["A1"] = "Rate"
        sheet["B1"] = 0.08
        sheet.defined_names.add(DefinedName("TheRate", attr_text="Model!$B$1"))
        path = Path(folder) / "host.xlsx"
        book.save(path)

        copy = marked_up_copy(
            path.read_bytes(),
            [MarkupFinding(severity="error", sheet="Model", ref="B1", text="x")],
        )
        marked = load_workbook(io.BytesIO(copy))
        #: The name is still scoped to Model — not to Findings, which
        #: now sits at the index the name used to carry.
        assert "TheRate" in marked["Model"].defined_names
        assert "TheRate" not in marked["Findings"].defined_names


def test_a_sheet_that_already_carries_comments_is_refused() -> None:
    """Merging into an existing comments part is not built; the
    refusal says so rather than dropping notes silently."""
    with tempfile.TemporaryDirectory() as folder:
        host = _host(folder)
        with zipfile.ZipFile(host) as archive:
            order = archive.infolist()
            members = {i.filename: archive.read(i.filename) for i in order}
        xml = members["xl/worksheets/sheet1.xml"].decode("utf-8")
        members["xl/worksheets/sheet1.xml"] = xml.replace(
            "</worksheet>", '<legacyDrawing r:id="rId9"/></worksheet>'
        ).encode("utf-8")
        commented = Path(folder) / "commented.xlsx"
        with zipfile.ZipFile(commented, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in order:
                archive.writestr(item, members[item.filename])

        with pytest.raises(MarkupRefused, match="Model"):
            marked_up_copy(commented.read_bytes(), FINDINGS)


def test_the_copy_gets_a_different_filename() -> None:
    assert marked_up_name("Project Alpha v14.xlsx") == (
        "Project Alpha v14 — marked up.xlsx"
    )
    assert marked_up_name("model") == "model — marked up.xlsx"
