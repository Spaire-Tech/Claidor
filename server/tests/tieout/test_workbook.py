"""Reading a model as labelled rows rather than as coordinates."""

from pathlib import Path

import pytest

from polar.tieout.model import read_outputs
from polar.tieout.provenance import chain, outputs_from_workbook, verify_outputs
from polar.tieout.workbook import precedents_of, read_workbook, references_of

CASCADE = Path(__file__).resolve().parents[2] / "scripts" / "cascade"
MODEL = str(CASCADE / "cascade_model.xlsx")


@pytest.fixture(scope="module")
def book():
    return read_workbook(MODEL)


def test_a_cell_is_named_by_its_row_and_its_period(book) -> None:
    """« Model!D26 » means nothing; « FY2025A Adjusted EBITDA » is the
    same shape of name an Outputs row carries, which is what lets one
    matcher serve a published interface and a raw workbook alike."""
    assert book.get("Model!D26").name == "FY2025A Adjusted EBITDA"
    assert book.get("Model!D16").name == "FY2025A Reported EBITDA"
    assert book.get("Comps!F13").name == "EV / EBITDA Median"


def test_a_bridge_under_a_forecast_grid_does_not_inherit_its_years(book) -> None:
    """DCF rows 6-9 are a five-year series under FY2026E…FY2030E; rows
    13-21 are a valuation bridge in one column. Reading the header down
    the page makes the enterprise value claim to be about FY2026E."""
    assert book.get("DCF!B9").name == "FY2026E Present value of FCF"
    assert book.get("DCF!B16").name == "Enterprise value"
    assert book.get("DCF!B21").name == "Implied value per share"


def test_an_indented_row_belongs_to_the_one_above(book) -> None:
    assert book.get("Model!D27").name == "FY2025A Adjusted EBITDA margin %"


def test_a_cell_that_only_restates_another_is_an_alias(book) -> None:
    """`Comps!B18` is `=Model!D26`. Left in the candidate set it gives
    « FY2025A adjusted EBITDA » two homes with one value, which is the
    ambiguity the matcher is built to refuse."""
    assert book.get("Comps!B18").alias_of == "Model!D26"
    assert book.get("Model!D30").alias_of == "Model!D26"
    assert book.get("Model!D26").alias_of is None
    assert book.get("Comps!B20").alias_of is None


def test_precedents_come_from_the_formula() -> None:
    assert precedents_of("=D16+D24", "Model") == ("Model!D16", "Model!D24")
    assert precedents_of("=SUM(D20:D23)", "Model") == tuple(
        f"Model!D{row}" for row in range(20, 24)
    )
    assert precedents_of("=B16/Model!D26", "DCF") == ("DCF!B16", "Model!D26")
    assert precedents_of("=1/(1+Assumptions!$B$19)^B7", "DCF") == (
        "Assumptions!B19",
        "DCF!B7",
    )


def test_a_hardcoded_input_is_the_end_of_the_chain(book) -> None:
    assert book.get("Assumptions!B19").hardcoded is True
    assert book.get("Model!D26").hardcoded is False
    assert chain(book, "Assumptions!B19") == "Assumptions!B19 = 0.098 (input)"


def test_the_chain_reads_as_a_sentence(book) -> None:
    assert chain(book, "Model!D26") == (
        "Model!D26 = 48.9, =D16+D24 = Reported EBITDA 41.2, Total adjustments 7.7"
    )


def test_the_chain_drops_what_the_reader_already_knows(book) -> None:
    """Model!E37's five precedents are all « FY2026E unlevered free cash
    flow » something. Printed in full that is the same eight words six
    times over, which buries the four that differ."""
    assert chain(book, "Model!E37").endswith(
        "= EBIT 40.69, Taxes on EBIT -10.17, Add back D&A 10.49, "
        "Capital expenditure -9.21, and 1 more"
    )


def test_the_outputs_tab_is_a_link_in_the_chain_not_the_end_of_it(book) -> None:
    """Four of Cascade's twenty-three source references point one row
    above the figure they name — the values are right, the pointers are
    stale, and a banker sent to Model!D25 finds an empty cell."""
    problems = {p.ref: p for p in verify_outputs(read_outputs(MODEL), book)}
    assert set(problems) == {"O5", "O6", "O8", "O10"}
    assert problems["O5"].claimed == "Model!D25"
    assert problems["O5"].found is None
    assert problems["O5"].actual == "Model!D26"
    assert problems["O6"].claimed == "Model!D26"
    assert problems["O6"].actual == "Model!D27"


def test_a_reference_citing_several_cells_is_not_a_claim_about_one(book) -> None:
    """O11's source is « Model!D6, I6 » — a CAGR derived from two cells.
    Reading the first of them and finding revenue rather than a growth
    rate reports a correct row as broken."""
    assert "O11" not in {p.ref for p in verify_outputs(read_outputs(MODEL), book)}


def test_a_repair_needs_the_name_as_well_as_the_value(book) -> None:
    """Four cells in the model hold 48.9. Only one of them is called
    adjusted EBITDA, and a repair proposed on the value alone would be the
    same mistake as a checker that links on values."""
    problem = next(
        p for p in verify_outputs(read_outputs(MODEL), book) if p.ref == "O5"
    )
    assert "named for it" in problem.how


def test_the_workbook_offers_far_more_than_the_interface_does(book) -> None:
    assert len(read_outputs(MODEL)) == 23
    assert len(outputs_from_workbook(book)) > 250
    assert not any(
        candidate.ref.startswith("Outputs!")
        for candidate in outputs_from_workbook(book)
    )


# --- what a formula reads that is not a plain cell reference -------------
#
# Every case below was found by running the parser over financial models
# published by economic regulators — `scripts/model_corpus.py` fetches
# them, `scripts/formula_coverage.py` measures. Before this, 2,602 of
# 59,705 formulas in those files lost a precedent and said nothing: the
# chain came back short by exactly the input that decided the answer, and
# looked complete. On the 2015 Ofgem transmission model that was 11.7% of
# every formula in the file.


@pytest.fixture
def names():
    from polar.tieout.workbook import Names

    return Names(
        book={
            "Tax_Rate": "Assumptions!$B$4",
            "Revenue": "Model!$D$6:$F$6",
            "Deleted": "#REF!",
            "Switch": "=OFFSET(Model!$A$1,1,1)",
        },
        sheet={("DCF", "Tax_Rate"): "DCF!$Z$9"},
        extent={"Model": (20, 8), "Inflation": (352, 14)},
    )


def test_a_defined_name_is_a_precedent(names) -> None:
    """`=B4*Tax_Rate` used to return `('Model!B4',)` — a chain missing an
    input, presented as the whole answer."""
    read = references_of("=B4*Tax_Rate", "Model", names)
    assert read.refs == ("Model!B4", "Assumptions!B4")
    assert read.unresolved == ()


def test_sheet_scope_beats_workbook_scope(names) -> None:
    """The same name means different cells on different sheets, and Excel
    resolves the sheet's own first. One lookup table would answer
    confidently and wrongly, which is worse than the silence it replaced —
    the 2026 Ofgem distribution model carries 594 sheet-scoped names."""
    assert references_of("=Tax_Rate", "Model", names).refs == ("Assumptions!B4",)
    assert references_of("=Tax_Rate", "DCF", names).refs == ("DCF!Z9",)


def test_a_name_pointing_at_a_deleted_row_is_a_finding(names) -> None:
    """Nine of these sit in a published Ofgem model. Not a parse failure —
    a defect in the workbook, and a person wants to be told."""
    read = references_of("=Deleted+1", "Model", names)
    assert read.refs == ()
    assert read.unresolved == (("Deleted", "a defined name pointing at #REF!"),)


def test_a_name_that_is_a_formula_says_so(names) -> None:
    read = references_of("=Switch", "Model", names)
    assert read.refs == ()
    assert "formula rather than a cell" in read.unresolved[0][1]


def test_a_whole_column_expands_against_the_sheet_it_names(names) -> None:
    """`AVERAGEIFS('Monthly Inflation'!$M:$M, ...)` appears 570 times in
    one real model. Expanded against the used extent, not the format's
    1,048,576 rows."""
    read = references_of("=SUM(Model!A:A)", "Other", names)
    assert read.refs == tuple(f"Model!A{row}" for row in range(1, 21))


def test_the_range_cap_says_when_it_bit(names) -> None:
    """A cap that truncates silently is the same failure this file is
    about, moved from « dropped » to « quietly shortened »."""
    read = references_of("=SUM(Inflation!M:M)", "Other", names)
    assert len(read.refs) == 200
    assert read.unresolved == (
        ("Inflation!M:M", "352 cells, of which the first 200 were followed"),
    )


def test_a_range_may_carry_its_sheet_on_both_ends(names) -> None:
    """`InputSummary!AR61:'InputSummary'!AR67` is legal and appears in the
    Ofgem distribution model. Read as a name, it produced the sentence
    « a name this workbook does not define » — which is not merely
    unresolved, it is false about the workbook."""
    read = references_of("=SUM(Model!A1:'Model'!A3)", "Other", names)
    assert read.refs == ("Model!A1", "Model!A2", "Model!A3")
    assert read.unresolved == ()


def test_another_workbook_is_named_rather_than_guessed_at(names) -> None:
    """Both spellings. The quoted form used to produce a reference to a
    sheet that does not exist in this workbook, and nothing downstream was
    told — a dangling precedent is worse than a missing one."""
    for formula in ("=[1]Sheet1!A1", "='[1]Cash Flow'!B4"):
        read = references_of(formula, "Model", names)
        assert read.refs == (), formula
        assert "another workbook" in read.unresolved[0][1], formula


def test_a_table_reference_is_named_rather_than_dropped(names) -> None:
    read = references_of("=SUM(Tbl[Amount])", "Model", names)
    assert read.refs == ()
    assert "table reference" in read.unresolved[0][1]


def test_without_a_name_map_a_name_is_reported_not_swallowed() -> None:
    """A legacy `.xls` has no defined names to give. The honest answer is
    a sentence saying so, not an empty tuple."""
    read = references_of("=B4*Tax_Rate", "Model")
    assert read.refs == ("Model!B4",)
    assert read.unresolved == (
        ("Tax_Rate", "a defined name, and this workbook's names were not read"),
    )


# --- Reading a model laid out the way a regulator lays one out ----------
#
# Every case below was found by running the reader against Ofgem's price
# control financial model and its published direction document — real
# files, neither written here. Each one silently produced *nothing* rather
# than something wrong, which is why none of them showed up in a test until
# somebody measured a pair.


def _sheet(rows: list[list[object]], formulas: dict[str, str] | None = None):
    """One sheet's grid — formulas and cached values — from row lists."""
    from openpyxl import Workbook as Book

    from polar.tieout.workbook import _grid_of

    values = Book()
    written = Book()
    for index, row in enumerate(rows, start=1):
        for column, cell in enumerate(row, start=1):
            values.active.cell(index, column).value = cell
            written.active.cell(index, column).value = cell
    for ref, formula in (formulas or {}).items():
        written.active[ref] = formula
    return _grid_of(written.active, values.active)


def test_a_label_may_sit_past_column_d() -> None:
    """Ofgem's model indents through B, C and D and names in **E**.

    With the search stopping at D, all 26,392 cells in that workbook came
    back unnamed, which empties the tie-out and the grounding both.
    """
    from polar.tieout.workbook import _label_column

    grid = _sheet(
        [
            ["Depn", None, None, None, None, None],
            [None, "NGET TO", None, None, None, None],
            [None, None, None, None, "Slow money", "£m 09/10 prices"],
            [None, None, None, None, "Net adjustments", "£m 09/10 prices"],
            [None, None, None, None, "Pre-RIIO net RAV additions", "£m 09/10 prices"],
        ]
    )
    assert _label_column(grid, grid.last_row, grid.last_column) == 5


def test_the_label_column_is_the_one_with_the_most_different_things_to_say() -> None:
    """The units column has as much text in it and says one thing.

    On one sheet of the real model `£m 09/10 prices` appears on 247 rows
    against 251 parameter names beside it — indistinguishable by volume,
    and 10 distinct values against 143.
    """
    from polar.tieout.workbook import _label_column

    grid = _sheet(
        [["Allowed revenue", "£m", None], ["Actual opex", "£m", None]] * 6
    )
    assert _label_column(grid, grid.last_row, grid.last_column) == 1


def test_a_label_that_is_a_formula_still_names_its_row() -> None:
    """A model that mirrors its input sheet writes `=Input!E31` where the
    name goes. The formula text is refused; the words it produced are not.
    """
    from polar.tieout.workbook import _shown

    grid = _sheet(
        [["Legacy price control adjustments", 95.5]],
        formulas={"A1": "=Input!A31"},
    )
    assert _shown(grid, 1, 1) == "Legacy price control adjustments"


def test_a_numeric_formula_is_still_not_a_label() -> None:
    """The reason `_label` refuses formulas in the first place: a column of
    arithmetic must not name every figure beside it."""
    from polar.tieout.workbook import _shown

    grid = _sheet([[42, 7]], formulas={"A1": "=B1*6"})
    assert _shown(grid, 1, 1) is None


def test_a_date_header_names_the_period_it_heads() -> None:
    """Ofgem writes `2017-03-31` where a banker writes `FY2017A`.

    Unread, every one of eight year columns on a row carried the same name,
    and a figure naming the row matched whichever column happened to hold a
    typed value.
    """
    import datetime

    from polar.tieout.workbook import _period_label

    assert _period_label(datetime.datetime(2017, 3, 31)) == "FY2017"
    assert _period_label(datetime.date(2015, 12, 31)) == "FY2015"
    assert _period_label("not a date") is None
    assert _period_label(42) is None


def test_the_words_beside_a_name_are_kept_apart_from_it() -> None:
    """A licence term three columns along can be the only thing telling two
    identically-named rows apart — `LAR` against `SOLAR`."""
    import openpyxl

    from polar.tieout.workbook import read_workbook

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "NGET TO"
    sheet["A1"], sheet["B1"], sheet["C1"] = "Period", "Term", "Units"
    sheet["A2"], sheet["B2"], sheet["C2"] = "Legacy adjustment", "LAR", "£m"
    sheet["D2"] = 95.5
    sheet["A3"], sheet["B3"], sheet["C3"] = "Actual opex", "SOACO", "£m"
    sheet["D3"] = 15.7
    sheet["A4"], sheet["B4"], sheet["C4"] = "Actual capex", "SOANC", "£m"
    sheet["D4"] = 32.2

    import io

    payload = io.BytesIO()
    book.save(payload)
    path = "/tmp/claude-0/_tags.xlsx"
    with open(path, "wb") as handle:
        handle.write(payload.getvalue())

    read = read_workbook(path)
    cell = read.get("NGET TO!D2")
    assert cell is not None
    assert cell.row_label == "Legacy adjustment"
    assert "LAR" in cell.row_tags
