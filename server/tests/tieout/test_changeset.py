"""swens-plan Track F, F2: the changeset, tested.

The contract: every correction recorded (before, after, why, who);
applied atomically or not at all; apply → undo byte-identical at the
member level; every write followed by a full re-read compare and a
re-audit as a hard gate — a write that changed anything it was not
asked to, or that would create new damage, is refused and nothing is
released; and a repair that leaves siblings wrong is « incomplete » —
its own state, not success, not failure.
"""

import io
import tempfile
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook as Book
from openpyxl import load_workbook

from polar.tieout import writer as writer_module
from polar.tieout.changeset import Correction, apply_corrections


def _typed_over_host(folder: str) -> Path:
    """B3 holds a typed 150 where the model's own growth formula
    belongs — the classic determined correction."""
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Rate"
    sheet["B1"] = 0.08
    sheet["A2"] = "Base"
    sheet["B2"] = 100
    sheet["A3"] = "Grown"
    sheet["B3"] = 150
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return path


def _frozen_pair_host(folder: str) -> Path:
    """Two frozen references in one row; fixing only one is the
    incomplete-repair case."""
    book = Book()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Costs"
    sheet["B1"] = "#REF!"
    sheet["C1"] = "#REF!"
    sheet["A2"] = "Rate"
    sheet["B2"] = 0.05
    path = Path(folder) / "host.xlsx"
    book.save(path)
    return path


def _members(data: bytes) -> tuple[dict[str, bytes], list[str]]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        order = [item.filename for item in archive.infolist()]
        return {name: archive.read(name) for name in order}, order


def test_f2_an_applied_correction_carries_its_full_record() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over; the block's own pattern declares the formula",
                )
            ],
            who="founder",
        )

        assert result.state == "applied"
        assert result.payload is not None
        assert result.who == "founder"
        assert result.edits[0].before_value == "150"
        assert result.edits[0].formula == "=B2*(1+B1)"
        assert "declares the formula" in result.corrections[0].why

        seen = load_workbook(io.BytesIO(result.payload))
        assert seen["Model"]["B3"].value == "=B2*(1+B1)"
        assert seen["Model"]["B2"].value == 100


def test_f2_apply_then_undo_is_byte_identical() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over",
                )
            ],
            who="founder",
        )
        assert result.state == "applied"
        assert _members(result.undo()) == _members(payload)


def test_f2_a_created_cells_undo_is_removal() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="D1",
                    formula=None,
                    value="see note",
                    why="the marked-up copy's note",
                    create=True,
                )
            ],
            who="founder",
        )
        assert result.state == "applied"
        assert result.edits[0].created is True
        seen = load_workbook(io.BytesIO(result.payload or b""))
        assert seen["Model"]["D1"].value == "see note"
        assert _members(result.undo()) == _members(payload)


def test_f2_refusal_is_atomic_and_releases_nothing() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over",
                ),
                Correction(
                    sheet="Missing",
                    ref="A1",
                    formula=None,
                    value="1",
                    why="a typo in the address",
                ),
            ],
            who="founder",
        )
        assert result.state == "refused"
        assert result.payload is None
        assert "no sheet" in result.reason
        with pytest.raises(ValueError, match="nothing to undo"):
            result.undo()


def test_f2_a_sabotaged_write_is_caught_by_its_own_gate(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The standing argument for the gate: a construct the writer
    cannot see will exist again. Simulated here as a save that also
    corrupts a cell nobody asked it to touch."""
    real_save = writer_module.WorkbookWriter.save

    def crooked(self: writer_module.WorkbookWriter, out: Path) -> Path:
        path = real_save(self, out)
        with zipfile.ZipFile(path) as archive:
            order = archive.infolist()
            members = {i.filename: archive.read(i.filename) for i in order}
        xml = members["xl/worksheets/sheet1.xml"].decode("utf-8")
        members["xl/worksheets/sheet1.xml"] = xml.replace(
            "<v>100</v>", "<v>999</v>", 1
        ).encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in order:
                archive.writestr(item, members[item.filename])
        return path

    monkeypatch.setattr(writer_module.WorkbookWriter, "save", crooked)

    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over",
                )
            ],
            who="founder",
        )
        assert result.state == "refused"
        assert result.payload is None
        assert "Model!B2" in result.reason


def test_f2_incomplete_repair_is_its_own_state() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _frozen_pair_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B1",
                    formula=None,
                    value="0",
                    why="frozen reference; the confirmed source says zero",
                    rule="error-value",
                )
            ],
            who="founder",
        )

        #: Applied — the fix itself landed — but not success: the same
        #: rule still fires beside it, and the state says so.
        assert result.state == "incomplete"
        assert result.payload is not None
        assert any("C1" in ref for ref in result.still_wrong)
        seen = load_workbook(io.BytesIO(result.payload))
        assert seen["Model"]["B1"].value == 0
        assert seen["Model"]["C1"].value == "#REF!"


def test_f2_the_record_is_serializable_for_the_changes_ui() -> None:
    with tempfile.TemporaryDirectory() as folder:
        payload = _typed_over_host(folder).read_bytes()
        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over",
                )
            ],
            who="founder",
        )
        record = result.record()
        assert record["state"] == "applied"
        assert record["who"] == "founder"
        assert record["edits"][0]["before"] == {"formula": None, "value": "150"}
        assert record["edits"][0]["after"] == {
            "formula": "=B2*(1+B1)",
            "value": "108",
        }
        assert record["edits"][0]["created"] is False
        import json

        json.dumps(record)  # the UI contract: plain JSON, no surprises


def test_f2_undo_restores_the_dropped_calc_chain() -> None:
    """openpyxl writes no calcChain; Excel does, and a formula edit
    drops it. Injected here so undo's restoration is actually
    exercised: the original — chain, content types and all — must
    come back member-for-member."""
    with tempfile.TemporaryDirectory() as folder:
        host = _typed_over_host(folder)
        with zipfile.ZipFile(host) as archive:
            order = archive.infolist()
            members = {i.filename: archive.read(i.filename) for i in order}
        members["xl/calcChain.xml"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<calcChain xmlns="http://schemas.openxmlformats.org/'
            b'spreadsheetml/2006/main"><c r="B3" i="1"/></calcChain>'
        )
        kinds = members["[Content_Types].xml"].decode("utf-8")
        members["[Content_Types].xml"] = kinds.replace(
            "</Types>",
            '<Override PartName="/xl/calcChain.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
            "</Types>",
        ).encode("utf-8")
        chained = Path(folder) / "chained.xlsx"
        with zipfile.ZipFile(chained, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in order:
                archive.writestr(item, members[item.filename])
            archive.writestr("xl/calcChain.xml", members["xl/calcChain.xml"])
        payload = chained.read_bytes()

        result = apply_corrections(
            payload,
            [
                Correction(
                    sheet="Model",
                    ref="B3",
                    formula="=B2*(1+B1)",
                    value="108",
                    why="typed over",
                )
            ],
            who="founder",
        )
        assert result.state == "applied"
        corrected_members, _ = _members(result.payload or b"")
        assert "xl/calcChain.xml" not in corrected_members
        assert _members(result.undo()) == _members(payload)


def test_f2_an_array_formula_elsewhere_does_not_refuse_the_gate() -> None:
    """openpyxl hands back a fresh ArrayFormula object on every load;
    compared by identity it looks changed forever — the false alarm a
    real RIIO-3 file exposed. The gate compares by content."""
    from openpyxl.worksheet.formula import ArrayFormula

    with tempfile.TemporaryDirectory() as folder:
        book = Book()
        sheet = book.active
        sheet.title = "Model"
        sheet["A1"] = 1
        sheet["A2"] = 2
        sheet["B1"] = ArrayFormula("B1", "=SUM(A1:A2*1)")
        sheet["C1"] = 150
        path = Path(folder) / "host.xlsx"
        book.save(path)

        result = apply_corrections(
            path.read_bytes(),
            [
                Correction(
                    sheet="Model",
                    ref="C1",
                    formula=None,
                    value="151",
                    why="the confirmed source moved",
                )
            ],
            who="founder",
        )
        assert result.state == "applied"
        assert _members(result.undo()) == _members(path.read_bytes())
