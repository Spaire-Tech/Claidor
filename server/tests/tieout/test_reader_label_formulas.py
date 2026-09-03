"""Formulas in the label column reach the prescan, the text rules and
the counts (`docs/pierce/reader-label-formulas.md`): the reader keeps
them as `label_cells`, apart from the numeric grid."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as XlsxWorkbook

from polar.tieout.audit import audit
from polar.tieout.recalc import Route
from polar.tieout.recalc.denylist import prescan, route_for
from polar.tieout.workbook import read_workbook


def _model(path: Path, label_formula: str) -> Path:
    """A three-year sheet whose second row label is built by formula."""
    book = XlsxWorkbook()
    sheet = book.active
    sheet.title = "Model"
    sheet["A1"] = "Line"
    sheet["B1"] = "FY2024"
    sheet["C1"] = "FY2025"
    sheet["D1"] = "FY2026"
    sheet["A2"] = "Revenue"
    sheet["B2"] = 100
    sheet["C2"] = "=B2*1.1"
    sheet["D2"] = "=C2*1.1"
    sheet["A3"] = label_formula
    sheet["B3"] = "=B2*0.4"
    sheet["C3"] = "=C2*0.4"
    sheet["D3"] = "=D2*0.4"
    sheet["A4"] = "Margin"
    sheet["B4"] = "=B3/B2"
    sheet["C4"] = "=C3/C2"
    sheet["D4"] = "=D3/D2"
    book.save(path)
    return path


def test_a_label_formula_is_kept_apart_from_the_grid(tmp_path: Path) -> None:
    book = read_workbook(str(_model(tmp_path / "m.xlsx", '="Cost of "&A2')))
    assert "Model!A3" not in book.cells
    assert "Model!A3" in book.label_cells
    assert book.label_cells["Model!A3"].formula == '="Cost of "&A2'
    # Every formula in the file: eight in the grid, one in the label column.
    assert len(book.formulas()) == 9
    assert sum(1 for cell in book.cells.values() if cell.formula) == 8


def test_typed_text_that_looks_like_a_formula_is_not_one(tmp_path: Path) -> None:
    """A « Key » sheet listing where each name points, as words: the
    FHWA tool holds 61 such cells, and a string is not arithmetic."""
    path = _model(tmp_path / "m.xlsx", "Cost")
    book = XlsxWorkbook()
    book.remove(book.active)
    # Reopen the saved model and add a typed string beside the grid.
    from openpyxl import load_workbook

    saved = load_workbook(path)
    sheet = saved["Model"]
    sheet["F2"] = "='Project Inputs'!$F$71"
    sheet["F2"].data_type = "s"
    saved.save(path)
    read = read_workbook(str(path))
    assert "Model!F2" not in read.formulas()
    assert read.cells["Model!F2"].formula is None if "Model!F2" in read.cells else True


def test_a_refused_function_in_a_label_formula_refuses_the_file(
    tmp_path: Path,
) -> None:
    book = read_workbook(str(_model(tmp_path / "m.xlsx", '=RTD("feed",,"name")')))
    assert route_for(prescan(book.cells)) is None
    assert route_for(prescan(book.formulas())) is Route.REFUSE


def test_an_external_link_in_a_label_formula_is_a_finding(tmp_path: Path) -> None:
    book = read_workbook(str(_model(tmp_path / "m.xlsx", "='[1]Names'!$A$3")))
    result = audit(book)
    links = [one for one in result.findings if one.rule == "external-link"]
    assert len(links) == 1
    assert links[0].ref == "Model!A3"
    # And it never becomes a numeric finding: the grid is untouched.
    assert not [one for one in result.findings if one.rule == "typed-over-formula"]
