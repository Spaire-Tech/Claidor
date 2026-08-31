"""The routes, through HTTP, where the promises are actually kept.

``test_spine.py`` proves the engine survives the trip into the tables.
What is left to prove is the part that only exists once there is a request
attached: that a deal is closed to everyone not on it, that an upload the
reader cannot use is refused with a sentence rather than a stack trace,
and that the shapes the founder is building screens against are the shapes
that come back.

The access tests are the ones worth reading twice. Every route here takes
a deal id or something reachable from one, and a single route that checks
the scope and forgets the membership is the widest hole this product could
have.
"""

import json
import os
from collections import Counter
from pathlib import Path
from typing import Any

import pytest
import sqlalchemy as sa
from httpx import AsyncClient

from polar.auth.scope import Scope
from polar.kit.db.postgres import AsyncSession
from polar.models import (
    Artifact,
    ArtifactKind,
    CheckStatus,
    Dossier,
    DossierMember,
    DossierRole,
    User,
    UserOrganization,
)
from polar.tieout.ingest import SUFFIXES
from polar.tieout.recalc.uno_calc import find_install
from polar.tieout.repository import TieOutRepository
from polar.tieout.service import tieout
from tests.fixtures.database import SaveFixture
from tests.fixtures.random_objects import create_organization, create_user

CASCADE = Path(__file__).resolve().parents[2] / "scripts" / "cascade"
CLEAN = CASCADE / "cascade_deck.pptx"
MODEL = CASCADE / "cascade_model.xlsx"
PREAPP = CASCADE / "example_preapp_model.xlsx"

DECK_MEDIA = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NOWHERE = "00000000-0000-0000-0000-000000000000"


async def _deal_for(
    session: AsyncSession, save_fixture: SaveFixture, owner: User
) -> Dossier:
    organization = await create_organization(save_fixture)
    deal = Dossier(
        organization_id=organization.id,
        name="Project Cascade",
        client_name="Cascade Industrial Holdings",
        created_by_id=owner.id,
    )
    session.add(deal)
    await session.flush()
    session.add(
        DossierMember(dossier_id=deal.id, user_id=owner.id, role=DossierRole.lead)
    )
    # The dashboard only reaches an organization its user belongs to, so
    # every real caller has this row; the settings routes check it.
    session.add(UserOrganization(user_id=owner.id, organization_id=organization.id))
    await session.flush()
    return deal


async def _loaded(
    session: AsyncSession, save_fixture: SaveFixture, owner: User
) -> Dossier:
    deal = await _deal_for(session, save_fixture, owner)
    for path, kind in ((MODEL, ArtifactKind.model), (CLEAN, ArtifactKind.deck)):
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=kind,
            filename=path.name,
            payload=path.read_bytes(),
            user_id=owner.id,
        )
    await tieout.run_tieout(session, dossier_id=deal.id, user_id=owner.id)
    await tieout.run_audit(session, dossier_id=deal.id, user_id=owner.id)
    await session.flush()
    return deal


@pytest.mark.asyncio
class TestAccess:
    async def test_anonymous_is_refused(self, client: AsyncClient) -> None:
        response = await client.get(f"/v1/tieout/deals/{NOWHERE}")
        assert response.status_code == 401

    @pytest.mark.auth
    async def test_a_deal_you_are_not_on_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        """404, not 403.

        A banker outside the deal should not learn that it exists, which
        is the same rule the rest of the workspace follows and the reason
        organization membership grants nothing here.
        """
        stranger = await create_user(save_fixture)
        deal = await _loaded(session, save_fixture, stranger)

        assert (await client.get(f"/v1/tieout/deals/{deal.id}")).status_code == 404
        assert (
            await client.get(f"/v1/tieout/deals/{deal.id}/findings")
        ).status_code == 404
        assert (
            await client.get(f"/v1/tieout/deals/{deal.id}/links")
        ).status_code == 404
        assert (
            await client.post(f"/v1/tieout/deals/{deal.id}/check")
        ).status_code == 404

    @pytest.mark.auth
    async def test_a_file_in_someone_elses_deal_does_not_exist_either(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        """The routes that take an artifact id reach a deal too.

        Checking the deal on the deal routes and forgetting it on these
        would leave every figure in the workspace readable by id.
        """
        stranger = await create_user(save_fixture)
        deal = await _loaded(session, save_fixture, stranger)
        repository = TieOutRepository.from_session(session)
        deck = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.deck
        )

        assert (await client.get(f"/v1/tieout/artifacts/{deck.id}")).status_code == 404
        assert (
            await client.get(f"/v1/tieout/artifacts/{deck.id}/figures")
        ).status_code == 404
        assert (
            await client.delete(f"/v1/tieout/artifacts/{deck.id}")
        ).status_code == 404

    @pytest.mark.auth
    async def test_a_finding_and_its_chain_are_closed_as_well(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _loaded(session, save_fixture, stranger)
        repository = TieOutRepository.from_session(session)
        finding = (await repository.findings_of(deal.id))[0]
        link = (await repository.links_of(deal.id))[0]

        assert (
            await client.get(f"/v1/tieout/findings/{finding.id}/chain")
        ).status_code == 404
        assert (await client.get(f"/v1/tieout/links/{link.id}")).status_code == 404
        assert (
            await client.patch(
                f"/v1/tieout/findings/{finding.id}", json={"state": "dismissed"}
            )
        ).status_code == 404


@pytest.mark.asyncio
class TestTheDataRoom:
    """The room pages, folds versions, and says how much it is not showing."""

    @pytest.mark.auth
    async def test_it_pages_and_says_how_many_there_are(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        for index in range(7):
            await tieout.ingest(
                session,
                dossier_id=deal.id,
                kind=ArtifactKind.source,
                filename=f"note {index}.txt",
                payload=b"not readable, and that is a state of the deal",
                user_id=user.id,
            )
        await session.flush()

        first = await client.get(
            f"/v1/tieout/deals/{deal.id}/artifacts", params={"limit": 3}
        )
        assert first.status_code == 200
        page = first.json()
        assert len(page["items"]) == 3
        # The total rides with the page, so the screen can write « 3 of 7 »
        # rather than quietly implying it has them all.
        assert page["total"] == 7

        rest = await client.get(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            params={"limit": 100, "offset": 3},
        )
        assert len(rest.json()["items"]) == 4

    @pytest.mark.auth
    async def test_a_re_upload_is_one_row_not_two(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        for _ in range(3):
            await tieout.ingest(
                session,
                dossier_id=deal.id,
                kind=ArtifactKind.model,
                filename=MODEL.name,
                payload=MODEL.read_bytes(),
                user_id=user.id,
            )
        await session.flush()

        page = (await client.get(f"/v1/tieout/deals/{deal.id}/artifacts")).json()
        assert page["total"] == 1
        assert len(page["items"]) == 1
        # Folded on the server, so the count and the rows agree — a client
        # that pages and folds can do neither.
        assert page["items"][0]["version"] == 3

    @pytest.mark.auth
    async def test_it_searches_on_the_name(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        page = (
            await client.get(
                f"/v1/tieout/deals/{deal.id}/artifacts", params={"q": "deck"}
            )
        ).json()
        assert page["total"] == 1
        assert "deck" in page["items"][0]["filename"]


@pytest.mark.asyncio
class TestTheDealPage:
    @pytest.mark.auth
    async def test_it_answers_in_one_request(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        response = await client.get(f"/v1/tieout/deals/{deal.id}")

        assert response.status_code == 200
        body = response.json()
        assert body["name"] == "Project Cascade"
        # The deal's spine — the documents it is built on — and counts
        # for the room, rather than the room itself.
        assert len(body["documents"]) == 2
        assert {one["kind"] for one in body["documents"]} == {"model", "deck"}
        assert body["files"] == 2
        assert body["lineages"] == 2
        # 106 with the cross-sheet parameter collapse (`link._timeless`);
        # 113 with the fraction gate letting a plain fraction figure claim
        # fraction-sized cells — gained links all agree, drifts unchanged.
        assert body["coverage"]["reconciled"] == 113
        assert body["coverage"]["drifting"] == 8
        # The coverage line carries its own misses, with reasons.
        assert body["coverage"]["unlinked"] > 0
        assert body["coverage"]["reasons"][0]["count"] > 0
        # Eight drifts from the tie-out and one smell from the audit. Both
        # checks ran, and their findings live in one list.
        assert body["findings"]["open"] == 9
        assert body["last_tieout"]["status"] == "done"
        assert body["last_audit"]["status"] == "done"

    @pytest.mark.auth
    async def test_an_empty_deal_is_not_an_error(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """A deal with no files is the first screen anybody sees."""
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()
        response = await client.get(f"/v1/tieout/deals/{deal.id}")

        assert response.status_code == 200
        body = response.json()
        assert body["documents"] == []
        assert body["files"] == 0
        assert body["coverage"]["reconciled"] == 0
        assert body["last_tieout"] is None


@pytest.mark.asyncio
class TestUpload:
    @pytest.mark.auth
    async def test_a_deck_is_read_on_the_way_in(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
        )

        assert response.status_code == 200
        body = response.json()
        # The kind was not passed; it came from the extension.
        assert body["kind"] == "deck"
        assert body["status"] == "ready"
        assert body["version"] == 1
        assert body["counts"]["figures"] > 0
        assert body["uploaded_by"]["id"] == str(user.id)

    @pytest.mark.auth
    async def test_a_file_nobody_can_read_is_refused_in_words(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={"file": ("notes.txt", b"hello", "text/plain")},
        )

        assert response.status_code == 415
        detail = response.json()["detail"]
        #: The intent, not the prose: refused in words, naming what the
        #: reader does take. Pinning the sentence's casing made this red
        #: when the list started being derived from `SUFFIXES`.
        assert "not a file this can read" in detail
        for suffix in SUFFIXES[ArtifactKind.model]:
            assert suffix in detail

    @pytest.mark.auth
    async def test_a_file_that_fails_to_parse_is_kept_with_its_reason(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """200, not 400. « This one did not work » is a state of the deal.

        The screen has to be able to show the file and what to do about it,
        which is impossible if the request failed and nothing was written.
        """
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={"file": ("model.xlsx", b"not a workbook", "application/zip")},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "failed"
        assert "Excel" in body["error"]


@pytest.mark.asyncio
class TestFindingsAndTheChain:
    @pytest.mark.auth
    async def test_the_findings_carry_where_and_what_the_model_says(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        response = await client.get(
            f"/v1/tieout/deals/{deal.id}/findings", params={"kind": "drift"}
        )

        assert response.status_code == 200
        findings = response.json()
        assert len(findings) == 8
        first = findings[0]
        assert first["printed"]
        assert first["expected"]
        assert first["where"]["filename"] == "cascade_deck.pptx"
        assert first["where"]["label"].startswith("slide ")
        assert first["source"]["ref"]
        # A one-tick difference is shown and ranked last, never hidden.
        assert [one["one_tick"] for one in findings] == sorted(
            one["one_tick"] for one in findings
        )

    @pytest.mark.auth
    async def test_the_panel_can_ask_for_one_document(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )
        response = await client.get(
            f"/v1/tieout/deals/{deal.id}/findings",
            params={"artifact_id": str(model.id)},
        )

        assert response.status_code == 200
        # The drifts belong to the deck, so the model's list is its audit's.
        assert all(one["kind"] != "drift" for one in response.json())

    @pytest.mark.auth
    async def test_dismissing_is_visible_and_reversible(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        findings = (await client.get(f"/v1/tieout/deals/{deal.id}/findings")).json()
        target = findings[0]["id"]

        # A dismissal without a reason is refused: dismissing says the
        # check is wrong about this one, which is exactly the decision
        # somebody questions three weeks later.
        bare = await client.patch(
            f"/v1/tieout/findings/{target}", json={"state": "dismissed"}
        )
        assert bare.status_code == 422

        dismissed = await client.patch(
            f"/v1/tieout/findings/{target}",
            json={"state": "dismissed", "note": "different basis, agreed"},
        )
        assert dismissed.status_code == 200
        assert dismissed.json()["state"] == "dismissed"
        assert dismissed.json()["note"] == "different basis, agreed"

        back = await client.patch(
            f"/v1/tieout/findings/{target}", json={"state": "open"}
        )
        assert back.json()["state"] == "open"
        # The note goes with the dismissal it explained.
        assert back.json()["note"] == ""

    @pytest.mark.auth
    async def test_the_chain_reaches_the_model(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        findings = (await client.get(f"/v1/tieout/deals/{deal.id}/findings")).json()

        response = await client.get(f"/v1/tieout/findings/{findings[0]['id']}/chain")

        assert response.status_code == 200
        steps = response.json()["steps"]
        assert steps[0]["kind"] == "figure"
        assert any(step["kind"] in {"cell", "input"} for step in steps)
        assert response.json()["summary"]
        # Present on every step, empty on a clean one. The field existing
        # is what stops the screen having to assume; see the next test for
        # why an absent field and an empty one are different claims.
        assert all("unresolved" in step for step in steps)

    @pytest.mark.auth
    async def test_an_input_the_chain_could_not_follow_reaches_the_screen(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The other half of the answer, all the way through.

        A cell reading something the parser cannot resolve to a cell — a
        reference into another workbook, a name pointing at `#REF!` — used
        to produce a chain short by that input and identical to a chain
        that never had one. Across two real Ofgem models that was 4.4% of
        formulas, and on one of them the dropped input was the switch
        deciding what the whole model computed.

        This walks it end to end: engine, database, route, wire.
        """
        from polar.models import ModelCell

        deal = await _loaded(session, save_fixture, user)
        findings = (await client.get(f"/v1/tieout/deals/{deal.id}/findings")).json()
        chain = (
            await client.get(f"/v1/tieout/findings/{findings[0]['id']}/chain")
        ).json()
        ref = next(
            step["ref"] for step in chain["steps"] if step["kind"] in {"cell", "input"}
        )

        cell = (
            await session.execute(
                sa.select(ModelCell).where(ModelCell.ref == ref).limit(1)
            )
        ).scalar_one()
        cell.unresolved = [["[1]Group.xlsx!B4", "in another workbook"]]
        session.add(cell)
        await session.flush()

        again = (
            await client.get(f"/v1/tieout/findings/{findings[0]['id']}/chain")
        ).json()
        step = next(one for one in again["steps"] if one["ref"] == ref)
        assert step["unresolved"] == ["in another workbook"]


@pytest.mark.asyncio
class TestTheConfirmationQueue:
    @pytest.mark.auth
    async def test_a_link_carries_both_sides(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        response = await client.get(
            f"/v1/tieout/deals/{deal.id}/links", params={"state": "proposed"}
        )

        assert response.status_code == 200
        links = response.json()
        assert links
        first = links[0]
        assert first["figure"]["printed"]
        assert first["cell"]["ref"]
        assert first["transformation"] == "identity"
        # Ordered by confidence, so the easy ones clear first.
        assert links[0]["confidence"] >= links[-1]["confidence"]

    @pytest.mark.auth
    async def test_one_link_offers_the_runners_up(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        links = (await client.get(f"/v1/tieout/deals/{deal.id}/links")).json()

        response = await client.get(f"/v1/tieout/links/{links[0]['id']}")

        assert response.status_code == 200
        alternatives = response.json()["alternatives"]
        # Never the cell already chosen — that is not an alternative.
        assert all(one["cell_id"] != links[0]["cell"]["id"] for one in alternatives)

    @pytest.mark.auth
    async def test_confirming_records_who_and_when(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        links = (await client.get(f"/v1/tieout/deals/{deal.id}/links")).json()

        response = await client.patch(
            f"/v1/tieout/links/{links[0]['id']}", json={"state": "confirmed"}
        )

        assert response.status_code == 200
        body = response.json()
        assert body["state"] == "confirmed"
        assert body["confirmed_by"]["id"] == str(user.id)
        assert body["confirmed_at"] is not None

    @pytest.mark.auth
    async def test_pointing_it_somewhere_else_stops_being_a_score(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """A banker's answer is not a 0.62. It is the answer."""
        deal = await _loaded(session, save_fixture, user)
        links = (await client.get(f"/v1/tieout/deals/{deal.id}/links")).json()
        target = next(one for one in links if one["confidence"] < 1)

        cells = (
            await client.get(
                f"/v1/tieout/artifacts/{target['cell']['artifact_id']}/cells",
                params={"q": "EBITDA"},
            )
        ).json()
        elsewhere = next(one for one in cells if one["id"] != target["cell"]["id"])

        response = await client.patch(
            f"/v1/tieout/links/{target['id']}",
            json={"state": "confirmed", "cell_id": elsewhere["id"]},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["cell"]["ref"] == elsewhere["ref"]
        assert body["confidence"] == 1.0

    @pytest.mark.auth
    async def test_a_cell_from_another_deal_is_refused(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        mine = await _loaded(session, save_fixture, user)
        theirs = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        other_model = next(
            one
            for one in await repository.current_artifacts(theirs.id)
            if one.kind is ArtifactKind.model
        )
        foreign = (await repository.cells_of(other_model.id))[0]
        link = (await repository.links_of(mine.id))[0]

        response = await client.patch(
            f"/v1/tieout/links/{link.id}",
            json={"state": "confirmed", "cell_id": str(foreign.id)},
        )

        assert response.status_code == 422


@pytest.mark.asyncio
class TestTheFigureMap:
    @pytest.mark.auth
    async def test_every_figure_appears_and_the_misses_say_why(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        deck = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.deck
        )

        response = await client.get(f"/v1/tieout/artifacts/{deck.id}/figures")

        assert response.status_code == 200
        body = response.json()
        shown = [one for slide in body["slides"] for one in slide["figures"]]
        assert len(shown) == deck.counts["figures"]
        assert {one["state"] for one in shown} <= {
            "agreeing",
            "drifting",
            "confirmed",
            "unlinked",
        }
        unlinked = [one for one in shown if one["state"] == "unlinked"]
        assert unlinked
        assert all(one["reason"] for one in unlinked)


@pytest.mark.asyncio
class TestThePanel:
    @pytest.mark.auth
    async def test_a_stamped_document_needs_no_guessing(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The lineage id travels inside the file and is not a guess."""
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        deck = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.deck
        )

        response = await client.post(
            "/v1/tieout/identify", json={"lineage_id": str(deck.lineage_id)}
        )

        assert response.status_code == 200
        body = response.json()
        assert body["matched_by"] == "stamp"
        assert body["artifact"]["id"] == str(deck.id)
        assert body["dossier_name"] == "Project Cascade"
        # Already stamped, so nothing to write back.
        assert body["stamp_lineage_id"] is None

    @pytest.mark.auth
    async def test_a_filename_inside_a_chosen_deal_is_a_guess_worth_stamping(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)

        response = await client.post(
            "/v1/tieout/identify",
            json={"filename": "cascade_deck.pptx", "dossier_id": str(deal.id)},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["matched_by"] == "filename"
        assert body["stamp_lineage_id"] is not None

    @pytest.mark.auth
    async def test_a_stamp_from_someone_elses_deal_reveals_nothing(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        """Not « found but forbidden ». Not found, and not a hint either."""
        stranger = await create_user(save_fixture)
        deal = await _loaded(session, save_fixture, stranger)
        repository = TieOutRepository.from_session(session)
        deck = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.deck
        )

        response = await client.post(
            "/v1/tieout/identify", json={"lineage_id": str(deck.lineage_id)}
        )

        assert response.status_code == 200
        body = response.json()
        assert body["matched_by"] == "none"
        assert body["artifact"] is None
        assert body["dossier_id"] is None

    @pytest.mark.auth
    async def test_an_unknown_document_says_so(
        self, client: AsyncClient, save_fixture: SaveFixture, user: User
    ) -> None:
        response = await client.post(
            "/v1/tieout/identify", json={"filename": "someone_elses.pptx"}
        )
        assert response.status_code == 200
        assert response.json()["matched_by"] == "none"

    @pytest.mark.auth
    async def test_the_deal_list_is_scoped_to_this_person(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        mine = await _deal_for(session, save_fixture, user)
        stranger = await create_user(save_fixture)
        await _deal_for(session, save_fixture, stranger)
        await session.flush()

        response = await client.get("/v1/tieout/deals")

        assert response.status_code == 200
        ids = [one["id"] for one in response.json()]
        assert str(mine.id) in ids
        assert len(ids) == 1

    @pytest.mark.auth
    async def test_a_deal_can_be_removed_and_stays_removed(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The route that did not exist: the panel's picker was listing
        every pre-pivot test deal with no way anywhere to be rid of
        them. Soft delete, membership as the whole permission."""
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        assert (await client.delete(f"/v1/tieout/deals/{deal.id}")).status_code == 204
        ids = [one["id"] for one in (await client.get("/v1/tieout/deals")).json()]
        assert str(deal.id) not in ids

    @pytest.mark.auth
    async def test_a_stranger_cannot_remove_a_deal(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _deal_for(session, save_fixture, stranger)
        await session.flush()

        assert (await client.delete(f"/v1/tieout/deals/{deal.id}")).status_code == 404

    @pytest.mark.auth
    async def test_a_deal_nobody_has_checked_does_not_look_clear(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The distinction Projects is built on.

        « No open findings » on a deal nobody ran reads exactly like « no
        open findings » on one checked this morning, and a product whose
        argument is that what was *not* checked stays on screen cannot let
        those two look the same. `checked_at` is null until a run finishes.
        """
        await _deal_for(session, save_fixture, user)
        await session.flush()

        never = (await client.get("/v1/tieout/deals")).json()[0]
        assert never["checked_at"] is None
        assert never["open_findings"] == 0
        assert never["failing_checks"] == 0

        checked = await _loaded(session, save_fixture, user)
        await session.flush()

        rows = {one["id"]: one for one in (await client.get("/v1/tieout/deals")).json()}
        assert rows[str(checked.id)]["checked_at"] is not None
        assert rows[str(checked.id)]["open_findings"] > 0

        # « 6 checks fail », not « 41 findings »: the row counts distinct
        # *checks* among the open findings, so it must agree with the
        # findings list's own rules — and can never exceed the findings.
        findings = (
            await client.get(f"/v1/tieout/deals/{checked.id}/findings?state=open")
        ).json()
        assert rows[str(checked.id)]["failing_checks"] == len(
            {one["rule"] for one in findings}
        )
        assert 0 < rows[str(checked.id)]["failing_checks"] <= len(findings)

    @pytest.mark.auth
    async def test_a_finding_carries_coordinates_a_host_can_act_on(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """« Click a finding and PowerPoint goes to it » is a lookup.

        Prose in `location` cannot be parsed back into a selection, so the
        same position rides along in numbers.
        """
        deal = await _loaded(session, save_fixture, user)
        findings = (
            await client.get(
                f"/v1/tieout/deals/{deal.id}/findings", params={"kind": "drift"}
            )
        ).json()

        for finding in findings:
            anchor = finding["where"]["anchor"]
            assert anchor["shape_id"] > 0
            assert anchor["kind"] in {"text", "table", "chart"}

        audit = (
            await client.get(
                f"/v1/tieout/deals/{deal.id}/findings", params={"kind": "audit"}
            )
        ).json()
        # An audit finding sits at a cell, which Excel selects as it stands.
        assert all(one["where"]["anchor"]["kind"] == "cell" for one in audit)
        assert all(one["where"]["anchor"]["ref"] for one in audit)


@pytest.mark.asyncio
class TestThePanelToken:
    async def test_anonymous_gets_nothing(self, client: AsyncClient) -> None:
        response = await client.post("/v1/tieout/panel/token")
        assert response.status_code == 401

    @pytest.mark.auth
    async def test_a_browser_session_mints_a_narrow_token(
        self, client: AsyncClient, user: User
    ) -> None:
        """The panel holds a credential, and only these two scopes.

        Fixed here rather than taken from the request: a caller that could
        name its own scopes would make this endpoint a way to widen any
        session into anything.
        """
        response = await client.post("/v1/tieout/panel/token")

        assert response.status_code == 201
        body = response.json()
        assert body["token"].startswith("claidor_pat_")
        assert body["scopes"] == ["tieout:read", "tieout:write"]
        assert body["expires_in"] == 30 * 24 * 3600

    @pytest.mark.auth
    async def test_the_token_it_mints_is_a_real_resolvable_credential(
        self, client: AsyncClient, session: AsyncSession, user: User
    ) -> None:
        """Resolve the plaintext the way the authenticator does.

        Not through a request: the `auth` marker replaces the
        authentication dependency wholesale, so a bearer header in a test
        is ignored and a route would answer 200 whatever the token said.
        An « end to end » assertion there would pass with the token
        deleted, which is worse than no test.

        So the token is looked up by its own hash — the same call the real
        authenticator makes — and checked for the two scopes and the
        expiry.
        """
        from polar.personal_access_token.service import (
            personal_access_token as tokens,
        )

        plaintext = (await client.post("/v1/tieout/panel/token")).json()["token"]
        found = await tokens.get_by_token(session, plaintext)

        assert found is not None
        assert found.user_id == user.id
        assert set(found.scopes) == {Scope.tieout_read, Scope.tieout_write}
        assert found.expires_at is not None

    @pytest.mark.auth
    async def test_it_carries_nothing_the_panel_does_not_need(
        self, client: AsyncClient, session: AsyncSession, user: User
    ) -> None:
        """Two scopes, and no way for a caller to ask for more.

        The scopes are fixed in the endpoint rather than read from the
        request. A body naming its own would turn this into a way to widen
        any browser session into anything the account can do.
        """
        from polar.personal_access_token.service import (
            personal_access_token as tokens,
        )

        plaintext = (
            await client.post(
                "/v1/tieout/panel/token",
                json={"scopes": ["organizations:write"]},
            )
        ).json()["token"]
        found = await tokens.get_by_token(session, plaintext)

        assert found is not None
        assert Scope.organizations_write not in found.scopes
        assert set(found.scopes) == {Scope.tieout_read, Scope.tieout_write}


@pytest.mark.asyncio
class TestCheckAFile:
    """The one-off check: a loose file, no deal required.

    The engine's own rules are pinned in ``test_solo.py``; what is proved
    here is the route — that the file is checked and dropped, that the
    answer is kept as a recent and replays whole, and that a recent is
    its owner's alone.
    """

    async def test_anonymous_is_refused(self, client: AsyncClient) -> None:
        response = await client.post(
            "/v1/tieout/check-file",
            files={"file": ("deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
        )
        assert response.status_code == 401

    @pytest.mark.auth
    async def test_a_deck_alone_is_checked_against_itself(
        self, client: AsyncClient, user: User
    ) -> None:
        response = await client.post(
            "/v1/tieout/check-file",
            files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
        )

        assert response.status_code == 201
        body = response.json()
        assert body["kind"] == "deck"
        assert body["against"] == ""
        assert body["dossier_id"] is None
        assert body["counts"]["figures"] > 0
        assert body["counts"]["repeated"] >= 2
        # The slide-3 chart against its table — the drift the clean deck
        # genuinely carries, and nothing else (see test_solo.py).
        assert body["counts"]["differences"] == 2
        assert [one["label"] for one in body["disagreements"]] == [
            "Adjusted EBITDA FY2023A",
            "Adjusted EBITDA FY2024A",
        ]
        first = body["disagreements"][0]
        assert first["first"]["printed"] == "37.8"
        assert first["other"]["printed"] == "30.8"
        assert first["first"]["page"] == 3
        assert body["drifts"] == []
        assert body["defects"] == []

    @pytest.mark.auth
    async def test_against_a_deal_it_also_meets_the_model(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)

        response = await client.post(
            f"/v1/tieout/check-file?dossier_id={deal.id}",
            files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
        )

        assert response.status_code == 201
        body = response.json()
        assert body["against"] == "Project Cascade"
        assert body["dossier_id"] == str(deal.id)
        assert [one["filename"] for one in body["models"]] == ["cascade_model.xlsx"]
        assert body["counts"]["reconciled"] > 0
        # The clean deck carries real drifts against the model — the
        # pre-existing set test_cascade.py documents — so this is not
        # allowed to come back empty.
        assert body["drifts"]
        one = body["drifts"][0]
        assert one["printed"]
        assert one["expected"]
        assert one["model_artifact_id"] == body["models"][0]["artifact_id"]
        # And the file still disagrees with itself, deal or no deal.
        assert len(body["disagreements"]) == 2

    @pytest.mark.auth
    async def test_a_deal_you_are_not_on_cannot_be_picked(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _deal_for(session, save_fixture, stranger)

        response = await client.post(
            f"/v1/tieout/check-file?dossier_id={deal.id}",
            files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
        )
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_model_routes_to_the_audit_and_a_deal_is_ignored(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """A workbook checked by itself is the model audit.

        Model-against-deal would be the grounding, which needs the deal's
        source documents and is not a one-off — so the deal is ignored
        and the stored answer says « on its own », because that is what
        ran.
        """
        deal = await _deal_for(session, save_fixture, user)

        response = await client.post(
            f"/v1/tieout/check-file?dossier_id={deal.id}",
            files={
                "file": (
                    "cascade_model.xlsx",
                    MODEL.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert response.status_code == 201
        body = response.json()
        assert body["kind"] == "model"
        assert body["against"] == ""
        assert body["dossier_id"] is None
        assert body["counts"]["cells"] > 0
        assert body["disagreements"] == []
        assert body["drifts"] == []
        # Errors and smells are never added into one number; the lists
        # and counts carry them apart.
        assert len(body["defects"]) == (
            body["counts"]["errors"] + body["counts"]["smells"]
        )

    @pytest.mark.auth
    async def test_a_pdf_is_refused_as_a_source(self, client: AsyncClient) -> None:
        response = await client.post(
            "/v1/tieout/check-file",
            files={"file": ("accounts.pdf", b"%PDF-1.4", "application/pdf")},
        )
        assert response.status_code == 415
        assert "a source" in response.json()["detail"]

    @pytest.mark.auth
    async def test_an_unreadable_file_is_a_sentence_and_no_recent(
        self, client: AsyncClient
    ) -> None:
        response = await client.post(
            "/v1/tieout/check-file",
            files={"file": ("deck.pptx", b"not a deck", DECK_MEDIA)},
        )
        assert response.status_code == 422
        assert "PowerPoint" in response.json()["detail"]

        recents = await client.get("/v1/tieout/check-file/recents")
        assert recents.json() == []

    @pytest.mark.auth
    async def test_a_recent_replays_whole_and_is_private(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        checked = (
            await client.post(
                "/v1/tieout/check-file",
                files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
            )
        ).json()

        recents = (await client.get("/v1/tieout/check-file/recents")).json()
        assert [one["id"] for one in recents] == [checked["id"]]
        assert recents[0]["filename"] == "cascade_deck.pptx"
        assert recents[0]["against"] == ""

        replayed = (await client.get(f"/v1/tieout/check-file/{checked['id']}")).json()
        assert replayed == checked

        # Another person's recent does not exist, rather than being
        # forbidden — same posture as a deal.
        other = await create_user(save_fixture)
        theirs = await tieout.check_file(
            session,
            user_id=other.id,
            kind=ArtifactKind.deck,
            filename="cascade_deck.pptx",
            payload=CLEAN.read_bytes(),
        )
        await session.flush()
        response = await client.get(f"/v1/tieout/check-file/{theirs.id}")
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_accepting_a_rule_marks_every_place_and_survives_replay(
        self, client: AsyncClient, user: User
    ) -> None:
        """The bench's « Accept with a note » — the ruling is about the
        check, so every place the rule fails is accepted together, and
        it stands when the recent is reopened."""
        checked = (
            await client.post(
                "/v1/tieout/check-file",
                files={
                    "file": (
                        "cascade_model.xlsx",
                        MODEL.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        ).json()
        assert len(checked["defects"]) > 0
        rule = checked["defects"][0]["rule"]
        of_rule = [one for one in checked["defects"] if one["rule"] == rule]

        response = await client.post(
            f"/v1/tieout/check-file/{checked['id']}/accept",
            json={"rule": rule, "note": "Known and priced in."},
        )
        assert response.status_code == 200
        accepted = [one for one in response.json()["defects"] if one["accepted"]]
        assert len(accepted) == len(of_rule)
        assert all(one["rule"] == rule for one in accepted)
        assert all(one["accepted_note"] == "Known and priced in." for one in accepted)

        #: The ruling is in the stored answer, not the response alone.
        replayed = (await client.get(f"/v1/tieout/check-file/{checked['id']}")).json()
        assert replayed == response.json()

    @pytest.mark.auth
    async def test_an_acceptance_without_a_reason_or_a_rule_is_refused(
        self, client: AsyncClient, user: User
    ) -> None:
        checked = (
            await client.post(
                "/v1/tieout/check-file",
                files={
                    "file": (
                        "cascade_model.xlsx",
                        MODEL.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        ).json()
        rule = checked["defects"][0]["rule"]

        bare = await client.post(
            f"/v1/tieout/check-file/{checked['id']}/accept",
            json={"rule": rule, "note": "   "},
        )
        assert bare.status_code == 422
        assert "reason" in bare.json()["detail"]

        unknown = await client.post(
            f"/v1/tieout/check-file/{checked['id']}/accept",
            json={"rule": "no-such-rule", "note": "still no"},
        )
        assert unknown.status_code == 422
        assert "no failing check named no-such-rule" in unknown.json()["detail"]


@pytest.mark.asyncio
class TestOpenArtifact:
    """« Open the cell »'s destination: the real document."""

    @pytest.mark.auth
    async def test_an_uploaded_model_answers_with_a_download(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """No SharePoint identity means no live document anywhere — the
        stored bytes are the truth, and the answer says so."""
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )

        response = await client.get(
            f"/v1/tieout/artifacts/{model.id}/open?ref=Model%21C6"
        )
        assert response.status_code == 200
        body = response.json()
        assert body["kind"] == "download"
        assert body["url"]
        assert body["filename"] == "cascade_model.xlsx"

    @pytest.mark.auth
    async def test_someone_elses_model_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _loaded(session, save_fixture, stranger)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )
        response = await client.get(f"/v1/tieout/artifacts/{model.id}/open")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestHouseRules:
    """The firm's rules: stored per organization, and actually obeyed."""

    async def _org_of(self, session: AsyncSession, deal: Dossier):
        return deal.organization_id

    @pytest.mark.auth
    async def test_defaults_before_anybody_decided(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        response = await client.get(
            f"/v1/tieout/house-rules?organization_id={deal.organization_id}"
        )
        assert response.status_code == 200
        body = response.json()
        assert body["rounding"] == "together"
        assert body["grounding"] is True
        # The audit's own catalogue, all on — never a list a screen
        # invented, and never a count this file pins by hand: a pinned
        # 17 went red the day Sentinel's adoptions made it 19, telling
        # nobody anything true. The invariant is that the endpoint
        # serves exactly the engine's catalogue — both families — with
        # every rule on by default, the newly adopted ones included.
        from polar.tieout.analytics import ANALYTIC_RULE_NAMES
        from polar.tieout.audit import RULE_NAMES

        assert {rule["key"] for rule in body["rules"]} == (
            set(RULE_NAMES) | set(ANALYTIC_RULE_NAMES)
        )
        assert all(rule["on"] for rule in body["rules"])
        assert {rule["key"] for rule in body["rules"] if rule["analytical"]} == set(
            ANALYTIC_RULE_NAMES
        )

    @pytest.mark.auth
    async def test_a_stranger_finds_no_organization(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _deal_for(session, save_fixture, stranger)
        response = await client.get(
            f"/v1/tieout/house-rules?organization_id={deal.organization_id}"
        )
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_an_unknown_rule_key_is_refused_whole(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        response = await client.put(
            f"/v1/tieout/house-rules?organization_id={deal.organization_id}",
            json={"audit_rules_off": ["skipped-cell", "made-up-rule"]},
        )
        assert response.status_code == 422
        assert "made-up-rule" in str(response.json())

    @pytest.mark.auth
    async def test_a_rule_switched_off_is_skipped_and_on_the_record(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The Cascade model carries hardcode smells; switch the rule off
        and the audit stays quiet about them — while its summary names
        the switch, because a rule turned off is a decision, never a
        silence."""
        deal = await _loaded(session, save_fixture, user)
        with_rule = (
            await client.get(f"/v1/tieout/deals/{deal.id}/findings?kind=audit")
        ).json()

        response = await client.put(
            f"/v1/tieout/house-rules?organization_id={deal.organization_id}",
            json={"audit_rules_off": ["hardcode-in-formula"]},
        )
        assert response.status_code == 200
        assert not [
            rule
            for rule in response.json()["rules"]
            if rule["key"] == "hardcode-in-formula" and rule["on"]
        ]

        runs = (await client.post(f"/v1/tieout/deals/{deal.id}/check")).json()
        audit = next(one for one in runs if one["kind"] == "audit")
        assert audit["summary"]["rules_off"] == ["hardcode-in-formula"]

        without_rule = (
            await client.get(f"/v1/tieout/deals/{deal.id}/findings?kind=audit")
        ).json()
        assert [one for one in with_rule if one["rule"] == "hardcode-in-formula"]
        assert not [one for one in without_rule if one["rule"] == "hardcode-in-formula"]

    @pytest.mark.auth
    async def test_grounding_off_means_two_runs_not_a_failed_third(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        await client.put(
            f"/v1/tieout/house-rules?organization_id={deal.organization_id}",
            json={"grounding": False},
        )
        runs = (await client.post(f"/v1/tieout/deals/{deal.id}/check")).json()
        assert sorted(one["kind"] for one in runs) == ["audit", "tieout"]

    @pytest.mark.auth
    async def test_two_firms_one_model_two_reports(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Piece 13's demonstration, criteria frozen first in
        docs/pierce/house-rules-demo.md § 2.

        Two firms, the founder's own pre-app model in both, one firm on
        shipped defaults and one with five rules off — spanning both
        catalogue families and all three layers of the report (findings,
        pass tallies, abstentions). The two reports must differ by
        exactly what the configuration says, and nothing else: every
        finding the stricter firm keeps is byte-for-byte the other
        firm's, the counters count the filtered list, and the switched
        rules are named on the record under ``rules_off``.
        """
        off = [
            "balance-sheet",
            "broken-name",
            "gapped-test",
            "long-formula",
            "model-own-check",
        ]

        async def firm() -> Dossier:
            deal = await _deal_for(session, save_fixture, user)
            await tieout.ingest(
                session,
                dossier_id=deal.id,
                kind=ArtifactKind.model,
                filename=PREAPP.name,
                payload=PREAPP.read_bytes(),
                user_id=user.id,
            )
            await session.flush()
            return deal

        firm_a, firm_b = await firm(), await firm()

        #: Criterion 2 — the settings surface serves every rule the
        #: engine runs (the two adopted rules included), and accepts
        #: switching them off. Before the catalogue carried them, this
        #: PUT was refused with a 422 calling `broken-name` « not a
        #: rule the audit runs » — criterion 1, kept on the record in
        #: the round document.
        response = await client.put(
            f"/v1/tieout/house-rules?organization_id={firm_b.organization_id}",
            json={"audit_rules_off": off},
        )
        assert response.status_code == 200, response.json()
        assert {one["key"] for one in response.json()["rules"] if not one["on"]} == set(
            off
        )
        for rule in ("gapped-test", "broken-name"):
            served = (
                await client.get(
                    f"/v1/tieout/house-rules?organization_id={firm_a.organization_id}"
                )
            ).json()["rules"]
            assert any(one["key"] == rule and one["on"] for one in served)

        reports = {}
        for deal in (firm_a, firm_b):
            runs = (await client.post(f"/v1/tieout/deals/{deal.id}/check")).json()
            audit = next(one for one in runs if one["kind"] == "audit")
            findings = (
                await client.get(f"/v1/tieout/deals/{deal.id}/findings?kind=audit")
            ).json()
            reports[deal.id] = (audit["summary"], findings)

        summary_a, found_a = reports[firm_a.id]
        summary_b, found_b = reports[firm_b.id]

        #: The two reports, verbatim, for the round document's evidence
        #: section — written only when asked for, never during CI, and
        #: before the criteria below so a failed run is still evidence.
        out = os.environ.get("HOUSE_RULES_DEMO_OUT")
        if out:
            Path(out).write_text(
                json.dumps(
                    {
                        "model": PREAPP.name,
                        "firm_a": {"summary": summary_a, "findings": found_a},
                        "firm_b": {
                            "rules_off": off,
                            "summary": summary_b,
                            "findings": found_b,
                        },
                    },
                    indent=2,
                )
            )

        #: Criterion 3 — Firm A's report is the conscience test's,
        #: through HTTP: twelve findings, the counters counting them,
        #: nothing switched off, the pass row and all four abstentions.
        assert Counter(one["rule"] for one in found_a) == {
            "broken-name": 2,
            "gapped-test": 1,
            "long-formula": 3,
            "hardcode-in-formula": 2,
            "hidden-sheet": 1,
            "inconsistent-row": 2,
            "skipped-cell": 1,
        }
        assert (summary_a["errors"], summary_a["smells"]) == (4, 8)
        assert summary_a["rules_off"] == []
        assert "model-own-check" in summary_a["tallies"]
        assert [one["rule"] for one in summary_a["abstentions"]] == [
            "balance-sheet",
            "cash-continuity",
            "debt-terminal",
            "interest-consistency",
        ]

        #: Criteria 4 and 5 — Firm B's report is Firm A's minus exactly
        #: the switched-off rules. Identity is checked field-for-field:
        #: a rule off must not reword, re-grade or reorder what remains.
        def carried(finding: dict[str, Any]) -> dict[str, Any]:
            #: Everything but the per-deal identifiers: `id` and
            #: `created_at` are the row's own, and the artifact id
            #: inside `where` differs per upload of the same bytes.
            return {
                key: (
                    {at: it for at, it in value.items() if at != "artifact_id"}
                    if key == "where"
                    else value
                )
                for key, value in finding.items()
                if key not in ("id", "created_at")
            }

        assert [carried(one) for one in found_b] == [
            carried(one) for one in found_a if one["rule"] not in off
        ]
        assert (summary_b["errors"], summary_b["smells"]) == (3, 3)
        assert summary_b["rules_off"] == off
        assert summary_b["tallies"] == {}
        assert [one["rule"] for one in summary_b["abstentions"]] == [
            "cash-continuity",
            "debt-terminal",
            "interest-consistency",
        ]


@pytest.mark.asyncio
class TestTheTeam:
    @pytest.mark.auth
    async def test_everyone_appears_and_no_deal_is_named(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The founder's decision (26 August): counts, never names.

        A colleague not on the deal must not learn what it is called
        from the team screen — the response carries no deal name
        anywhere, only how many deals each person is on.
        """
        deal = await _deal_for(session, save_fixture, user)
        colleague = await create_user(save_fixture)
        session.add(
            UserOrganization(user_id=colleague.id, organization_id=deal.organization_id)
        )
        await session.flush()

        response = await client.get(
            f"/v1/tieout/team?organization_id={deal.organization_id}"
        )
        assert response.status_code == 200
        body = response.json()
        assert body["total_deals"] == 1
        by_id = {one["id"]: one for one in body["members"]}
        mine = by_id[str(user.id)]
        assert mine["you"] is True
        assert mine["deal_count"] == 1
        theirs = by_id[str(colleague.id)]
        assert theirs["you"] is False
        # In the organization and on no deal — zero, never a borrowed
        # count.
        assert theirs["deal_count"] == 0
        # And the name itself appears nowhere in the payload: the
        # count is the most this screen may say.
        assert "Project Cascade" not in response.text
        assert "deals" not in mine


@pytest.mark.asyncio
class TestTheChat:
    """The three chat scopes, with a scripted model.

    The loop itself is proven in `tests/dossier/test_agent_loop.py`;
    what is proven here is the wiring — that a finding's context and the
    conversation reach the prompt, that the one-off chat holds only its
    file's tools, and that a stranger's check stays closed.
    """

    def _client(self, *script):
        from tests.dossier.test_agent_loop import FakeClient

        return FakeClient(*script)

    def _configure(self, monkeypatch: pytest.MonkeyPatch, client) -> None:
        import polar.tieout.endpoints as endpoints

        monkeypatch.setattr(endpoints, "agent_client", lambda: client)

    @pytest.mark.auth
    async def test_the_finding_and_the_conversation_reach_the_prompt(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from tests.dossier.test_agent_loop import says

        deal = await _loaded(session, save_fixture, user)
        finding = (await client.get(f"/v1/tieout/deals/{deal.id}/findings")).json()[0]

        fake = self._client(says("Looked up, not composed."))
        self._configure(monkeypatch, fake)
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/ask",
            json={
                "prompt": "What else reads this cell?",
                "finding_id": finding["id"],
                "history": [
                    {"who": "you", "text": "Where does it come from?"},
                    {"who": "pierce", "text": "Model!B26, a formula."},
                ],
            },
        )

        assert response.status_code == 201
        assert response.json()["answer"] == "Looked up, not composed."
        sent = fake.messages.calls[0]["messages"][0]["content"]
        assert finding["title"] in sent
        assert "The banker said: Where does it come from?" in sent
        assert "You answered: Model!B26, a formula." in sent
        assert sent.endswith("The banker now asks: What else reads this cell?")

    @pytest.mark.auth
    async def test_a_finding_from_another_deal_is_refused(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from tests.dossier.test_agent_loop import says

        deal = await _loaded(session, save_fixture, user)
        other = await _loaded(session, save_fixture, user)
        stray = (await client.get(f"/v1/tieout/deals/{other.id}/findings")).json()[0]

        self._configure(monkeypatch, self._client(says("never reached")))
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/ask",
            json={"prompt": "About that finding?", "finding_id": stray["id"]},
        )
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_the_file_chat_holds_only_its_files_tools(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from tests.dossier.test_agent_loop import calls, says

        checked = (
            await client.post(
                "/v1/tieout/check-file",
                files={"file": ("cascade_deck.pptx", CLEAN.read_bytes(), DECK_MEDIA)},
            )
        ).json()

        fake = self._client(
            calls("list_findings"),
            says("The chart and the table disagree on slide 3."),
        )
        self._configure(monkeypatch, fake)
        response = await client.post(
            f"/v1/tieout/check-file/{checked['id']}/ask",
            json={"prompt": "What disagrees inside this file?"},
        )

        assert response.status_code == 201
        body = response.json()
        assert body["answer"] == "The chart and the table disagree on slide 3."
        assert [step["tool"] for step in body["steps"]] == ["list_findings"]
        # The toolset offered is the file's own two tools and nothing
        # else — no deal tools to reach with.
        offered = {tool["name"] for tool in fake.messages.calls[0]["tools"]}
        assert offered == {"file_summary", "list_findings"}
        # And the system prompt carries the boundary.
        assert "one file" in fake.messages.calls[0]["system"]

    @pytest.mark.auth
    async def test_someone_elses_check_cannot_be_asked_about(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from tests.dossier.test_agent_loop import says

        other = await create_user(save_fixture)
        theirs = await tieout.check_file(
            session,
            user_id=other.id,
            kind=ArtifactKind.deck,
            filename="cascade_deck.pptx",
            payload=CLEAN.read_bytes(),
        )
        await session.flush()

        self._configure(monkeypatch, self._client(says("never reached")))
        response = await client.post(
            f"/v1/tieout/check-file/{theirs.id}/ask",
            json={"prompt": "What is in it?"},
        )
        assert response.status_code == 404


@pytest.mark.asyncio
class TestSinceYouLooked:
    """« Since you looked » — the watch's voice, derived from a visit row.

    The background watch re-syncs and re-checks without anybody looking,
    which clears *stale* silently; these two counts are what keep that
    from being invisible. Derived from timestamps against the person's
    last visit, never stored — so the test's assertions are about what
    the records imply, not about notification rows.
    """

    @pytest.mark.auth
    async def test_opening_the_deal_is_the_seen_event(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        # Never looked: no visit, and zeros — not « everything is new ».
        row = (await client.get("/v1/tieout/deals")).json()[0]
        assert row["visited_at"] is None
        assert row["arrived_since_visit"] == 0
        assert row["findings_since_visit"] == 0

        marked = await client.post(f"/v1/tieout/deals/{deal.id}/visit")
        assert marked.status_code == 204

        row = (await client.get("/v1/tieout/deals")).json()[0]
        assert row["visited_at"] is not None

    @pytest.mark.auth
    async def test_what_arrived_after_the_visit_is_counted(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()
        assert (
            await client.post(f"/v1/tieout/deals/{deal.id}/visit")
        ).status_code == 204

        # The watch's work arrives after the visit: files and findings.
        broken = CASCADE / "cascade_deck_broken.pptx"
        for path, kind in ((MODEL, ArtifactKind.model), (broken, ArtifactKind.deck)):
            await tieout.ingest(
                session,
                dossier_id=deal.id,
                kind=kind,
                filename=path.name,
                payload=path.read_bytes(),
                user_id=user.id,
            )
        await tieout.run_tieout(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()

        row = (await client.get("/v1/tieout/deals")).json()[0]
        assert row["arrived_since_visit"] == 2
        assert row["findings_since_visit"] > 0

        # Looking again clears it — for this person, and nobody else.
        assert (
            await client.post(f"/v1/tieout/deals/{deal.id}/visit")
        ).status_code == 204
        row = (await client.get("/v1/tieout/deals")).json()[0]
        assert row["arrived_since_visit"] == 0
        assert row["findings_since_visit"] == 0

    @pytest.mark.auth
    async def test_a_strangers_deal_cannot_be_visited(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        stranger = await create_user(save_fixture)
        theirs = await _deal_for(session, save_fixture, stranger)
        await session.flush()

        refused = await client.post(f"/v1/tieout/deals/{theirs.id}/visit")
        assert refused.status_code == 404


@pytest.mark.asyncio
class TestTheVersionAudit:
    """The version dropdown's re-scoping: one stored version, checked now.

    The route computes the audit from the picked version's stored cells
    and persists nothing — the deal's findings, runs and rulings belong
    to the current version, and looking at history must never move them.
    """

    @pytest.mark.auth
    async def test_an_old_version_answers_with_its_own_findings(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        first = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )
        # A second version of the same lineage — different bytes, same
        # filename, which is how a banker re-uploads « the model ».
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=(CASCADE / "audit_fixture.xlsx").read_bytes(),
            user_id=user.id,
        )
        await session.flush()

        response = await client.get(f"/v1/tieout/artifacts/{first.id}/audit")
        assert response.status_code == 200
        body = response.json()
        assert body["version"] == 1
        assert body["filename"] == "cascade_model.xlsx"
        # The summary carries the same record a stored run keeps, and the
        # findings are the audit's own shapes — rule, cell, evidence.
        assert body["summary"]["cells"] > 0
        for finding in body["findings"]:
            assert finding["kind"] == "audit"
            assert finding["state"] == "open"
            assert finding["where"]["anchor"]["kind"] == "cell"

    @pytest.mark.auth
    async def test_nothing_is_persisted_by_looking(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )
        before = {one.id for one in await repository.findings_of(deal.id)}

        response = await client.get(f"/v1/tieout/artifacts/{model.id}/audit")
        assert response.status_code == 200
        # The response's finding ids exist nowhere: they cannot be ruled
        # on, and the stored findings are exactly what they were.
        after = {one.id for one in await repository.findings_of(deal.id)}
        assert after == before
        for finding in response.json()["findings"]:
            ruled = await client.patch(
                f"/v1/tieout/findings/{finding['id']}",
                json={"state": "dismissed", "note": "should not exist"},
            )
            assert ruled.status_code == 404

    @pytest.mark.auth
    async def test_a_deck_has_no_version_audit(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        deck = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.deck
        )
        response = await client.get(f"/v1/tieout/artifacts/{deck.id}/audit")
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_strangers_version_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        theirs = await _loaded(session, save_fixture, stranger)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(theirs.id)
            if one.kind is ArtifactKind.model
        )
        response = await client.get(f"/v1/tieout/artifacts/{model.id}/audit")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestTheVersionDelta:
    """The Watch, served: what a revision did, in review language.

    The engine's own semantics are proven in `test_watch*`; what is
    tested here is the route — deal posture, the honest null for a
    first version, and that the wire shape carries the report's counts
    and ranked items as the engine produced them.
    """

    async def _two_versions(
        self, session: AsyncSession, save_fixture: SaveFixture, owner: User
    ) -> tuple[Dossier, Artifact, Artifact]:
        deal = await _deal_for(session, save_fixture, owner)
        first = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=owner.id,
        )
        second = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=(CASCADE / "audit_fixture.xlsx").read_bytes(),
            user_id=owner.id,
        )
        await session.flush()
        return deal, first, second

    @pytest.mark.auth
    async def test_a_revision_answers_in_review_language(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        _, first, second = await self._two_versions(session, save_fixture, user)

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/delta")
        assert response.status_code == 200
        body = response.json()
        assert body["old_version"] == 1
        assert body["new_version"] == 2
        # The study's counts are integers, never blurred; the nameless
        # are counted apart rather than guessed at.
        for key in (
            "new_defects",
            "repaired_defects",
            "persistent_defects",
            "unmatched_old",
            "unmatched_new",
        ):
            assert isinstance(body[key], int)
        # Items are the Watch's eight classes, in the engine's own rank.
        kinds = [item["kind"] for item in body["items"]]
        #: The engine's own ranking, not a copy of it. Pinning the list
        #: here made this test red the moment the Watch registered two
        #: new classes (`emptied_cell`, `filled_cell`) — a green test
        #: that fails on the engine growing is testing the wrong thing.
        from polar.tieout.watch.delta import _KIND_ORDER as order

        assert all(kind in order for kind in kinds)
        assert kinds == sorted(kinds, key=order.index)
        # Two entirely different workbooks under one lineage: the
        # report must have found *something* to say.
        assert body["items"]

    @pytest.mark.auth
    async def test_a_first_version_is_null_not_an_error(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        only = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.get(f"/v1/tieout/artifacts/{only.id}/delta")
        assert response.status_code == 200
        assert response.json() is None

    @pytest.mark.auth
    async def test_against_outside_the_lineage_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal, _, second = await self._two_versions(session, save_fixture, user)
        deck = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.deck,
            filename="cascade_deck.pptx",
            payload=CLEAN.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.get(
            f"/v1/tieout/artifacts/{second.id}/delta?against={deck.id}"
        )
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_strangers_revision_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        _, _, second = await self._two_versions(session, save_fixture, stranger)
        response = await client.get(f"/v1/tieout/artifacts/{second.id}/delta")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestTheSourcePage:
    """The source viewer's ground: a stored PDF's page, as pixels.

    The Chain's facts cite pages and boxes; this route serves the
    pixels those citations sit on. What matters at the route: the deal
    posture, the honest out-of-range sentence, and that a non-PDF
    reads as not found rather than half-rendering.
    """

    async def _source(
        self, session: AsyncSession, save_fixture: SaveFixture, owner: User
    ) -> tuple[Dossier, Artifact]:
        deal = await _deal_for(session, save_fixture, owner)
        source = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.source,
            filename="cascade_accounts.pdf",
            payload=(CASCADE / "cascade_accounts.pdf").read_bytes(),
            user_id=owner.id,
        )
        await session.flush()
        return deal, source

    @pytest.mark.auth
    async def test_a_page_renders_as_png(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        _, source = await self._source(session, save_fixture, user)
        response = await client.get(f"/v1/tieout/artifacts/{source.id}/page/1")
        assert response.status_code == 200
        assert response.headers["content-type"] == "image/png"
        assert response.content.startswith(b"\x89PNG")

    @pytest.mark.auth
    async def test_a_page_outside_the_document_says_the_range(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        _, source = await self._source(session, save_fixture, user)
        response = await client.get(f"/v1/tieout/artifacts/{source.id}/page/999")
        assert response.status_code == 404
        assert "no page 999" in response.json()["detail"]

    @pytest.mark.auth
    async def test_a_model_is_not_a_page_source(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.get(f"/v1/tieout/artifacts/{model.id}/page/1")
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_strangers_page_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        _, source = await self._source(session, save_fixture, stranger)
        response = await client.get(f"/v1/tieout/artifacts/{source.id}/page/1")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestTheDealsList:
    """« Last checked » must mean what the column says.

    A deal with no deck has nothing to reconcile, so its tie-out never
    runs — and the list read « Not checked yet » beside the findings
    its own audit had just produced. Most of the real corpus is exactly
    that shape: a model, no deliverables.
    """

    @pytest.mark.auth
    async def test_an_audited_deal_is_not_called_unchecked(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        run = await tieout.run_audit(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()
        assert run.finished_at is not None

        response = await client.get("/v1/tieout/deals")
        assert response.status_code == 200
        row = next(one for one in response.json() if one["id"] == str(deal.id))
        #: The audit ran and finished; the column says « Last checked ».
        assert row["checked_at"] is not None


@pytest.mark.asyncio
class TestWhatIntakeWillNotRead:
    """The refusal a person meets when the format is not ours.

    Found by the corpus rather than by a customer: two of eleven
    eligible models arrived as `.xlsb` (`population-proof.md`,
    27 Aug). **That one is no longer refused** — since 28 Aug
    `polar.tieout.binary` converts it and the reader takes it — so what
    is tested here is that the formats still named say what to do, and
    that `.xlsb` is not among them.
    """

    @pytest.mark.auth
    async def test_the_binary_workbook_is_no_longer_refused_on_format(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """A `.xlsb` is read now, so it may not be turned away for being
        one. These bytes are not a workbook, so this upload still fails —
        the point is *why*: for its contents, not for its extension."""
        deal = await _deal_for(session, save_fixture, user)
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={
                "file": ("model.xlsb", b"PK\x03\x04binary", "application/octet-stream")
            },
        )
        assert response.status_code != 415
        assert "Save As" not in str(response.json().get("detail", ""))

    @pytest.mark.auth
    async def test_a_format_we_really_cannot_read_says_what_to_do(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={"file": ("model.numbers", b"whatever", "application/octet-stream")},
        )
        assert response.status_code == 415
        # The fix, not just the fact: a person can act on this.
        assert ".xlsx" in response.json()["detail"]

    @pytest.mark.auth
    async def test_the_formats_named_are_the_formats_read(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={"file": ("sketch.psd", b"8BPS", "application/octet-stream")},
        )
        assert response.status_code == 415
        detail = response.json()["detail"]
        #: Derived from the reader's own table, never a copy of it: the
        #: sentence said « models are .xlsx or .xls » while `.xlsm` —
        #: the format most project-finance models arrive in — had been
        #: accepted all along.
        for suffixes in SUFFIXES.values():
            for suffix in suffixes:
                assert suffix in detail, f"{suffix} is read but not named"

    @pytest.mark.auth
    async def test_a_macro_enabled_model_is_taken(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        response = await client.post(
            f"/v1/tieout/deals/{deal.id}/artifacts",
            files={
                "file": (
                    "macro_model.xlsm",
                    MODEL.read_bytes(),
                    "application/vnd.ms-excel.sheet.macroEnabled.12",
                )
            },
        )
        assert response.status_code == 200
        assert response.json()["status"] == "ready"


@pytest.mark.asyncio
class TestTheRecalculation:
    """The mark: « validated by recalculation », and its honest refusal face.

    The gate's own semantics are proven in the engine's tests; what is
    tested here is the route — deal posture, that only a model earns a
    mark, that a refused file answers in words without any engine
    running, and that the mark persists onto the artifact so every
    later read repeats it. The full-engine pass runs only where an
    adequate LibreOffice exists, and is skipped honestly elsewhere.
    """

    @staticmethod
    def _rtd_model() -> bytes:
        """A tiny real workbook carrying one RTD call — refused, no engine."""
        import io

        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet["A1"] = "Revenue"
        sheet["B1"] = 100
        sheet["A2"] = "Live price"
        sheet["B2"] = '=RTD("feed.prog",,"topic")'
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    @pytest.mark.auth
    async def test_a_denylisted_model_is_refused_in_words(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="live_feed_model.xlsx",
            payload=self._rtd_model(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.post(f"/v1/tieout/artifacts/{model.id}/recalculate")
        assert response.status_code == 200
        body = response.json()
        assert body["verdict"] == "refused"
        # No engine ran, and none is claimed.
        assert body["engine"] is None
        assert body["compared"] == 0
        # The refusal is words a person can act on: the construct, the
        # cell, and where it routes — RTD is gone the moment the file
        # is saved, so nothing we run could recompute it.
        assert body["refusal_count"] >= 1
        refusal = body["refusals"][0]
        assert refusal["category"] == "rtd"
        assert refusal["ref"] == "Sheet!B2"
        assert body["route"] == "refuse"

    @pytest.mark.auth
    async def test_the_mark_persists_onto_the_artifact(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="live_feed_model.xlsx",
            payload=self._rtd_model(),
            user_id=user.id,
        )
        await session.flush()
        marked = await client.post(f"/v1/tieout/artifacts/{model.id}/recalculate")
        assert marked.status_code == 200
        again = await client.get(f"/v1/tieout/artifacts/{model.id}")
        assert again.status_code == 200
        kept = again.json()["counts"]["recalc"]
        assert kept["verdict"] == "refused"
        assert kept["refusals"][0]["ref"] == "Sheet!B2"

    @pytest.mark.skipif(
        find_install() is None,
        reason="no LibreOffice >= 25.8 on this machine (dev/setup-libreoffice)",
    )
    @pytest.mark.auth
    async def test_a_clean_model_is_validated_by_the_real_engine(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.post(f"/v1/tieout/artifacts/{model.id}/recalculate")
        assert response.status_code == 200
        body = response.json()
        # The fixture's cached values were computed by a real engine, so
        # a real recalculation reproduces them — and the mark names the
        # engine so a fake can never be mistaken for a machine result.
        assert body["verdict"] == "pass"
        assert body["engine"] is not None
        assert "LibreOffice" in body["engine"]
        assert body["compared"] > 0
        assert body["matched"] == body["compared"]
        assert body["match_rate"] == 1.0

    @pytest.mark.auth
    async def test_a_deck_cannot_be_recalculated(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        deck = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.deck,
            filename="cascade_deck.pptx",
            payload=CLEAN.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        response = await client.post(f"/v1/tieout/artifacts/{deck.id}/recalculate")
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_strangers_model_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        deal = await _deal_for(session, save_fixture, stranger)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=stranger.id,
        )
        await session.flush()
        response = await client.post(f"/v1/tieout/artifacts/{model.id}/recalculate")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestTheMarkedUpModel:
    """« Download the marked-up model », through HTTP.

    ``test_markup.py`` proves the surgery — colour and notes only,
    nothing altered, verified before release. What is left to prove is
    the route: that the download is closed exactly as the deal is, that
    a deal with nothing to mark says so instead of handing over an
    untouched copy, and that what comes back is the workbook the
    founder's card promises — a Findings sheet in front, the model's
    own sheets behind it, under a filename that is not the original's.
    """

    @pytest.mark.auth
    async def test_the_download_is_the_marked_copy(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        import io
        import zipfile

        from openpyxl import load_workbook

        deal = await _loaded(session, save_fixture, user)

        response = await client.get(f"/v1/tieout/deals/{deal.id}/markup")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        # A different filename, so the original is never at risk.
        disposition = response.headers["content-disposition"]
        assert "attachment" in disposition
        assert "marked" in disposition
        assert "cascade_model.xlsx" not in disposition

        copy = load_workbook(io.BytesIO(response.content))
        # The findings list rides in front; the model rides behind it,
        # whole — every original sheet, in its order.
        assert copy.sheetnames[0] == "Findings"
        assert copy.sheetnames[1:] == [
            "Assumptions",
            "Model",
            "DCF",
            "Comps",
            "Outputs",
        ]
        # The listing carries one row per marked finding, under a header.
        listing = copy["Findings"]
        assert listing.max_row >= 2

        # And it is genuinely a different file from the one uploaded —
        # never the original handed back under a new name.
        original = MODEL.read_bytes()
        assert response.content != original
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert archive.testzip() is None

    @pytest.mark.auth
    async def test_a_stranger_gets_the_same_404_as_everywhere(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
    ) -> None:
        stranger = await create_user(save_fixture)
        theirs = await _loaded(session, save_fixture, stranger)

        response = await client.get(f"/v1/tieout/deals/{theirs.id}/markup")
        assert response.status_code == 404

    @pytest.mark.auth
    async def test_a_deal_with_no_model_says_so(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await session.flush()

        response = await client.get(f"/v1/tieout/deals/{deal.id}/markup")
        assert response.status_code == 404
        assert "no model" in response.json()["detail"]

    @pytest.mark.auth
    async def test_nothing_open_means_nothing_to_hand_over(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Every model finding ruled on ⇒ 404, not an unmarked copy.

        The card exists to hand over marked problems; a clean handover
        of an untouched file would look like the product vouching for
        the model, which it is not doing.
        """
        deal = await _loaded(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        model = next(
            one
            for one in await repository.current_artifacts(deal.id)
            if one.kind is ArtifactKind.model
        )
        listed = (
            await client.get(
                f"/v1/tieout/deals/{deal.id}/findings",
                params={"artifact_id": str(model.id)},
            )
        ).json()
        assert listed  # the model has findings, or this test proves nothing
        for finding in listed:
            ruled = await client.patch(
                f"/v1/tieout/findings/{finding['id']}",
                json={"state": "dismissed", "note": "ruled on for this test"},
            )
            assert ruled.status_code == 200

        response = await client.get(f"/v1/tieout/deals/{deal.id}/markup")
        assert response.status_code == 404
        assert "nothing to mark up" in response.json()["detail"]


class TestTheCategoryMap:
    """Every rule the engine emits has a family on a partner's report.

    The report groups findings by family, and a rule with no family
    falls to « Other findings » — the heading a partner reads over a
    defect the engine called an error. The map lives in the frontend
    (`clients/apps/web/src/components/Workspace/files.ts`), so nothing
    on the server would notice it going stale.

    It went stale twice, and both times for the same reason: the map was
    re-checked against `RULE_NAMES` and `ANALYTIC_RULE_NAMES` rather
    than against the rules the engine actually writes into a finding.
    A rule missing from the catalogues is precisely the rule most likely
    to be missing from the map, and reading the catalogues finds nothing
    wrong with it. `broken-name` fires on eight of the nine readable
    corpus models and was unmapped the whole time.

    So this test reads the rule literals out of the engine's own source.
    That is coarse — it cannot see a rule composed at runtime — but it
    is the check that would have caught what the careful one missed.
    """

    ROOT = Path(__file__).resolve().parents[3]
    ENGINE = ROOT / "server" / "polar" / "tieout"
    FILES_TS = (
        ROOT
        / "clients"
        / "apps"
        / "web"
        / "src"
        / "components"
        / "Workspace"
        / "files.ts"
    )

    def _emitted(self) -> set[str]:
        import re

        rules: set[str] = set()
        for source in self.ENGINE.rglob("*.py"):
            rules |= set(re.findall(r'rule="([a-z0-9-]+)"', source.read_text()))
        return rules

    def _mapped(self) -> set[str]:
        import re

        text = self.FILES_TS.read_text()
        start = text.index("const CATEGORY_OF")
        body = text[start : text.index("}", start)]
        return set(re.findall(r"^\s*'?([a-z0-9-]+)'?:", body, re.MULTILINE))

    def test_the_engine_emits_rules(self) -> None:
        """The literal scan works — a green suite over nothing proves nothing."""
        emitted = self._emitted()
        assert "balance-sheet" in emitted
        assert "typed-over-formula" in emitted
        assert len(emitted) > 15

    def test_every_emitted_rule_has_a_family(self) -> None:
        unmapped = sorted(self._emitted() - self._mapped())
        assert unmapped == [], (
            "these rules reach a report with no family and read as "
            f"« Other findings »: {', '.join(unmapped)} — map them in "
            "files.ts"
        )

    #: Rules mapped **ahead of the engine**, because a merge is
    #: sequenced behind the mapping. The lead held Sentinel's unit
    #: checks at the tip until this map covered them, so for one sweep
    #: the map legitimately knows two rules the merged engine does not.
    #:
    #: The exemption clears itself: `test_the_pending_map_is_still
    #: _pending` fails the moment the engine *does* emit one of these,
    #: which is the sweep it has to be deleted. An exemption that
    #: outlives its reason is how a guard quietly stops guarding.
    #: Emptied on 28 Aug at the sweep that merged Sentinel's unit checks,
    #: which is what this exemption was built to provoke. Both rules now
    #: appear in the engine's source and both are mapped in `files.ts`
    #: to « Units that do not agree », so the full guard applies again.
    #:
    #: Worth knowing, because it reads oddly: `_unit_mismatch` is present
    #: and unit-tested but **deliberately not wired into `audit()`** —
    #: measured on the closed-deal corpus it raised 103 findings and every
    #: one was a false alarm. So the engine cannot emit these today. The
    #: map having a home for them is right either way: the family is
    #: ready for the round that arms them, and until then nothing reaches
    #: it. Put a name back here only for a rule mapped *ahead* of its
    #: code, which is the case this list exists for.
    AHEAD_OF_THE_ENGINE: set[str] = set()

    def test_the_map_invents_nothing(self) -> None:
        """A family for a rule the engine cannot emit is dead prose."""
        invented = sorted(self._mapped() - self._emitted() - self.AHEAD_OF_THE_ENGINE)
        assert invented == [], (
            f"files.ts files rules the engine never emits: {', '.join(invented)}"
        )

    def test_the_pending_map_is_still_pending(self) -> None:
        """Delete the exemption the sweep its rule lands.

        This is the half that keeps the guard honest. Sentinel's unit
        checks are mapped before they merge, on the lead's sequencing;
        the moment they arrive this test goes red and the only way to
        make it green is to remove the name from
        `AHEAD_OF_THE_ENGINE`, which restores the full check.
        """
        landed = sorted(self.AHEAD_OF_THE_ENGINE & self._emitted())
        assert landed == [], (
            f"{', '.join(landed)} is emitted by the engine now — delete it "
            "from AHEAD_OF_THE_ENGINE so the map is fully guarded again"
        )


@pytest.mark.asyncio
class TestWhichModelAnAnswerIsAbout:
    """A deal-scoped answer names the model it read.

    Three places narrowed a deal to one model with
    `next(one for one in current if one.kind is model)` — the deals
    list's model column, the marked-up download, and the assistant's
    workspace. That took whichever lineage `current_artifacts` happened
    to return first, which follows `list_artifacts`' ordering and
    promises nothing. It was the newest model most of the time, by
    luck, and silently another one the rest.

    The assistant is where it hurt: it answers in prose, so a wrong
    pick is a confident paragraph about the wrong workbook with nothing
    on the screen to say so. Found on a demo deal carrying three model
    lineages, where chat answered about a scratch file while the deal's
    subject sat next to it.

    The fix is `service.subject_model` — the most recently uploaded,
    said out loud — and these tests hold both halves: that all three
    callers agree, and that the answer names what it read.
    """

    async def _two_models(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> tuple[Dossier, str]:
        """A deal carrying two models, the second uploaded later."""
        deal = await _deal_for(session, save_fixture, user)
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="lenders_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        return deal, "cascade_model.xlsx"

    async def test_the_subject_is_the_most_recently_uploaded(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        from polar.tieout.service import models_of, subject_model

        deal, newest = await self._two_models(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        current = await repository.current_artifacts(deal.id)

        chosen = subject_model(current)
        assert chosen is not None
        assert chosen.filename == newest
        #: And the ordering is total, not « whatever came back first ».
        assert [one.filename for one in models_of(current)] == [
            newest,
            "lenders_model.xlsx",
        ]

    async def test_a_deal_with_no_model_has_no_subject(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        from polar.tieout.service import subject_model

        deal = await _deal_for(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        assert subject_model(await repository.current_artifacts(deal.id)) is None

    @pytest.mark.auth
    async def test_the_deals_list_names_the_same_model(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The row and the assistant must not disagree about the file."""
        deal, newest = await self._two_models(session, save_fixture, user)

        response = await client.get("/v1/tieout/deals")
        assert response.status_code == 200
        row = next(one for one in response.json() if one["id"] == str(deal.id))
        assert row["model_name"] == newest

    async def test_the_assistants_workspace_takes_the_subject_and_names_the_rest(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        from polar.tieout.agent.service import load_model_workspace

        deal, newest = await self._two_models(session, save_fixture, user)
        workspace = await load_model_workspace(session, deal.id, "the deal")

        assert workspace is not None
        assert workspace.filename == newest
        #: The files this answer is *not* about, named — a count would
        #: not let a reader tell whether the right one was read.
        assert workspace.others == ["lenders_model.xlsx (v1)"]

    async def test_one_model_leaves_the_scope_line_unsaid(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The ordinary deal must not grow a paragraph about itself."""
        from polar.tieout.agent.service import load_model_workspace

        deal = await _deal_for(session, save_fixture, user)
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        workspace = await load_model_workspace(session, deal.id, "the deal")
        assert workspace is not None
        assert workspace.others == []

    async def test_the_prompt_tells_the_assistant_what_it_cannot_see(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Without this the assistant cannot know the other file exists.

        Its tools only ever hold one model, so asked about a line that
        lives in the deal's *other* workbook it would answer « that is
        not in this model » — true, and read as « your deal does not
        contain it ».
        """
        from polar.tieout.endpoints import _scope_line
        from polar.tieout.service import models_of

        deal, newest = await self._two_models(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)
        models = models_of(await repository.current_artifacts(deal.id))

        said = _scope_line(models[0], models[1:])
        assert newest in said
        assert "lenders_model.xlsx (v1)" in said
        assert "cannot see" in said
        #: One model, and the prompt says nothing at all about scope.
        assert _scope_line(models[0], []) == ""


@pytest.mark.asyncio
class TestWhatARevisionDidToTheDeck:
    """The failure this product exists for, served.

    Not one typo: a model revision the deck never caught up with,
    because nobody knows which of its hundred printed figures the
    revision touched. The Watch's C5 comparison has been in the engine
    with no endpoint and no screen behind it — the same deck tied out
    against both versions, and every break attributed to the model
    change underneath it.

    The engine's semantics are proven in `test_watch*`. What is tested
    here is the route: deal posture, the honest null, and — the part
    that carries the promise — that the four lists come back as four
    different sentences and are never summed into one.
    """

    async def _revision_with_a_deck(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        owner: User,
        *,
        second: bytes | None = None,
    ) -> tuple[Dossier, Artifact, Artifact]:
        deal = await _deal_for(session, save_fixture, owner)
        first = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=owner.id,
        )
        later = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=second if second is not None else MODEL.read_bytes(),
            user_id=owner.id,
        )
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.deck,
            filename="cascade_deck.pptx",
            payload=CLEAN.read_bytes(),
            user_id=owner.id,
        )
        await session.flush()
        return deal, first, later

    @pytest.mark.auth
    async def test_a_revision_that_changed_nothing_breaks_nothing(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The property that makes the rest trustworthy.

        The same workbook re-uploaded cannot have broken a figure. The
        deck disagrees with the model in eight places either way, and
        every one of them belongs to « still drifting » — the list
        whose whole job is to keep a deck's pre-existing quarrels off
        this revision's account.
        """
        _, _, second = await self._revision_with_a_deck(session, save_fixture, user)

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/deck-delta")
        assert response.status_code == 200
        body = response.json()
        assert body["broken"] == []
        assert body["repaired"] == []
        assert body["still_drifting"] != []
        #: Coverage did not move either — the same model reconciles the
        #: same figures.
        assert body["coverage_changed"] == []
        assert body["checked_old"] == body["checked_new"]

    @pytest.mark.auth
    async def test_a_real_revision_separates_broken_from_lost_sight_of(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """« I lost sight of it » is not « it broke ».

        Folding the two together is how a checker earns a reputation
        for crying wolf: this revision reconciles far fewer of the
        deck's figures than the version before, and every figure it can
        no longer reach would read as a break if the lists were summed.
        """
        _, _, second = await self._revision_with_a_deck(
            session,
            save_fixture,
            user,
            second=(CASCADE / "audit_fixture.xlsx").read_bytes(),
        )

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/deck-delta")
        assert response.status_code == 200
        body = response.json()
        assert body["broken"], "a real revision moved figures the deck prints"
        assert body["coverage_changed"], "and lost sight of others"
        #: The counts are the honest reason coverage moved.
        assert body["checked_new"] < body["checked_old"]
        #: A break names what the deck prints and what the model now
        #: says — both, because one without the other is not checkable.
        first = body["broken"][0]
        assert first["printed"]
        assert first["expected"]
        assert first["slide"] > 0

    @pytest.mark.auth
    async def test_a_cause_is_never_guessed(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """An unattributed break says nothing rather than the nearest change.

        `cause` carries the model change underneath a break in the
        Watch's own words. Where the Watch could not attribute it the
        field is empty and the screen says so in words — a guess
        printed as a cause is the one claim a banker would repeat.
        """
        _, _, second = await self._revision_with_a_deck(
            session,
            save_fixture,
            user,
            second=(CASCADE / "audit_fixture.xlsx").read_bytes(),
        )

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/deck-delta")
        assert response.status_code == 200
        for one in response.json()["broken"]:
            assert "cause" in one
            #: Empty is allowed; invented is not — every non-empty cause
            #: has to be a sentence the Watch wrote, never a bare ref.
            assert one["cause"] == "" or len(one["cause"]) > 3

    @pytest.mark.auth
    async def test_a_first_version_is_null_not_an_error(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        only = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()

        response = await client.get(f"/v1/tieout/artifacts/{only.id}/deck-delta")
        assert response.status_code == 200
        assert response.json() is None

    @pytest.mark.auth
    async def test_a_deal_with_no_deck_is_null_not_an_error(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Nothing was sent out, so nothing can have gone stale."""
        deal = await _deal_for(session, save_fixture, user)
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        second = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=user.id,
        )
        await session.flush()

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/deck-delta")
        assert response.status_code == 200
        assert response.json() is None

    @pytest.mark.auth
    async def test_a_strangers_revision_does_not_exist(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Closed by default, like every other route reachable from a deal."""
        stranger = await create_user(save_fixture)
        _, _, second = await self._revision_with_a_deck(session, save_fixture, stranger)

        response = await client.get(f"/v1/tieout/artifacts/{second.id}/deck-delta")
        assert response.status_code == 404


@pytest.mark.asyncio
class TestTheTwoWaysToReadCells:
    """The light read must be the same read.

    The assistant waited 28 seconds before it could think on a real
    model, and 16 of those were the ORM building an instrumented object
    for each of 470,594 cells that `_workbook_of` reads once and throws
    away. `cells_for_graph` returns the same rows as columns — 4.0
    seconds against 16.0, medians of three alternating runs — and the
    audit and the assistant now take it.

    A timing test would be a flake, so what is held here is the
    property that made the switch safe: **the workbook built from
    either read is the same workbook**. If it ever is not, every audit
    finding on every model is suspect, which is a far worse failure
    than a slow one.
    """

    async def _model(
        self, session: AsyncSession, save_fixture: SaveFixture, owner: User
    ) -> Artifact:
        deal = await _deal_for(session, save_fixture, owner)
        made = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=owner.id,
        )
        await session.flush()
        return made

    async def test_both_reads_build_the_same_workbook(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        from polar.tieout.service import _workbook_of

        model = await self._model(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)

        heavy = _workbook_of(await repository.cells_of(model.id))
        light = _workbook_of(await repository.cells_for_graph(model.id))

        assert set(heavy.cells) == set(light.cells)
        assert heavy.sheets == light.sheets
        assert heavy.cells, "the fixture has to have cells for this to mean anything"
        for ref, was in heavy.cells.items():
            now = light.cells[ref]
            assert (was.sheet, was.row, was.column) == (now.sheet, now.row, now.column)
            assert was.value == now.value
            assert was.formula == now.formula
            assert (was.row_label, was.column_label) == (
                now.row_label,
                now.column_label,
            )
            assert was.precedents == now.precedents
            assert was.unresolved == now.unresolved
            assert was.alias_of == now.alias_of

    async def test_the_light_read_carries_no_identity(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The one difference, and it is the point.

        A link points at a cell *row*, so grounding and the tie-out need
        `cells_of`. Leaving `id` off the light read is what stops a
        caller reaching for the fast one and quietly losing the ability
        to say which row it meant.
        """
        model = await self._model(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)

        light = await repository.cells_for_graph(model.id)
        assert light
        assert not hasattr(light[0], "id")
        heavy = await repository.cells_of(model.id)
        assert heavy[0].id is not None

    async def test_an_audit_reads_the_same_findings_either_way(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The switch's real blast radius, checked end to end.

        `run_audit` is what changed; this proves the findings it would
        have produced from the ORM read are the findings it produces now.
        """
        model = await self._model(session, save_fixture, user)
        repository = TieOutRepository.from_session(session)

        heavy_findings, heavy_record = tieout._audit_one(
            model,
            await repository.cells_of(model.id),
            dossier_id=model.dossier_id,
            check_run_id=None,
            rules_off=set(),
        )
        light_findings, light_record = tieout._audit_one(
            model,
            await repository.cells_for_graph(model.id),
            dossier_id=model.dossier_id,
            check_run_id=None,
            rules_off=set(),
        )

        assert heavy_record == light_record
        assert [(one.rule, one.location, one.title) for one in heavy_findings] == [
            (one.rule, one.location, one.title) for one in light_findings
        ]


@pytest.mark.asyncio
class TestTheFactsTheCellsCannotSay:
    """The product must audit the same workbook the engine does.

    It never audits a file: `_workbook_of` rebuilds one from stored
    rows, and every field the reader filled at *open* time was empty by
    the time a rule read it. `hidden_sheets` was the one anybody
    noticed, and it was fixed alone. Over the nine readable corpus
    models the rest cost **41 of 116 findings** — `error-value` ×28,
    thirteen at error severity, `broken-name` ×12, one `hidden-sheet`
    downgraded — and took four models to « nothing failing ». One of
    them prints `#N/A` across forty-eight cells of a live repayment
    column while its report says nothing is failing.

    Ingest now keeps those facts on the artifact and `_audit_one` puts
    them back. Restoring them closes the corpus gap exactly: nothing
    missing, nothing invented, on all nine.
    """

    PREAPP = CASCADE / "example_preapp_model.xlsx"

    async def _audited(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        owner: User,
        *,
        payload: bytes,
        filename: str,
    ) -> tuple[Dossier, Artifact]:
        deal = await _deal_for(session, save_fixture, owner)
        model = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename=filename,
            payload=payload,
            user_id=owner.id,
        )
        await session.flush()
        return deal, model

    async def test_broken_defined_names_reach_a_finding(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The repo's own judged model carries fifty of them.

        `broken-name` fires on eight of the nine readable corpus models
        and appeared in no deal in the database, because the rule reads
        `book.broken_names` and the rebuilt workbook had none.
        """
        deal, _ = await self._audited(
            session,
            save_fixture,
            user,
            payload=self.PREAPP.read_bytes(),
            filename="example_preapp_model.xlsx",
        )
        run = await tieout.run_audit(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()

        repository = TieOutRepository.from_session(session)
        rules = {one.rule for one in await repository.findings_of(deal.id)}
        assert "broken-name" in rules
        assert run.status is CheckStatus.done

    async def test_cached_error_values_reach_a_finding(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """`error-value` was the biggest loss of the four — and dead.

        It reads `book.errors`, which the reader fills from the values
        Excel cached. A cell showing `#N/A` has no number, so it is not
        even stored as a cell: without the fact carried over, the
        product could not know the cell existed.
        """
        deal, _ = await self._audited(
            session,
            save_fixture,
            user,
            payload=(CASCADE / "audit_fixture.xlsx").read_bytes(),
            filename="audit_fixture.xlsx",
        )
        await tieout.run_audit(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()

        repository = TieOutRepository.from_session(session)
        rules = {one.rule for one in await repository.findings_of(deal.id)}
        assert "error-value" in rules

    async def test_the_stored_facts_come_back_as_they_went_in(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """JSON has no integer keys and no tuples; the restore knows."""
        from polar.tieout.service import _restore_file_facts
        from polar.tieout.workbook import Workbook, read_workbook

        _, model = await self._audited(
            session,
            save_fixture,
            user,
            payload=self.PREAPP.read_bytes(),
            filename="example_preapp_model.xlsx",
        )
        from_file = read_workbook(str(self.PREAPP))

        rebuilt = Workbook()
        _restore_file_facts(rebuilt, model.counts)

        assert rebuilt.broken_names == from_file.broken_names
        assert rebuilt.foreign_names == from_file.foreign_names
        assert rebuilt.errors == from_file.errors
        assert rebuilt.unparseable == from_file.unparseable
        assert rebuilt.populated == from_file.populated
        assert rebuilt.iterative == from_file.iterative
        #: The row numbers went to JSON as strings and have to come back
        #: as integers, or every lookup the audit makes misses.
        assert rebuilt.row_words == from_file.row_words
        assert all(
            isinstance(row, int) for rows in rebuilt.row_words.values() for row in rows
        )

    async def test_a_model_stored_before_the_fix_still_reads(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """No migration, and no pretending an old artifact knows more.

        Every model already in a database was ingested without these
        facts. It has to audit exactly as it did — poorer, and without
        raising — and uploading it again is what teaches it.
        """
        from polar.tieout.service import _restore_file_facts
        from polar.tieout.workbook import Workbook

        deal, model = await self._audited(
            session,
            save_fixture,
            user,
            payload=self.PREAPP.read_bytes(),
            filename="example_preapp_model.xlsx",
        )
        counts = dict(model.counts)
        counts.pop("workbook", None)
        model.counts = counts
        session.add(model)
        await session.flush()

        rebuilt = Workbook()
        _restore_file_facts(rebuilt, model.counts)
        assert rebuilt.broken_names == []
        assert rebuilt.errors == {}
        assert rebuilt.row_words == {}

        run = await tieout.run_audit(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()
        assert run.status is CheckStatus.done
        repository = TieOutRepository.from_session(session)
        rules = {one.rule for one in await repository.findings_of(deal.id)}
        #: The old, poorer answer — which is the honest one for an
        #: artifact that never carried the fact.
        assert "broken-name" not in rules


@pytest.mark.asyncio
class TestAFindingWithNoCell:
    """A finding about the workbook still says where it is.

    Most findings sit at a cell. Some are about the *file* — the
    defined names pointing into workbooks that are not here, which on
    the repo's judged model is fifty of them and on a corpus model
    three hundred and thirty-eight. Those carry no sheet and no ref,
    and every screen drew an empty grey pill beside them: the reader
    sees a rendering fault where the truth is « the whole workbook ».

    The engine already names what such a finding is about. Saying that
    is both shorter and truer than a blank.
    """

    @pytest.mark.auth
    async def test_a_workbook_finding_names_what_it_is_about(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        deal = await _deal_for(session, save_fixture, user)
        await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="example_preapp_model.xlsx",
            payload=(CASCADE / "example_preapp_model.xlsx").read_bytes(),
            user_id=user.id,
        )
        await session.flush()
        await tieout.run_audit(session, dossier_id=deal.id, user_id=user.id)
        await session.flush()

        response = await client.get(f"/v1/tieout/deals/{deal.id}/findings")
        assert response.status_code == 200
        findings = response.json()
        workbook_wide = [one for one in findings if one["rule"] == "broken-name"]
        assert workbook_wide, "the judged model carries fifty broken names"
        for one in workbook_wide:
            #: No cell, by the rule's nature — and still not blank.
            assert one["where"]["anchor"].get("ref") in (None, "")
            assert one["where"]["label"], "an empty pill is a rendering fault"
            assert one["where"]["label"] == "defined names"


@pytest.mark.asyncio
class TestTheSameFileUploadedAgain:
    """Two uploads with one digest are one file, and nothing to compare.

    Refusing « compute the delta from stored cells » was right — it lost
    an `unmatched_new` on a real pair — and a refusal is half a turn.
    This is the successor, and it differs in kind: not a faster
    comparison, but not comparing at all.

    Measured, that matters: the Watch spends **158 seconds** on a
    432,596-cell model to conclude a re-upload changed nothing — 58 of
    them reading the two files, 5 auditing them, and most of the rest
    aligning two large sheets against themselves. Identical bytes are
    identical workbooks, so the same answer comes out of a string
    comparison and comes out *exact*.

    Ingest has recorded the digest since 28 Aug. A version stored before
    that has none, and is compared as before: two absences are not a
    match, which is the case that would otherwise turn « we know
    nothing about either file » into « they are the same ».
    """

    async def _pair(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        owner: User,
        *,
        second: bytes | None = None,
    ) -> tuple[Artifact, Artifact]:
        deal = await _deal_for(session, save_fixture, owner)
        first = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=MODEL.read_bytes(),
            user_id=owner.id,
        )
        later = await tieout.ingest(
            session,
            dossier_id=deal.id,
            kind=ArtifactKind.model,
            filename="cascade_model.xlsx",
            payload=second if second is not None else MODEL.read_bytes(),
            user_id=owner.id,
        )
        await session.flush()
        return first, later

    async def test_a_re_upload_is_recognised_without_reading_anything(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        first, later = await self._pair(session, save_fixture, user)
        assert first.counts["sha256"] == later.counts["sha256"]
        assert tieout.identical_upload(first, later) is True

    async def test_a_real_revision_is_not(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        first, later = await self._pair(
            session,
            save_fixture,
            user,
            second=(CASCADE / "audit_fixture.xlsx").read_bytes(),
        )
        assert first.counts["sha256"] != later.counts["sha256"]
        assert tieout.identical_upload(first, later) is False

    async def test_two_absences_are_not_a_match(
        self,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """Every version stored before 28 Aug carries no digest.

        Reading « both have none » as « both are the same » would tell
        a reader two different files are one, which is the worst
        answer this fast path could give.
        """
        first, later = await self._pair(session, save_fixture, user)
        for one in (first, later):
            counts = dict(one.counts)
            counts.pop("sha256", None)
            one.counts = counts
            session.add(one)
        await session.flush()

        assert tieout.identical_upload(first, later) is False

    @pytest.mark.auth
    async def test_the_digest_reaches_the_screen_that_uses_it(
        self,
        client: AsyncClient,
        session: AsyncSession,
        save_fixture: SaveFixture,
        user: User,
    ) -> None:
        """The Versions tab decides this itself, off the versions it has.

        No new endpoint and no new field: `VersionRead.counts` already
        carries whatever ingest kept, so the screen can prove the two
        uploads match and never ask for a comparison.
        """
        _, later = await self._pair(session, save_fixture, user)

        response = await client.get(f"/v1/tieout/artifacts/{later.id}/versions")
        assert response.status_code == 200
        digests = {one["counts"].get("sha256") for one in response.json()}
        assert len(digests) == 1
        assert next(iter(digests))
