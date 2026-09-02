"""The regularity check, on constructed blocks.

Registered in `docs/pierce/regularity-check.md`: a break inside a
large block of one shape outranks the same break beside a ragged
run. The registration's proposed new kind (an island enclosed by a
block) was measured and not adopted — the row detector already
accuses every such island — so the tests below hold that no such
kind is raised and that the weight term reaches the row finding.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

from openpyxl import Workbook as Book

from polar.tieout.audit import _shape, audit
from polar.tieout.regularity import ISLAND_MAX, REGION_MIN, bump_for, islands
from polar.tieout.workbook import read_workbook


def _block(rows: int, columns: int, island: tuple[int, int] | None, edge: bool = False):
    """A `rows` × `columns` block of `=B{r}*$A$1` starting at C3, with one
    island of a different shape at `island` (row, column offsets)."""
    book = Book()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Model"
    sheet["A1"] = 1.1
    for r in range(3, 3 + rows):
        sheet.cell(row=r, column=2, value=100 + r)
        sheet.cell(row=r, column=1, value=f"Line {r}")
        for c in range(3, 3 + columns):
            sheet.cell(
                row=r,
                column=c,
                value=f"={sheet.cell(row=r, column=c - 1).coordinate}*$A$1",
            )
    if island is not None:
        r, c = 3 + island[0], 3 + island[1]
        sheet.cell(
            row=r, column=c, value=f"={sheet.cell(row=r, column=c - 1).coordinate}+$A$1"
        )
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        path = handle.name
    book.save(path)
    return path


def test_an_enclosed_island_is_found_with_its_region() -> None:
    path = _block(6, 8, island=(2, 4))
    book = read_workbook(path)
    found = islands(book, _shape)
    assert [one.ref for one in found] == ["Model!G5"]
    one = found[0]
    assert one.region_area >= REGION_MIN
    assert one.enclosed == 4
    assert one.own_shape != one.region_shape
    Path(path).unlink(missing_ok=True)


def test_an_edge_island_is_seen_but_not_enclosed() -> None:
    path = _block(6, 8, island=(0, 4))
    book = read_workbook(path)
    found = islands(book, _shape)
    assert [one.ref for one in found] == ["Model!G3"]
    assert found[0].enclosed < 4
    Path(path).unlink(missing_ok=True)


def test_no_new_kind_is_raised_and_the_row_detector_keeps_the_island() -> None:
    """Measured 2 September (regularity-check.md, results): an island
    enclosed by a one-shape block is always inside a row family, so the
    row detector already accuses it, and the « island in a block » kind
    never fired on 49 real models. The kind is not adopted; the audit
    carries exactly one finding at the island, from the row detector,
    and that finding carries the block term."""
    enclosed = _block(6, 8, island=(2, 4))
    book = read_workbook(enclosed)
    at_island = [f for f in audit(book).findings if f.ref == "Model!G5"]
    assert len(at_island) == 1
    assert at_island[0].rule == "inconsistent-row"
    assert at_island[0].kind != "island in a block"
    assert "breaks a block of" in at_island[0].basis
    assert not any(f.kind == "island in a block" for f in audit(book).findings)
    Path(enclosed).unlink(missing_ok=True)


def test_a_break_inside_a_large_block_outranks_one_in_a_small_one() -> None:
    big = _block(12, 12, island=(5, 5))
    small = _block(3, 4, island=(1, 1))
    heavy = [f for f in audit(read_workbook(big)).findings if f.ref == "Model!H8"]
    light = [f for f in audit(read_workbook(small)).findings if f.ref == "Model!D4"]
    assert heavy
    assert light
    assert heavy[0].weight >= light[0].weight
    assert "block of" in heavy[0].basis
    for path in (big, small):
        Path(path).unlink(missing_ok=True)


def test_the_bump_scales_with_the_region_and_caps() -> None:
    from polar.tieout.regularity import Island
    from polar.tieout.workbook import Cell

    cell = Cell(
        sheet="M",
        ref="M!A1",
        row=1,
        column=1,
        value=None,
        formula="=B1",
        row_label="",
        column_label="",
    )
    small = Island("M!B2", "M", REGION_MIN, cell, 4, "a", "b")
    huge = Island("M!B2", "M", 10_000, cell, 4, "a", "b")
    assert 0 < bump_for(small) < bump_for(huge) <= 0.1
    assert ISLAND_MAX == 2
