"""C3 — the delta report where the right answer is by construction.

V2 of the registration: each review-language class provoked by
exactly one synthetic edit, and only the right class may speak. The
real-pair verification (V1 parity with the study matcher, V3 on the
PR24 pairs) lives under the registration in the lane log.
"""

from decimal import Decimal
from pathlib import Path

import openpyxl

from polar.tieout.audit import Finding
from polar.tieout.watch import delta_of
from polar.tieout.watch.delta import keyed_findings
from polar.tieout.workbook import Cell, Workbook


def cell(
    ref: str,
    row: int,
    column: int,
    *,
    value: str | None = None,
    formula: str | None = None,
    label: str = "",
    period: str = "FY2025",
) -> Cell:
    return Cell(
        sheet="M",
        ref=f"M!{ref}",
        row=row,
        column=column,
        value=None if value is None else Decimal(value),
        formula=formula,
        row_label=label,
        column_label=period,
    )


def book(cells: list[Cell]) -> Workbook:
    built = Workbook(sheets=["M"])
    for one in cells:
        built.cells[one.ref] = one
    return built


BASE = [
    cell("B2", 2, 2, value="100", label="Revenue"),
    cell("B3", 3, 2, value="40", formula="=B2*0.4", label="Cost"),
    cell("B4", 4, 2, value="60", formula="=B2-B3", label="EBITDA"),
    cell("B5", 5, 2, value="0.6", formula="=B4/B2", label="Margin"),
]


def kinds(report) -> dict[str, int]:
    counted: dict[str, int] = {}
    for item in report.items:
        counted[item.kind] = counted.get(item.kind, 0) + 1
    return counted


class TestOneEditOneClass:
    def test_a_typed_over_formula_is_a_class_change_and_nothing_else(self) -> None:
        revised = [
            BASE[0],
            BASE[1],
            cell("B4", 4, 2, value="42", label="EBITDA"),  # typed over
            BASE[3],
        ]
        report = delta_of(book(BASE), book(revised))
        counted = kinds(report)
        assert counted.get("class_change") == 1
        assert "moved_assumption" not in counted
        assert "methodology_change" not in counted
        item = next(i for i in report.items if i.kind == "class_change")
        assert (item.sheet, item.first_row) == ("M", 4)
        assert "typed constant" in item.detail

    def test_a_retyped_input_is_a_moved_assumption_and_nothing_else(self) -> None:
        revised = [
            cell("B2", 2, 2, value="120", label="Revenue"),
            *BASE[1:],
        ]
        report = delta_of(book(BASE), book(revised))
        counted = kinds(report)
        assert counted.get("moved_assumption") == 1
        assert "class_change" not in counted
        assert "methodology_change" not in counted
        item = next(i for i in report.items if i.kind == "moved_assumption")
        assert item.detail == "100 → 120"

    def test_a_rewritten_formula_is_a_methodology_change(self) -> None:
        revised = [
            BASE[0],
            cell("B3", 3, 2, value="40", formula="=B2+7", label="Cost"),
            *BASE[2:],
        ]
        report = delta_of(book(BASE), book(revised))
        assert kinds(report).get("methodology_change") == 1

    def test_an_output_move_speaks_only_past_the_registered_line(self) -> None:
        past = [
            BASE[0],
            BASE[1],
            cell("B4", 4, 2, value="66", formula="=B2-B3", label="EBITDA"),
            BASE[3],
        ]
        under = [
            BASE[0],
            BASE[1],
            cell("B4", 4, 2, value="60.05", formula="=B2-B3", label="EBITDA"),
            BASE[3],
        ]
        loud = delta_of(book(BASE), book(past))
        quiet = delta_of(book(BASE), book(under))
        assert kinds(loud).get("material_output") == 1
        assert "material_output" not in kinds(quiet)

    def test_identical_books_report_nothing(self) -> None:
        report = delta_of(book(BASE), book(BASE))
        assert report.items == []
        assert report.new_defects == 0
        assert report.repaired_defects == 0


class TestFold:
    def test_adjacent_retyped_rows_fold_to_one_block(self) -> None:
        base = BASE + [
            cell("B6", 6, 2, value="1", label="Rate one"),
            cell("B7", 7, 2, value="2", label="Rate two"),
        ]
        revised = BASE + [
            cell("B6", 6, 2, value="3", label="Rate one"),
            cell("B7", 7, 2, value="4", label="Rate two"),
        ]
        report = delta_of(book(base), book(revised))
        moved = [i for i in report.items if i.kind == "moved_assumption"]
        assert len(moved) == 1
        assert (moved[0].first_row, moved[0].last_row) == (6, 7)
        assert moved[0].columns == ("B",)


class TestStudyKey:
    """The defect-matching semantics, exactly the study's."""

    @staticmethod
    def finding(rule: str, name: str, ref: str = "B2") -> Finding:
        return Finding(
            rule=rule, severity="error", ref=ref, sheet="M", name=name, detail=""
        )

    def test_empty_names_bucket_as_unmatched_never_guessed(self) -> None:
        keys, _, unmatched = keyed_findings(
            [self.finding("hardcode", ""), self.finding("hardcode", "  ")]
        )
        assert unmatched == 2
        assert not keys

    def test_multisets_match_by_count(self) -> None:
        old_keys, _, _ = keyed_findings(
            [self.finding("hardcode", "EBITDA"), self.finding("hardcode", "EBITDA")]
        )
        new_keys, _, _ = keyed_findings([self.finding("hardcode", "EBITDA")])
        assert sum((old_keys & new_keys).values()) == 1
        assert sum((old_keys - new_keys).values()) == 1
        assert sum((new_keys - old_keys).values()) == 0


class TestRelabelledLine:
    """Class 7, the written amendment: EPN's « Spare » row became
    « Connections Reform Costs » with every value untouched, and the
    report must say so."""

    def test_a_renamed_row_is_a_relabelled_line_and_nothing_else(self) -> None:
        revised = [
            BASE[0],
            BASE[1],
            BASE[2],
            cell("B5", 5, 2, value="0.6", formula="=B4/B2", label="EBITDA margin"),
        ]
        report = delta_of(book(BASE), book(revised))
        counted = kinds(report)
        assert counted.get("relabelled_line") == 1
        assert "moved_assumption" not in counted
        assert "methodology_change" not in counted
        assert "structure" not in counted
        item = next(i for i in report.items if i.kind == "relabelled_line")
        assert item.first_row == 5
        assert item.detail == "« margin » → « ebitda margin »"


class TestCellsInsideMatchedStructure:
    """The C3 deferral, now due: a cell that appears or goes at a
    *matched* position — the new period's actual, typed into a row
    that already existed. An inserted row is `structure` and must
    never be counted here twice."""

    #: Two periods, so the column carrying the change is matched on
    #: its own label rather than colliding with its neighbour.
    GRID = [
        cell("B2", 2, 2, value="100", label="Revenue"),
        cell("C2", 2, 3, value="110", label="Revenue", period="FY2026"),
        cell("B3", 3, 2, value="40", formula="=B2*0.4", label="Cost"),
        cell("C3", 3, 3, value="44", formula="=C2*0.4", label="Cost", period="FY2026"),
        cell("B4", 4, 2, value="60", formula="=B2-B3", label="EBITDA"),
        cell("C4", 4, 3, value="66", formula="=C2-C3", label="EBITDA", period="FY2026"),
        cell("B5", 5, 2, value="0.6", formula="=B4/B2", label="Margin"),
    ]
    LATER = cell(
        "C5", 5, 3, value="0.6", formula="=C4/C2", label="Margin", period="FY2026"
    )

    def test_a_cell_filled_in_an_existing_row_is_its_own_item(self) -> None:
        report = delta_of(book(self.GRID), book([*self.GRID, self.LATER]))
        counted = kinds(report)
        assert counted.get("filled_cell") == 1
        assert "structure" not in counted
        assert "emptied_cell" not in counted
        item = next(i for i in report.items if i.kind == "filled_cell")
        assert (item.first_row, item.columns) == (5, ("C",))
        assert "=C4/C2" in item.detail

    def test_a_cell_emptied_in_an_existing_row_is_its_own_item(self) -> None:
        report = delta_of(book([*self.GRID, self.LATER]), book(self.GRID))
        counted = kinds(report)
        assert counted.get("emptied_cell") == 1
        assert "filled_cell" not in counted
        assert "structure" not in counted
        assert (
            "=C4/C2" in next(i for i in report.items if i.kind == "emptied_cell").detail
        )

    def test_an_inserted_row_stays_structure_and_is_not_counted_twice(self) -> None:
        """The scope line: cells of a row with no counterpart are the
        row's own story, and this class must stay silent about them."""
        shifted = [
            c
            if c.row < 5
            else cell(
                f"{'B' if c.column == 2 else 'C'}6",
                6,
                c.column,
                value=str(c.value),
                formula=c.formula,
                label=c.row_label,
                period=c.column_label,
            )
            for c in self.GRID
        ]
        inserted = [
            *shifted,
            cell("B5", 5, 2, value="5", label="One-offs"),
            cell("C5", 5, 3, value="5", label="One-offs", period="FY2026"),
        ]
        report = delta_of(book(self.GRID), book(inserted))
        counted = kinds(report)
        assert counted.get("structure", 0) >= 1
        assert "filled_cell" not in counted
        assert "emptied_cell" not in counted


class TestTheIdenticalPairFastPath:
    """Routed from Atelier: when two uploads share a SHA-256 there is
    nothing to compare, and the true report is knowable exactly."""

    def finding(self, rule: str, sheet: str, name: str) -> Finding:
        return Finding(
            rule=rule,
            severity="error",
            sheet=sheet,
            ref=f"{sheet}!B2",
            name=name,
            detail="",
        )

    def test_persistent_is_what_the_one_workbook_keys_to(self) -> None:
        from polar.tieout.watch import unchanged_report

        findings = [
            self.finding("hardcode", "Model", "Revenue"),
            self.finding("hardcode", "Model", "Costs"),
            self.finding("long-formula", "Model", "Revenue"),
        ]
        report = unchanged_report(findings, "book.xlsx")
        assert (report.new_defects, report.repaired_defects) == (0, 0)
        assert report.persistent_defects == 3
        assert report.items == []
        assert report.old == report.new == "book.xlsx"

    def test_it_agrees_with_running_a_delta_of_a_book_against_itself(self) -> None:
        """The premise check: the fast path must give what the slow
        path gives, or it is not a fast path."""
        from polar.tieout.watch import keyed_findings, unchanged_report

        findings = [
            self.finding("hardcode", "Model", "Revenue"),
            self.finding("stale-value", "Other", ""),
        ]
        keys, _behind, unmatched = keyed_findings(findings)
        report = unchanged_report(findings)
        assert report.persistent_defects == sum(keys.values())
        assert report.unmatched_old == report.unmatched_new == unmatched


class TestAuditsCanBeHandedIn:
    """Measured on the GD3 pair: the two audits are 448.7 s of a
    1,119.6 s delta_of — 40% of the report. The product audits every
    upload as it arrives, so recomputing them here is the most
    expensive quarter of the work done twice."""

    def test_passing_findings_gives_the_same_report_as_running_them(
        self, tmp_path: Path
    ) -> None:
        from polar.tieout.audit import audit
        from polar.tieout.structure import period_axes
        from polar.tieout.watch import delta_of
        from polar.tieout.workbook import read_workbook

        old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Model"
        sheet["A2"], sheet["B2"] = "Revenue", 100
        sheet["A3"], sheet["B3"] = "Costs", 40
        sheet["A4"], sheet["B4"] = "Margin", "=B2-B3"
        book.save(old)
        sheet["B3"] = 45
        book.save(new)

        old_book, new_book = read_workbook(str(old)), read_workbook(str(new))
        computed = delta_of(old_book, new_book)
        handed = delta_of(
            old_book,
            new_book,
            old_findings=audit(old_book, axes=period_axes(old_book)).findings,
            new_findings=audit(new_book, axes=period_axes(new_book)).findings,
        )
        assert handed.new_defects == computed.new_defects
        assert handed.repaired_defects == computed.repaired_defects
        assert handed.persistent_defects == computed.persistent_defects
        assert [(i.kind, i.sheet, i.detail) for i in handed.items] == [
            (i.kind, i.sheet, i.detail) for i in computed.items
        ]

    def test_an_empty_list_is_not_the_same_as_no_argument(self) -> None:
        """`None` means « run the audit »; `[]` means « this book has
        no findings ». Conflating them would silently drop defects."""
        from polar.tieout.watch.delta import delta_of
        from polar.tieout.workbook import Workbook

        empty = Workbook(sheets=["M"])
        report = delta_of(empty, empty, old_findings=[], new_findings=[])
        assert report.persistent_defects == 0


class TestRetainedParseCaches:
    """The delta pipeline's fix for paying the cold parse pass three
    times: inside a retained scope `audit()` leaves the content-keyed
    caches warm for the next phase; outside one, nothing changes."""

    def formulas(self) -> Workbook:
        return book(
            [
                cell("B2", 2, 2, value="100", label="Revenue"),
                cell("B3", 3, 2, value="40", formula="=B2*0.4", label="Cost"),
                cell("B4", 4, 2, value="60", formula="=B2-B3", label="EBITDA"),
            ]
        )

    def test_an_unretained_audit_clears_as_it_always_has(self) -> None:
        from polar.tieout.audit import _shape_of, audit
        from polar.tieout.workbook import tokens_of

        audit(self.formulas())
        assert tokens_of.cache_info().currsize == 0
        assert _shape_of.cache_info().currsize == 0

    def test_a_retained_scope_keeps_the_caches_then_clears_on_exit(self) -> None:
        from polar.tieout.audit import _shape_of, audit, retained_parse_caches
        from polar.tieout.workbook import tokens_of

        with retained_parse_caches():
            audit(self.formulas())
            assert tokens_of.cache_info().currsize > 0
            assert _shape_of.cache_info().currsize > 0
        assert tokens_of.cache_info().currsize == 0
        assert _shape_of.cache_info().currsize == 0

    def test_nested_scopes_clear_once_when_the_outermost_closes(self) -> None:
        from polar.tieout.audit import audit, retained_parse_caches
        from polar.tieout.workbook import tokens_of

        with retained_parse_caches():
            with retained_parse_caches():
                audit(self.formulas())
            #: The inner exit must not clear — delta_of retains and
            #: delta_report retains around it; the pair must nest.
            assert tokens_of.cache_info().currsize > 0
        assert tokens_of.cache_info().currsize == 0

    def test_the_scope_clears_even_when_the_work_inside_raises(self) -> None:
        import gc

        from polar.tieout.audit import retained_parse_caches
        from polar.tieout.workbook import tokens_of

        try:
            with retained_parse_caches():
                tokens_of("=1+1")
                raise RuntimeError("mid-pipeline failure")
        except RuntimeError:
            pass
        assert tokens_of.cache_info().currsize == 0
        assert gc.isenabled()

    def test_the_scope_holds_the_collector_and_restores_it(self) -> None:
        """The measured 100 s of a cold comparison was the cyclic
        collector re-scanning the growing caches; the scope holds it
        and puts back exactly the state it found — nested scopes
        included."""
        import gc

        from polar.tieout.audit import retained_parse_caches

        assert gc.isenabled()
        with retained_parse_caches():
            assert not gc.isenabled()
            with retained_parse_caches():
                assert not gc.isenabled()
            assert not gc.isenabled()
        assert gc.isenabled()

    def test_a_collector_already_off_stays_off(self) -> None:
        import gc

        from polar.tieout.audit import retained_parse_caches

        gc.disable()
        try:
            with retained_parse_caches():
                assert not gc.isenabled()
            assert not gc.isenabled()
        finally:
            gc.enable()


class TestTierOneOnRewrites:
    """A methodology change carries tier 1's answer: proved the same
    function, or the assignment where the two differ, or an honest
    refusal by construct."""

    def test_an_equivalent_rewrite_is_proved(self) -> None:
        revised = [
            BASE[0],
            cell("B3", 3, 2, value="40", formula="=B2*2/5", label="Cost"),
            *BASE[2:],
        ]
        report = delta_of(book(BASE), book(revised))
        item = next(i for i in report.items if i.kind == "methodology_change")
        assert "proved the same function" in item.detail

    def test_a_different_rewrite_is_refuted_with_the_assignment(self) -> None:
        revised = [
            BASE[0],
            cell("B3", 3, 2, value="40", formula="=B2+7", label="Cost"),
            *BASE[2:],
        ]
        report = delta_of(book(BASE), book(revised))
        item = next(i for i in report.items if i.kind == "methodology_change")
        assert "differs at" in item.detail
        assert "M!B2=" in item.detail

    def test_a_rewrite_outside_the_fragment_is_refused_by_name(self) -> None:
        revised = [
            BASE[0],
            cell("B3", 3, 2, value="40", formula="=INDEX(B2:B2,1)*0.4", label="Cost"),
            *BASE[2:],
        ]
        report = delta_of(book(BASE), book(revised))
        item = next(i for i in report.items if i.kind == "methodology_change")
        assert "not provable (lookup_or_selection: INDEX)" in item.detail
