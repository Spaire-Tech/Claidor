"""The metadata checker.

**These are unit tests, and they are not the evidence.** Every file below
is built here, which makes it worth exactly as much as any fixture whose
author also wrote the checker — it proves the rule fires and the parsing
path holds, and nothing about what real files contain. The evidence is
`scripts/metadata_survey.py`, which runs the same code over public Office
files from people who have never heard of us, and whose numbers are in
`docs/pierce/accuracy-backlog.md`.

Where a fixture can be produced by a library that also writes files Office
opens — `python-pptx`, `openpyxl` — it is, rather than by hand-writing the
XML the checker is about to read. A test that writes the same XML it
asserts on tests nothing.
"""

import io
import zipfile

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from polar.tieout.metadata import (
    NotAnOfficeFile,
    read_metadata_bytes,
)

CONTENT_TYPES = (
    '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
    'package/2006/content-types">'
    '<Default Extension="xml" ContentType="application/xml"/></Types>'
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _package(parts: dict[str, str | bytes]) -> bytes:
    """A zip with the parts given, plus the content types every one has."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", CONTENT_TYPES)
        for name, body in parts.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def _word(body: str, **extra: str) -> bytes:
    return _package(
        {
            "word/document.xml": (
                f'<?xml version="1.0"?><w:document xmlns:w="{W}">'
                f"<w:body>{body}</w:body></w:document>"
            ),
            **extra,
        }
    )


def _rules(report: object) -> set[str]:
    return {one.rule for one in report.findings}  # type: ignore[attr-defined]


def _of(report: object, rule: str) -> list:
    return [one for one in report.findings if one.rule == rule]  # type: ignore[attr-defined]


class TestRefusal:
    def test_not_a_zip(self) -> None:
        with pytest.raises(NotAnOfficeFile, match="not an Office file"):
            read_metadata_bytes(b"just some bytes, not a file")

    def test_legacy_ole_says_which_problem_it_is(self) -> None:
        """A `.doc` and an encrypted `.docx` look identical from outside.

        Both are OLE compound files, so the message names both rather than
        guessing at one. Guessing here means telling somebody their file is
        corrupt when it is password protected.
        """
        with pytest.raises(NotAnOfficeFile, match="legacy Office file"):
            read_metadata_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    def test_a_plain_zip_is_not_a_document(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("notes.txt", "hello")
        with pytest.raises(NotAnOfficeFile, match="no content types"):
            read_metadata_bytes(buffer.getvalue())

    def test_too_many_parts_is_refused_before_reading(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", CONTENT_TYPES)
            for index in range(20_001):
                archive.writestr(f"part{index}.xml", "<a/>")
        with pytest.raises(NotAnOfficeFile, match="more than an Office document"):
            read_metadata_bytes(buffer.getvalue())


class TestProperties:
    def test_who_made_it_and_who_saved_it(self) -> None:
        report = read_metadata_bytes(
            _package(
                {
                    "docProps/core.xml": (
                        '<?xml version="1.0"?><cp:coreProperties '
                        'xmlns:cp="http://schemas.openxmlformats.org/package/'
                        '2006/metadata/core-properties" '
                        'xmlns:dc="http://purl.org/dc/elements/1.1/">'
                        "<dc:creator>A. Analyst</dc:creator>"
                        "<cp:lastModifiedBy>MD, Coverage</cp:lastModifiedBy>"
                        "<cp:revision>47</cp:revision>"
                        "</cp:coreProperties>"
                    )
                }
            )
        )
        said = {one.detail: one.evidence for one in _of(report, "document-properties")}
        assert said["the author recorded in the file"] == "A. Analyst"
        assert said["the last person to save it"] == "MD, Coverage"
        assert said["how many times it has been saved"] == "47"
        # Who touched a file is not content leaving the building.
        assert all(one.severity == "trace" for one in report.findings)

    def test_an_empty_property_is_not_a_finding(self) -> None:
        report = read_metadata_bytes(
            _package(
                {
                    "docProps/core.xml": (
                        '<?xml version="1.0"?><cp:coreProperties '
                        'xmlns:cp="http://schemas.openxmlformats.org/package/'
                        '2006/metadata/core-properties" '
                        'xmlns:dc="http://purl.org/dc/elements/1.1/">'
                        "<dc:creator></dc:creator><dc:title>  </dc:title>"
                        "</cp:coreProperties>"
                    )
                }
            )
        )
        assert report.findings == []


class TestWord:
    def test_deleted_text_is_still_readable(self) -> None:
        """The finding that matters most in Word.

        `w:del` does not remove the words; it marks them. They are in the
        file, and « show markup » off is a display setting, not a redaction.
        """
        report = read_metadata_bytes(
            _word(
                '<w:p><w:del w:author="Jo Reviewer" w:date="2026-08-01">'
                "<w:r><w:delText>the vendor will accept 6.5x</w:delText>"
                "</w:r></w:del></w:p>"
            )
        )
        found = _of(report, "tracked-changes")
        assert any("Jo Reviewer" in one.detail for one in found)
        assert any("the vendor will accept 6.5x" in one.evidence for one in found)
        assert all(one.severity == "leak" for one in found)

    def test_a_formatting_change_carries_a_name_with_no_text(self) -> None:
        report = read_metadata_bytes(
            _word(
                '<w:p><w:pPr><w:pPrChange w:author="Someone Else" '
                'w:date="2026-08-01"><w:pPr/></w:pPrChange></w:pPr></w:p>'
            )
        )
        found = _of(report, "tracked-changes")
        assert len(found) == 1
        assert "Someone Else" in found[0].detail

    def test_a_clean_document_reports_no_revisions(self) -> None:
        report = read_metadata_bytes(_word("<w:p><w:r><w:t>Clean.</w:t></w:r></w:p>"))
        assert "tracked-changes" not in _rules(report)

    def test_hidden_text(self) -> None:
        report = read_metadata_bytes(
            _word(
                "<w:p><w:r><w:rPr><w:vanish/></w:rPr>"
                "<w:t>internal only — do not circulate</w:t></w:r></w:p>"
            )
        )
        assert _of(report, "hidden-text")[0].evidence == (
            "internal only — do not circulate"
        )

    def test_vanish_switched_off_is_not_hidden(self) -> None:
        """`<w:vanish w:val="0"/>` is how Word turns hidden text back on.

        Read as presence-means-hidden — which is how the element usually
        works — every un-hidden run in the document becomes a finding.
        """
        report = read_metadata_bytes(
            _word(
                '<w:p><w:r><w:rPr><w:vanish w:val="0"/></w:rPr>'
                "<w:t>ordinary text</w:t></w:r></w:p>"
            )
        )
        assert "hidden-text" not in _rules(report)

    def test_comments_carry_their_author(self) -> None:
        report = read_metadata_bytes(
            _word(
                "<w:p/>",
                **{
                    "word/comments.xml": (
                        f'<?xml version="1.0"?><w:comments xmlns:w="{W}">'
                        f'<w:comment w:author="Jo Reviewer" w:date="2026-08-02">'
                        f"<w:p><w:r><w:t>is this the right multiple?</w:t>"
                        f"</w:r></w:p></w:comment></w:comments>"
                    )
                },
            )
        )
        found = _of(report, "comment")
        assert "Jo Reviewer" in found[0].detail
        assert found[0].evidence == "is this the right multiple?"


class TestExcel:
    def _book(self, **kwargs: object) -> Workbook:
        book = Workbook()
        book.active.title = "Summary"  # type: ignore[union-attr]
        return book

    def _bytes(self, book: Workbook) -> bytes:
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    def test_very_hidden_and_hidden_are_different_findings(self) -> None:
        """And graded differently, because they are different problems.

        Hidden is one right-click from visible and everyone knows it.
        Very hidden does not appear in Excel's unhide list at all, which is
        why people believe those sheets were deleted.
        """
        book = self._book()
        book.create_sheet("Workings").sheet_state = "hidden"
        book.create_sheet("Old bid model").sheet_state = "veryHidden"
        report = read_metadata_bytes(self._bytes(book))

        very = _of(report, "very-hidden-sheet")
        assert len(very) == 1
        assert very[0].where == "sheet « Old bid model »"
        assert very[0].severity == "leak"

        hidden = _of(report, "hidden-sheet")
        assert len(hidden) == 1
        assert hidden[0].where == "sheet « Workings »"
        assert hidden[0].severity == "trace"

    def test_a_visible_workbook_reports_no_sheets(self) -> None:
        book = self._book()
        book.create_sheet("Detail")
        report = read_metadata_bytes(self._bytes(book))
        assert not _of(report, "hidden-sheet")
        assert not _of(report, "very-hidden-sheet")

    def test_an_external_link_carries_its_path_and_its_numbers(self) -> None:
        """Two findings from one part, and the second is the surprise.

        `externalLink1.xml` caches the values the link last returned, so
        another deal's numbers travel inside this workbook even when the
        recipient cannot reach the file they came from.
        """
        book = self._book()
        payload = self._bytes(book)
        buffer = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(payload)) as source:
            with zipfile.ZipFile(buffer, "w") as target:
                for item in source.infolist():
                    target.writestr(item, source.read(item.filename))
                target.writestr(
                    "xl/externalLinks/externalLink1.xml",
                    '<?xml version="1.0"?><externalLink xmlns="http://schemas.'
                    'openxmlformats.org/spreadsheetml/2006/main"><externalBook>'
                    "<sheetDataSet><sheetData><row><cell><v>228.9</v></cell>"
                    "<cell><v>41.2</v></cell></row></sheetData></sheetDataSet>"
                    "</externalBook></externalLink>",
                )
                target.writestr(
                    "xl/externalLinks/_rels/externalLink1.xml.rels",
                    '<?xml version="1.0"?><Relationships xmlns="http://schemas.'
                    'openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" Type="http://schemas.openxmlformats'
                    '.org/officeDocument/2006/relationships/externalLinkPath" '
                    'Target="file:///\\\\fs01\\Projects\\Falcon\\Model v12.xlsx" '
                    'TargetMode="External"/></Relationships>',
                )
        report = read_metadata_bytes(buffer.getvalue())

        path = _of(report, "external-workbook")
        assert len(path) == 1
        assert "Falcon" in path[0].evidence

        cached = _of(report, "external-cached-values")
        assert len(cached) == 1
        assert "2 values" in cached[0].detail
        assert "228.9" in cached[0].evidence

        # The general local-path rule sees the same relationship. Two
        # findings about one fact is how a report stops being read, and
        # `external-workbook` is the one that also knows about the cache,
        # so it is the one that keeps the path.
        assert not _of(report, "local-path")


class TestPowerPoint:
    def _deck(self) -> Presentation:
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6])
        return deck

    def _bytes(self, deck: Presentation) -> bytes:
        buffer = io.BytesIO()
        deck.save(buffer)
        return buffer.getvalue()

    def test_speaker_notes(self) -> None:
        deck = self._deck()
        notes = deck.slides[0].notes_slide.notes_text_frame
        notes.text = "Do not show them the downside case."
        report = read_metadata_bytes(self._bytes(deck))
        found = _of(report, "speaker-notes")
        assert len(found) == 1
        assert found[0].where == "slide 1"
        assert found[0].evidence == "Do not show them the downside case."

    def test_an_untyped_notes_page_is_not_a_finding(self) -> None:
        """Every slide has a notes page and every notes page has a slide
        number on it. Counting those reports notes on every deck ever made.
        """
        deck = self._deck()
        deck.slides[0].notes_slide  # brings the part into being, empty
        report = read_metadata_bytes(self._bytes(deck))
        assert not _of(report, "speaker-notes")

    def test_a_shape_dragged_off_the_canvas(self) -> None:
        deck = self._deck()
        slide = deck.slides[0]
        parked = slide.shapes.add_textbox(-Inches(6), Inches(1), Inches(4), Inches(1))
        parked.text_frame.text = "last week: 4.1x"
        report = read_metadata_bytes(self._bytes(deck))
        found = _of(report, "off-canvas-shape")
        assert len(found) == 1
        assert found[0].where == "slide 1"
        assert "to the left of" in found[0].detail
        assert found[0].evidence == "last week: 4.1x"

    def test_a_shape_bleeding_over_the_edge_is_a_design_decision(self) -> None:
        """A logo half off the edge is on every professional deck.

        Only a shape with no overlap at all is reported, because the
        alternative is a finding on every slide of every deck, which is the
        same as no findings.
        """
        deck = self._deck()
        deck.slides[0].shapes.add_textbox(-Inches(1), Inches(1), Inches(4), Inches(1))
        report = read_metadata_bytes(self._bytes(deck))
        assert not _of(report, "off-canvas-shape")

    def test_a_hidden_slide(self) -> None:
        deck = self._deck()
        deck.slides[0]._element.set("show", "0")
        report = read_metadata_bytes(self._bytes(deck))
        assert _of(report, "hidden-slide")[0].where == "slide 1"


class TestCapping:
    def test_a_long_list_says_how_long_it_was(self) -> None:
        """The rule from the roadmap: no silent caps.

        A deck with two hundred notes pages is one fact, not two hundred
        findings — but a report that lists forty and stops looks exactly
        like a file with forty.
        """
        deck = Presentation()
        for index in range(45):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.notes_slide.notes_text_frame.text = f"note {index}"
        buffer = io.BytesIO()
        deck.save(buffer)
        report = read_metadata_bytes(buffer.getvalue())

        found = _of(report, "speaker-notes")
        assert len(found) == 41
        assert found[-1].where == "the rest of the file"
        assert "5 more" in found[-1].detail
        assert "45 in total" in found[-1].detail


class TestGrading:
    def test_leaks_and_traces_are_never_one_number(self) -> None:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = "internal"
        buffer = io.BytesIO()
        deck.save(buffer)
        report = read_metadata_bytes(buffer.getvalue())

        assert report.kind == "powerpoint"
        assert {one.severity for one in report.leaks} == {"leak"}
        assert {one.severity for one in report.traces} == {"trace"}
        assert len(report.leaks) + len(report.traces) == len(report.findings)
        assert report.parts > 0
