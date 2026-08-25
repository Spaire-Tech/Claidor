"""C2 — the DP alignment on lines built to test it, then end to end.

The planted-edit harness measures recovery on real corpus sheets
under its own registration; these tests pin the algorithm's
semantics where the right answer is known by construction.
"""

from pathlib import Path

import openpyxl
import pytest

from polar.tieout.watch import (
    LITERAL,
    Line,
    align_lines,
    align_sheet,
    sheet_grids,
    structural_changes,
)
from polar.tieout.watch.align import similarity
from polar.tieout.workbook import read_workbook


def line(index: int, label: str, *signatures: str) -> Line:
    return Line(index, label, tuple(enumerate(signatures, start=1)))


ROWS = [
    line(2, "revenue", LITERAL, LITERAL, LITERAL),
    line(3, "growth", "D3/C3-#", "D3/C3-#"),
    line(4, "ebitda", "C2*#", "C2*#", "C2*#"),
    line(5, "margin", "C4/C2", "C4/C2", "C4/C2"),
]


def shifted(rows: list[Line], at: int, by: int = 1) -> list[Line]:
    return [
        Line(row.index + by, row.label, row.cells) if row.index >= at else row
        for row in rows
    ]


class TestAlignment:
    def test_one_inserted_row_is_one_structural_change(self) -> None:
        new_rows = shifted(ROWS, at=4)
        new_rows.insert(2, line(4, "one-offs", LITERAL, LITERAL))
        result = align_lines(ROWS, new_rows)
        assert result.mapping == {2: 2, 3: 3, 4: 5, 5: 6}
        assert result.inserted == (4,)
        assert result.deleted == ()

    def test_a_deleted_row_is_reported_and_the_rest_hold(self) -> None:
        result = align_lines(ROWS, [ROWS[0], *shifted(ROWS[2:], at=0, by=-1)])
        assert result.deleted == (3,)
        assert result.mapping == {2: 2, 4: 3, 5: 4}

    def test_a_changed_row_stays_matched_through_its_label(self) -> None:
        new_rows = list(ROWS)
        new_rows[2] = line(4, "ebitda", "C2*#", "C2*#-D4", "C2*#")
        result = align_lines(ROWS, new_rows)
        match = {m.old: m for m in result.matched}[4]
        assert match.new == 4
        assert 0.5 <= match.similarity < 1.0

    def test_relabelled_rows_match_only_on_identical_shapes(self) -> None:
        renamed = [
            Line(row.index, row.label + " (restated)", row.cells) for row in ROWS
        ]
        result = align_lines(ROWS, renamed)
        assert result.mapping == {2: 2, 3: 3, 4: 4, 5: 5}
        assert all(m.similarity == 0.5 for m in result.matched)
        different = line(3, "renamed", "X9+#")
        assert similarity(ROWS[1], different) == 0.0

    def test_values_never_enter_a_signature(self) -> None:
        retyped = [
            ROWS[0],
            line(3, "growth", "D3/C3-#", "D3/C3-#"),
            ROWS[2],
            ROWS[3],
        ]
        result = align_lines(ROWS, retyped)
        assert result.inserted == () == result.deleted
        assert all(m.similarity == 1.0 for m in result.matched)


class TestEndToEnd:
    """Through the engine's frozen reader surface and shape hash."""

    @staticmethod
    def _write(path: Path, inserted: bool) -> str:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Model"
        sheet["B1"], sheet["C1"], sheet["D1"] = "FY2024", "FY2025", "FY2026"
        rows = [
            ("Revenue", [100, 110, 121]),
            ("Cost", [40, 44, 48]),
            ("EBITDA", ["=B2-B3", "=C2-C3", "=D2-D3"]),
            ("Margin", ["=B4/B2", "=C4/C2", "=D4/D2"]),
        ]
        if inserted:
            # Excel-correct insertion above the EBITDA row: the two
            # rows below shift down and their formulas' references to
            # rows >= 4 shift with them.
            rows = [
                rows[0],
                rows[1],
                ("One-offs", [5, 5, 5]),
                ("EBITDA", ["=B2-B3", "=C2-C3", "=D2-D3"]),
                ("Margin", ["=B5/B2", "=C5/C2", "=D5/D2"]),
            ]
        for offset, (label, cells) in enumerate(rows):
            sheet.cell(row=2 + offset, column=1, value=label)
            for column, value in enumerate(cells, start=2):
                sheet.cell(row=2 + offset, column=column, value=value)
        book.save(path)
        return str(path)

    @pytest.fixture
    def books(self, tmp_path: Path) -> tuple[dict, dict]:
        old = read_workbook(self._write(tmp_path / "old.xlsx", inserted=False))
        new = read_workbook(self._write(tmp_path / "new.xlsx", inserted=True))
        return sheet_grids(old), sheet_grids(new)

    def test_one_planted_insert_reads_as_one_structural_change(
        self, books: tuple[dict, dict]
    ) -> None:
        old_grids, new_grids = books
        alignment = align_sheet(old_grids["Model"], new_grids["Model"])
        structure = [
            change
            for change in structural_changes(alignment)
            if not change["kind"].startswith("changed")
        ]
        assert structure == [{"kind": "inserted_rows", "first": 4, "last": 4}]
        assert alignment.rows.mapping == {2: 2, 3: 3, 4: 5, 5: 6}
        assert alignment.columns.inserted == ()
        assert alignment.columns.deleted == ()


class TestCrossSheetAbsolute:
    """Round 3's amendment: a piece on another sheet does not move
    with the cell, so its signature must not either."""

    def test_a_pull_through_survives_its_cells_shift(self) -> None:
        from polar.tieout.watch.signature import cell_signature

        old = cell_signature(
            "=SelectedInputs!E64", "SelectedInputs!R[-307]C[+0]", "InputSummary", 371, 5
        )
        new = cell_signature(
            "=SelectedInputs!E64", "SelectedInputs!R[-308]C[+0]", "InputSummary", 372, 5
        )
        assert old == new == "SelectedInputs!R64CE"

    def test_a_same_sheet_piece_stays_relative(self) -> None:
        from polar.tieout.watch.signature import cell_signature

        assert (
            cell_signature("=AM424", "R[+53]C[-1]", "InputSummary", 371, 40)
            == "R[+53]C[-1]"
        )
        assert (
            cell_signature(
                "=InputSummary!A1",
                "InputSummary!R[-370]C[-39]",
                "InputSummary",
                371,
                40,
            )
            == "InputSummary!R[-370]C[-39]"
        )

    def test_a_cross_sheet_range_tail_goes_absolute_too(self) -> None:
        from polar.tieout.watch.signature import cell_signature

        got = cell_signature(
            "=SUM('Annual Inflation'!A1:A5)",
            "SUM('Annual Inflation'!R[-1]C[+0]:R[+3]C[+0])",
            "Model",
            2,
            1,
        )
        assert got == "SUM('Annual Inflation'!R1CA:R5CA)"

    def test_anchored_pieces_pass_untouched(self) -> None:
        from polar.tieout.watch.signature import cell_signature

        assert (
            cell_signature("=Other!$B$19", "Other!R19CB", "Model", 7, 3)
            == "Other!R19CB"
        )
