"""The flagship check, as it reaches a report.

`test_period_flow_stock.py` tests the judgement — what a row's kind is
and where it breaks. This tests the wiring: that a real workbook laid
out the way a model is laid out produces a finding pointing at the
**cell**, that a file with nothing to read against says so instead of
printing silence, and that a clean file stays clean.

The round is `docs/pierce/e3c-flow-stock.md`.
"""

import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

from openpyxl import Workbook as Book
from openpyxl.utils import get_column_letter

from polar.tieout.audit import Audit, audit
from polar.tieout.workbook import read_workbook

#: 1 January 2030, as Excel counts.
START = 47484.0
DAYS = 365.25


def _lay_out(
    sheet: Any, label: str, values: list[float], *, months: int, first: int = 2
) -> None:
    """A model's own layout: a header row, a date axis, a labelled row.

    `months` is the step between columns — 6 for half-years, 12 for
    years — so the granularity is read off the dates exactly as the
    check reads it off a real file, by arithmetic and never by a word.

    **The text header row above the dates is not decoration.** The
    reader treats a sheet's header row as words rather than content and
    keeps its cells out of the audit, so a file whose *only* row of
    dates is its header row hands this check no axis at all. Real
    models put a period number or a label line above the dates, which
    is why the corpus rounds saw axes; a fixture without one measures
    the reader's header rule instead of the check.
    """
    sheet.cell(row=1, column=1).value = "Period"
    sheet.cell(row=2, column=1).value = f"{months} month ending"
    sheet.cell(row=3, column=1).value = label
    for index, value in enumerate(values):
        column = first + index
        sheet.cell(row=1, column=column).value = f"P{index + 1}"
        date = sheet.cell(row=2, column=column)
        date.value = START + index * months * DAYS / 12
        date.number_format = "yyyy-mm-dd"
        sheet.cell(row=3, column=column).value = value


def _audit_of(build: Callable[[Book], None]) -> Audit:
    book = Book()
    build(book)
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / "built.xlsx"
        book.save(path)
        return audit(read_workbook(str(path)))


def _findings(result: Audit) -> list[Any]:
    return [f for f in result.findings if f.rule == "broken-aggregation"]


def _abstention(result: Audit) -> str | None:
    for one in result.abstentions:
        if one.rule == "broken-aggregation":
            return one.why
    return None


#: Twenty-four half-years and the twelve years they roll into.
FINE = [float(n) for n in range(1, 25)]
COARSE = [FINE[i * 2] + FINE[i * 2 + 1] for i in range(12)]


def _two_blocks(book: Book, coarse: list[float], label: str = "Revenue") -> None:
    half = book.active
    half.title = "Half"
    _lay_out(half, label, FINE, months=6)
    year = book.create_sheet("Year")
    _lay_out(year, label, coarse, months=12)


class TestTheFindingReachesTheReport:
    def test_a_broken_year_names_the_cell_it_broke_in(self) -> None:
        broken = list(COARSE)
        #: The eighth year takes the second half alone.
        broken[7] = FINE[15]

        found = _findings(_audit_of(lambda b: _two_blocks(b, broken)))

        assert len(found) == 1
        #: The eighth of twelve columns starting at B is I — the cell,
        #: not the row, so « Open the cell » has somewhere to go.
        assert found[0].ref == f"Year!{get_column_letter(2 + 7)}3"
        assert found[0].severity == "error"
        assert found[0].name == "revenue"
        assert "adds up one half-year" in found[0].detail
        assert found[0].source == "the row's own behaviour across its time axis"

    def test_the_detail_says_what_the_row_does_everywhere_else(self) -> None:
        broken = list(COARSE)
        broken[7] = FINE[15]

        found = _findings(_audit_of(lambda b: _two_blocks(b, broken)))

        #: « adds up one half-year where the row adds all 2 » — the
        #: pattern the break departs from, stated in the same breath
        #: as the break, in the period words the model itself uses.
        assert "the row adds all 2" in found[0].detail

    def test_a_clean_pair_of_blocks_reports_nothing(self) -> None:
        result = _audit_of(lambda b: _two_blocks(b, list(COARSE)))

        assert _findings(result) == []
        #: And it looked: a tally, not an abstention.
        assert _abstention(result) is None
        assert result.tallies["broken-aggregation"]["total"] == 1
        assert result.tallies["broken-aggregation"]["raised"] == 0

    def test_the_tally_counts_the_finding_it_raised(self) -> None:
        broken = list(COARSE)
        broken[7] = FINE[15]

        result = _audit_of(lambda b: _two_blocks(b, broken))

        assert result.tallies["broken-aggregation"] == {"total": 1, "raised": 1}


class TestSilenceIsSaidOutLoud:
    """Six of twenty-two real close models give this check nothing to
    read. A report that printed nothing there would be claiming a clean
    bill it never earned, so the abstention carries the reason."""

    def test_one_dated_sheet_abstains_with_its_reason(self) -> None:
        def build(book: Book) -> None:
            _lay_out(book.active, "Revenue", FINE, months=6)

        result = _audit_of(build)

        assert _findings(result) == []
        assert _abstention(result) == (
            "the file lays out no two sheets of dated columns to "
            "read one row against the other"
        )

    def test_two_blocks_sharing_no_row_abstains_too(self) -> None:
        def build(book: Book) -> None:
            _two_blocks(book, list(COARSE))
            #: Same blocks, different labels — nothing to compare.
            book["Year"].cell(row=3, column=1).value = "Something else"

        result = _audit_of(build)

        assert _findings(result) == []
        assert _abstention(result) is not None
        assert "broken-aggregation" not in result.tallies

    def test_a_frozen_row_is_not_a_pattern_and_says_so(self) -> None:
        def build(book: Book) -> None:
            half = book.active
            half.title = "Half"
            _lay_out(half, "Revenue", [0.0] * 24, months=6)
            _lay_out(book.create_sheet("Year"), "Revenue", [0.0] * 12, months=12)

        result = _audit_of(build)

        assert _findings(result) == []
        assert _abstention(result) == (
            "no row on the file's dated sheets holds still long "
            "enough across its periods to establish a pattern"
        )


class TestTheRuleIsAFirstClassRule:
    def test_it_can_be_switched_off_and_has_a_headline(self) -> None:
        #: The map comment on `files.ts` records five rules that reach a
        #: report the engine cannot name and a firm cannot switch off.
        #: This one is not going to be the sixth.
        from polar.tieout.audit import HEADLINES, RULE_NAMES

        assert "broken-aggregation" in RULE_NAMES
        assert HEADLINES["broken-aggregation"] == "Broken aggregation"
