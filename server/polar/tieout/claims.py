"""The prover behind the investigator loop
(`docs/pierce/investigator-loop.md`).

The investigator — a language model reading a sheet — proposes
claims in a fixed shape: a cell, a kind, and the facts the claim
rests on. This module checks every cited fact against the cells,
deterministically, and returns a verdict:

- **confirmed** — every fact holds in the workbook;
- **refuted** — a cited fact is false (the range does include the
  cell, the referenced cell is not empty, the neighbours do not
  share a shape, the alternative formula gives a different number);
- **unverifiable** — the kind is `other`, or the alternative formula
  uses something the small evaluator here does not know.

A confirmed verdict proves the facts, never the judgment. « Reads
D5, and D10 is populated » is a fact the engine can stand behind;
« and it should read D10 » is the reviewer's call, and the verdict's
note says so. Nothing here guesses: a claim that cannot be checked
is returned as such, in words.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl.utils import column_index_from_string, get_column_letter

from .workbook import ERROR_VALUES, Cell, Workbook, tokens_of

KINDS = frozenset(
    {
        "range-omits",
        "wrong-reference",
        "empty-reference",
        "differs-from-neighbours",
        "malformed",
        "error-value",
        "arithmetic",
        "value-contradiction",
        "other",
    }
)

_A1 = re.compile(
    r"^(?:(?P<sheet>'[^']+'|[^'!]+)!)?\$?(?P<col>[A-Za-z]{1,3})\$?(?P<row>\d+)$"
)


@dataclass(frozen=True)
class Claim:
    sheet: str
    cell: str
    kind: str
    reason: str = ""
    confidence: str = ""
    reads: str = ""
    should_read: str = ""
    omits: tuple[str, ...] = ()
    neighbours: tuple[str, ...] = ()
    alternative_formula: str = ""
    expected_value: str = ""
    other_cell: str = ""

    @classmethod
    def from_json(cls, sheet: str, data: dict[str, Any]) -> Claim:
        def text(key: str) -> str:
            value = data.get(key)
            return "" if value is None else str(value).strip()

        def many(key: str) -> tuple[str, ...]:
            value = data.get(key) or ()
            if isinstance(value, str):
                value = [piece for piece in re.split(r"[,;\s]+", value) if piece]
            return tuple(str(one).strip() for one in value if str(one).strip())

        kind = text("kind") or "other"
        return cls(
            sheet=sheet,
            cell=text("cell").upper(),
            kind=kind if kind in KINDS else "other",
            reason=text("reason"),
            confidence=text("confidence"),
            reads=text("reads").upper(),
            should_read=text("should_read").upper(),
            omits=tuple(one.upper() for one in many("omits")),
            neighbours=tuple(one.upper() for one in many("neighbours")),
            alternative_formula=text("alternative_formula"),
            expected_value=text("expected_value"),
            other_cell=text("other_cell").upper(),
        )


@dataclass
class Verdict:
    status: str  # confirmed | refuted | unverifiable
    #: Every fact checked, as « holds: … » or « fails: … » sentences.
    facts: list[str] = field(default_factory=list)
    #: What the verdict does not settle — the judgment left to a person.
    note: str = ""


# --- helpers ---------------------------------------------------------------


def _parse(ref: str, sheet: str) -> tuple[str, int, int] | None:
    match = _A1.match(ref.strip())
    if match is None:
        return None
    target = (match.group("sheet") or sheet).strip("'")
    try:
        column = column_index_from_string(match.group("col").upper())
    except ValueError:
        return None
    return target, int(match.group("row")), column


def _cell(book: Workbook, sheet: str, ref: str) -> Cell | None:
    parsed = _parse(ref, sheet)
    if parsed is None:
        return None
    target, row, column = parsed
    return book.cells.get(f"{target}!{get_column_letter(column)}{row}")


def _populated(book: Workbook, sheet: str, ref: str) -> bool:
    cell = _cell(book, sheet, ref)
    return cell is not None and (cell.formula is not None or cell.value is not None)


def _references(formula: str, sheet: str) -> list[tuple[str, int, int, int, int]]:
    """Every reference in the formula as (sheet, r1, c1, r2, c2)."""
    out = []
    try:
        tokens = tokens_of(formula)
    except Exception:
        return out
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        pieces = token.value.split(":")
        first = _parse(pieces[0], sheet)
        if first is None:
            continue
        last = _parse(pieces[-1], first[0]) if len(pieces) > 1 else first
        if last is None:
            continue
        out.append(
            (
                first[0],
                min(first[1], last[1]),
                min(first[2], last[2]),
                max(first[1], last[1]),
                max(first[2], last[2]),
            )
        )
    return out


def _covers(refs: list[tuple[str, int, int, int, int]], sheet: str, ref: str) -> bool:
    parsed = _parse(ref, sheet)
    if parsed is None:
        return False
    target, row, column = parsed
    return any(
        s == target and r1 <= row <= r2 and c1 <= column <= c2
        for s, r1, c1, r2, c2 in refs
    )


def _shape(cell: Cell) -> str:
    """References made relative to the cell, numbers erased — enough to
    say whether two cells are the same formula filled across."""
    if cell.formula is None:
        return ""
    out = []
    try:
        tokens = tokens_of(cell.formula)
    except Exception:
        return cell.formula
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            pieces = []
            for piece in token.value.split(":"):
                parsed = _parse(piece, cell.sheet)
                if parsed is None:
                    pieces.append(piece)
                    continue
                target, row, column = parsed
                absolute_col = "$" in piece.split("!")[-1].split(str(row))[0]
                absolute_row = piece.rstrip().endswith(f"${row}")
                col_part = (
                    f"C{column}" if absolute_col else f"C[{column - cell.column}]"
                )
                row_part = f"R{row}" if absolute_row else f"R[{row - cell.row}]"
                prefix = f"{target}!" if target != cell.sheet else ""
                pieces.append(f"{prefix}{row_part}{col_part}")
            out.append(":".join(pieces))
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            out.append("#")
        else:
            out.append(token.value.replace(" ", ""))
    return "".join(out)


def _value(cell: Cell | None) -> Decimal | None:
    if cell is None or cell.value is None:
        return None
    if isinstance(cell.value, Decimal):
        return cell.value
    try:
        return Decimal(str(cell.value))
    except InvalidOperation:
        return None


class _Evaluator:
    """Sums, minimums, maximums, averages, rounding and the four
    operations over cells and ranges — and nothing else. Anything else
    raises, and the caller says « unverifiable »."""

    FUNCTIONS = {"SUM", "MIN", "MAX", "AVERAGE", "ROUND", "ABS"}

    def __init__(self, book: Workbook, sheet: str) -> None:
        self.book, self.sheet = book, sheet

    def evaluate(self, formula: str) -> Decimal:
        text = formula.strip()
        if text.startswith("="):
            text = text[1:]
        tokens = [t for t in tokens_of("=" + text) if t.type != "WHITE-SPACE"]
        self.tokens, self.at = tokens, 0
        value = self._expression()
        if self.at != len(self.tokens):
            raise ValueError("trailing tokens")
        return value

    def _peek(self):
        return self.tokens[self.at] if self.at < len(self.tokens) else None

    def _take(self):
        token = self._peek()
        self.at += 1
        return token

    def _expression(self) -> Decimal:
        value = self._term()
        while (
            (t := self._peek()) is not None
            and t.type == "OPERATOR-INFIX"
            and t.value in "+-"
        ):
            self._take()
            other = self._term()
            value = value + other if t.value == "+" else value - other
        return value

    def _term(self) -> Decimal:
        value = self._factor()
        while (
            (t := self._peek()) is not None
            and t.type == "OPERATOR-INFIX"
            and t.value in "*/"
        ):
            self._take()
            other = self._factor()
            value = value * other if t.value == "*" else value / other
        return value

    def _factor(self) -> Decimal:
        token = self._take()
        if token is None:
            raise ValueError("unexpected end")
        if token.type == "OPERATOR-PREFIX":
            inner = self._factor()
            return -inner if token.value == "-" else inner
        if token.type == "PAREN" and token.subtype == "OPEN":
            value = self._expression()
            close = self._take()
            if close is None or close.type != "PAREN":
                raise ValueError("missing )")
            return value
        if token.type == "OPERAND" and token.subtype == "NUMBER":
            return Decimal(token.value)
        if token.type == "OPERAND" and token.subtype == "RANGE":
            values = self._range_values(token.value)
            if len(values) != 1:
                raise ValueError("a range where one cell was expected")
            return values[0]
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value.rstrip("(").upper()
            if name not in self.FUNCTIONS:
                raise ValueError(f"function {name} is outside the evaluator")
            args: list[list[Decimal]] = []
            while True:
                nxt = self._peek()
                if nxt is not None and nxt.type == "FUNC" and nxt.subtype == "CLOSE":
                    self._take()
                    break
                if (
                    nxt is not None
                    and nxt.type == "OPERAND"
                    and nxt.subtype == "RANGE"
                    and ":" in nxt.value
                ):
                    self._take()
                    args.append(self._range_values(nxt.value))
                else:
                    args.append([self._expression()])
                sep = self._peek()
                if sep is not None and sep.type == "SEP":
                    self._take()
            flat = [v for arg in args for v in arg]
            if name == "SUM":
                return sum(flat, Decimal(0))
            if name == "MIN":
                return min(flat)
            if name == "MAX":
                return max(flat)
            if name == "AVERAGE":
                return sum(flat, Decimal(0)) / Decimal(len(flat))
            if name == "ABS":
                return abs(flat[0])
            if name == "ROUND":
                places = int(flat[1]) if len(flat) > 1 else 0
                return round(flat[0], places)
        raise ValueError(f"token {token.value!r} is outside the evaluator")

    def _range_values(self, text: str) -> list[Decimal]:
        pieces = text.split(":")
        first = _parse(pieces[0], self.sheet)
        if first is None:
            raise ValueError(f"unreadable reference {text}")
        last = _parse(pieces[-1], first[0]) if len(pieces) > 1 else first
        if last is None:
            raise ValueError(f"unreadable reference {text}")
        out = []
        for row in range(min(first[1], last[1]), max(first[1], last[1]) + 1):
            for column in range(min(first[2], last[2]), max(first[2], last[2]) + 1):
                cell = self.book.cells.get(
                    f"{first[0]}!{get_column_letter(column)}{row}"
                )
                value = _value(cell)
                if value is not None:
                    out.append(value)
                elif len(pieces) == 1:
                    #: A single empty cell reads as zero, as Excel does.
                    out.append(Decimal(0))
        return out


# --- the checks ------------------------------------------------------------


def check(book: Workbook, claim: Claim) -> Verdict:
    """Every cited fact against the cells; the verdict and its facts."""
    facts: list[str] = []
    cell = _cell(book, claim.sheet, claim.cell)

    def holds(text: str) -> None:
        facts.append(f"holds: {text}")

    def fails(text: str) -> Verdict:
        facts.append(f"fails: {text}")
        return Verdict("refuted", facts)

    if claim.kind == "other":
        return Verdict("unverifiable", facts, "no checkable fact was cited")
    if cell is None:
        return fails(f"{claim.cell} holds nothing on {claim.sheet}")

    if claim.kind == "error-value":
        shown = book.errors.get(cell.ref)
        if shown in ERROR_VALUES:
            holds(f"{claim.cell} shows {shown}")
            return Verdict("confirmed", facts)
        return fails(f"{claim.cell} does not show an error value")

    if claim.kind == "malformed":
        if cell.formula is None:
            return fails(f"{claim.cell} holds no formula")
        try:
            tokens = tokens_of(cell.formula)
        except Exception:
            holds(f"{claim.cell}'s formula does not parse")
            return Verdict("confirmed", facts)
        depth = 0
        for token in tokens:
            if token.type in ("FUNC", "PAREN"):
                depth += 1 if token.subtype == "OPEN" else -1
            elif (
                depth == 0
                and token.type == "OPERATOR-INFIX"
                and token.value in ("=", "<>", "<", ">", "<=", ">=")
            ):
                holds(
                    f"{claim.cell}'s formula compares with « {token.value} » at its top level, so it yields TRUE or FALSE"
                )
                return Verdict("confirmed", facts)
        return fails(f"{claim.cell}'s formula parses and has no top-level comparison")

    if claim.kind in (
        "range-omits",
        "wrong-reference",
        "empty-reference",
        "differs-from-neighbours",
    ):
        if cell.formula is None:
            return fails(f"{claim.cell} holds no formula")
    refs = _references(cell.formula or "", claim.sheet)

    if claim.kind == "range-omits":
        if not claim.omits:
            return Verdict("unverifiable", facts, "no omitted cell was named")
        for ref in claim.omits:
            if _covers(refs, claim.sheet, ref):
                return fails(f"{claim.cell}'s formula does cover {ref}")
            holds(f"{claim.cell}'s formula leaves {ref} out")
            if not _populated(book, claim.sheet, ref):
                return fails(f"{ref} holds nothing")
            holds(f"{ref} is populated")
        return Verdict(
            "confirmed",
            facts,
            "whether the range should cover those cells is the reviewer's call",
        )

    if claim.kind == "wrong-reference":
        if not claim.reads:
            return Verdict("unverifiable", facts, "the cell it reads was not named")
        if not _covers(refs, claim.sheet, claim.reads):
            return fails(f"{claim.cell}'s formula does not read {claim.reads}")
        holds(f"{claim.cell}'s formula reads {claim.reads}")
        if claim.should_read:
            if not _populated(book, claim.sheet, claim.should_read):
                return fails(f"{claim.should_read} holds nothing")
            holds(f"{claim.should_read} is populated")
        return Verdict(
            "confirmed", facts, "which cell it should read is the reviewer's call"
        )

    if claim.kind == "empty-reference":
        if not claim.reads:
            return Verdict("unverifiable", facts, "the cell it reads was not named")
        if not _covers(refs, claim.sheet, claim.reads):
            return fails(f"{claim.cell}'s formula does not read {claim.reads}")
        holds(f"{claim.cell}'s formula reads {claim.reads}")
        if _populated(book, claim.sheet, claim.reads):
            return fails(f"{claim.reads} is not empty")
        holds(f"{claim.reads} is empty")
        return Verdict("confirmed", facts)

    if claim.kind == "differs-from-neighbours":
        if not claim.neighbours:
            return Verdict("unverifiable", facts, "no neighbour was named")
        mine = _shape(cell)
        theirs = []
        for ref in claim.neighbours:
            other = _cell(book, claim.sheet, ref)
            if other is None or other.formula is None:
                return fails(f"{ref} holds no formula")
            theirs.append(_shape(other))
        if any(shape == mine for shape in theirs):
            return fails(f"{claim.cell} is built the same way as a named neighbour")
        holds(f"{claim.cell} is built unlike {', '.join(claim.neighbours)}")
        if len(set(theirs)) > 1:
            return fails(
                "the named neighbours are not built the same way as each other"
            )
        if len(theirs) > 1:
            holds("the named neighbours are built the same way as each other")
        return Verdict(
            "confirmed",
            facts,
            "whether the row should be uniform is the reviewer's call",
        )

    if claim.kind == "arithmetic":
        if not claim.alternative_formula:
            return Verdict("unverifiable", facts, "no alternative formula was given")
        try:
            computed = _Evaluator(book, claim.sheet).evaluate(claim.alternative_formula)
        except Exception as error:
            return Verdict(
                "unverifiable",
                facts,
                f"the alternative could not be evaluated: {error}",
            )
        holds(f"{claim.alternative_formula} evaluates to {computed}")
        if claim.expected_value:
            try:
                expected = Decimal(claim.expected_value.replace(",", ""))
            except InvalidOperation:
                return Verdict(
                    "unverifiable", facts, "the expected value is not a number"
                )
            scale = max(abs(expected), abs(computed), Decimal(1))
            if abs(expected - computed) / scale > Decimal("0.000001"):
                return fails(
                    f"the claim expected {expected}, the alternative gives {computed}"
                )
            holds(f"which matches the {expected} the claim expects")
        current = _value(cell)
        if current is None:
            return fails(f"{claim.cell} holds no number to compare")
        if abs(current - computed) <= Decimal("0.000001") * max(
            abs(current), Decimal(1)
        ):
            return fails(
                f"{claim.cell} already holds {current}, the same as the alternative"
            )
        holds(f"{claim.cell} holds {current}, which differs")
        return Verdict(
            "confirmed",
            facts,
            "whether the alternative is the right calculation is the reviewer's call",
        )

    if claim.kind == "value-contradiction":
        if not claim.other_cell:
            return Verdict("unverifiable", facts, "the other cell was not named")
        mine, theirs = _value(cell), _value(_cell(book, claim.sheet, claim.other_cell))
        if mine is None or theirs is None:
            return fails("one of the two cells holds no number")
        if mine == theirs:
            return fails(f"{claim.cell} and {claim.other_cell} hold the same value")
        holds(f"{claim.cell} holds {mine} and {claim.other_cell} holds {theirs}")
        return Verdict(
            "confirmed", facts, "whether they should agree is the reviewer's call"
        )

    return Verdict("unverifiable", facts, f"kind {claim.kind} is not checked")


__all__ = ["KINDS", "Claim", "Verdict", "check"]
