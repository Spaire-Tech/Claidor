"""Ingesting a file, running a check, and answering the deal page.

The engine's own modules read files. This one reads **rows**, and that is
the whole difference between a library and a product: once a deck and a
model have been ingested, every re-check is a query and a comparison. No
file is opened, nothing is re-parsed, and a link a banker confirmed is
re-tested by arithmetic that cannot come out differently on Tuesday.

The adapters below — :func:`_workbook_of`, :func:`_figures_of` — exist for
exactly that. They rebuild the engine's own shapes out of what was stored,
so one implementation of the linker, the comparison and the chain serves
both the offline scripts and the running product.
"""

import hashlib
import re
from collections.abc import Sequence
from decimal import Decimal
from typing import Any
from uuid import UUID

import structlog

from polar.kit.db.postgres import AsyncReadSession, AsyncSession
from polar.models import (
    Artifact,
    ArtifactKind,
    ArtifactStatus,
    CheckKind,
    CheckRun,
    CheckStatus,
    Dossier,
    FindingKind,
    FindingSeverity,
    LinkState,
    OneOffCheck,
)
from polar.models import Figure as FigureRow
from polar.models import FigureLink as LinkRow
from polar.models import Finding as FindingRow
from polar.models import ModelCell as CellRow

from . import figures as engine_figures
from . import storage
from .check import compare
from .ingest import Ingested, Unreadable, read_artifact
from .link import link as propose_links
from .link import rank as rank_outputs
from .message import read_message
from .model import Output
from .numbers import show
from .provenance import chain as render_chain
from .provenance import inputs_from_workbook, outputs_from_workbook
from .repository import TieOutRepository
from .solo import Statement, disagreements, repeated
from .workbook import Cell as EngineCell
from .workbook import Workbook

log = structlog.get_logger()


class TieOutService:
    # --- ingestion ------------------------------------------------------

    async def ingest(
        self,
        session: AsyncSession,
        *,
        dossier_id: UUID,
        kind: ArtifactKind,
        filename: str,
        payload: bytes,
        user_id: UUID,
        file_id: UUID | None = None,
        external_id: str | None = None,
        external_version: str | None = None,
        lineage_of: UUID | None = None,
    ) -> Artifact:
        """Read a file into rows, or record why it could not be read.

        A failed artifact is kept rather than discarded: the deal page has
        to be able to show *« this one did not work, and here is what to do
        about it »*, which is impossible if the row is gone.

        **The bytes are kept too, and a file that fails to read keeps them
        as well.** Every check runs off the rows and never opens a document
        again; writing a correction is the one thing that cannot, because a
        correction is a new version of a real file. Storing is best-effort
        by design — see :mod:`polar.tieout.storage`.
        """
        repository = TieOutRepository.from_session(session)
        artifact = await repository.create_artifact(
            dossier_id=dossier_id,
            kind=kind,
            filename=filename,
            uploaded_by_id=user_id,
            file_id=file_id,
            external_id=external_id,
            external_version=external_version,
            lineage_of=lineage_of,
        )
        artifact.storage_path = storage.keep(artifact, payload)

        try:
            ingested = read_artifact(payload, filename, kind)
        except Unreadable as error:
            artifact.status = ArtifactStatus.failed
            artifact.error = str(error)
            session.add(artifact)
            await session.flush()
            log.info(
                "tieout.ingest.unreadable",
                artifact=str(artifact.id),
                reason=str(error)[:120],
            )
            return artifact

        await self._persist(session, artifact, ingested)
        artifact.status = ArtifactStatus.ready
        artifact.counts = ingested.counts
        artifact.outputs = ingested.outputs
        artifact.error = None
        session.add(artifact)
        await session.flush()
        return artifact

    async def ingest_message(
        self,
        session: AsyncSession,
        *,
        dossier_id: UUID,
        subject: str,
        body: str,
        html: bool,
        user_id: UUID,
        external_id: str,
        external_version: str,
        lineage_of: UUID | None = None,
    ) -> Artifact:
        """Read an email into rows, so it can be checked like anything else.

        Sibling to :meth:`ingest` rather than a branch inside it, because a
        message is not a file: there are no bytes, no suffix and nothing to
        keep. What comes out is identical — figures, named, located — and
        from there the tie-out cannot tell it from a memo, which is the
        point.

        **Nothing is stored.** :meth:`ingest` keeps the document because a
        correction has to write a new version of a real file; a message is
        corrected in Outlook and never here, so there is nothing a stored
        copy of somebody's mail would make possible. What is kept is what
        is kept for every document: the figures, and the sentence each one
        was found in.
        """
        repository = TieOutRepository.from_session(session)
        artifact = await repository.create_artifact(
            dossier_id=dossier_id,
            kind=ArtifactKind.message,
            # The subject *is* the name of this document, and it is the
            # name a person would use for it.
            filename=subject.strip() or "(no subject)",
            uploaded_by_id=user_id,
            external_id=external_id,
            external_version=external_version,
            lineage_of=lineage_of,
        )

        extraction = read_message(subject, body, html=html)
        ingested = Ingested(
            figures=extraction.figures,
            counts={
                "figures": len(extraction.figures),
                "paragraphs_with_figures": len(
                    {one.location for one in extraction.figures}
                ),
            },
        )
        await self._persist(session, artifact, ingested)
        artifact.status = ArtifactStatus.ready
        artifact.counts = ingested.counts
        artifact.error = None
        session.add(artifact)
        await session.flush()
        return artifact

    async def _persist(
        self, session: AsyncSession, artifact: Artifact, ingested: Ingested
    ) -> None:
        repository = TieOutRepository.from_session(session)
        if ingested.figures:
            await repository.replace_figures(
                artifact.id,
                [
                    FigureRow(
                        artifact_id=artifact.id,
                        printed=figure.printed,
                        value=figure.value,
                        decimals=figure.decimals,
                        kind=figure.kind,
                        page=figure.slide,
                        label=figure.label,
                        location=figure.location,
                        context=figure.context,
                        section=figure.section,
                        range_endpoint=figure.range_endpoint,
                        parenthesised=figure.parenthesised,
                        subject=figure.subject,
                        anchor=figure.anchor,
                    )
                    for figure in ingested.figures
                ],
            )
        if ingested.cells:
            await repository.replace_cells(
                artifact.id,
                [
                    CellRow(
                        artifact_id=artifact.id,
                        ref=cell.ref,
                        sheet=cell.sheet,
                        row=cell.row,
                        column=cell.column,
                        value=cell.value,
                        formula=cell.formula,
                        row_label=cell.row_label,
                        column_label=cell.column_label,
                        number_format=cell.number_format,
                        name=cell.name,
                        precedents=list(cell.precedents),
                        unresolved=[list(one) for one in cell.unresolved],
                        alias_of=cell.alias_of,
                    )
                    for cell in ingested.cells
                ],
            )

    # --- checking -------------------------------------------------------

    async def run_tieout(
        self, session: AsyncSession, *, dossier_id: UUID, user_id: UUID | None
    ) -> CheckRun:
        """Reconcile every deck in the deal against every model in it.

        Reads rows. The files are never opened again, which is what makes
        a confirmed link re-checkable forever and what lets the documents
        be dropped while the chain is kept.
        """
        repository = TieOutRepository.from_session(session)
        current = await repository.current_artifacts(dossier_id)
        # A memo is checked exactly as a deck is: printed figures, named
        # by the words around them, reconciled against the model. The only
        # difference is that it has paragraphs where a deck has slides,
        # and that difference lives in the reader, not here. A message is
        # the same sentence again — it asserts figures, so it is checked.
        decks = [
            one
            for one in current
            if one.kind in (ArtifactKind.deck, ArtifactKind.memo, ArtifactKind.message)
        ]
        models = [one for one in current if one.kind is ArtifactKind.model]

        run = await repository.start_run(
            dossier_id=dossier_id,
            kind=CheckKind.tieout,
            artifact_ids=[one.id for one in decks + models],
            requested_by_id=user_id,
        )
        if not decks or not models:
            missing = "model" if not models else "deck or memo"
            return await repository.finish_run(
                run,
                error=f"nothing to reconcile — this deal has no {missing} yet",
            )

        findings: list[FindingRow] = []
        links: list[LinkRow] = []
        totals = {"reconciled": 0, "agreeing": 0, "drifting": 0, "unlinked": 0}
        reasons: dict[str, int] = {}
        decided = await repository.decided_pairs(dossier_id)

        for model in models:
            cells = await repository.cells_of(model.id)
            book = _workbook_of(cells)
            by_ref = {cell.ref: cell for cell in cells}
            candidates = _candidates(book, cells, model.outputs)

            for deck in decks:
                rows = await repository.figures_of(deck.id)
                engine, back = _figures_of(rows)

                # Two passes, merged. The published pass reaches figures
                # the workbook has no cell for — a CAGR computed on the
                # Outputs tab itself — and the workbook pass reaches
                # everything the tab never published. Neither subsumes the
                # other, and running only one costs seven reconciled
                # figures on the Cascade deck.
                proposed, unlinked = _both_passes(engine, candidates)
                drifts, agreed = compare(proposed)

                totals["reconciled"] += len(proposed)
                totals["agreeing"] += len(agreed)
                totals["drifting"] += len(drifts)
                totals["unlinked"] += len(unlinked)
                for item in unlinked:
                    reasons[_reason(item.reason)] = (
                        reasons.get(_reason(item.reason), 0) + 1
                    )

                for item in proposed:
                    figure_row = back[id(item.figure)]
                    cell_row = candidates.cells.get(item.output.source) or (
                        candidates.cells.get(item.output.ref)
                    )
                    if cell_row is None or (figure_row.id, cell_row.id) in decided:
                        continue
                    links.append(
                        LinkRow(
                            dossier_id=dossier_id,
                            figure_id=figure_row.id,
                            cell_id=cell_row.id,
                            state=LinkState.proposed,
                            confidence=round(item.score, 3),
                            transformation="identity",
                            basis=item.output.basis,
                            cell_name=item.output.name,
                            figure_label=item.figure.label,
                        )
                    )

                for drift in drifts:
                    source = by_ref.get(drift.source)
                    findings.append(
                        FindingRow(
                            dossier_id=dossier_id,
                            check_run_id=run.id,
                            artifact_id=deck.id,
                            kind=FindingKind.drift,
                            severity=FindingSeverity.error,
                            fingerprint=_fingerprint(
                                "drift", deck.lineage_id, drift.location, drift.printed
                            ),
                            page=drift.slide,
                            printed=drift.printed,
                            expected=drift.expected,
                            one_tick=drift.one_tick,
                            title=f"{drift.printed} where the model says {drift.expected}",
                            detail=drift.context,
                            location=drift.location,
                            anchor=drift.anchor,
                            evidence={
                                "ref": drift.ref,
                                "name": drift.name,
                                "source": drift.source,
                                "basis": drift.basis,
                                "confidence": round(drift.confidence, 3),
                                "model_artifact_id": str(model.id),
                                "chain": render_chain(book, drift.source)
                                if source
                                else "",
                            },
                        )
                    )

        await repository.replace_findings(dossier_id, CheckKind.tieout, findings)
        await repository.replace_proposals(
            dossier_id, links, figures_on=[one.id for one in decks]
        )

        summary = {
            **totals,
            "confirmed": sum(
                1
                for one in await repository.links_of(dossier_id)
                if one.state is LinkState.confirmed
            ),
            "reasons": [
                {"reason": reason, "count": count}
                for reason, count in sorted(reasons.items(), key=lambda pair: -pair[1])
            ],
        }
        return await repository.finish_run(run, summary=summary)

    async def run_audit(
        self, session: AsyncSession, *, dossier_id: UUID, user_id: UUID | None
    ) -> CheckRun:
        """Check every model in the deal against itself.

        The firm's house rules apply here: an audit rule the
        organization switched off is skipped, and the run's summary
        names what was skipped — a rule turned off is a decision on the
        record, never a silence.

        The statement checks run with the mechanical rules and land in
        the same findings table: whether the balance sheet balances,
        cash carries forward, debt repays, the time axis holds, and
        what the model's own check rows say. They read cached values,
        so they still speak on values-pasted close copies — and the
        summary says when a model is such a copy, what each check
        examined, and where one abstained rather than guess.
        """
        from .analytics import (
            ANALYTIC_RULE_NAMES,
            ANALYTIC_STANDARD_SENTENCES,
            ANALYTIC_STANDARDS,
            run_analytics,
        )
        from .audit import HEADLINES, plain_words
        from .audit import audit as run_rules
        from .structure import read_structure

        repository = TieOutRepository.from_session(session)
        rules_off = await self._rules_off(repository, dossier_id)
        models = [
            one
            for one in await repository.current_artifacts(dossier_id)
            if one.kind is ArtifactKind.model
        ]
        run = await repository.start_run(
            dossier_id=dossier_id,
            kind=CheckKind.audit,
            artifact_ids=[one.id for one in models],
            requested_by_id=user_id,
        )
        if not models:
            return await repository.finish_run(
                run, error="this deal has no model to audit yet"
            )

        findings: list[FindingRow] = []
        errors = smells = 0
        values_only = False
        abstentions: list[dict[str, str]] = []
        tallies: dict[str, dict[str, int]] = {}
        statement_keys = set(ANALYTIC_RULE_NAMES) - rules_off
        for model in models:
            cells = await repository.cells_of(model.id)
            book = _workbook_of(cells)
            #: The cells cannot say what the workbook hides — that fact
            #: was kept on the artifact at ingest, and the audit needs
            #: it back before it runs.
            book.hidden_sheets = tuple(model.counts.get("hidden_sheets", []))
            book.very_hidden_sheets = tuple(model.counts.get("very_hidden_sheets", []))
            structure = read_structure(book)
            result = run_rules(book, axes=structure.axes)
            result.findings = [
                one for one in result.findings if one.rule not in rules_off
            ]
            errors += len(result.errors)
            smells += len(result.smells)
            for defect in result.findings:
                findings.append(
                    FindingRow(
                        dossier_id=dossier_id,
                        check_run_id=run.id,
                        artifact_id=model.id,
                        kind=FindingKind.audit,
                        severity=FindingSeverity(defect.severity),
                        fingerprint=_fingerprint(
                            "audit", model.lineage_id, defect.ref, defect.rule
                        ),
                        rule=defect.rule,
                        standard=defect.source,
                        printed=defect.ref,
                        #: The plain sentence is the title — what a
                        #: person reads first; the formula stays in the
                        #: detail as evidence beneath it.
                        title=plain_words(defect, structure.axes),
                        detail=defect.detail,
                        location=defect.ref,
                        # An audit finding already sits at a cell, which
                        # is a coordinate Excel selects as it stands.
                        anchor={
                            "kind": "cell",
                            "ref": defect.ref,
                            "sheet": defect.sheet,
                        },
                        evidence={
                            "sheet": defect.sheet,
                            "name": defect.name,
                            "headline": HEADLINES.get(defect.rule, ""),
                            "chain": render_chain(book, defect.ref),
                            "figure": defect.figure,
                            "figure_unit": defect.figure_unit,
                            "flow": defect.flow,
                            "fix": defect.fix,
                            "fix_before": defect.fix_before,
                            "grid": _neighbourhood(
                                book,
                                defect.sheet,
                                defect.ref,
                                axes=structure.axes,
                            ),
                        },
                    )
                )

            #: The statement checks, on the same workbook. They read
            #: values, not formulas, so a values-pasted close copy —
            #: where the rules above are nearly blind — is exactly
            #: where they earn their keep.
            values_only = values_only or structure.values_pasted
            if statement_keys:
                told = run_analytics(book, structure)
                for claim in told.findings:
                    if claim.rule not in statement_keys:
                        continue
                    if claim.severity == "smell":
                        smells += 1
                    else:
                        errors += 1
                    findings.append(
                        FindingRow(
                            dossier_id=dossier_id,
                            check_run_id=run.id,
                            artifact_id=model.id,
                            kind=FindingKind.audit,
                            severity=FindingSeverity(claim.severity),
                            fingerprint=_fingerprint(
                                "audit", model.lineage_id, claim.ref, claim.rule
                            ),
                            rule=claim.rule,
                            standard=ANALYTIC_STANDARDS.get(claim.rule, ""),
                            printed=claim.figure,
                            title=(f"{ANALYTIC_RULE_NAMES[claim.rule]} at {claim.ref}"),
                            detail=claim.detail,
                            location=claim.ref,
                            anchor={
                                "kind": "cell",
                                "ref": claim.ref,
                                "sheet": claim.sheet,
                            },
                            evidence={
                                "sheet": claim.sheet,
                                "name": claim.row_label,
                                "headline": ANALYTIC_RULE_NAMES[claim.rule],
                                "period": claim.period,
                                "value": claim.value,
                                "figure": claim.figure,
                                "figure_unit": claim.figure_unit,
                                "grid": _neighbourhood(
                                    book,
                                    claim.sheet,
                                    claim.ref,
                                    axes=structure.axes,
                                ),
                                "standard_sentence": (
                                    ANALYTIC_STANDARD_SENTENCES.get(claim.rule, "")
                                ),
                            },
                        )
                    )
                abstentions.extend(
                    {"rule": one.rule, "why": one.why}
                    for one in told.abstentions
                    if one.rule in statement_keys
                )
                for rule, tally in told.tallies.items():
                    if rule not in statement_keys:
                        continue
                    merged = tallies.setdefault(rule, {"total": 0, "clean": 0})
                    merged["total"] += tally["total"]
                    merged["clean"] += tally["clean"]

        await repository.replace_findings(dossier_id, CheckKind.audit, findings)
        return await repository.finish_run(
            run,
            summary={
                "errors": errors,
                "smells": smells,
                "models": len(models),
                "cells": sum(int(one.counts.get("cells", 0)) for one in models),
                "rules_off": sorted(rules_off),
                #: The statement checks' own record: what a pass row may
                #: claim, why a check stayed silent, and whether this is
                #: a copy the construction rules could not read.
                "values_only": values_only,
                "abstentions": abstentions,
                "tallies": tallies,
            },
        )

    async def _rules_off(
        self, repository: TieOutRepository, dossier_id: UUID
    ) -> set[str]:
        """The audit rules this deal's organization switched off."""
        organization_id = await repository.organization_of(dossier_id)
        if organization_id is None:
            return set()
        rules = await repository.house_rules_for(organization_id)
        return set(rules.audit_rules_off) if rules else set()

    async def grounding_on(self, session: AsyncSession, *, dossier_id: UUID) -> bool:
        """Whether the firm runs the grounding pass with the others."""
        repository = TieOutRepository.from_session(session)
        organization_id = await repository.organization_of(dossier_id)
        if organization_id is None:
            return True
        rules = await repository.house_rules_for(organization_id)
        return rules.grounding if rules else True

    async def run_crosscheck(
        self, session: AsyncSession, *, dossier_id: UUID, user_id: UUID | None
    ) -> CheckRun:
        """Ground the model's typed inputs in the documents behind them.

        **The chain's last hop, and the only check that leaves the deal's
        own arithmetic.** Everything else here asks whether two things the
        team produced agree with each other. This asks the question that
        was underneath all of them: the model says revenue was 228.9 —
        *says who?*

        The matcher is the one that links a deck to a model, used in the
        other direction: a figure printed in the accounts is the printed
        thing, and a typed input is the candidate it might be the origin
        of. That is not a convenience. The gates that make the tie-out
        refuse rather than guess — the label must name it, the period must
        not contradict, the margin over the runner-up must be clear — are
        exactly the gates a hundred pages of somebody else's prose needs,
        and a second matcher written for this would have to earn them
        again.

        A disagreement here is a `contradiction`: two documents in the deal
        say different things. It is deliberately not a `drift`, which is a
        deliverable disagreeing with the model and is a mistake somebody
        made this week. « The signed accounts restated cost of sales and
        the model still carries the draft figure » is a different sentence
        and a different fix.
        """
        repository = TieOutRepository.from_session(session)
        current = await repository.current_artifacts(dossier_id)
        sources = [one for one in current if one.kind is ArtifactKind.source]
        models = [one for one in current if one.kind is ArtifactKind.model]

        run = await repository.start_run(
            dossier_id=dossier_id,
            kind=CheckKind.crosscheck,
            artifact_ids=[one.id for one in sources + models],
            requested_by_id=user_id,
        )
        if not sources or not models:
            missing = "model" if not models else "source document"
            return await repository.finish_run(
                run,
                error=(
                    f"nothing to ground — this deal has no {missing} yet. "
                    "A source is the audited accounts, a term sheet, "
                    "anything a typed input came out of"
                ),
            )

        findings: list[FindingRow] = []
        links: list[LinkRow] = []
        totals = {"grounded": 0, "agreeing": 0, "contradicting": 0, "unlinked": 0}
        reasons: dict[str, int] = {}
        decided = await repository.decided_pairs(dossier_id)

        for model in models:
            cells = await repository.cells_of(model.id)
            book = _workbook_of(cells)
            by_ref = {cell.ref: cell for cell in cells}
            inputs = inputs_from_workbook(book)

            for source in sources:
                rows = await repository.figures_of(source.id)
                engine, back = _figures_of(rows)
                proposed, unlinked = propose_links(
                    engine,
                    inputs,
                    year=(source.counts or {}).get("document_year"),
                )
                contradictions, agreed = compare(proposed)

                totals["grounded"] += len(proposed)
                totals["agreeing"] += len(agreed)
                totals["contradicting"] += len(contradictions)
                totals["unlinked"] += len(unlinked)
                for miss in unlinked:
                    reasons[_reason(miss.reason)] = (
                        reasons.get(_reason(miss.reason), 0) + 1
                    )

                for item in proposed:
                    figure_row = back[id(item.figure)]
                    cell_row = by_ref.get(item.output.ref)
                    if cell_row is None or (figure_row.id, cell_row.id) in decided:
                        continue
                    links.append(
                        LinkRow(
                            dossier_id=dossier_id,
                            figure_id=figure_row.id,
                            cell_id=cell_row.id,
                            state=LinkState.proposed,
                            confidence=round(item.score, 3),
                            transformation="identity",
                            basis=item.output.basis,
                            cell_name=item.output.name,
                            figure_label=item.figure.label,
                        )
                    )

                for gap in contradictions:
                    findings.append(
                        FindingRow(
                            dossier_id=dossier_id,
                            check_run_id=run.id,
                            #: The finding is *on the source*, because that
                            #: is where a reader has to go to settle it —
                            #: page 2 of the accounts, not a cell that is
                            #: only doing what it was told.
                            artifact_id=source.id,
                            kind=FindingKind.contradiction,
                            severity=FindingSeverity.error,
                            fingerprint=_fingerprint(
                                "contradiction",
                                source.lineage_id,
                                gap.location,
                                gap.printed,
                            ),
                            page=gap.slide,
                            printed=gap.printed,
                            expected=gap.expected,
                            one_tick=gap.one_tick,
                            title=(
                                f"{source.filename} says {gap.printed} "
                                f"where the model has {gap.expected}"
                            ),
                            detail=gap.context,
                            location=gap.location,
                            anchor=gap.anchor,
                            evidence={
                                "ref": gap.ref,
                                "name": gap.name,
                                "source": gap.source,
                                "basis": gap.basis,
                                "confidence": round(gap.confidence, 3),
                                "model_artifact_id": str(model.id),
                                "chain": render_chain(book, gap.source),
                                #: Which way round the disagreement is.
                                #: The model is the thing that would be
                                #: corrected, and it is not on the screen
                                #: this finding points at.
                                "grounded_in": source.filename,
                            },
                        )
                    )

        await repository.replace_proposals(
            dossier_id, links, figures_on=[one.id for one in sources]
        )
        await repository.replace_findings(dossier_id, CheckKind.crosscheck, findings)
        return await repository.finish_run(
            run,
            summary={
                **totals,
                "sources": len(sources),
                "reasons": [
                    {"reason": reason, "count": count}
                    for reason, count in sorted(
                        reasons.items(), key=lambda pair: -pair[1]
                    )
                ],
            },
        )

    async def check_file(
        self,
        session: AsyncSession,
        *,
        user_id: UUID,
        kind: ArtifactKind,
        filename: str,
        payload: bytes,
        dossier: Dossier | None = None,
    ) -> OneOffCheck:
        """Check a loose file, keep the answer, drop the file.

        The one check that runs outside any deal's data room. What runs
        depends on what the file is:

        - **A deck or a memo alone** is checked against itself — the solo
          check, one name carrying two figures.
        - **A deck or a memo against a deal** additionally reconciles its
          figures with every current model in that deal, through the same
          two-pass linker and comparison as the deal's own tie-out. The
          solo findings still ride along: the file disagreeing with
          itself is worth saying whatever it was checked against.
        - **A model** is audited — a workbook checked by itself *is* the
          model audit, which has its own rules and cites its standards.
          A deal picked alongside a model is deliberately not a check:
          model-against-deal is the grounding, which needs the deal's
          source documents and is not a one-off — the stored row says
          « on its own », because that is what ran.

        Nothing is uploaded anywhere: the bytes are read in a temporary
        file and discarded, and the row keeps the counts and findings
        exactly as the screen received them, so a recent reopens to the
        same answer. `Unreadable` propagates to the caller with its
        reason in words.
        """
        ingested = read_artifact(payload, filename, kind)
        counts: dict[str, Any] = dict(ingested.counts)
        # A sheet listing is layout for the model page, not a tally.
        counts.pop("sheet_order", None)
        result: dict[str, Any] = {
            "disagreements": [],
            "drifts": [],
            "defects": [],
            "models": [],
        }
        checked_against = ""

        if kind is ArtifactKind.model:
            #: The workbook, rebuilt once for everything below: the
            #: plain sentences, the little grids, the statement checks.
            from .analytics import (
                ANALYTIC_RULE_NAMES,
                ANALYTIC_STANDARDS,
                run_analytics,
            )
            from .audit import HEADLINES, plain_words
            from .structure import read_structure
            from .workbook import Workbook as EngineWorkbook

            book = EngineWorkbook()
            for cell in ingested.cells:
                book.cells[cell.ref] = cell
            book.sheets = list(ingested.counts.get("sheet_order", []))
            structure = read_structure(book)

            result["defects"] = [
                {
                    "rule": defect.rule,
                    "severity": defect.severity,
                    "ref": defect.ref,
                    "sheet": defect.sheet,
                    "name": defect.name,
                    #: What a person reads first; the formula is
                    #: evidence beneath it, never the headline.
                    "plain": plain_words(defect, structure.axes),
                    "headline": HEADLINES.get(defect.rule, ""),
                    "detail": defect.detail,
                    "standard": defect.source,
                    "figure": defect.figure,
                    "figure_unit": defect.figure_unit,
                    "flow": defect.flow,
                    "fix": defect.fix,
                    "fix_before": defect.fix_before,
                    "grid": _neighbourhood(
                        book, defect.sheet, defect.ref, axes=structure.axes
                    ),
                }
                for defect in ingested.defects
            ]

            #: The statement checks run here too — same engine as the
            #: deal audit, on the cells just read, before the file is
            #: dropped. They read values, so a values-pasted close copy
            #: still gets a verdict about whether its accounts add up.
            told = run_analytics(book, structure)
            result["defects"].extend(
                {
                    "rule": claim.rule,
                    "severity": claim.severity,
                    "ref": claim.ref,
                    "sheet": claim.sheet,
                    "name": claim.row_label,
                    "plain": claim.detail,
                    "headline": ANALYTIC_RULE_NAMES.get(claim.rule, ""),
                    "detail": claim.detail,
                    "standard": ANALYTIC_STANDARDS.get(claim.rule, ""),
                    "analytical": True,
                    "figure": claim.figure,
                    "figure_unit": claim.figure_unit,
                    "period": claim.period,
                    "grid": _neighbourhood(
                        book, claim.sheet, claim.ref, axes=structure.axes
                    ),
                }
                for claim in told.findings
            )
            result["values_only"] = structure.values_pasted
            result["abstentions"] = [
                {"rule": one.rule, "why": one.why} for one in told.abstentions
            ]
            result["tallies"] = told.tallies
        else:
            extraction = engine_figures.Extraction(figures=list(ingested.figures))
            found = disagreements(extraction)
            counts["repeated"] = repeated(extraction)
            result["disagreements"] = [
                {
                    "label": one.label,
                    "statements": one.statements,
                    "first": _statement_json(one.first),
                    "other": _statement_json(one.other),
                }
                for one in found
            ]

            if dossier is not None:
                repository = TieOutRepository.from_session(session)
                models = [
                    one
                    for one in await repository.current_artifacts(dossier.id)
                    if one.kind is ArtifactKind.model
                ]
                totals = {"reconciled": 0, "agreeing": 0, "unlinked": 0}
                for model in models:
                    cells = await repository.cells_of(model.id)
                    book = _workbook_of(cells)
                    candidates = _candidates(book, cells, model.outputs)
                    proposed, unlinked = _both_passes(
                        list(ingested.figures), candidates
                    )
                    drifts, agreed = compare(proposed)
                    totals["reconciled"] += len(proposed)
                    totals["agreeing"] += len(agreed)
                    totals["unlinked"] += len(unlinked)
                    result["models"].append(
                        {
                            "artifact_id": str(model.id),
                            "filename": model.filename,
                            "version": model.version,
                            "read_at": model.created_at.isoformat(),
                        }
                    )
                    for drift in drifts:
                        result["drifts"].append(
                            {
                                "printed": drift.printed,
                                "expected": drift.expected,
                                "label": drift.name,
                                "page": drift.slide,
                                "location": drift.location,
                                "context": drift.context,
                                "ref": drift.ref,
                                "name": drift.name,
                                "basis": drift.basis,
                                "confidence": round(drift.confidence, 3),
                                "one_tick": drift.one_tick,
                                "model_artifact_id": str(model.id),
                            }
                        )
                counts.update(totals)
                counts["drifting"] = len(result["drifts"])
                checked_against = dossier.name

        counts["differences"] = len(result["disagreements"]) + len(result["drifts"])

        repository = TieOutRepository.from_session(session)
        return await repository.save_one_off(
            OneOffCheck(
                user_id=user_id,
                dossier_id=dossier.id if dossier and checked_against else None,
                against=checked_against,
                filename=filename,
                kind=kind,
                counts=counts,
                result=result,
            )
        )

    # --- reading --------------------------------------------------------

    async def coverage_of(
        self, session: AsyncSession | AsyncReadSession, *, dossier_id: UUID
    ) -> dict[str, Any]:
        """The coverage line, from the last tie-out that finished.

        Read off the stored run rather than recomputed, so the number on
        the screen is the number the check produced. A deal that has never
        been checked reports zeroes, which is honest: nothing has been
        looked at.
        """
        repository = TieOutRepository.from_session(session)
        run = await repository.latest_run(dossier_id, CheckKind.tieout)
        summary = dict(run.summary) if run and run.status is CheckStatus.done else {}
        return {
            "reconciled": int(summary.get("reconciled", 0)),
            "agreeing": int(summary.get("agreeing", 0)),
            "drifting": int(summary.get("drifting", 0)),
            "unlinked": int(summary.get("unlinked", 0)),
            "confirmed": int(summary.get("confirmed", 0)),
            "reasons": list(summary.get("reasons", [])),
        }

    async def figure_map(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        dossier_id: UUID,
        artifact_id: UUID,
    ) -> list[tuple[int, list[dict[str, Any]]]]:
        """Every figure in a deck, slide by slide, and what became of it.

        The unlinked ones are the point. They are what the tool did *not*
        check, and the reason each one gives is the difference between an
        honest screen and an impressive one.

        The states come from re-running the same two passes the check runs,
        over the same rows, rather than from reading the stored links. Two
        reasons, and the second is the one that bit:

        1. The reason a figure was skipped is a property of *that check*
           against *that model*, so it is not on the figure's row.
        2. Not every reconciled figure has a link row. A figure the
           Outputs tab publishes but the workbook has no cell for — a CAGR
           computed on the tab itself — is checked and agrees, and there is
           no cell to point a link at. Reading state off the links alone
           called seven such figures « unlinked », which is the exact lie
           this screen exists to prevent.
        """
        repository = TieOutRepository.from_session(session)
        rows = await repository.figures_of(artifact_id)
        engine, back = _figures_of(rows)

        models = [
            one
            for one in await repository.current_artifacts(dossier_id)
            if one.kind is ArtifactKind.model
        ]
        reasons: dict[UUID, str] = {}
        states: dict[UUID, str] = {}
        if models:
            model = models[0]
            cells = await repository.cells_of(model.id)
            candidates = _candidates(_workbook_of(cells), cells, model.outputs)
            proposed, unlinked = _both_passes(engine, candidates)
            _, agreed = compare(proposed)
            for item in agreed:
                states[back[id(item.figure)].id] = "agreeing"
            # A `Drift` carries the slide and the location but not the
            # figure it came from, so drifting is « proposed and not
            # agreed » rather than a second lookup.
            for item in proposed:
                states.setdefault(back[id(item.figure)].id, "drifting")
            for item in unlinked:
                identifier = back[id(item.figure)].id
                states[identifier] = "unlinked"
                reasons[identifier] = item.reason

        links = {
            link.figure_id: link
            for link in await repository.links_of(dossier_id)
            if link.state is not LinkState.rejected
        }

        pages: dict[int, list[dict[str, Any]]] = {}
        for row in rows:
            link = links.get(row.id)
            state = states.get(row.id, "unlinked")
            # A person's decision outranks a score, in both directions.
            if link is not None and link.state is LinkState.confirmed:
                state = "confirmed"
            reason = reasons.get(row.id)
            if state == "unlinked" and reason is None:
                reason = "no model in this deal yet"
            pages.setdefault(row.page, []).append(
                {
                    "id": row.id,
                    "printed": row.printed,
                    "label": row.label,
                    "location": row.location,
                    "state": state,
                    "link_id": link.id if link else None,
                    "reason": reason if state == "unlinked" else None,
                    "anchor": row.anchor or {},
                }
            )
        return sorted(pages.items())

    async def alternatives_for(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        dossier_id: UUID,
        figure_id: UUID,
    ) -> list[dict[str, Any]]:
        """What else this figure could be, best first.

        « Or did you mean this one » — the third action on the confirmation
        queue, and the only one that turns a wrong guess into a right fact
        rather than throwing it away.
        """
        repository = TieOutRepository.from_session(session)
        figure_row = await repository.get_figure(figure_id)
        if figure_row is None:
            return []
        engine, _ = _figures_of([figure_row])

        models = [
            one
            for one in await repository.current_artifacts(dossier_id)
            if one.kind is ArtifactKind.model
        ]
        if not models:
            return []
        model = models[0]
        cells = await repository.cells_of(model.id)
        candidates = _candidates(_workbook_of(cells), cells, model.outputs)

        pool = candidates.published + candidates.workbook
        ranked = rank_outputs(engine[0], pool, limit=6)
        seen: set[str] = set()
        out: list[dict[str, Any]] = []
        for output, score in ranked:
            if output.source in seen:
                continue
            seen.add(output.source)
            cell = candidates.cells.get(output.source) or candidates.cells.get(
                output.ref
            )
            out.append(
                {
                    "cell_id": cell.id if cell else None,
                    "ref": output.source or output.ref,
                    "name": output.name,
                    "value": _text(output.value),
                    "confidence": round(score, 3),
                }
            )
        return out

    async def model_diff(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        dossier_id: UUID,
        artifact_id: UUID,
    ) -> dict[str, Any] | None:
        """What moved since the version before, and what went stale with it.

        The realistic failure is not one typo. It is a model revision the
        deck never caught up with, which is why the count that matters on
        this screen is not « cells changed » but « deck figures now wrong
        because of it ».
        """
        repository = TieOutRepository.from_session(session)
        current = await repository.get_artifact(artifact_id)
        if current is None:
            return None
        previous = await repository.previous_version(current)
        if previous is None:
            return None

        now = {cell.ref: cell for cell in await repository.cells_of(current.id)}
        was = {cell.ref: cell for cell in await repository.cells_of(previous.id)}
        links = await repository.links_of(dossier_id)
        stale: dict[str, int] = {}
        by_id = {cell.id: cell for cell in was.values()}
        for link in links:
            cell = by_id.get(link.cell_id)
            if cell is not None:
                stale[cell.ref] = stale.get(cell.ref, 0) + 1

        changed: list[dict[str, Any]] = []
        for ref, cell in now.items():
            before = was.get(ref)
            if before is None or before.value == cell.value:
                continue
            changed.append(
                {
                    "ref": ref,
                    "name": cell.name or cell.row_label,
                    "was": _text(before.value),
                    "now": _text(cell.value),
                    "stale_figures": stale.get(ref, 0),
                }
            )
        changed.sort(key=lambda one: -one["stale_figures"])
        return {
            "artifact_id": current.id,
            "from_version": previous.version,
            "to_version": current.version,
            "changed": changed,
            "added": len(set(now) - set(was)),
            "removed": len(set(was) - set(now)),
        }

    async def model_grid(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        dossier_id: UUID,
        artifact_id: UUID,
        rows_per_sheet: int = 240,
    ) -> dict[str, Any] | None:
        """The model laid out as it is laid out — sheets, rows, columns.

        Every screen this has had until now asked a person to know already
        what they were looking for: a search box, or a cell reference
        carried by a finding. A banker opening a model opens *a sheet* and
        reads down it, and until that is possible the model is the one
        document in the deal nobody can actually look at.

        **Only cells with a row label.** A number with no words beside it
        cannot be named, cannot be linked, and is not what anyone opens a
        model to read. The same rule the linker uses, for the same reason.

        **`linked` is the column that earns the screen.** Any grid can
        print a workbook back. What this one adds is which cells a
        deliverable is standing on, which is the difference between a
        spreadsheet viewer and a tie-out engine.
        """
        repository = TieOutRepository.from_session(session)
        artifact = await repository.get_artifact(artifact_id)
        if artifact is None:
            return None

        cells = [
            cell for cell in await repository.cells_of(artifact_id) if cell.row_label
        ]
        by_id = {cell.id: cell for cell in cells}
        linked_refs = {
            by_id[link.cell_id].ref
            for link in await repository.links_of(dossier_id)
            if link.state is not LinkState.rejected and link.cell_id in by_id
        }

        #: Sheets in the workbook's own tab order — the order the person
        #: who built the model chose. Recorded at ingest because it cannot
        #: be recovered afterwards: the cells carry a sheet name and
        #: nothing about where that tab sat, so ordering by anything they
        #: hold gives alphabetical (Assumptions before Model, on every
        #: model ever written) or whatever the database happened to
        #: return, which is not stable between two loads of the same page.
        tabs: list[str] = list(artifact.counts.get("sheet_order") or [])
        present = {cell.sheet for cell in cells}
        # A model ingested before the tab order was recorded, and any sheet
        # that appeared since: appended in name order, so the answer is at
        # least the same every time.
        order = [name for name in tabs if name in present]
        order += sorted(present - set(order))

        sheets: list[dict[str, Any]] = []
        for name in order:
            here = [cell for cell in cells if cell.sheet == name]
            #: Columns left to right, keyed by heading rather than index:
            #: two columns headed FY2024A are one column to a reader, and a
            #: model doing that is saying they mean the same thing.
            columns: dict[str, int] = {}
            for cell in sorted(here, key=lambda one: one.column):
                columns.setdefault(cell.column_label, cell.column)

            order = list(columns)
            grouped: dict[str, dict[str, Any]] = {}
            for cell in sorted(here, key=lambda one: (one.row, one.column)):
                row = grouped.setdefault(
                    cell.row_label,
                    {"label": cell.row_label, "at": cell.row, "cells": {}},
                )
                row["cells"].setdefault(
                    cell.column_label,
                    {
                        "ref": cell.ref,
                        "value": _text(cell.value),
                        "display": show(cell.value, cell.number_format),
                        "linked": cell.ref in linked_refs,
                    },
                )

            rows = sorted(grouped.values(), key=lambda one: one["at"])
            sheets.append(
                {
                    "name": name,
                    "columns": order,
                    "rows": [
                        {
                            "label": row["label"],
                            "cells": [row["cells"].get(heading) for heading in order],
                        }
                        for row in rows[:rows_per_sheet]
                    ],
                    "rows_total": len(rows),
                }
            )

        return {
            "artifact_id": artifact.id,
            "filename": artifact.filename,
            "version": artifact.version,
            "uploaded_at": artifact.created_at,
            "sheets": sheets,
        }

    async def chain_of_finding(
        self, session: AsyncSession | AsyncReadSession, *, finding: FindingRow
    ) -> dict[str, Any]:
        """The whole path behind one finding, as steps a screen can render."""
        repository = TieOutRepository.from_session(session)
        evidence = finding.evidence or {}
        steps: list[dict[str, Any]] = []

        if finding.kind is FindingKind.drift:
            steps.append(
                {
                    "kind": "figure",
                    "label": finding.location,
                    "name": finding.detail,
                    "printed": finding.printed,
                }
            )

        model_id = evidence.get("model_artifact_id") or (
            str(finding.artifact_id) if finding.kind is FindingKind.audit else None
        )
        ref = evidence.get("source") or evidence.get("ref") or finding.location
        if model_id and ref:
            cells = await repository.cells_of(UUID(str(model_id)))
            book = _workbook_of(cells)
            steps.extend(_steps_from(book, str(ref), evidence.get("basis")))
            steps.extend(
                await self._grounding(
                    session,
                    dossier_id=finding.dossier_id,
                    cells=cells,
                    ends_at=steps[-1] if steps else None,
                )
            )

        return {
            "finding_id": finding.id,
            "steps": steps,
            "summary": str(evidence.get("chain") or finding.detail or ""),
        }

    async def _grounding(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        dossier_id: UUID,
        cells: Sequence[CellRow],
        ends_at: dict[str, Any] | None,
    ) -> list[dict[str, Any]]:
        """« Audited accounts FY24 · p.42 », when there is one.

        **The step the design draws at the end of every chain, and the
        first one that leaves this deal's own arithmetic.** A chain that
        stops at a typed input has answered « where in the model » and not
        « says who »; this answers the second, or says nothing at all.

        Nothing is inferred. There is a step here only when a source
        document was read, the linker matched it to *this* cell, and — if
        anybody has ruled on that link — they did not reject it. A guess
        at provenance is the worst thing this product could invent: it is
        the one claim a banker would repeat to a client without checking.
        """
        if ends_at is None or ends_at.get("kind") != "input":
            return []
        ref = str(ends_at.get("ref") or "")
        cell = next((one for one in cells if one.ref == ref), None)
        if cell is None:
            return []

        repository = TieOutRepository.from_session(session)
        grounded = await repository.grounding_for(dossier_id, cell.id)
        if grounded is None:
            return []
        link, figure, artifact = grounded

        return [
            {
                "kind": "source",
                "ref": figure.location,
                #: The sentence **as printed**, not the label the matcher
                #: used. Those differ by one deliberate rewrite — accounts
                #: say « 31 December 2025 » and the linker needs
                #: « FY2025A » — and showing a reader a year this product
                #: invented, on the one screen whose job is to say where a
                #: number came from, would be the wrong place to be clever.
                "name": figure.context or figure.label,
                "printed": figure.printed,
                "value": f"{figure.value.normalize():f}",
                # « Audited accounts FY24 · p.42 » — the document and the
                # page, which is the whole of what a person needs to check
                # it themselves.
                "label": f"{artifact.filename} · {figure.location}",
                "basis": link.basis or None,
                "note": (
                    "confirmed"
                    if link.state is LinkState.confirmed
                    else "proposed — nobody has confirmed this yet"
                ),
                "inputs": [],
            }
        ]

    async def chain_for(
        self,
        session: AsyncSession | AsyncReadSession,
        *,
        artifact_id: UUID,
        ref: str,
        dossier_id: UUID | None = None,
    ) -> list[dict[str, Any]]:
        """The steps behind one cell, for the screen that sells the product.

        Given the deal, a typed input carries the document it came out of —
        which is the whole question this screen is asked: *where did this
        number come from.* Without it the honest answer stops one hop
        short, at « somebody typed this ».
        """
        repository = TieOutRepository.from_session(session)
        cells = await repository.cells_of(artifact_id)
        book = _workbook_of(cells)
        cell = book.get(ref)
        if cell is None:
            return []

        steps: list[dict[str, Any]] = [
            {
                "kind": "cell",
                "ref": cell.ref,
                "name": cell.name,
                "value": _text(cell.value),
                "formula": cell.formula,
            }
        ]
        for precedent in cell.precedents[:6]:
            source = book.get(precedent)
            if source is None:
                continue
            steps.append(
                {
                    "kind": "input" if source.formula is None else "cell",
                    "ref": source.ref,
                    "name": source.name,
                    "value": _text(source.value),
                    "formula": source.formula,
                    "note": None if source.formula else "typed, not calculated",
                }
            )

        if dossier_id is not None:
            steps.extend(
                await self._grounding(
                    session,
                    dossier_id=dossier_id,
                    cells=cells,
                    #: The cell asked about, when it is itself typed. A
                    #: precedent that is grounded belongs on the chain of
                    #: *that* cell, one press further on, and putting it
                    #: here would say this figure came off a page it did
                    #: not come off.
                    ends_at={"kind": "input", "ref": cell.ref}
                    if cell.formula is None
                    else None,
                )
            )
        return steps


# --- adapters ------------------------------------------------------------


def _workbook_of(cells: Sequence[CellRow]) -> Workbook:
    """The engine's `Workbook`, rebuilt from stored rows.

    Everything downstream — the audit, the chain, the candidate list —
    takes a `Workbook`, so rebuilding one is what lets a check run without
    the file it came from.
    """
    book = Workbook()
    for row in cells:
        book.cells[row.ref] = EngineCell(
            sheet=row.sheet,
            ref=row.ref,
            row=row.row,
            column=row.column,
            value=row.value,
            formula=row.formula,
            row_label=row.row_label,
            column_label=row.column_label,
            precedents=tuple(row.precedents or ()),
            unresolved=tuple(
                (str(one[0]), str(one[1]))
                for one in (row.unresolved or [])
                if len(one) == 2
            ),
            alias_of=row.alias_of,
        )
    book.sheets = list(dict.fromkeys(row.sheet for row in cells))
    return book


class _Candidates:
    """What a deck figure may be, and the row each candidate came from.

    `outputs_from_workbook` keys each candidate by the cell's own
    reference, so the way back to the stored row is that reference — which
    is what a link has to record, since a link points at a row and not at
    an engine object.
    """

    def __init__(
        self,
        published: list[Output],
        workbook: list[Output],
        cells: dict[str, CellRow],
    ) -> None:
        #: What the model says it publishes. Named by a person, carrying a
        #: stated basis, and preferred wherever both passes fire.
        self.published = published
        #: Every named cell. Wider, and the only pass that works on a model
        #: with no Outputs tab, which is most of them.
        self.workbook = workbook
        self.cells = cells


def _statement_json(statement: Statement) -> dict[str, Any]:
    return {
        "printed": statement.printed,
        "location": statement.location,
        "page": statement.page,
        "section": statement.section,
        "context": statement.context,
    }


def _candidates(
    book: Workbook, cells: Sequence[CellRow], published: list[dict[str, Any]]
) -> _Candidates:
    return _Candidates(
        [
            Output(
                ref=row["ref"],
                name=row["name"],
                value=Decimal(row["value"]),
                source=row["source"],
                basis=row["basis"],
            )
            for row in published or []
        ],
        outputs_from_workbook(book),
        {cell.ref: cell for cell in cells},
    )


def _both_passes(
    figures: list[engine_figures.Figure], candidates: _Candidates
) -> tuple[list[Any], list[Any]]:
    """Link against the published figures, then against the workbook.

    The published pass wins where both fire, because a name a human chose
    and a stated basis read better in a finding than `Model!D26`. A figure
    the first pass reconciled is not offered to the second.
    """
    published, _ = (
        propose_links(figures, candidates.published)
        if candidates.published
        else ([], [])
    )
    settled = {id(item.figure) for item in published}
    remaining = [figure for figure in figures if id(figure) not in settled]
    workbook, unlinked = propose_links(remaining, candidates.workbook)
    return published + workbook, unlinked


def _figures_of(
    rows: Sequence[FigureRow],
) -> tuple[list[engine_figures.Figure], dict[int, FigureRow]]:
    """Engine figures, and the way back to the rows that produced them."""
    engine: list[engine_figures.Figure] = []
    back: dict[int, FigureRow] = {}
    for row in rows:
        figure = engine_figures.Figure(
            printed=row.printed,
            value=row.value,
            decimals=row.decimals,
            kind=row.kind,
            slide=row.page,
            label=row.label,
            location=row.location,
            context=row.context,
            section=row.section,
            range_endpoint=row.range_endpoint,
            parenthesised=row.parenthesised,
            subject=row.subject,
            anchor=row.anchor or {},
        )
        engine.append(figure)
        back[id(figure)] = row
    return engine, back


def _steps_from(
    book: Workbook, ref: str, basis: str | None, depth: int = 3
) -> list[dict[str, Any]]:
    """Walk back from a cell to the typed inputs behind it.

    Stops at the first cell with no formula, because that is the edge of
    the model: a number somebody typed, and the beginning of the next
    question — *where did that come from*. Later the answer is a page in a
    PDF, which is why a step carries a `kind` rather than always being a
    cell.
    """
    steps: list[dict[str, Any]] = []
    seen: set[str] = set()
    current: str | None = ref

    while current and current not in seen and len(steps) < depth:
        seen.add(current)
        cell = book.get(current)
        if cell is None:
            steps.append({"kind": "cell", "ref": current, "name": "", "value": None})
            break

        inputs = []
        for precedent in cell.precedents[:6]:
            source = book.get(precedent)
            if source is None:
                continue
            inputs.append(
                {
                    "ref": source.ref,
                    "name": source.name or source.row_label,
                    "value": _text(source.value),
                }
            )
        steps.append(
            {
                "kind": "cell" if cell.formula else "input",
                "ref": cell.ref,
                "name": cell.name or cell.row_label,
                "value": _text(cell.value),
                "formula": cell.formula,
                "basis": basis if len(steps) == 0 else None,
                "note": None if cell.formula else "typed, not calculated",
                "inputs": inputs,
                # Carried to the screen rather than counted. « Four inputs
                # and one we could not follow » is a different sentence
                # from « four inputs », and a banker deciding whether to
                # trust a number needs the second half of it.
                "unresolved": [why for _, why in cell.unresolved],
            }
        )
        if not cell.formula:
            break
        # Follow the precedent that is itself calculated; a chain that
        # walks into a constant has already ended.
        current = next(
            (
                one
                for one in cell.precedents
                if (nxt := book.get(one)) is not None and nxt.formula
            ),
            None,
        )
    return steps


def _neighbourhood(
    book: Workbook,
    sheet: str,
    ref: str,
    *,
    rows: int = 4,
    columns: int = 4,
    axes: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    """The finding's cell with its neighbours — the design's little
    Excel grid, composed where the cells are in scope so every screen
    (workspace modal, panel, a stored one-off) can draw it without a
    second request. The window is small on purpose: enough to see the
    row break its pattern, never a spreadsheet viewer.

    Shape follows the design's `xl` block: the sheet, the selected
    coordinate, its formula (or its printed value), the column letters,
    and the rows — first cell of each row is the model's own label for
    it, `hot` marks the finding's cell.
    """
    from openpyxl.utils import column_index_from_string, get_column_letter

    coordinate = ref.split("!", 1)[-1]
    match = re.match(r"^([A-Z]+)(\d+)$", coordinate)
    if match is None:
        return None
    column = column_index_from_string(match.group(1))
    row = int(match.group(2))

    per_sheet = [c for c in book.cells.values() if c.sheet == sheet]
    if not per_sheet:
        return None
    first_row = max(1, row - (rows - 2))
    first_column = max(1, column - (columns - 2))
    span_rows = list(range(first_row, first_row + rows))
    span_columns = list(range(first_column, first_column + columns))

    by_place = {(c.row, c.column): c for c in per_sheet}
    labels = {
        c.row: c.row_label for c in per_sheet if c.row in span_rows and c.row_label
    }
    target = by_place.get((row, column))

    def printed(cell: EngineCell | None) -> str:
        if cell is None or cell.value is None:
            return ""
        text = f"{cell.value.normalize():f}"
        try:
            number = float(text)
        except ValueError:
            return text
        #: What Excel shows: digits with thousands separators, never
        #: scientific notation — `,.6g` quietly turns 512,500,000 into
        #: « 5.125e+08 » once `g` runs out of significant figures.
        if abs(number) >= 1e5:
            return f"{number:,.0f}"
        return f"{number:,.6g}"

    #: The model's own year for each column, from the structure layer —
    #: the founder's grid shows « FY2032 », not « N ».
    periods: dict[int, str] = {}
    axis = (axes or {}).get(sheet)
    if axis is not None:
        periods = dict(axis.columns)

    #: The sheet-tab strip, windowed around this sheet, as the design
    #: draws beneath the grid.
    tabs: list[str] = []
    if book.sheets and sheet in book.sheets:
        at = book.sheets.index(sheet)
        start = max(0, min(at - 2, len(book.sheets) - 5))
        tabs = book.sheets[start : start + 5]

    return {
        "sheet": sheet,
        "sel": coordinate,
        "formula": ((target.formula or printed(target)) if target is not None else ""),
        "sheets": tabs,
        "cols": [
            {"l": get_column_letter(c), "p": periods.get(c, "")} for c in span_columns
        ],
        "rows": [
            {
                "n": r,
                "label": labels.get(r, ""),
                "cells": [
                    {
                        "v": printed(by_place.get((r, c))),
                        "hot": r == row and c == column,
                    }
                    for c in span_columns
                ],
            }
            for r in span_rows
        ],
    }


def _fingerprint(kind: str, lineage: UUID, where: str, what: str) -> str:
    """What survives a re-run.

    Keyed on the *lineage* rather than the artifact, so re-uploading the
    deck does not resurrect a finding somebody dismissed on the version
    before. Keyed on where and what rather than on a row id, because the
    row is new every time.
    """
    raw = f"{kind}|{lineage}|{where}|{what}"
    return hashlib.sha256(raw.encode()).hexdigest()[:32]


def _reason(text: str) -> str:
    """Group the linker's reasons so the coverage line stays readable."""
    for prefix in (
        "no output fits the label",
        "two outputs fit equally well",
    ):
        if text.startswith(prefix):
            return prefix
    return text


def _text(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return f"{value.normalize():f}"


tieout = TieOutService()

__all__ = ["TieOutService", "tieout"]
