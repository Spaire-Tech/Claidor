"""The marked-up model — the founder's spec (swens-product-plan.md §4).

Your own model handed back to you, with every problem cell coloured
in and a note stuck on it saying what's wrong, plus a first sheet
listing all findings so you can sort them and tick them off. The
promise that matters: **nothing in the model is altered. Not one
formula, not one number. Only colour and notes.** It's a copy, with
a different filename, so the original is never at risk.

The promise is kept the writer's way — surgery on the zip, every
member the markup does not need copied through byte-identical — and
then *verified* the changeset's way: the copy is re-read beside the
original in an independent library and every cell's formula and
value must match, or the markup refuses itself. Colour is a style
index; a note is a comment part; neither touches content.

What the surgery adds, and nothing else:

- a « Findings » sheet in first position (severity, sheet, cell,
  what's wrong, and blank Notes/Done columns, with an autofilter so
  the list sorts and ticks);
- two solid fills in the style table (Excel's own review red and
  amber) and, per coloured cell, a clone of that cell's existing
  style with only the fill changed — number formats, fonts and
  borders keep their exact style;
- per touched sheet, a legacy comments part and its VML twin, the
  note anchored to the cell, authored « Swens ».

A finding may point at a cell the file never wrote (a skipped cell);
the cell is materialized empty and styled — still not one number. A
sheet that already carries its own comments part is refused in
words: merging into existing comments is not built, and dropping the
notes silently is not an option.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from .changeset import comparable
from .writer import (
    CELL,
    REF,
    _column_index,
    _widen_dimension,
    place_cell,
    sheet_map,
)

#: Excel's own review colours — the light red and amber every model
#: reviewer already knows from conditional formatting.
COLOURS = {"error": "FFFFC7CE", "smell": "FFFFEB9C"}

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

FILLS = re.compile(r'<fills count="(\d+)">(.*?)</fills>', re.DOTALL)
CELL_XFS = re.compile(r'<cellXfs count="(\d+)">(.*?)</cellXfs>', re.DOTALL)
XF = re.compile(r"<xf\b[^>]*(?:/>|>.*?</xf>)", re.DOTALL)
STYLE = re.compile(r'\ss="(\d+)"')


class MarkupRefused(Exception):
    """The markup will not produce a copy it cannot vouch for."""


@dataclass(frozen=True)
class MarkupFinding:
    """One row of the list, one coloured cell, one note."""

    severity: str
    sheet: str
    ref: str
    text: str


def marked_up_name(filename: str) -> str:
    """A different filename, so the original is never at risk."""
    stem, dot, _ = filename.rpartition(".xlsx")
    base = stem if dot else filename
    return f"{base} — marked up.xlsx"


def marked_up_copy(payload: bytes, findings: list[MarkupFinding]) -> bytes:
    """The copy: coloured, noted, listed — and verified unaltered."""
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        order = [item.filename for item in archive.infolist()]
        members = {name: archive.read(name) for name in order}
    sheets = sheet_map(members)

    by_sheet: dict[str, dict[str, list[MarkupFinding]]] = {}
    for finding in findings:
        if finding.sheet not in sheets:
            raise MarkupRefused(f"no sheet named « {finding.sheet} »")
        if REF.fullmatch(finding.ref) is None:
            raise MarkupRefused(
                f"{finding.sheet}!{finding.ref} is not a cell reference"
            )
        by_sheet.setdefault(finding.sheet, {}).setdefault(finding.ref, []).append(
            finding
        )
    for sheet in by_sheet:
        if b"<legacyDrawing " in members[sheets[sheet]]:
            raise MarkupRefused(
                f"« {sheet} » already carries its own comments; merging into "
                f"them is not built yet, and dropping the notes silently is "
                f"not an option"
            )

    style_of = _add_styles(members, by_sheet)
    for sheet, cells in by_sheet.items():
        _colour_cells(members, sheets[sheet], cells, style_of)
        _add_notes(members, sheets[sheet], cells, order)
    listing_name = _add_findings_sheet(members, findings, order, set(sheets))

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name in order:
            archive.writestr(name, members[name])
    copy = out.getvalue()
    _verify_unaltered(payload, copy, listing_name)
    return copy


def _worst(group: list[MarkupFinding]) -> str:
    return "error" if any(f.severity == "error" for f in group) else "smell"


# --- styles: two fills, and a clone of each touched style with only
# --- the fill changed


def _add_styles(
    members: dict[str, bytes],
    by_sheet: dict[str, dict[str, list[MarkupFinding]]],
) -> dict[tuple[int, str], int]:
    """Append the fills and the per-(style, severity) xf clones.

    Returns (original style index, severity) → new style index.
    """
    xml = members["xl/styles.xml"].decode("utf-8")
    fills = FILLS.search(xml)
    xfs = CELL_XFS.search(xml)
    if fills is None or xfs is None:
        raise MarkupRefused("the style table has no fills or cellXfs block")

    fill_count = int(fills.group(1))
    fill_of = {
        severity: fill_count + i for i, severity in enumerate(("error", "smell"))
    }
    new_fills = "".join(
        f'<fill><patternFill patternType="solid">'
        f'<fgColor rgb="{COLOURS[severity]}"/><bgColor indexed="64"/>'
        f"</patternFill></fill>"
        for severity in ("error", "smell")
    )
    xml = xml.replace(
        fills.group(0),
        f'<fills count="{fill_count + 2}">{fills.group(2)}{new_fills}</fills>',
        1,
    )

    existing = XF.findall(xfs.group(2))
    needed: list[tuple[int, str]] = []
    for sheet, cells in by_sheet.items():
        member = sheet_map(members)[sheet]
        sheet_xml = members[member].decode("utf-8")
        found = {m.group(1): m.group(0) for m in CELL.finditer(sheet_xml)}
        for ref, group in cells.items():
            old = found.get(ref, "")
            style = STYLE.search(old.split(">", 1)[0]) if old else None
            key = (int(style.group(1)) if style else 0, _worst(group))
            if key not in needed:
                needed.append(key)

    clones = []
    style_of: dict[tuple[int, str], int] = {}
    for i, (base, severity) in enumerate(needed):
        if base >= len(existing):
            raise MarkupRefused(f"a cell points at style {base}, which is not there")
        clone = existing[base]
        clone = (
            re.sub(r'fillId="\d+"', f'fillId="{fill_of[severity]}"', clone, count=1)
            if 'fillId="' in clone
            else clone.replace("<xf ", f'<xf fillId="{fill_of[severity]}" ', 1)
        )
        clone = (
            re.sub(r'applyFill="[^"]*"', 'applyFill="1"', clone, count=1)
            if 'applyFill="' in clone
            else clone.replace("<xf ", '<xf applyFill="1" ', 1)
        )
        clones.append(clone)
        style_of[(base, severity)] = len(existing) + i

    xml = xml.replace(
        xfs.group(0),
        f'<cellXfs count="{len(existing) + len(clones)}">'
        f"{xfs.group(2)}{''.join(clones)}</cellXfs>",
        1,
    )
    members["xl/styles.xml"] = xml.encode("utf-8")
    return style_of


def _colour_cells(
    members: dict[str, bytes],
    member: str,
    cells: dict[str, list[MarkupFinding]],
    style_of: dict[tuple[int, str], int],
) -> None:
    xml = members[member].decode("utf-8")
    for ref, group in cells.items():
        found = {m.group(1): m.group(0) for m in CELL.finditer(xml)}
        old = found.get(ref)
        if old is None:
            parsed = REF.fullmatch(ref)
            assert parsed is not None  # checked at the door
            column, row_number = (
                _column_index(parsed.group(1)),
                int(parsed.group(2)),
            )
            base_style = style_of[(0, _worst(group))]
            xml = place_cell(
                xml, f'<c r="{ref}" s="{base_style}"/>', column, row_number
            )
            xml = _widen_dimension(xml, column, row_number)
            continue
        head, rest = old.split(">", 1)
        style = STYLE.search(head)
        new_index = style_of[(int(style.group(1)) if style else 0, _worst(group))]
        if style:
            head = head.replace(style.group(0), f' s="{new_index}"', 1)
        else:
            head = head.replace(f'<c r="{ref}"', f'<c r="{ref}" s="{new_index}"', 1)
        xml = xml.replace(old, f"{head}>{rest}", 1)
    members[member] = xml.encode("utf-8")


# --- notes: a legacy comments part and its VML twin per touched sheet


def _add_notes(
    members: dict[str, bytes],
    member: str,
    cells: dict[str, list[MarkupFinding]],
    order: list[str],
) -> None:
    comments_name = _fresh(members, "xl/comments{}.xml")
    vml_name = _fresh(members, "xl/drawings/vmlDrawing{}.vml")

    comment_rows = []
    shapes = []
    for i, (ref, group) in enumerate(sorted(cells.items())):
        text = "\n\n".join(f.text for f in group)
        comment_rows.append(
            f'<comment ref="{ref}" authorId="0"><text><r>'
            f'<t xml:space="preserve">{escape(text)}</t></r></text></comment>'
        )
        parsed = REF.fullmatch(ref)
        assert parsed is not None
        column, row_number = _column_index(parsed.group(1)), int(parsed.group(2))
        shapes.append(
            f'<v:shape id="_x0000_s{1025 + i}" type="#_x0000_t202" '
            f'style="position:absolute;margin-left:80pt;margin-top:2pt;'
            f'width:180pt;height:64pt;z-index:{i + 1};visibility:hidden" '
            f'fillcolor="#ffffe1" o:insetmode="auto">'
            f'<v:fill color2="#ffffe1"/>'
            f'<v:shadow on="t" color="black" obscured="t"/>'
            f'<x:ClientData ObjectType="Note"><x:MoveWithCells/>'
            f"<x:SizeWithCells/><x:AutoFill>False</x:AutoFill>"
            f"<x:Row>{row_number - 1}</x:Row>"
            f"<x:Column>{column - 1}</x:Column></x:ClientData></v:shape>"
        )

    members[comments_name] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<comments xmlns="{MAIN}"><authors><author>Swens</author></authors>'
        f"<commentList>{''.join(comment_rows)}</commentList></comments>"
    ).encode()
    members[vml_name] = (
        f'<xml xmlns:v="urn:schemas-microsoft-com:vml" '
        f'xmlns:o="urn:schemas-microsoft-com:office:office" '
        f'xmlns:x="urn:schemas-microsoft-com:office:excel">'
        f'<o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/>'
        f"</o:shapelayout>"
        f'<v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" '
        f'path="m,l,21600r21600,l21600,xe"><v:stroke joinstyle="miter"/>'
        f'<v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>'
        f"{''.join(shapes)}</xml>"
    ).encode()
    order.extend([comments_name, vml_name])

    rels_name = f"xl/worksheets/_rels/{member.rsplit('/', 1)[1]}.rels"
    comment_rid, vml_rid = _fresh_rids(members.get(rels_name), 2)
    additions = (
        f'<Relationship Id="{comment_rid}" Type="{DOC_RELS}/comments" '
        f'Target="../{comments_name.removeprefix("xl/")}"/>'
        f'<Relationship Id="{vml_rid}" Type="{DOC_RELS}/vmlDrawing" '
        f'Target="../{vml_name.removeprefix("xl/")}"/>'
    )
    if rels_name in members:
        members[rels_name] = (
            members[rels_name]
            .decode("utf-8")
            .replace("</Relationships>", f"{additions}</Relationships>", 1)
            .encode("utf-8")
        )
    else:
        members[rels_name] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<Relationships xmlns="{RELS}">{additions}</Relationships>'
        ).encode()
        order.append(rels_name)

    #: The r prefix is declared on the element itself — openpyxl's
    #: worksheet root does not declare it; Excel's does. Both accept
    #: the local declaration.
    sheet_xml = members[member].decode("utf-8")
    members[member] = sheet_xml.replace(
        "</worksheet>",
        f'<legacyDrawing xmlns:r="{DOC_RELS}" r:id="{vml_rid}"/></worksheet>',
        1,
    ).encode("utf-8")

    _register(
        members,
        "Default",
        "vml",
        "application/vnd.openxmlformats-officedocument.vmlDrawing",
    )
    _register(
        members,
        "Override",
        f"/{comments_name}",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml",
    )


# --- the first sheet: the list


def _add_findings_sheet(
    members: dict[str, bytes],
    findings: list[MarkupFinding],
    order: list[str],
    taken_names: set[str],
) -> str:
    listing_name = "Findings"
    while listing_name in taken_names:
        listing_name = "Swens " + listing_name.lower()

    sheet_member = _fresh(members, "xl/worksheets/sheet{}.xml")
    headers = ["Severity", "Sheet", "Cell", "What's wrong", "Notes", "Done"]
    rows = [
        "".join(
            f'<c r="{chr(65 + i)}1" t="inlineStr"><is><t>{escape(h)}</t></is></c>'
            for i, h in enumerate(headers)
        )
    ]
    for n, f in enumerate(findings, start=2):
        values = [f.severity.capitalize(), f.sheet, f.ref, f.text]
        rows.append(
            "".join(
                f'<c r="{chr(65 + i)}{n}" t="inlineStr">'
                f'<is><t xml:space="preserve">{escape(v)}</t></is></c>'
                for i, v in enumerate(values)
            )
        )
    body = "".join(f'<row r="{n + 1}">{row}</row>' for n, row in enumerate(rows))
    last = len(rows)
    widths = "".join(
        f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>'
        for i, w in ((1, 10), (2, 18), (3, 10), (4, 80), (5, 28), (6, 8))
    )
    members[sheet_member] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{MAIN}" xmlns:r="{DOC_RELS}">'
        f'<dimension ref="A1:F{last}"/>'
        f'<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
        f"<cols>{widths}</cols>"
        f"<sheetData>{body}</sheetData>"
        f'<autoFilter ref="A1:F{last}"/>'
        f"</worksheet>"
    ).encode()
    order.append(sheet_member)

    rels = members["xl/_rels/workbook.xml.rels"].decode("utf-8")
    (rid,) = _fresh_rids(members["xl/_rels/workbook.xml.rels"], 1)
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{rid}" Type="{DOC_RELS}/worksheet" '
        f'Target="{sheet_member.removeprefix("xl/")}"/></Relationships>',
        1,
    )
    members["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    book = members["xl/workbook.xml"].decode("utf-8")
    #: localSheetId is an index into sheet order; a new first sheet
    #: shifts every scoped name by one, so they are all re-pointed.
    book = re.sub(
        r'localSheetId="(\d+)"',
        lambda m: f'localSheetId="{int(m.group(1)) + 1}"',
        book,
    )
    #: The copy opens on the list — that is what it is for.
    book = re.sub(r'activeTab="\d+"', 'activeTab="0"', book)
    ids = [int(m) for m in re.findall(r'sheetId="(\d+)"', book)]
    #: The r prefix is declared on the element itself: Excel declares
    #: it on the workbook root, openpyxl on each sheet element — this
    #: works in both files.
    book = book.replace(
        "<sheets>",
        f'<sheets><sheet xmlns:r="{DOC_RELS}" name="{escape(listing_name)}" '
        f'sheetId="{max(ids) + 1}" r:id="{rid}"/>',
        1,
    )
    members["xl/workbook.xml"] = book.encode("utf-8")

    _register(
        members,
        "Override",
        f"/{sheet_member}",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
    )
    return listing_name


# --- shared small parts


def _fresh(members: dict[str, bytes], pattern: str) -> str:
    n = 1
    while pattern.format(n) in members:
        n += 1
    return pattern.format(n)


def _fresh_rids(rels: bytes | None, count: int) -> list[str]:
    taken = (
        {int(m) for m in re.findall(r'Id="rId(\d+)"', rels.decode("utf-8"))}
        if rels
        else set()
    )
    start = max(taken, default=0) + 1
    return [f"rId{start + i}" for i in range(count)]


def _register(members: dict[str, bytes], kind: str, key: str, content: str) -> None:
    kinds = members["[Content_Types].xml"].decode("utf-8")
    if kind == "Default":
        if f'Extension="{key}"' in kinds:
            return
        addition = f'<Default Extension="{key}" ContentType="{content}"/>'
    else:
        if f'PartName="{key}"' in kinds:
            return
        addition = f'<Override PartName="{key}" ContentType="{content}"/>'
    members["[Content_Types].xml"] = kinds.replace(
        "</Types>", f"{addition}</Types>", 1
    ).encode("utf-8")


# --- the promise, verified


def _verify_unaltered(original: bytes, copy: bytes, listing_name: str) -> None:
    """Not one formula, not one number — or no copy at all."""
    a = load_workbook(io.BytesIO(original))
    b = load_workbook(io.BytesIO(copy))
    if b.sheetnames != [listing_name, *a.sheetnames]:
        raise MarkupRefused(
            "the copy's sheets are not the original's plus the list — refused"
        )
    for name in a.sheetnames:
        held = {
            cell.coordinate: comparable(cell.value)
            for row in a[name].iter_rows()
            for cell in row
            if cell.value is not None
        }
        now = {
            cell.coordinate: comparable(cell.value)
            for row in b[name].iter_rows()
            for cell in row
            if cell.value is not None
        }
        if held != now:
            different = sorted(
                ref for ref in set(held) | set(now) if held.get(ref) != now.get(ref)
            )
            raise MarkupRefused(
                f"the copy altered « {name} » at {', '.join(different[:5])} — "
                f"the promise is broken, so there is no copy"
            )
