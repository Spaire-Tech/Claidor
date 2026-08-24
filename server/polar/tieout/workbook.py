"""Reading a model the way a banker reads it: labels, not coordinates.

`Model!D26` means nothing. « FY2025A adjusted EBITDA » means something,
and the workbook already says so — the words are in column A and row 4,
which is where every financial model in the world puts them. A model is a
grid with its own names written down the side and along the top, and
reading those names turns two hundred anonymous cells into two hundred
named figures.

This matters for two separate reasons.

**The Outputs tab is a link in the chain, not the end of it.** Cascade's
Outputs tab says FY2025A adjusted EBITDA lives at `Model!D25`. It lives at
`Model!D26`; D25 is blank. Four of its twenty-three references are off by
one row, all in the adjustments block, which is what happens when a row is
inserted after the tab was written. The values are right and the pointers
are stale — so a checker that trusts them sends a banker to an empty cell,
and a checker that cannot read the workbook has no way to know.

**Most models have no Outputs tab.** Cascade has one because it was built
well. If a cell can be named from its own labels, then the same matcher
that links a deck figure to an output row links it to a *cell*, and the
Outputs tab becomes an optimisation rather than a requirement.

**Precedents come from the formulas.** `=D16+D24` says adjusted EBITDA is
reported EBITDA plus total adjustments, and openpyxl already ships
Microsoft's own published formula tokenizer, so the arithmetic behind a
figure is readable without a calculation engine and without taking on a
copyleft dependency. What is *not* here is evaluation: Excel stores its
last computed values in the file, and reading them is both sufficient and
honest — the model's own answer, not a re-derivation of it.
"""

import datetime
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from .legacy import read_legacy

#: How many text cells a row must have, outside the label column, before
#: it is read as the header naming the columns. Two is enough to tell a
#: period header — « FY2023A FY2024A FY2025A » — from a stray note, and
#: low enough that a two-year table still qualifies.
HEADER_TEXTS = 2

#: How far down a sheet the header row can be. Models open with a title, a
#: units note and a blank line; past this it is data, not a heading.
HEADER_SEARCH = 12

#: Row labels that name a derivation of the row above rather than a line
#: item of their own — the workbook's version of the same rule the deck
#: reader applies to `% margin`. Indentation says the same thing and is
#: also read, but not every model indents.
DERIVED_ROW = re.compile(
    r"^\s*(?:%\s*)?(?:as\s+%\s+of\s+)?(margin|growth|change|yield|%)\b",
    re.IGNORECASE,
)

#: How many cells a range in a formula is expanded to. `SUM(A1:A20)` is
#: worth knowing cell by cell; `SUM(A1:IV65536)` is sixteen million
#: strings and says nothing the first two hundred do not. The cap bounds
#: the work on a workbook built by somebody who selected whole columns,
#: which is most workbooks in the wild and none of the models seen so far.
MAX_RANGE = 200

#: A whole column — `$M:$M`, `A:C`. Real models are full of them: the
#: 2026 Ofgem distribution model averages over `'Monthly Inflation'!$M:$M`
#: 570 times. Expanded against the sheet's used extent rather than to the
#: format's 1,048,576 rows.
WHOLE_COLUMN = re.compile(r"^\$?(?P<first>[A-Z]{1,3}):\$?(?P<last>[A-Z]{1,3})$")
WHOLE_ROW = re.compile(r"^\$?(?P<first>\d+):\$?(?P<last>\d+)$")

#: A defined name, as it appears where a reference could. Excel's rules:
#: starts with a letter, underscore or backslash, no spaces, not a cell
#: address.
DEFINED_NAME = re.compile(r"^[A-Za-z_\\][A-Za-z0-9_.\\]*$")

#: A cell or range inside a formula, with an optional sheet. Quoted sheet
#: names — 'Cash Flow'!B4 — are the reason this is not a two-line split.
REFERENCE = re.compile(
    r"""
    (?:(?P<sheet>'[^']+'|[A-Za-z0-9_.]+)!)?
    \$?(?P<column>[A-Z]{1,3})\$?(?P<row>\d+)
    (?:
      :
      (?:(?P<sheet2>'[^']+'|[A-Za-z0-9_.]+)!)?
      \$?(?P<column2>[A-Z]{1,3})\$?(?P<row2>\d+)
    )?
    """,
    re.VERBOSE,
)


@dataclass(frozen=True)
class Cell:
    """One numeric cell, and the words the model uses for it."""

    sheet: str
    #: « Model!D26 ».
    ref: str
    row: int
    column: int
    #: The cached value, as Excel last computed it — or None when there
    #: is none. A workbook written by a generator and never opened in
    #: Excel has formulas and no values at all, which makes it invisible
    #: to anything that keys on values and perfectly readable to the
    #: audit, which only needs the formulas.
    value: Decimal | None
    #: The formula, when there is one. `None` for a hardcoded input — which
    #: is itself worth knowing, and is most of the model audit.
    formula: str | None
    #: Column A: « Adjusted EBITDA ».
    row_label: str
    #: Row 4: « FY2025A ». Empty on sheets laid out as label/value pairs.
    column_label: str
    #: The workbook's own format code — `0.0%`, `#,##0.0`, `"$"#,##0`.
    #: Presentation rather than data, and the only thing that can tell a
    #: screen that `0.1222587719` is meant to read `12.2%`. Read here
    #: rather than reconstructed later, because it is in the file and a
    #: guess from the value never can be. See `numbers.show`.
    number_format: str | None = None
    #: The cells this one is computed from, in the order they appear.
    precedents: tuple[str, ...] = ()
    #: The subset of `precedents` read only through a lookup table —
    #: INDEX's first argument. Excel resolves the pick before hunting
    #: circular references, so the cycle hunter must not walk these;
    #: everything else about a precedent still applies to them.
    lookup_reads: tuple[str, ...] = ()
    #: What this cell reads that could not be resolved to a cell, each with
    #: a sentence saying why — an external workbook, a defined name left
    #: pointing at `#REF!`, a table reference. Kept because a chain missing
    #: an input while looking complete is the one failure this product
    #: cannot afford, and because on a real model it is not rare: 11.7% of
    #: formulas in the 2015 Ofgem transmission model were in that state
    #: before names were read.
    unresolved: tuple[tuple[str, str], ...] = ()
    #: The other short words printed on this row, left of the data: a
    #: units column, a licence-condition reference, a mnemonic. Kept apart
    #: from `row_label` on purpose — they say what kind of thing the row is
    #: rather than what it is called, so they belong with the sheet name in
    #: the matcher's `basis` and are weighted below the name there.
    #:
    #: **They can be the only thing that tells two rows apart.** Ofgem's
    #: model holds `Legacy price control adjustments to allowed revenue` on
    #: the transmission owner's sheet and on the system operator's, with
    #: the same words; the direction document prints the same row and
    #: distinguishes them by the licence term beside it, `LAR` against
    #: `SOLAR`. The model has that term too, three columns along.
    row_tags: tuple[str, ...] = ()
    #: Set when the whole formula is a single reference — `=Model!D26`,
    #: `=F13`. Such a cell restates a figure rather than being one, and
    #: Cascade has six of them: the comps bridge pulls adjusted EBITDA from
    #: the Model tab, the DCF pulls free cash flow, and so on. Left in the
    #: candidate set they produce several cells with one name and one
    #: value, which is exactly the ambiguity the matcher refuses to
    #: resolve — so a figure that is a pointer is followed to what it
    #: points at.
    alias_of: str | None = None

    @property
    def name(self) -> str:
        """« FY2025A Adjusted EBITDA » — the same shape as an output row's
        name, deliberately, so one matcher serves both."""
        if not self.column_label:
            return self.row_label
        if self.column_label.lower() in self.row_label.lower():
            return self.row_label
        return f"{self.column_label} {self.row_label}".strip()

    @property
    def hardcoded(self) -> bool:
        """A number typed into a row that is otherwise computed. The single
        most common real defect in a banker's model."""
        return self.formula is None


#: Excel's own error values. `#REF!` and `#NAME?` are always defects —
#: a deleted row, a typo'd function. `#N/A` and `#DIV/0!` are frequently
#: deliberate in a template with empty inputs, which is why the audit
#: grades them differently rather than counting them together.
ERROR_VALUES = frozenset(
    {"#REF!", "#NAME?", "#VALUE!", "#NULL!", "#NUM!", "#N/A", "#DIV/0!"}
)


@dataclass(frozen=True)
class Names:
    """The workbook's defined names and the extent of each sheet.

    **Scope is not decoration.** Excel lets the same name mean different
    cells on different sheets, and sheet scope beats workbook scope. The
    2026 Ofgem distribution model carries 594 sheet-scoped names beside
    791 workbook-scoped ones, so a single name-to-reference map would not
    merely miss things — it would resolve a name to the wrong cell and
    report the answer with confidence. That is worse than the silent drop
    it replaced, which is why the two maps are kept apart.
    """

    #: name → what it points at, for names visible everywhere.
    book: dict[str, str] = field(default_factory=dict)
    #: (sheet, name) → what it points at. Wins over `book`.
    sheet: dict[tuple[str, str], str] = field(default_factory=dict)
    #: sheet → (last row, last column) actually used. What a whole-column
    #: reference is expanded against.
    extent: dict[str, tuple[int, int]] = field(default_factory=dict)

    def lookup(self, name: str, sheet: str) -> str | None:
        return self.sheet.get((sheet, name)) or self.book.get(name)


@dataclass(frozen=True)
class Precedents:
    """What a formula reads, and what could not be read.

    Both halves, always. A formula that loses one of its precedents
    produces a chain short by exactly the input that decided the answer
    and looks complete — measured at 11.7% of formulas on a real Ofgem
    transmission model, where the dropped name was the switch selecting
    which company the whole model was calculating for. Rule 3: what was
    not checked is part of the answer.
    """

    refs: tuple[str, ...] = ()
    #: (what it said, why it could not be resolved). Phrased for a reader.
    unresolved: tuple[tuple[str, str], ...] = ()
    #: The subset of `refs` read *only* through a lookup table — INDEX's
    #: first argument — where Excel resolves the pick before hunting
    #: circular references. Real precedents for flows and coverage;
    #: edges the cycle hunter must not walk.
    via_lookup: tuple[str, ...] = ()

    def __bool__(self) -> bool:
        return bool(self.refs or self.unresolved)


@dataclass
class Workbook:
    cells: dict[str, Cell] = field(default_factory=dict)
    sheets: list[str] = field(default_factory=list)
    #: Cells whose cached value is one of Excel's error values.
    errors: dict[str, str] = field(default_factory=dict)
    #: Sheets Excel is hiding. Concealment is a repeated cause in every
    #: published catalogue of spreadsheet disasters, so it is recorded
    #: even though hiding a working sheet is often perfectly innocent.
    hidden_sheets: tuple[str, ...] = ()
    #: The subset set to *very hidden* — absent from Excel's own unhide
    #: menu, reachable only through the VBA editor. A different fact
    #: from hidden, and reported as one.
    very_hidden_sheets: tuple[str, ...] = ()
    #: True when the workbook has iterative calculation switched on, which
    #: is a model saying its circular references are deliberate.
    iterative: bool = False
    #: Raw populated cells per sheet, counted at read time — `cells`
    #: holds only what the reader could name, and « is this sheet
    #: empty » must be answered from the file, not from what survived
    #: labelling. Absent on hand-built books.
    populated: dict[str, int] = field(default_factory=dict)
    #: Every row's label text, per sheet — including rows that carry
    #: no numeric cell of their own, which is exactly where a block
    #: header lives (« Asset beta at 0.075 debt beta » two rows above
    #: the formulas it documents). The audit reads it to honour
    #: numbers the sheet's own words already state.
    row_words: dict[str, dict[int, str]] = field(default_factory=dict)
    #: Cells whose formula the tokenizer rejected. One malformed
    #: formula must cost one finding, never the whole workbook —
    #: Round 4's planting run found the reader dying on a
    #: partial-range #REF! and taking the file with it.
    unparseable: list[str] = field(default_factory=list)
    #: Defined names whose target was deleted (`#REF!`) and names that
    #: point into other workbooks (`'[2]Control Panel'!$D$144`). No
    #: live formula needs to use them: they still raise Excel's update
    #: prompts, and a new formula written against a broken name breaks
    #: on arrival — a judged model carried dozens of both.
    broken_names: list[str] = field(default_factory=list)
    foreign_names: list[tuple[str, str]] = field(default_factory=list)

    def get(self, ref: str) -> Cell | None:
        return self.cells.get(ref)

    def named(self, name: str) -> list[Cell]:
        wanted = name.strip().lower()
        return [cell for cell in self.cells.values() if cell.name.lower() == wanted]


@dataclass
class _Grid:
    """One sheet's cells, in plain dictionaries.

    **openpyxl's ordinary mode is what made a real model take half an
    hour.** `load_workbook` without `read_only` builds a Python object
    for every cell in every sheet's rectangle — on a company financial
    model from the 2024 water price review that is nine million objects,
    twice over, since the formulas and the values are two loads of the
    file. Profiled at 150 seconds in, the reader had finished seven of
    the sixty-four sheets, and most of the time was openpyxl
    materialising empty cells so this module could ask them questions
    they answer with `None`.

    So each sheet is now **streamed once per load** (`read_only=True`)
    into dictionaries holding only the cells that carry something, and
    every later pass — labels, headers, tags, the numeric sweep — asks
    the dictionaries. The extent comes from the cells actually present
    rather than from the file's declared dimension, which real files
    misstate.
    """

    #: Cached values, from the `data_only` load.
    values: dict[tuple[int, int], Any] = field(default_factory=dict)
    #: What is written in the cell: a formula string or the typed value.
    written: dict[tuple[int, int], Any] = field(default_factory=dict)
    #: Number formats, for the cells that have written content.
    formats: dict[tuple[int, int], str | None] = field(default_factory=dict)
    #: Columns present per row, across both loads, for the numeric sweep.
    by_row: dict[int, set[int]] = field(default_factory=dict)
    last_row: int = 0
    last_column: int = 0

    def _saw(self, row: int, column: int) -> None:
        self.by_row.setdefault(row, set()).add(column)
        if row > self.last_row:
            self.last_row = row
        if column > self.last_column:
            self.last_column = column


def _grid_of(written_sheet: Any, values_sheet: Any) -> _Grid:
    """One sheet from both loads, streamed into a :class:`_Grid`."""
    grid = _Grid()
    cells = getattr(written_sheet, "cells", None)
    if isinstance(cells, dict):
        # A legacy `.xls` sheet already is a dictionary of cells.
        for (row, column), value in cells.items():
            if value is None:
                continue
            grid.written[(row, column)] = value
            grid._saw(row, column)
        for (row, column), value in getattr(values_sheet, "cells", {}).items():
            if value is None:
                continue
            grid.values[(row, column)] = value
            grid._saw(row, column)
        return grid

    for row in written_sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            at = (cell.row, cell.column)
            grid.written[at] = cell.value
            grid.formats[at] = getattr(cell, "number_format", None)
            grid._saw(*at)
    for row in values_sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            grid.values[(cell.row, cell.column)] = cell.value
            grid._saw(cell.row, cell.column)
    return grid


def read_workbook(path: str) -> Workbook:
    """Every numeric cell in the model, named and with its precedents.

    Reads `.xlsx` through openpyxl and `.xls` through
    :mod:`polar.tieout.legacy`, which presents the same surface. Which
    one a model is in decides nothing else: the labels, the precedents
    and every rule downstream are the same either way, and they have to
    be, because the corpora with real spreadsheets in them are all the
    old format and the models a bank sends are all the new one.
    """
    if path.lower().endswith((".xls", ".xlt")):
        formulas, values = read_legacy(path)
        close = None
    else:
        formulas = load_workbook(path, data_only=False, read_only=True)
        values = load_workbook(path, data_only=True, read_only=True)
        close = (formulas, values)

    try:
        book = Workbook(
            sheets=list(formulas.sheetnames),
            hidden_sheets=tuple(
                name
                for name in formulas.sheetnames
                if getattr(formulas[name], "sheet_state", "visible") != "visible"
            ),
            very_hidden_sheets=tuple(
                name
                for name in formulas.sheetnames
                if getattr(formulas[name], "sheet_state", "visible") == "veryHidden"
            ),
            iterative=bool(getattr(formulas.calculation, "iterate", False)),
        )
        grids: dict[str, _Grid] = {}
        for name in formulas.sheetnames:
            sheet = formulas[name]
            # A chart sheet is a sheet in the file format and a picture to
            # everybody else: no cells, no grid, no `max_row`. Real models
            # have them and the Cascade test pair does not, which is the
            # sort of thing only a real model tells you.
            if not hasattr(sheet, "max_row"):
                continue
            grids[name] = _grid_of(sheet, values[name])
            book.populated[name] = len(grids[name].written)
        names = _names_of(formulas, grids)
        every_name = list(names.book.items()) + [
            (scoped, target) for (_, scoped), target in names.sheet.items()
        ]
        book.broken_names = sorted(
            {name for name, target in every_name if "#REF!" in target}
        )
        book.foreign_names = sorted(
            {
                (name, target)
                for name, target in every_name
                if re.search(r"\[\d+\]", target)
            }
        )
        for name, grid in grids.items():
            _read_sheet(book, name, grid, names)
        return book
    finally:
        # A read-only load keeps the archive open behind the streaming
        # readers, and an unclosed one holds the temporary upload's file
        # handle for as long as the objects live.
        if close is not None:
            for one in close:
                one.close()


def _names_of(formulas: Any, grids: dict[str, "_Grid"]) -> Names:
    """Defined names and sheet extents, in one pass.

    Read from the *formula* workbook. A legacy `.xls` read through
    :mod:`polar.tieout.legacy` has no defined names to give, and asks for
    nothing: the maps come back empty and every name falls through to
    `unresolved` with a sentence saying so, which is the honest state
    rather than a silent one.

    Extents come from the grids — the cells actually present — because a
    file's declared dimension is a claim, and whole-column expansion
    against an overstated one is exactly the waste the cap exists for.
    """
    names = Names()
    for name, one in (getattr(formulas, "defined_names", {}) or {}).items():
        names.book[str(name)] = str(getattr(one, "value", one))
    for title in formulas.sheetnames:
        sheet = formulas[title]
        for name, one in (getattr(sheet, "defined_names", {}) or {}).items():
            names.sheet[(title, str(name))] = str(getattr(one, "value", one))
        grid = grids.get(title)
        if grid is not None:
            names.extent[title] = (grid.last_row, grid.last_column)
    return names


def _read_sheet(
    book: Workbook, name: str, grid: _Grid, names: Names | None = None
) -> None:
    last_row = grid.last_row
    last_column = grid.last_column
    header_row = _header_row(grid, last_column)
    label_column = _label_column(grid, last_row, last_column)

    #: The last row label that named a line item rather than a derivation
    #: of one, so « % growth » reads as « growth of that ».
    subject = ""
    labels: dict[int, str] = {}
    for row in range(1, last_row + 1):
        named = _shown(grid, row, label_column)
        text = named.strip() if named else ""
        if not text:
            continue
        # Indentation is how a model marks a derived row, and it is a
        # property of the words shown rather than of the formula that
        # produced them.
        indented = named is not None and named[:1].isspace()
        derived = DERIVED_ROW.match(text) is not None or (indented and subject)
        if derived and subject and subject.lower() not in text.lower():
            labels[row] = f"{subject} {text}"
        elif derived and subject:
            labels[row] = text
        else:
            subject = text
            labels[row] = text

    book.row_words[name] = labels

    #: The short descriptors printed beside the name, per row. Long text is
    #: left out: columns C and D of a regulator's model carry whole
    #: paragraphs beginning « Note: », and a paragraph is commentary rather
    #: than a descriptor of the row.
    #:
    #: Only for rows that have a name. A descriptor describes something,
    #: and a row with nothing to describe has none — which on a 22,004-row
    #: allocations table is most of the sheet.
    tags: dict[int, tuple[str, ...]] = {}
    for row in labels:
        if row == header_row:
            continue
        found = []
        for column in range(1, min(last_column, LABEL_COLUMNS) + 1):
            if column == label_column:
                continue
            beside = _shown(grid, row, column)
            beside = beside.strip() if beside else ""
            if beside and len(beside) <= TAG_LENGTH:
                found.append(beside)
        if found:
            tags[row] = tuple(found)

    headers: dict[int, str] = {}
    if header_row is not None:
        for column in range(1, last_column + 1):
            named = _shown(grid, header_row, column)
            if named:
                headers[column] = named.strip()

    for (row, column), shown in grid.values.items():
        if isinstance(shown, str) and shown.strip() in ERROR_VALUES:
            at = f"{name}!{get_column_letter(column)}{row}"
            book.errors[at] = shown.strip()

    for row in sorted(grid.by_row):
        columns = grid.by_row[row]
        numeric = [
            column
            for column in sorted(columns)
            if column != label_column
            and (
                _decimal(grid.values.get((row, column))) is not None
                or _formula(grid.written.get((row, column))) is not None
            )
        ]
        if row == header_row:
            #: The header row names periods, and its cells are words the
            #: audit must not read as content — except a formula that
            #: reaches other rows. A year walker (`=C7+1`) is part of
            #: the header; a warning banner testing the cash row lives
            #: wherever its author parked it, and a judged model parked
            #: one on the header row, where the old whole-row skip made
            #: the engine blind to a coverage gap the banner carried.
            def _reaches_out(column: int) -> bool:
                formula = _formula(grid.written.get((row, column)))
                if formula is None:
                    return False
                for target in references_of(formula, name, names).refs:
                    digits = re.search(r"\d+", target.rsplit("!", 1)[-1])
                    if digits and int(digits.group(0)) != row:
                        return True
                return False

            numeric = [column for column in numeric if _reaches_out(column)]
        # A row carrying one number is a label and a value — « Enterprise
        # value | 489.5 » in a valuation bridge. A row carrying several is
        # a series, and only then does the header above a column name
        # anything about it. Without this, every figure in the bridge under
        # a DCF's forecast grid inherits « FY2026E » and claims to be about
        # a year it has nothing to do with.
        series = len(numeric) > 1

        for column in numeric:
            number = _decimal(grid.values.get((row, column)))
            formula = _formula(grid.written.get((row, column)))
            read = references_of(formula, name, names) if formula else Precedents()
            references = read.refs
            ref = f"{name}!{get_column_letter(column)}{row}"
            if any(
                why == "the formula could not be parsed" for _, why in read.unresolved
            ):
                book.unparseable.append(ref)
            book.cells[ref] = Cell(
                sheet=name,
                ref=ref,
                row=row,
                column=column,
                value=number,
                formula=formula,
                row_label=labels.get(row, ""),
                row_tags=tags.get(row, ()),
                column_label=headers.get(column, "") if series else "",
                # Off the *formula* load: the value load is `data_only`,
                # and a legacy `.xls` read through xlrd has no format on
                # its cells at all, which is why this may be absent.
                number_format=grid.formats.get((row, column)),
                precedents=references,
                lookup_reads=read.via_lookup,
                unresolved=read.unresolved,
                alias_of=_alias(formula, references),
            )


def precedents_of(
    formula: str, sheet: str, names: Names | None = None
) -> tuple[str, ...]:
    """The cells a formula reads, expanded from ranges, in order."""
    return references_of(formula, sheet, names).refs


# Not every reference argument is a read. `ROW(A1)` is 1 whatever A1
# holds — the argument names a place, and the function answers a
# question about the place, never about the value in it. An edge from
# such an argument is a false precedent, and with the dependency graph
# feeding cycle detection a false edge can invent a loop that is not
# there — `=CELL("filename",$A$1)` sitting in A1 read as A1 depending
# on itself. These functions' reference arguments never become edges.
METADATA_ARGS = frozenset(
    {"SHEET", "SHEETS", "ISREF", "ROW", "COLUMN", "ROWS", "COLUMNS", "AREAS"}
)

# CELL is the split case: the first argument decides whether the second
# is read. These info types ask about the cell's location or dressing —
# never its value — so the reference stays out of the graph. The other
# three ("contents", "type", "prefix") do look at what the cell holds,
# and an unrecognisable first argument is treated as if it might, which
# keeps the edge; a doubtful edge is a smaller lie than a missing one.
LOCATION_INFO = frozenset(
    {
        "filename",
        "address",
        "row",
        "col",
        "width",
        "format",
        "color",
        "protect",
        "parentheses",
    }
)

# The two whose real read happens at run time. The references they show
# are genuine dependencies — `INDIRECT("S"&D4)` reads D4 to build the
# string, OFFSET's anchor decides where the walk starts — so those stay
# edges. What neither can show a static read is where the walk *lands*,
# and pretending otherwise is how a builder gets these exactly half
# right. The landing place is declared unfollowable instead.
RUNTIME_TARGET = {
    "INDIRECT": "a reference assembled while the model runs, which a read of the file cannot follow",
    "OFFSET": "a range measured out at run time from the anchor it names",
}

# INDEX's first argument is a table the function picks *one* cell out
# of at run time, and Excel resolves the pick before it hunts circular
# references — which is why `=INDEX(A1:A10,5)` written inside its own
# table calculates instead of warning, and why modellers reach for
# INDEX instead of OFFSET to break a deliberate cycle without
# volatility. The whole table stays a precedent — changing any cell of
# it can change the answer, and flows and coverage are right to say so
# — but the cycle hunter must not walk it: two shipped regulator
# models (CAA's H7 final determination, Ofgem's GT3 business-plan
# model) carry chains that close *only* through INDEX tables, both
# with iterative calculation off, and both calculate cleanly in Excel.
# A loop invented from edges Excel does not walk is not a finding.
LOOKUP_TABLE = frozenset({"INDEX"})


def references_of(formula: str, sheet: str, names: Names | None = None) -> Precedents:
    """Everything a formula reads, and everything it reads that we cannot.

    Parsed through openpyxl's bundled tokenizer, which is a port of the
    grammar Microsoft published, so `SUM(D20:D23)` and `'Cash Flow'!B4`
    and a string literal containing a colon are all told apart properly
    rather than by a regex over the whole formula.

    Without `names` this resolves plain references only, which is what it
    did before there was anything else. With them it resolves defined
    names in scope and expands whole-column and whole-row references
    against the sheet's real extent.

    References are read in the context of the function holding them,
    against the tables above — the same operand is a precedent inside
    SUM and not one inside ROW. Nesting is honoured through a frame
    stack: in `SUM((A1), ROW(C3))` the parenthesised A1 still belongs
    to SUM and C3 to ROW.
    """
    found: list[str] = []
    seen: set[str] = set()
    missing: list[tuple[str, str]] = []
    said: set[str] = set()
    #: Which side of the lookup line each reference was met on. A cell
    #: read both inside an INDEX table and plainly is a plain read.
    table_reads: set[str] = set()
    plain_reads: set[str] = set()

    def keep(refs: list[str], in_table: bool = False) -> None:
        (table_reads if in_table else plain_reads).update(refs)
        for ref in refs:
            if ref not in seen:
                seen.add(ref)
                found.append(ref)

    def give_up(text: str, why: str) -> None:
        if text not in said:
            said.add(text)
            missing.append((text, why))

    # One frame per open call or parenthesis. A parenthesis frame has no
    # name and defers to the nearest named frame around it, so grouping
    # an argument does not change whose argument it is.
    frames: list[dict[str, Any]] = []

    def enclosing() -> dict[str, Any] | None:
        for frame in reversed(frames):
            if frame["name"]:
                return frame
        return None

    try:
        tokens = Tokenizer(formula).items
    except Exception:
        #: The grammar rejected the whole formula. Report the fact and
        #: keep reading the rest of the workbook.
        give_up(formula[:80], "the formula could not be parsed")
        return Precedents(unresolved=tuple(missing))

    for token in tokens:
        for frame in frames:
            frame["call"].append(token.value)
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper().removeprefix("_XLFN.")
            frames.append({"name": name, "arg": 0, "info": None, "call": [token.value]})
            continue
        if token.type == "PAREN" and token.subtype == "OPEN":
            frames.append({"name": "", "arg": 0, "info": None, "call": []})
            continue
        if token.type in ("FUNC", "PAREN") and token.subtype == "CLOSE":
            if frames:
                closed = frames.pop()
                if closed["name"] in RUNTIME_TARGET:
                    give_up("".join(closed["call"]), RUNTIME_TARGET[closed["name"]])
            continue
        if token.type == "SEP" and token.subtype == "ARG" and frames:
            frames[-1]["arg"] += 1
            continue
        if token.type != "OPERAND":
            continue

        here = frames[-1] if frames else None
        if here is not None and here["name"] == "CELL" and here["arg"] == 0:
            # The first argument, when it is one clean string literal,
            # is the info type; anything else leaves it unknown and the
            # unknown case keeps its edges.
            if here["info"] is None and token.subtype == "TEXT":
                here["info"] = token.value.strip('"').lower()
            else:
                here["info"] = "?"

        if token.subtype != "RANGE":
            continue
        owner = enclosing()
        if owner is not None:
            if owner["name"] in METADATA_ARGS:
                continue
            if owner["name"] == "CELL" and owner["info"] in LOCATION_INFO:
                continue
        text = token.value.strip()
        refs = _expand(text, sheet, names, give_up)
        if refs:
            keep(
                refs,
                in_table=owner is not None
                and owner["name"] in LOOKUP_TABLE
                and owner["arg"] == 0,
            )

    return Precedents(
        refs=tuple(found),
        unresolved=tuple(missing),
        via_lookup=tuple(
            ref for ref in found if ref in table_reads and ref not in plain_reads
        ),
    )


def _expand(
    text: str, sheet: str, names: Names | None, give_up: Any, depth: int = 0
) -> list[str]:
    """One operand, as the cells it stands for."""
    where, rest = _split_sheet(text)

    # A reference into another workbook. The path is in the file and the
    # cells are not, so this can never resolve here — but saying so is the
    # point: an external link is the one precedent most likely to be stale,
    # because nobody re-opens the workbook it points at.
    if where.strip("'").startswith("[") or rest.startswith("["):
        if "]" in rest and not rest.startswith("["):
            give_up(text, "a table reference, which this does not read yet")
        else:
            give_up(text, "in another workbook, which is not in this deal")
        return []

    if "[" in rest:
        give_up(text, "a table reference, which this does not read yet")
        return []

    match = REFERENCE.fullmatch(text)
    if match is not None:
        return _capped(
            text,
            _cells(
                (match.group("sheet") or sheet).strip("'"),
                match.group("column"),
                int(match.group("row")),
                match.group("column2"),
                int(match.group("row2") or 0),
            ),
            give_up,
        )

    here = where.strip("'") or sheet

    column = WHOLE_COLUMN.match(rest)
    if column:
        rows = names.extent.get(here, (0, 0))[0] if names else 0
        if not rows:
            give_up(text, "a whole column, and that sheet is not in this workbook")
            return []
        return _capped(
            text,
            _cells(here, column.group("first"), 1, column.group("last"), rows),
            give_up,
        )

    row = WHOLE_ROW.match(rest)
    if row:
        columns = names.extent.get(here, (0, 0))[1] if names else 0
        if not columns:
            give_up(text, "a whole row, and that sheet is not in this workbook")
            return []
        return _capped(
            text,
            _cells(
                here,
                "A",
                int(row.group("first")),
                get_column_letter(max(columns, 1)),
                int(row.group("last")),
            ),
            give_up,
        )

    if DEFINED_NAME.fullmatch(rest) and names is not None:
        points_at = names.lookup(rest, sheet)
        if points_at is None:
            give_up(text, "a name this workbook does not define")
            return []
        if points_at.startswith("#"):
            # A defined name left pointing at a deleted row. Nine of them
            # sit in a published Ofgem model. Not a parse failure — a
            # defect in the workbook, and worth saying so.
            give_up(text, f"a defined name pointing at {points_at}")
            return []
        if depth >= 2:
            give_up(text, "a name defined in terms of other names")
            return []
        # A name can point at anything, including another name or a
        # formula. Following it once more covers the real cases and the
        # depth cap stops a name defined in terms of itself.
        #
        # The recursion reports nothing of its own. What a person needs to
        # be told is « Switch could not be followed », not
        # « OFFSET(Model!$A$1,1,1) is a reference this does not understand » —
        # the second is this module talking about itself.
        resolved = _expand(points_at.lstrip("="), sheet, names, _quiet, depth + 1)
        if not resolved:
            give_up(text, "a name that is a formula rather than a cell")
        return resolved

    if DEFINED_NAME.fullmatch(rest):
        give_up(text, "a defined name, and this workbook's names were not read")
        return []

    give_up(text, "a reference this does not understand")
    return []


def _quiet(text: str, why: str) -> None:
    """Swallow a reason raised inside a name's own definition."""


def _capped(text: str, answer: tuple[list[str], int], give_up: Any) -> list[str]:
    """Say so when the cap bit."""
    refs, span = answer
    if span > len(refs):
        give_up(
            text,
            f"{span:,} cells, of which the first {len(refs):,} were followed",
        )
    return refs


def _split_sheet(token: str) -> tuple[str, str]:
    """« 'Live Results (SO)'!C15 » → the sheet and the rest.

    A quoted sheet name may contain brackets, spaces, even exclamation
    marks. Splitting naively is how `'Live Results (SO)'!$C$15` gets
    called a table reference, which is a mistake this module's own
    measurement script made first.
    """
    if token.startswith("'"):
        end = token.find("'", 1)
        while end != -1 and token[end : end + 2] == "''":
            end = token.find("'", end + 2)
        if end != -1 and token[end + 1 : end + 2] == "!":
            return token[: end + 1], token[end + 2 :]
    if "!" in token:
        where, _, rest = token.rpartition("!")
        return where, rest
    return "", token


def _cells(
    where: str, column: str, row: int, column2: str | None, row2: int
) -> tuple[list[str], int]:
    """The cells a reference stands for, and how many there were in all.

    The second number is not decoration. `MAX_RANGE` exists so that
    `SUM(A1:IV65536)` does not become sixteen million strings, and a cap
    that truncates without saying so is the same silent loss this module
    was just fixed for — moved from « dropped » to « quietly shortened ».
    """
    if not row2:
        return [f"{where}!{column}{row}"], 1
    first_column = _column_index(column)
    last_column = _column_index(column2 or "")
    span = (row2 - row + 1) * max(last_column - first_column + 1, 0)
    refs: list[str] = []
    for one in range(row, row2 + 1):
        for two in range(first_column, last_column + 1):
            if len(refs) >= MAX_RANGE:
                return refs, span
            refs.append(f"{where}!{get_column_letter(two)}{one}")
    return refs, span


def _alias(formula: str | None, references: tuple[str, ...]) -> str | None:
    """The cell this one merely restates, if that is all it does.

    `=Model!D26` and `=F13` are pointers. `=B18*B19` is a figure. The test
    is that stripping the one reference out of the formula leaves nothing
    but the equals sign — no arithmetic, no function, no second operand.
    """
    if formula is None or len(references) != 1:
        return None
    match = REFERENCE.fullmatch(formula[1:].strip())
    if match is not None and match.group("row2") is None:
        return references[0]
    return None


def _column_index(letters: str) -> int:
    index = 0
    for letter in letters:
        index = index * 26 + (ord(letter.upper()) - 64)
    return index


def _shown(grid: _Grid, row: int, column: int) -> str | None:
    """The words a cell shows, whether they were typed or computed.

    **A label can be a formula, and refusing every formula loses a whole
    entity.** Ofgem's price control model keeps one sheet per licensed
    business and builds each from the input sheet, so the name beside every
    row of `NGET TO` is `=Input!E31` rather than words. :func:`_label`
    declines a formula for a good reason — a column of arithmetic would
    otherwise name every figure beside it after the arithmetic — but the
    reason is about the *formula text*, not about what the formula
    produces. Measured on that model: 3,694 typed inputs found, and not one
    of them on the transmission-owner sheet, so every figure the direction
    document states for the transmission owner matched a cell belonging to
    the system operator instead. Seven false contradictions.

    So the formula text is still refused, and the **cached value** is used
    when the formula produced words. A formula that produced a number
    produces no label, which is the case `_label` was protecting against
    and is unaffected: numbers are not strings.
    """
    text = _label(grid.written.get((row, column)))
    if text is not None:
        return text
    value = grid.values.get((row, column))
    if isinstance(value, str):
        return value
    return _period_label(value)


def _period_label(value: Any) -> str | None:
    """A date cell read as the period it heads.

    **A model's column headers are very often real dates.** Ofgem's writes
    `2017-03-31` where a banker's writes `FY2017A`, and a date is not a
    string, so the header row came back empty and every one of the eight
    year columns on a row carried the same name. A figure naming the row
    and not the year then matched whichever column happened to hold a typed
    value — reported as the document contradicting the model, twice, on the
    first real pair.

    `FY` plus the calendar year, because that is the spelling the matcher
    already understands. **It is a convention and it can be wrong**: a
    31 March 2017 year end reads `FY2017` here, which is the British
    convention and not the American one, and a company whose year ends in
    January would call the same date `FY2016`. No basis letter is added —
    a date says when, not whether the number is an actual or an estimate,
    and `_same_period` treats an unmarked year as compatible with either.
    """
    year = getattr(value, "year", None)
    if not isinstance(year, int) or not 1900 <= year <= 2200:
        return None
    return f"FY{year}"


def _header_row(grid: _Grid, last_column: int) -> int | None:
    """The row whose text names the columns. Usually the period header."""
    best: tuple[int, int] | None = None
    for row in range(1, min(grid.last_row, HEADER_SEARCH) + 1):
        texts = sum(
            1 for column in range(2, last_column + 1) if _shown(grid, row, column)
        )
        if texts >= HEADER_TEXTS and (best is None or texts > best[1]):
            best = (row, texts)
    return best[0] if best else None


#: How far right the row names may sit. Four was enough for every model
#: this had seen and not for the one it had not: Ofgem's price control
#: financial model indents through columns B, C and D for section and
#: sub-section headings and puts the parameter names in **column E**. With
#: the search stopping at D, all 26,392 cells in that workbook came back
#: with no name at all — which silently empties the tie-out and the
#: grounding both, since each needs a named cell to have anything to
#: match. Eight covers that layout with room, and is still far to the left
#: of any data column.
LABEL_COLUMNS = 8

#: Longer than this and a cell beside the name is commentary, not a
#: descriptor. Ofgem's model puts whole `Note: …` paragraphs in the two
#: columns left of its parameter names.
TAG_LENGTH = 40

#: How many rows are read to decide *which* column holds the names.
#:
#: **Which column names the rows is a fact about a sheet's layout, and a
#: sheet does not change its layout half way down.** Reading all of a
#: 22,004-row allocations table to choose between eight columns is 352,000
#: cell reads that cannot change the answer the first few hundred gave —
#: and every one of them makes `openpyxl` materialise a cell object that
#: was not in the file. Found by a published schools funding workbook that
#: took tens of minutes to read.
LABEL_SCAN = 1000


def _label_column(grid: _Grid, last_row: int, last_column: int) -> int:
    """The column holding the row names.

    **Chosen by how many *different* things a column says, not how much it
    says.** The obvious rule — the column with the most text in it — picks
    the units column on a regulator's model: `£m 09/10 prices` appears on
    247 rows of one sheet against 251 parameter names beside it, so the
    two are indistinguishable by volume. They are not remotely
    indistinguishable by variety: 10 distinct values against 143. A column
    of row names is nearly all distinct by definition, because that is what
    naming a row is for.

    Ties go left, which keeps column A for the ordinary model that has its
    labels there and a `Notes` column somewhere off to the right.
    """
    best, most = 1, -1
    depth = min(last_row, LABEL_SCAN)
    for column in range(1, min(last_column, LABEL_COLUMNS) + 1):
        seen: set[str] = set()
        for row in range(1, depth + 1):
            text = _shown(grid, row, column)
            if text and text.strip():
                seen.add(text.strip())
        if len(seen) > most:
            best, most = column, len(seen)
    return best


def _formula(value: Any) -> str | None:
    #: Array formulas arrive as objects, not strings, and dropping them
    #: registered every array-calculated cell as a typed value — the
    #: single largest source of false typed-over findings in the
    #: usefulness audit (the RIIO-3 models array-enter whole blocks).
    if isinstance(value, ArrayFormula):
        text = value.text or ""
        if not text:
            return None
        return text if text.startswith("=") else f"={text}"
    return value if isinstance(value, str) and value.startswith("=") else None


def _label(value: Any) -> str | None:
    """A cell's text, when the text is a name and not a calculation.

    A formula is a string, so a column full of formulas counts as a column
    full of labels unless this says otherwise — and then every figure on
    the sheet is named after the arithmetic in the column beside it. The
    Cascade model never showed it because nothing in its column A is
    calculated; a real model with a formula down the left-hand side shows
    it immediately.
    """
    if not isinstance(value, str) or value.startswith("="):
        return None
    return value if value.strip() else None


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool | str):
        return None
    if isinstance(value, int | float):
        # Through `repr`: 0.098 as a float carries fifteen digits of noise
        # that mean nothing to a figure printed to one decimal place.
        return Decimal(repr(value))
    if isinstance(value, Decimal):
        return value
    #: A date IS a number to Excel — the serial it computes with. The
    #: reader used to drop date cells, which made every formula that
    #: reads one (`=InpC!F109` on Bertha Park) unevaluable and hid the
    #: whole timeline machinery of a model from the engine. Serial
    #: convention: days since 1899-12-30, fraction for time of day.
    if isinstance(value, datetime.datetime):
        delta = value - datetime.datetime(1899, 12, 30)
        return Decimal(delta.days) + (
            Decimal(delta.seconds) / Decimal(86400) if delta.seconds else 0
        )
    if isinstance(value, datetime.date):
        return Decimal((value - datetime.date(1899, 12, 30)).days)
    if isinstance(value, datetime.time):
        seconds = value.hour * 3600 + value.minute * 60 + value.second
        return Decimal(seconds) / Decimal(86400)
    return None


__all__ = [
    "DERIVED_ROW",
    "HEADER_TEXTS",
    "Cell",
    "Names",
    "Precedents",
    "Workbook",
    "precedents_of",
    "read_workbook",
    "references_of",
]
