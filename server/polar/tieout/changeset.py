"""The changeset — swens-plan Track F, F2.

Every correction recorded (before, after, why, who), applied
atomically, invertible — and every write followed by a full re-read
and re-audit as a hard gate: any unexpected change refuses the whole
set and releases nothing. Not a check that runs; a gate that blocks.
The self-closing-XML bug is the standing argument: a construct the
writer cannot see will exist again, and the re-read is what catches
it, because it compares the corrected file against the original cell
by cell in an independent library rather than trusting the writer's
account of itself.

The gate has three walls:

1. **The cell-exact compare.** Every sheet, every cell, both files,
   read through openpyxl. Anything different that was not asked for —
   refused. The targets must hold exactly what was asked — else
   refused.
2. **The re-audit.** The engine's own audit runs on both versions. A
   defect class appearing on a sheet where it did not exist before is
   damage the write created — refused.
3. **« Incomplete repair » as a first-class state** — not success,
   not failure. A correction names the rule that motivated it; when
   that rule still fires on the same sheet after the fix (the E41
   case: one cell repaired, its siblings still wrong), the changeset
   applies but says « incomplete » and names the cells still wrong.

Custody (F4) builds on this: the corrected bytes are a new version
held by Swens — the caller releases them on acceptance, or calls
`undo()` and gets the original back, member-for-member identical.
"""

import io
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .audit import audit
from .workbook import read_workbook
from .writer import Edit, WorkbookWriter, WriteRefused


@dataclass(frozen=True)
class Correction:
    """One determined fix, with its determination in words.

    `rule` names the audit rule this correction answers, so the
    re-audit can tell a finished repair from an incomplete one.
    """

    sheet: str
    ref: str
    formula: str | None
    value: str
    why: str
    rule: str = ""
    create: bool = False


@dataclass
class Changeset:
    """What one apply did — the record the Changes UI renders.

    `state` is one of « applied », « incomplete », « refused ».
    A refused changeset carries its reason and no payload; an
    incomplete one carries both the corrected payload and the cells
    still wrong, because the fix itself landed.
    """

    who: str
    corrections: list[Correction]
    state: str
    edits: list[Edit] = field(default_factory=list)
    reason: str = ""
    still_wrong: list[str] = field(default_factory=list)
    payload: bytes | None = None
    undo_members: dict[str, bytes | None] = field(default_factory=dict)
    member_order: list[str] = field(default_factory=list)

    def undo(self) -> bytes:
        """The original file back, member-for-member identical."""
        if self.payload is None:
            raise ValueError(
                "a refused changeset released nothing; there is nothing to undo"
            )
        with zipfile.ZipFile(io.BytesIO(self.payload)) as archive:
            members = {
                item.filename: archive.read(item.filename)
                for item in archive.infolist()
            }
        for name, data in self.undo_members.items():
            if data is None:
                members.pop(name, None)
            else:
                members[name] = data
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            for name in self.member_order:
                archive.writestr(name, members[name])
        return out.getvalue()

    def record(self) -> dict[str, Any]:
        """The plain-JSON contract for the founder's Changes UI."""
        return {
            "who": self.who,
            "state": self.state,
            "reason": self.reason,
            "still_wrong": list(self.still_wrong),
            "corrections": [
                {
                    "sheet": c.sheet,
                    "ref": c.ref,
                    "why": c.why,
                    "rule": c.rule,
                }
                for c in self.corrections
            ],
            "edits": [
                {
                    "sheet": e.sheet,
                    "ref": e.ref,
                    "before": {"formula": e.before_formula, "value": e.before_value},
                    "after": {"formula": e.formula, "value": e.value},
                    "created": e.created,
                }
                for e in self.edits
            ],
        }


def apply_corrections(
    payload: bytes, corrections: list[Correction], *, who: str
) -> Changeset:
    """All the corrections, or none of them, with the gate in between."""

    def refused(reason: str, edits: list[Edit] | None = None) -> Changeset:
        return Changeset(
            who=who,
            corrections=list(corrections),
            state="refused",
            reason=reason,
            edits=edits or [],
        )

    with tempfile.TemporaryDirectory() as folder:
        original = Path(folder) / "original.xlsx"
        original.write_bytes(payload)
        try:
            writer = WorkbookWriter(original)
        except (zipfile.BadZipFile, KeyError, UnicodeDecodeError) as broken:
            return refused(f"not a workbook this writer can open: {broken}")

        edits: list[Edit] = []
        try:
            for c in corrections:
                edits.append(
                    writer.set_cell(
                        c.sheet,
                        c.ref,
                        formula=c.formula,
                        value=c.value,
                        create=c.create,
                    )
                )
        except WriteRefused as refusal:
            return refused(str(refusal))
        corrected_path = writer.save(Path(folder) / "corrected.xlsx")

        problem = _unexpected_change(original, corrected_path, corrections)
        if problem:
            return refused(problem, edits)

        before = audit(read_workbook(str(original)))
        after = audit(read_workbook(str(corrected_path)))
        known = {(f.rule, f.sheet) for f in before.findings}
        damage = [f for f in after.findings if (f.rule, f.sheet) not in known]
        if damage:
            first = damage[0]
            return refused(
                f"the write would create new damage — {first.rule} at "
                f"{first.ref}: {first.detail}",
                edits,
            )

        still: list[str] = []
        for c in corrections:
            if not c.rule:
                continue
            for f in after.findings:
                if f.rule == c.rule and f.sheet == c.sheet and f.ref not in still:
                    still.append(f.ref)

        corrected = corrected_path.read_bytes()
        undo_members, order = _undo_map(payload, corrected)
        return Changeset(
            who=who,
            corrections=list(corrections),
            state="incomplete" if still else "applied",
            edits=edits,
            still_wrong=still,
            payload=corrected,
            undo_members=undo_members,
            member_order=order,
        )


def _unexpected_change(
    before_path: Path, after_path: Path, corrections: list[Correction]
) -> str:
    """The cell-exact compare: one sentence naming the problem, or ""."""
    targets = {(c.sheet, c.ref) for c in corrections}
    a = load_workbook(before_path)
    b = load_workbook(after_path)
    if a.sheetnames != b.sheetnames:
        return "the write changed the sheet list itself — rolled back"
    for name in a.sheetnames:
        held = {
            cell.coordinate: cell.value
            for row in a[name].iter_rows()
            for cell in row
            if cell.value is not None
        }
        now = {
            cell.coordinate: cell.value
            for row in b[name].iter_rows()
            for cell in row
            if cell.value is not None
        }
        for coordinate in set(held) | set(now):
            if (name, coordinate) in targets:
                continue
            if held.get(coordinate) != now.get(coordinate):
                return (
                    f"{name}!{coordinate} changed without being asked to — "
                    f"the write is rolled back"
                )
    for c in corrections:
        got = b[c.sheet][c.ref].value
        if c.formula is not None:
            if got != c.formula:
                return (
                    f"{c.sheet}!{c.ref} holds {got!r}, not the asked "
                    f"formula — rolled back"
                )
        elif str(got) != c.value and not _same_number(got, c.value):
            return f"{c.sheet}!{c.ref} holds {got!r}, not the asked value — rolled back"
    return ""


def _same_number(got: Any, value: str) -> bool:
    try:
        return float(got) == float(value)
    except (TypeError, ValueError):
        return False


def _undo_map(
    original: bytes, corrected: bytes
) -> tuple[dict[str, bytes | None], list[str]]:
    """What undo must restore: original bytes for every member the
    apply changed or dropped, None for members it added."""
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        order = [item.filename for item in archive.infolist()]
        held = {name: archive.read(name) for name in order}
    with zipfile.ZipFile(io.BytesIO(corrected)) as archive:
        now = {
            item.filename: archive.read(item.filename) for item in archive.infolist()
        }
    undo: dict[str, bytes | None] = {}
    for name, data in held.items():
        if now.get(name) != data:
            undo[name] = data
    for name in now:
        if name not in held:
            undo[name] = None
    return undo, order
