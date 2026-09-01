"""D3 round 10 — page selection by title, on the RMU FERC pair.

The registration is ``docs/pierce/chain-round10.md``, committed before
this harness produced a verdict. Two stages: the model row's section
header (nearest column-B string above it on Appendix A) is matched
against each document page's TITLE — the longest uppercase run in the
page's first six lines — and the row's label is then matched against
the selected pages' table lines only, prose excluded. Scoring within a
page is identical to ``corpus_d3_oracle.py``, which reproduced round 8
exactly: label-token coverage, FLOOR = 0.5, exact tie at the top means
abstain. Never-by-value holds throughout — ``propose.label_tokens``
drops purely numeric tokens from both sides, and values appear in the
report only.

    uv run python scripts/corpus_d3_round10.py <corpus-dir>

``<corpus-dir>`` is where ``corpus_ferc_fetch.py`` put the RMU pair.
Verdicts are written to ``docs/pierce/scribe-d3-round10-verdicts.json``.
"""

import importlib.util
import json
import re
import sys
import warnings
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

DOCS = Path(__file__).parent.parent.parent / "docs" / "pierce"

#: The FERC Form 1 page mark is the one in the page footer — the rule
#: that owns the answer, from ``corpus_d3_oracle.py``.
FOOTER = re.compile(r"FERC\s+FORM\s+NO\.\s*1.*?Page\s+(\d{3})\b", re.IGNORECASE)

#: A schedule prints its line number on the left of a left-hand page and
#: on the right of its continuation.
LINE_LEAD = re.compile(r"^\s*(\d{1,3})(?![0-9])")
LINE_TAIL = re.compile(r"(?:^|\s)(\d{1,3})\s*$")

#: A Form 1 page's title, frozen in the registration: the longest run of
#: capitals (with the punctuation a schedule title carries) over the
#: page's first six lines, needing at least 8 uppercase letters.
#: Calibrated on the 96 footer pages the sample does not cite: 93 titles,
#: 3 empty.
UPPER_RUN = re.compile(r"[A-Z][A-Z0-9 &/,\-\.()':]{7,}")
TITLE_LINES = 6
TITLE_MIN_UPPER = 8

#: A candidate line carrying more than this many alphabetic word tokens
#: is prose and ineligible. Round 8's truth lines carry 0-8; the
#: instruction sentence that drew both of its wrong answers carries 20.
PROSE_WORDS = 10

FLOOR = 0.5

SCORABLE = {10, 12, 13, 19, 21, 36, 41, 44, 57, 59, 88, 119, 121, 179, 187}

_WORD = re.compile(r"[A-Za-z]+")


def _chain() -> tuple[Any, Any]:
    chain = Path(__file__).parent.parent / "polar" / "tieout" / "chain"
    out = []
    for name in ("extract", "propose"):
        spec = importlib.util.spec_from_file_location(name, chain / f"{name}.py")
        assert spec is not None
        assert spec.loader is not None
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        out.append(module)
    return out[0], out[1]


extract, propose = _chain()


def title_of(text: str) -> str:
    best = ""
    for line in text.split("\n")[:TITLE_LINES]:
        for match in UPPER_RUN.findall(line):
            if sum(c.isupper() for c in match) >= TITLE_MIN_UPPER and len(match) > len(
                best
            ):
                best = match
    return best.strip()


def _anchors(text: str) -> set[int]:
    lead, tail = LINE_LEAD.match(text), LINE_TAIL.search(text)
    return {int(m.group(1)) for m in (lead, tail) if m}


def section_headers(model: Path) -> dict[int, str]:
    """Row -> nearest non-empty column-B string at or above it."""
    import openpyxl

    sheet = openpyxl.load_workbook(model, data_only=False)["Appendix A - TSRR Summary"]
    headers: dict[int, str] = {}
    current = ""
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=2).value
        if isinstance(value, str) and value.strip():
            current = value.strip()
        headers[row] = current
    return headers


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    corpus = Path(argv[0])
    import pdfplumber

    document = corpus / "rmu-2015-form1.pdf"
    sample = [
        s
        for s in json.loads((DOCS / "scribe-d3-round8-sample.json").read_text())
        if s["row"] in SCORABLE
    ]
    headers = section_headers(corpus / "rmu-2016-formula-rate.xlsx")

    printed: dict[int, int] = {}  # printed page -> pdf page (1-based)
    titles: dict[int, str] = {}  # pdf page (1-based) -> title
    with pdfplumber.open(document) as pdf:
        for index, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            found = FOOTER.findall(text)
            if found:
                printed[int(found[0])] = index + 1
            titles[index + 1] = title_of(text)

    numbers = extract.extract_pdf(document).numbers
    by_page: dict[int, list[Any]] = {}
    for number in numbers:
        by_page.setdefault(number.page, []).append(number)

    verdicts = []
    stage1_hits = 0
    counts = {"correct": 0, "wrong": 0, "abstained": 0, "unreachable-page": 0}
    for entry in sample:
        section = headers.get(entry["row"], "")
        wanted_page = propose.label_tokens(section)
        truth_page = printed[entry["page"]]

        # Stage 1 — pages tied at the maximum title score, if it is > 0.
        best = 0.0
        selected: list[int] = []
        for pdf_page, title in titles.items():
            tokens = propose.label_tokens(title)
            score = len(wanted_page & tokens) / len(wanted_page) if wanted_page else 0.0
            if score > best:
                best, selected = score, [pdf_page]
            elif score == best and best > 0:
                selected.append(pdf_page)
        page_hit = truth_page in selected
        stage1_hits += page_hit

        record: dict[str, Any] = {
            "row": entry["row"],
            "label": entry["label"],
            "section": section,
            "citation": entry["citation"],
            "truth_pdf_page": truth_page,
            "selected_pages": sorted(selected),
            "page_score": best,
            "truth_page_selected": page_hit,
        }

        if not selected:
            counts["unreachable-page"] += 1
            record["outcome"] = "unreachable-page"
            verdicts.append(record)
            continue

        # Stage 2 — the frozen round-8 scoring, restricted to the
        # selected pages, prose excluded.
        wanted = propose.label_tokens(entry["label"] or "")
        candidates = [
            n
            for p in selected
            for n in by_page.get(p, [])
            if len(_WORD.findall(n.line)) <= PROSE_WORDS
        ]
        scored = (
            sorted(
                (
                    (len(wanted & propose.label_tokens(n.line)) / len(wanted), n)
                    for n in candidates
                ),
                key=lambda pair: -pair[0],
            )
            if wanted
            else []
        )

        if not scored or scored[0][0] < FLOOR:
            counts["abstained"] += 1
            record["outcome"] = "abstained (floor)"
        elif len(scored) > 1 and scored[1][0] == scored[0][0]:
            counts["abstained"] += 1
            tied = sum(1 for score, _ in scored if score == scored[0][0])
            record["outcome"] = f"abstained (tie of {tied})"
        else:
            top = scored[0][1]
            record["proposal"] = {
                "pdf_page": top.page,
                "line": top.line,
                "text": top.text,
            }
            if top.page == truth_page and entry["line"] in _anchors(top.line):
                counts["correct"] += 1
                record["outcome"] = "correct"
            else:
                counts["wrong"] += 1
                record["outcome"] = "wrong"
        verdicts.append(record)

    out = DOCS / "scribe-d3-round10-verdicts.json"
    out.write_text(json.dumps(verdicts, indent=1))

    print(f"stage 1: truth page selected on {stage1_hits} of {len(sample)}")
    for key, value in counts.items():
        print(f"  {key:18} {value}")
    for record in verdicts:
        print(
            f"  r{record['row']:<4} page {'HIT ' if record['truth_page_selected'] else 'miss'}"
            f"  {record['outcome']:22} {record['label'][:44]}"
        )
    print(f"verdicts: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
