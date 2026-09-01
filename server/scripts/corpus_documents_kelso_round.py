"""D3 round 4: the deal team's own provenance tab as the answer sheet.

Registered in the Scribe log before this existed. The Scottish close
models publish alongside their signed contracts, and Kelso's model
carries a hand-written provenance tab — « clause → term → figure »,
73 rows, written by the deal team at close and owing nothing to us.
That is the only ground truth in this corpus nobody here influenced,
so the sample is every row of it, not a draw.

Two phases, same discipline as every round in this lane: the judging
sheet is produced with **no matcher output on it**, the truth is
recorded blind, and only then is the frozen matcher scored.

    uv run python -m scripts.corpus_documents_kelso_round sheet \\
        --model  scripts/corpus_sft/kelso-high-school-model.xlsx \\
        --contract scripts/corpus_documents/sft/kelso-high-school-agreement.pdf

    uv run python -m scripts.corpus_documents_kelso_round score TRUTH.json \\
        --model ... --contract ...

**Blindness is structural**: the provenance sheet is dropped from the
model side entirely — it contributes no cells, no labels and no
candidates. It exists here only as the answer sheet.

**The typed-cell convention, declared wherever a number from this
round is quoted**: the published models are formula-stripped, so
typed-versus-computed cannot be read from the file. The provenance
tab itself names the document-fed figures; the cell under test is
the one carrying a provenance row's figure. Locating it uses the
figure's value — locating is sampling, and sampling by value stays
legitimate exactly where scoring by value is not. The matcher sees
only labels.

**Conditions, counted as their own rows and never as failures**:
`unreachable-unpublished` (the row cites paper that was never
published — some loan agreements), `unreachable-ocr` (the contract
statement falls on a page the extractor refuses, or the number did
not survive the photocopy). The extractor's own coverage and
refusals for the contract print with the round.
"""

import importlib.util
import json
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

_CHAIN = Path(__file__).parent.parent / "polar" / "tieout" / "chain"

#: Words that mark a sheet as the provenance tab when one is not named.
PROVENANCE_WORDS = ("clause", "provenance", "source", "reference", "contract")


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, _CHAIN / f"{name}.py")
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


extract = _load("extract")
propose = _load("propose")


def find_provenance_sheet(path: Path) -> str | None:
    """The sheet that looks like the deal team's marking scheme.

    A guess, printed for a person to confirm or override with
    ``--provenance-sheet`` — never silently trusted, because getting
    this wrong would leak the answer sheet into the matcher's inputs.
    """
    from openpyxl import load_workbook

    book = load_workbook(str(path), read_only=True, data_only=True)
    best: tuple[int, str] | None = None
    for sheet in book.worksheets:
        if any(word in sheet.title.lower() for word in PROVENANCE_WORDS):
            return str(sheet.title)
        hits = 0
        for row in sheet.iter_rows(min_row=1, max_row=6, values_only=True):
            text = " ".join(str(v).lower() for v in row if isinstance(v, str))
            hits += sum(word in text for word in PROVENANCE_WORDS)
        if hits and (best is None or hits > best[0]):
            best = (hits, str(sheet.title))
    return best[1] if best else None


def provenance_rows(path: Path, sheet_name: str) -> list[dict]:
    """« clause → term → figure », one entry per row that has a figure."""
    from openpyxl import load_workbook

    book = load_workbook(str(path), read_only=True, data_only=True)
    sheet = book[sheet_name]
    rows = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        numbers = [
            v for v in row if isinstance(v, (int, float)) and not isinstance(v, bool)
        ]
        words = [str(v).strip() for v in row if isinstance(v, str) and str(v).strip()]
        if not numbers or not words:
            continue
        rows.append(
            {
                "row": index,
                "text": " | ".join(words),
                "figure": float(numbers[-1]),
            }
        )
    return rows


def model_cells(path: Path, excluded: frozenset[str]):
    """Every numeric cell EXCEPT the answer sheet's own — and, per the
    amendment registered in ``chain-kelso-round.md`` before the round
    ran, any sibling tab of the same hand-written extract class named
    with ``--also-exclude``: cells that exist only because a person
    already did the linking by hand are answer material, not model."""
    from polar.models.tieout import ArtifactKind
    from polar.tieout.ingest import read_artifact

    got = read_artifact(path.read_bytes(), path.name, ArtifactKind.model)
    return [
        cell
        for cell in got.cells
        if cell.sheet not in excluded and cell.value is not None
    ]


def contract_facts(path: Path):
    """D1 over the signed agreement: facts, and the pages refused."""
    extraction = extract.extract_pdf(path)
    facts = [
        (f"p{n.page}|{ordinal}", n) for ordinal, n in enumerate(extraction.numbers)
    ]
    return facts, extraction


def _close(a: float, b: float) -> bool:
    """The same number, allowing the model's own rounding."""
    if a == b:
        return True
    scale = max(abs(a), abs(b), 1.0)
    return abs(a - b) <= 0.005 * scale


def sheet(
    model: Path, contract: Path, provenance: str, excluded: frozenset[str]
) -> int:
    rows = provenance_rows(model, provenance)
    cells = model_cells(model, excluded)
    facts, extraction = contract_facts(contract)
    print(
        f"provenance rows: {len(rows)} | model cells (answer sheet excluded): "
        f"{len(cells):,} | contract facts: {len(facts):,} | pages refused: "
        f"{len(extraction.refusals)}",
        file=sys.stderr,
    )
    for refusal in extraction.refusals:
        print(f"  [refused] {refusal.reason}", file=sys.stderr)

    record = []
    for entry in rows:
        located = [
            {
                "ref": cell.ref,
                "labels": f"{cell.row_label} {cell.column_label}".strip(),
            }
            for cell in cells
            if _close(float(cell.value), entry["figure"])
        ]
        hits = [
            {"key": key, "text": number.text, "line": number.line[:160]}
            for key, number in facts
            if _close(number.value, entry["figure"])
        ]
        record.append(
            {
                **entry,
                "located": located[:10],
                "value_hits": hits[:20],
                # the judge fills both:
                "truth": None,  # list of fact keys the contract states it at
                "condition": None,  # ok | unreachable-unpublished | unreachable-ocr
            }
        )
        print(f"\n[row {entry['row']}] {entry['text'][:90]}  figure={entry['figure']}")
        print(f"   model cells with this value: {len(located)}")
        for cell in located[:3]:
            print(f"     {cell['ref']}: {cell['labels'][:70]!r}")
        for hit in hits[:5]:
            print(f"   ? {hit['key']}: {hit['text']!r} in {hit['line'][:80]!r}")
        if not hits:
            print("   no contract line prints this figure")

    out = contract.parent / "kelso-round-sheet.json"
    out.write_text(json.dumps(record, indent=1))
    print(f"\nsheet written: {out} — judge, then run score", file=sys.stderr)
    return 0


def score(
    truth_path: Path, model: Path, contract: Path, excluded: frozenset[str]
) -> int:
    entries = json.loads(truth_path.read_text())
    cells = model_cells(model, excluded)
    by_ref = {cell.ref: cell for cell in cells}
    facts, _ = contract_facts(contract)
    candidates = [(key, number.line, number.text) for key, number in facts]

    counts: dict[str, int] = {}
    verdicts = []
    for entry in entries:
        condition = entry.get("condition") or "unjudged"
        if condition != "ok":
            counts[condition] = counts.get(condition, 0) + 1
            continue
        located = entry.get("located") or []
        if not located:
            counts["no-cell-located"] = counts.get("no-cell-located", 0) + 1
            continue
        cell = by_ref.get(located[0]["ref"])
        labels = (
            f"{cell.row_label} {cell.column_label}".strip()
            if cell
            else located[0]["labels"]
        )
        stated = entry.get("truth") or []
        answer = propose.propose(labels, candidates)
        if isinstance(answer, propose.Proposed):
            picked = answer.candidate.fact_id
            verdict = "true proposal" if picked in stated else "false proposal"
        else:
            picked = None
            verdict = "true abstention" if not stated else "missed"
        counts[verdict] = counts.get(verdict, 0) + 1
        verdicts.append(
            {
                "row": entry["row"],
                "text": entry["text"],
                "figure": entry["figure"],
                "cell": located[0]["ref"],
                "labels": labels,
                "stated": stated,
                "proposed": picked,
                "verdict": verdict,
            }
        )
        print(f"row {entry['row']:4} {verdict:16} {located[0]['ref']:20} -> {picked}")

    reachable = sum(
        counts.get(name, 0)
        for name in ("true proposal", "false proposal", "true abstention", "missed")
    )
    print("\n--- the registered table (reachable rows) ---")
    for name in ("true proposal", "false proposal", "true abstention", "missed"):
        print(f"{name}: {counts.get(name, 0)}")
    proposals = counts.get("true proposal", 0) + counts.get("false proposal", 0)
    if proposals:
        print(
            f"proposal precision: {counts.get('true proposal', 0)}/{proposals} "
            f"= {counts.get('true proposal', 0) / proposals:.0%}"
        )
    print(f"reachable rows: {reachable} of {len(entries)}")
    for name in (
        "unreachable-unpublished",
        "unreachable-ocr",
        "no-cell-located",
        "unjudged",
    ):
        if counts.get(name):
            print(f"{name}: {counts[name]}")

    out = contract.parent / "kelso-round-verdicts.json"
    out.write_text(json.dumps(verdicts, indent=1))
    print(f"verdicts written: {out}", file=sys.stderr)
    return 0


def _argument(name: str, default: str | None = None) -> str | None:
    return sys.argv[sys.argv.index(name) + 1] if name in sys.argv else default


def main() -> int:
    if len(sys.argv) < 2 or sys.argv[1] not in ("sheet", "score"):
        print(__doc__)
        return 1
    model = Path(_argument("--model", "") or "")
    contract = Path(_argument("--contract", "") or "")
    if not model.exists() or not contract.exists():
        print(
            "the pair is not here yet — see the Scribe log: the hub's "
            "certificate is expired and wrong-host, and the archive holds "
            "no capture of the models. Nothing to run until the files land.",
            file=sys.stderr,
        )
        return 1
    provenance = _argument("--provenance-sheet") or find_provenance_sheet(model)
    if provenance is None:
        print(
            "no provenance tab found; name it with --provenance-sheet. "
            "Refusing to run: guessing this wrong would leak the answer "
            "sheet into the matcher's own inputs.",
            file=sys.stderr,
        )
        return 1
    also = _argument("--also-exclude") or ""
    excluded = frozenset(
        {provenance} | {s.strip() for s in also.split(",") if s.strip()}
    )
    print(
        f"answer sheet(s) excluded from the model side: {sorted(excluded)!r}",
        file=sys.stderr,
    )
    if sys.argv[1] == "sheet":
        return sheet(model, contract, provenance, excluded)
    return score(Path(sys.argv[2]), model, contract, excluded)


if __name__ == "__main__":
    raise SystemExit(main())
