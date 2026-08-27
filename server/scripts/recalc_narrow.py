"""B4's narrowing round — from a violated law to the responsible cell.

Runs `polar.tieout.recalc.narrow` over the plants already registered
in `scripts.recalc_behave` (same selector maps, same defects, so the
narrowing is measured on exactly the population the per-class table
was measured on).

    cd server && uv run python -m scripts.recalc_narrow FILTER OUT.json [--ddmin]

Per planted copy: recalculate it once for the baseline, once per
applicable perturbation, find the violated outputs with the same law
functions B4 uses, and narrow each violation:

- perturbation laws (zero-input, proportionality, scale invariance)
  → the **frontier walk**, free of extra recalculation;
- consolidation → the **structural** narrowing (which declared
  segment the total's own formula never reaches).

`--ddmin` adds Zeller's minimizing delta debugging on top, with a
real oracle: pin a candidate subset to the values the law predicts,
recalculate, ask whether the law holds. One recalculation per test,
so it is for registered subsamples only.

**Scored, per the registration**: a narrowing is *exact* when its
answer is the planted cell alone, *hit* when the planted cell is in
a larger set, *miss* otherwise — and the set's size is recorded
either way, because a « narrowing » that names forty cells has not
narrowed anything.
"""

import json
import sys
import time
from collections.abc import Mapping, Sequence
from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from polar.tieout.recalc.laws import (
    consolidation_violations,
    proportionality_violations,
    scale_invariance_violations,
    zero_input_violations,
)
from polar.tieout.recalc.narrow import (
    ddmin,
    formula_references,
    frontier,
    missing_segments,
)
from polar.tieout.recalc.uno_calc import UnoCalculator, find_install
from polar.tieout.workbook import read_workbook
from scripts.recalc_behave import (
    PILOTS,
    Plant,
    Selectors,
    _inputs_scaled,
    _inputs_zero,
    _split,
    perturb,
    plant,
)

#: The perturbation laws, with the factor the frontier walk classifies
#: against: zero-input has no scaled state, a doubling has 2.0, a
#: cents-for-pounds rescale has 100.0.
PERTURBATIONS = (
    ("zero-input", None),
    ("proportionality", 2.0),
    ("scale-invariance", 100.0),
)


def numeric(values: Mapping[str, Any]) -> dict[str, float]:
    return {ref: v for ref, v in values.items() if isinstance(v, float)}


def perturbation_inputs(
    law: str, base_file: Path, selectors: Selectors
) -> dict[str, float] | None:
    if law == "zero-input" and selectors.volume:
        return _inputs_zero(base_file, selectors.volume)
    if law == "proportionality" and selectors.price:
        return _inputs_scaled(base_file, selectors.price, 2.0)
    if law == "scale-invariance" and selectors.monetary:
        return _inputs_scaled(base_file, selectors.monetary, 100.0)
    return None


def violations_for(
    law: str,
    baseline: dict[str, float],
    perturbed: dict[str, float],
    selectors: Selectors,
) -> list[str]:
    if law == "zero-input":
        return [v.ref for v in zero_input_violations(perturbed, selectors.revenue)]
    if law == "proportionality":
        return [
            v.ref
            for v in proportionality_violations(
                baseline, perturbed, 2.0, selectors.revenue
            )
        ]
    return [
        v.ref
        for v in scale_invariance_violations(baseline, perturbed, selectors.ratios)
    ]


def expected_value(law: str, ref: str, baseline: dict[str, float]) -> float:
    """What the law predicts for one cell — the value ddmin pins."""
    if law == "zero-input":
        return 0.0
    if law == "proportionality":
        return baseline.get(ref, 0.0) * 2.0
    return baseline.get(ref, 0.0)


def patched_cells(cells: Mapping[str, Any], defect: Plant) -> dict[str, Any]:
    """The reader's cells with the planted cell's own formula in place.

    The workbook is read once per file (minutes on a 100k-cell
    model); the plant changes exactly one formula, and the narrowing
    must see *that* formula — its references are the planted cell's
    real precedents, and an omitted segment is precisely a reference
    that stopped being there. Patched here rather than re-read.
    """
    from openpyxl.utils import get_column_letter, range_boundaries

    sheet, _ = _split(defect.target)
    precedents: list[str] = []
    for operand in sorted(formula_references(defect.formula)):
        ref_sheet, _, ref = operand.rpartition("!")
        try:
            c1, r1, c2, r2 = range_boundaries(ref if ":" in ref else f"{ref}:{ref}")
        except Exception:
            continue
        if None in (c1, r1, c2, r2):
            continue
        for row in range(int(r1), int(r2) + 1):
            for column in range(int(c1), int(c2) + 1):
                precedents.append(
                    f"{ref_sheet or sheet}!{get_column_letter(column)}{row}"
                )
    original = cells.get(defect.target)
    patched = dict(cells)
    if original is not None:
        patched[defect.target] = replace(
            original, formula=defect.formula, precedents=tuple(precedents)
        )
    return patched


def score(culprits: list[str], planted: str) -> str:
    if culprits == [planted]:
        return "exact"
    return "hit" if planted in culprits else "miss"


def narrow_plant(
    calculator: UnoCalculator,
    work: Path,
    source: Path,
    defect: Plant,
    selectors: Selectors,
    cells: dict[str, Any],
    stem: str,
    use_ddmin: bool,
) -> dict[str, Any]:
    planted = work / f"{stem}.xlsx"
    plant(source, planted, defect)
    cells = patched_cells(cells, defect)
    baseline = numeric(calculator.recalculate(str(planted)).values)
    record: dict[str, Any] = {
        "kind": defect.kind,
        "planted": defect.target,
        "laws": {},
    }

    for law, factor in PERTURBATIONS:
        changes = perturbation_inputs(law, planted, selectors)
        if changes is None:
            continue
        copy = work / f"{stem}-{law}.xlsx"
        perturb(planted, copy, changes)
        perturbed = numeric(calculator.recalculate(str(copy)).values)
        broken = violations_for(law, baseline, perturbed, selectors)
        if not broken:
            continue
        entry: dict[str, Any] = {"violated": broken, "narrowed": {}}
        for output in broken:
            culprits = frontier(cells, baseline, perturbed, output, factor=factor)
            entry["narrowed"][output] = {
                "culprits": culprits,
                "size": len(culprits),
                "score": score(culprits, defect.target),
            }
            if use_ddmin and culprits:
                entry["narrowed"][output]["ddmin"] = run_ddmin(
                    calculator,
                    work,
                    copy,
                    law,
                    baseline,
                    selectors,
                    output,
                    culprits,
                    defect.target,
                    stem,
                )
        record["laws"][law] = entry

    if selectors.totals:
        broken = [
            v.ref
            for v in consolidation_violations(
                baseline, {t: list(s) for t, s in selectors.totals.items()}
            )
        ]
        if broken:
            structural = {}
            for total in broken:
                sheet, _ = _split(total)
                cell = cells.get(total)
                gaps = missing_segments(
                    getattr(cell, "formula", "") or "", selectors.totals[total], sheet
                )
                # An unreached declared segment *is* the authoring
                # decision, and it lives in the total's own formula.
                # With every segment reached, the identity is broken
                # somewhere among the parts and this method cannot
                # say where — the honest answer is the whole set.
                culprits = [total] if gaps else [total, *selectors.totals[total]]
                structural[total] = {
                    "unreached_segments": gaps,
                    "culprits": culprits,
                    "size": len(culprits),
                    "score": score(culprits, defect.target),
                }
            record["laws"]["consolidation"] = {
                "violated": broken,
                "narrowed": structural,
            }
    return record


def run_ddmin(
    calculator: UnoCalculator,
    work: Path,
    perturbed_file: Path,
    law: str,
    baseline: dict[str, float],
    selectors: Selectors,
    output: str,
    candidates: list[str],
    planted: str,
    stem: str,
) -> dict[str, Any]:
    """Zeller's ddmin with pin-and-recalculate as the oracle."""
    tests = {"count": 0}

    def holds(subset: Sequence[str]) -> bool:
        tests["count"] += 1
        pinned = work / f"{stem}-ddmin-{tests['count']}.xlsx"
        book = load_workbook(perturbed_file)
        for ref in subset:
            sheet, at = _split(ref)
            book[sheet][at] = expected_value(law, ref, baseline)
        book.save(pinned)
        values = numeric(calculator.recalculate(str(pinned)).values)
        return output not in violations_for(law, baseline, values, selectors)

    minimal = ddmin(candidates, holds)
    return {
        "minimal": minimal,
        "size": len(minimal),
        "tests": tests["count"],
        "score": score(minimal, planted),
    }


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    only, out = sys.argv[1], Path(sys.argv[2])
    use_ddmin = "--ddmin" in sys.argv
    if find_install() is None:
        print("no LibreOffice >= 25.8 here — refusing")
        return 1
    chosen = {path: entry for path, entry in PILOTS.items() if only in path}
    if not chosen:
        print("no registered pilot matches; nothing to narrow")
        return 1

    work = out.parent / f"{out.stem}-work"
    work.mkdir(parents=True, exist_ok=True)
    records = []
    for rel_path, (selectors, plants) in chosen.items():
        source = Path(rel_path)
        print(f"== {source.name}", flush=True)
        started = time.monotonic()
        cells = read_workbook(str(source)).cells
        calculator = UnoCalculator(document_timeout=3600)
        calculator.start()
        try:
            for i, defect in enumerate(plants):
                record = narrow_plant(
                    calculator,
                    work,
                    source,
                    defect,
                    selectors,
                    cells,
                    f"{source.stem}-p{i}",
                    use_ddmin,
                )
                record["file"] = source.name
                records.append(record)
                print(
                    f"   plant {i} ({defect.kind} at {defect.target}): "
                    + json.dumps(
                        {
                            law: {
                                ref: n["score"] for ref, n in entry["narrowed"].items()
                            }
                            for law, entry in record["laws"].items()
                        }
                    ),
                    flush=True,
                )
        except Exception as error:
            records.append(
                {
                    "file": source.name,
                    "unmeasurable": f"{type(error).__name__}: {error}",
                }
            )
            print(f"   UNMEASURABLE: {type(error).__name__}: {error}", flush=True)
        finally:
            calculator.stop()
        print(f"   ({round(time.monotonic() - started, 1)}s)", flush=True)
        out.write_text(json.dumps(records, indent=1))
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
