"""D3 round 11 — the floor the evidence produces, on the RMU FERC pair.

The registration is ``docs/pierce/chain-round11.md``, committed before
this harness ran. Round 10's two-stage design with two frozen changes:
a page's own title line is not a candidate (a heading names the table,
it does not state a value), and the stage-2 floor is DERIVED by the
founder's negative control — each model label is deliberately matched
against every wrong section, and the floor is the maximum score any
deliberately-wrong pairing achieves, so that at scoring time zero of
them pass. A wrong pairing whose selected pages overlap the true
section's is excluded from the control: a shuffle that lands on the
real pages is the real pairing wearing a shuffle.

Never-by-value holds throughout. The kill rule is strict: a proposal
on the label half of the cited line counts as wrong, and is also
counted under its own name so the report shows what the wrongs were.

    uv run python scripts/corpus_d3_round11.py <corpus-dir>

Verdicts and the control distribution are written to
``docs/pierce/scribe-d3-round11-verdicts.json``.
"""

import importlib.util
import json
import re
import sys
import warnings
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

DOCS = Path(__file__).parent.parent.parent / "docs" / "pierce"

FOOTER = re.compile(r"FERC\s+FORM\s+NO\.\s*1.*?Page\s+(\d{3})\b", re.IGNORECASE)
LINE_LEAD = re.compile(r"^\s*(\d{1,3})(?![0-9])")
LINE_TAIL = re.compile(r"(?:^|\s)(\d{1,3})\s*$")
UPPER_RUN = re.compile(r"[A-Z][A-Z0-9 &/,\-\.()':]{7,}")
TITLE_LINES = 6
TITLE_MIN_UPPER = 8
PROSE_WORDS = 10
SCORABLE = {10, 12, 13, 19, 21, 36, 41, 44, 57, 59, 88, 119, 121, 179, 187}
_WORD = re.compile(r"[A-Za-z]+")


def _chain() -> tuple[Any, Any]:
    chain = Path(__file__).parent.parent / "polar" / "tieout" / "chain"
    out = []
    for name in ("extract", "propose"):
        spec = importlib.util.spec_from_file_location(name, chain / f"{name}.py")
        assert spec is not None
        assert spec.loader is not None
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        out.append(module)
    return out[0], out[1]


extract, propose = _chain()


def title_of(text: str) -> str:
    best = ""
    for line in text.split("\n")[:TITLE_LINES]:
        for match in UPPER_RUN.findall(line):
            if sum(c.isupper() for c in match) >= TITLE_MIN_UPPER and len(match) > len(
                best
            ):
                best = match
    return best.strip()


def _anchors(text: str) -> set[int]:
    lead, tail = LINE_LEAD.match(text), LINE_TAIL.search(text)
    return {int(m.group(1)) for m in (lead, tail) if m}


def section_headers(model: Path) -> dict[int, str]:
    import openpyxl

    sheet = openpyxl.load_workbook(model, data_only=False)["Appendix A - TSRR Summary"]
    headers: dict[int, str] = {}
    current = ""
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=2).value
        if isinstance(value, str) and value.strip():
            current = value.strip()
        headers[row] = current
    return headers


class Round:
    def __init__(self, corpus: Path) -> None:
        import pdfplumber

        document = corpus / "rmu-2015-form1.pdf"
        self.printed: dict[int, int] = {}
        self.titles: dict[int, str] = {}
        with pdfplumber.open(document) as pdf:
            for index, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                found = FOOTER.findall(text)
                if found:
                    self.printed[int(found[0])] = index + 1
                self.titles[index + 1] = title_of(text)
        self.by_page: dict[int, list[Any]] = {}
        for number in extract.extract_pdf(document).numbers:
            self.by_page.setdefault(number.page, []).append(number)

    def select_pages(self, section: str) -> tuple[float, set[int]]:
        wanted = propose.label_tokens(section)
        best = 0.0
        selected: set[int] = set()
        for pdf_page, title in self.titles.items():
            tokens = propose.label_tokens(title)
            score = len(wanted & tokens) / len(wanted) if wanted else 0.0
            if score > best:
                best, selected = score, {pdf_page}
            elif score == best and best > 0:
                selected.add(pdf_page)
        return best, selected

    def candidates(self, pages: set[int]) -> list[Any]:
        """Prose excluded; a page's own title line excluded."""
        out = []
        for page in pages:
            title = self.titles.get(page, "")
            for n in self.by_page.get(page, []):
                if len(_WORD.findall(n.line)) > PROSE_WORDS:
                    continue
                if title and title in n.line:
                    continue
                out.append(n)
        return out

    def scored(self, label: str, pages: set[int]) -> list[tuple[float, Any]]:
        wanted = propose.label_tokens(label or "")
        if not wanted:
            return []
        return sorted(
            (
                (len(wanted & propose.label_tokens(n.line)) / len(wanted), n)
                for n in self.candidates(pages)
            ),
            key=lambda pair: -pair[0],
        )


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    corpus = Path(argv[0])
    sample = [
        s
        for s in json.loads((DOCS / "scribe-d3-round8-sample.json").read_text())
        if s["row"] in SCORABLE
    ]
    headers = section_headers(corpus / "rmu-2016-formula-rate.xlsx")
    for entry in sample:
        entry["section"] = headers.get(entry["row"], "")

    the_round = Round(corpus)
    selections = {
        entry["row"]: the_round.select_pages(entry["section"]) for entry in sample
    }

    # --- the control, before any scored verdict ------------------------
    pool = sorted({entry["section"] for entry in sample})
    control: list[dict[str, Any]] = []
    excluded_overlap = 0
    for entry in sample:
        _, true_pages = selections[entry["row"]]
        for wrong_section in pool:
            if wrong_section == entry["section"]:
                continue
            _, wrong_pages = the_round.select_pages(wrong_section)
            if not wrong_pages:
                continue
            if wrong_pages & true_pages:
                excluded_overlap += 1
                continue
            scored = the_round.scored(entry["label"], wrong_pages)
            top = scored[0][0] if scored else 0.0
            control.append(
                {
                    "row": entry["row"],
                    "label": entry["label"],
                    "wrong_section": wrong_section,
                    "top_score": top,
                }
            )
    tops = sorted((c["top_score"] for c in control), reverse=True)
    floor = tops[0] if tops else 0.0
    print(f"control: {len(control)} deliberately-wrong pairings scored")
    print(
        f"         {excluded_overlap} excluded for page overlap with the true section"
    )
    print(f"         top wrong scores: {[round(t, 3) for t in tops[:8]]}")
    print(f"DERIVED FLOOR: a proposal must score STRICTLY ABOVE {floor:.4f}")

    # --- the scored run ------------------------------------------------
    verdicts: list[dict[str, Any]] = []
    counts = {
        "correct": 0,
        "wrong": 0,
        "label-half-of-cited-line": 0,
        "abstained": 0,
        "unreachable-page": 0,
    }
    for entry in sample:
        best, selected = selections[entry["row"]]
        truth_page = the_round.printed[entry["page"]]
        record: dict[str, Any] = {
            "row": entry["row"],
            "label": entry["label"],
            "section": entry["section"],
            "citation": entry["citation"],
            "truth_pdf_page": truth_page,
            "selected_pages": sorted(selected),
            "truth_page_selected": truth_page in selected,
        }
        if not selected:
            counts["unreachable-page"] += 1
            record["outcome"] = "unreachable-page"
            verdicts.append(record)
            continue
        scored = the_round.scored(entry["label"], selected)
        if not scored or scored[0][0] <= floor:
            counts["abstained"] += 1
            record["outcome"] = "abstained (derived floor)"
        elif len(scored) > 1 and scored[1][0] == scored[0][0]:
            counts["abstained"] += 1
            tied = sum(1 for score, _ in scored if score == scored[0][0])
            record["outcome"] = f"abstained (tie of {tied})"
        else:
            score, top = scored[0]
            record["proposal"] = {
                "pdf_page": top.page,
                "line": top.line,
                "text": top.text,
                "score": score,
            }
            if top.page == truth_page and entry["line"] in _anchors(top.line):
                counts["correct"] += 1
                record["outcome"] = "correct"
            else:
                counts["wrong"] += 1
                record["outcome"] = "wrong"
                if top.page != truth_page and entry["line"] in _anchors(top.line):
                    counts["label-half-of-cited-line"] += 1
                    record["outcome"] = "wrong (label-half-of-cited-line)"
        verdicts.append(record)

    # --- COUNTERFACTUAL, not the round's result ------------------------
    # The registered floor is pooled: one number, the maximum any wrong
    # pairing reached anywhere. This block re-reads the same control
    # per row — each label barred by its own wrong pairings' maximum —
    # so the report can show what the pooling decision cost. It runs
    # after the registered result and changes nothing above.
    per_row_floor: dict[int, float] = {}
    for c in control:
        row = c["row"]
        per_row_floor[row] = max(per_row_floor.get(row, 0.0), c["top_score"])
    counterfactual: list[dict[str, Any]] = []
    for entry in sample:
        _, selected = selections[entry["row"]]
        if not selected:
            continue
        row_floor = per_row_floor.get(entry["row"], 0.0)
        truth_page = the_round.printed[entry["page"]]
        scored = the_round.scored(entry["label"], selected)
        if not scored or scored[0][0] <= row_floor:
            outcome = "abstained (own floor)"
        elif len(scored) > 1 and scored[1][0] == scored[0][0]:
            outcome = "abstained (tie)"
        else:
            score, top = scored[0]
            if top.page == truth_page and entry["line"] in _anchors(top.line):
                outcome = "correct"
            elif top.page != truth_page and entry["line"] in _anchors(top.line):
                outcome = "wrong (label-half-of-cited-line)"
            else:
                outcome = "wrong"
        counterfactual.append(
            {
                "row": entry["row"],
                "own_floor": row_floor,
                "outcome": outcome,
            }
        )

    out = DOCS / "scribe-d3-round11-verdicts.json"
    out.write_text(
        json.dumps(
            {
                "derived_floor": floor,
                "control": control,
                "verdicts": verdicts,
                "counterfactual_per_row_floor": counterfactual,
            },
            indent=1,
        )
    )
    dead = counts["wrong"] > counts["correct"]
    print()
    for key, value in counts.items():
        print(f"  {key:26} {value}")
    print(
        f"  kill rule (wrong > correct): {'FIRES — the round is dead' if dead else 'does not fire'}"
    )
    for record in verdicts:
        print(f"  r{record['row']:<4} {record['outcome']:34} {record['label'][:40]}")
    print("\nCOUNTERFACTUAL (per-row floors; NOT the round's registered result):")
    for c in counterfactual:
        print(f"  r{c['row']:<4} own floor {c['own_floor']:.3f}  {c['outcome']}")
    print(f"verdicts: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
