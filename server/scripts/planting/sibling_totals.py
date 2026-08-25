"""Plant sibling-total disagreements — the A3 candidate-1 recall test.

Protocol: docs/pierce/a3-sibling-totals.md, registered before any
result. Four defect classes from the mining round's evidence — plug,
off-by-one, bleed, mis-drag — planted into rows of same-signature
column totals by XML surgery on the xlsx zip, so every untouched
cell keeps its cached value. One plant per family, no two plants
sharing a row or a column on one sheet, sites drawn with the
registered seed. The ground truth is written before the engine sees
the planted file. One planting run per host — no re-rolls.

    uv run python -m scripts.planting.sibling_totals --scan CORPUS_DIR
    uv run python -m scripts.planting.sibling_totals HOST.xlsx OUT.xlsx TRUTH.json 20260825

The eligibility logic here is harness-local on purpose: it mirrors
the detector's membership rule in spirit but is written against the
registration's text, not the detector's code, so the two cannot
quietly agree by sharing a bug.
"""

import json
import random
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl.utils import column_index_from_string, get_column_letter

#: The registration's numbers, in one place.
SEED_CLASSES = ("plug", "off-by-one", "bleed", "mis-drag")
SITES_PER_CLASS = 5
MIN_FAMILY = 4
MAX_HOST_CELLS = 400_000
PLUGS = (250, 1000, 3100, 47500)

#: A totals cell the harness can mutate: nothing but one SUM call,
#: optionally signed. (The detector also admits a surround; the
#: harness plants only into clean families so the mutation is the
#: only difference.)
BARE = re.compile(r"^=\s*\+?\s*SUM\((?P<args>[^()]*)\)\s*$", re.IGNORECASE)

#: One argument: optional sheet, one reference or range, A1 style.
ARG = re.compile(
    r"^\s*(?:'(?P<qsheet>[^']+)'!|(?P<sheet>[A-Za-z0-9_.]+)!)?"
    r"\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)"
    r"(?::\$?(?P<col2>[A-Z]{1,3})\$?(?P<row2>\d+))?\s*$",
    re.IGNORECASE,
)


def _coverage(cell_sheet, cell_col, cell_row, formula):
    """(covered {(dc, row)}, own-column multi-row?) — or None when the
    formula is not a clean own-column total entirely above its cell."""
    m = BARE.match(formula)
    if m is None:
        return None
    covered: set[tuple[int, int]] = set()
    multi = False
    args = m.group("args").split(",")
    if not args or not m.group("args").strip():
        return None
    for arg in args:
        a = ARG.match(arg)
        if a is None:
            return None
        sheet = a.group("qsheet") or a.group("sheet")
        if sheet is not None and sheet != cell_sheet:
            return None
        c1 = column_index_from_string(a.group("col").upper())
        c2 = column_index_from_string((a.group("col2") or a.group("col")).upper())
        r1 = int(a.group("row"))
        r2 = int(a.group("row2") or a.group("row"))
        c1, c2 = min(c1, c2), max(c1, c2)
        r1, r2 = min(r1, r2), max(r1, r2)
        if not (c1 <= cell_col <= c2):
            return None
        if r2 >= cell_row:
            return None
        if c1 == c2 == cell_col and r2 > r1:
            multi = True
        for c in range(c1, c2 + 1):
            for r in range(r1, r2 + 1):
                covered.add((c - cell_col, r))
    return (covered, multi) if multi else None


def families(book):
    """Eligible row-direction families: (sheet, row) → members, where
    4+ clean own-column totals share one coverage signature. Members
    are (ref, column, formula, own-column rows sorted)."""
    by_row: dict[tuple[str, int], list] = {}
    for cell in book.cells.values():
        if not cell.formula:
            continue
        got = _coverage(cell.sheet, cell.column, cell.row, cell.formula)
        if got is None:
            continue
        covered, _ = got
        rows0 = tuple(sorted(r for dc, r in covered if dc == 0))
        signature = frozenset(covered)
        by_row.setdefault((cell.sheet, cell.row), []).append(
            (cell.ref, cell.column, cell.formula, rows0, signature)
        )
    out = {}
    for key, members in by_row.items():
        tally: dict[frozenset, list] = {}
        for member in members:
            tally.setdefault(member[4], []).append(member)
        best = max(tally.values(), key=len)
        if len(best) >= MIN_FAMILY:
            out[key] = sorted(best, key=lambda m: m[1])
    return out


def scan(corpus: Path) -> None:
    from polar.tieout.workbook import read_workbook

    ranked = []
    for path in sorted(corpus.rglob("*.xls[xm]")):
        name = str(path.relative_to(corpus))
        try:
            book = read_workbook(str(path))
        except Exception as problem:
            print(f"unreadable {name}: {problem}")
            continue
        n = len(families(book))
        print(f"{name}: {len(book.cells)} cells, {n} eligible families", flush=True)
        if len(book.cells) < MAX_HOST_CELLS:
            ranked.append((-n, name))
    ranked.sort()
    print("\nhosts (top 3 eligible under 400k cells, ties by name):")
    for minus_n, name in ranked[:3]:
        print(f"  {name}  ({-minus_n} families)")


# ---- XML surgery -----------------------------------------------------


def sheet_files(zf: zipfile.ZipFile) -> dict[str, str]:
    book = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rel_to_target = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    out = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', book):
        target = rel_to_target.get(rid, "")
        if target.startswith("/"):
            target = target[1:]
        elif not target.startswith("xl/"):
            target = "xl/" + target
        out[name.replace("&amp;", "&").replace("&quot;", '"').replace("&apos;", "'")] = (
            target
        )
    return out


def rewrite_formula(sheet_xml: str, ref: str, new_formula: str) -> str | None:
    """The cell's <f> replaced with a plain formula, cached <v> kept.

    Returns None when the cell is a shared-formula master or an array
    formula — mutating either would change more cells than the one
    planted — or when the cell or its <f> cannot be found.
    """
    cell_pattern = re.compile(
        rf'(<c r="{re.escape(ref)}"[^>]*>)(.*?)(</c>)', re.DOTALL
    )
    m = cell_pattern.search(sheet_xml)
    if m is None:
        return None
    body = m.group(2)
    f = re.search(r"<f(?:\s[^>]*)?(?:/>|>.*?</f>)", body, re.DOTALL)
    if f is None:
        return None
    attrs = re.match(r"<f\s([^>]*?)/?>", f.group(0))
    tag = attrs.group(1) if attrs else ""
    if 'ref="' in tag or 't="array"' in tag:
        return None  # shared master or array: the plant would spread
    plain = f"<f>{escape(new_formula.lstrip('='))}</f>"
    new_body = body[: f.start()] + plain + body[f.end() :]
    return sheet_xml[: m.start()] + m.group(1) + new_body + m.group(3) + sheet_xml[m.end():]


def _shift_start(formula: str) -> str | None:
    """The first own-column range's start row moved down one."""

    def bump(m: re.Match[str]) -> str:
        return (
            m.group(1)
            + str(int(m.group(2)) + 1)
            + m.group(3)
        )

    new, n = re.subn(
        r"(\(\s*\$?[A-Z]{1,3}\$?)(\d+)(\s*:)", bump, formula, count=1
    )
    return new if n else None


def _widen(formula: str, own_col: int, toward: int) -> str | None:
    """The range widened one column: SUM(C6:C13) → SUM(C6:D13) toward
    the right, SUM(B6:C13) toward the left."""
    m = re.search(
        r"\(\s*(\$?([A-Z]{1,3})\$?\d+)\s*:\s*(\$?)([A-Z]{1,3})(\$?\d+)\s*\)", formula
    )
    if m is None:
        return None
    start, end_dollar, end_col, end_rest = m.group(1), m.group(3), m.group(4), m.group(5)
    if toward > own_col:
        new_end_col = get_column_letter(own_col + 1)
        replaced = f"({start}:{end_dollar}{new_end_col}{end_rest})"
    else:
        new_start_col = get_column_letter(own_col - 1)
        new_start = re.sub(r"[A-Z]{1,3}", new_start_col, start, count=1)
        replaced = f"({new_start}:{end_dollar}{end_col}{end_rest})"
    return formula[: m.start()] + replaced + formula[m.end() :]


def plant(host: Path, out: Path, truth_path: Path, seed: int) -> None:
    from polar.tieout.workbook import read_workbook

    book = read_workbook(str(host))
    eligible = families(book)
    rng = random.Random(seed)
    used_rows: set[tuple[str, int]] = set()
    used_cols: set[tuple[str, int]] = set()
    plants: list[dict] = []

    ordered = sorted(eligible.items())
    for defect in SEED_CLASSES:
        shuffled = ordered[:]
        rng.shuffle(shuffled)
        made = 0
        for (sheet, row), members in shuffled:
            if made >= SITES_PER_CLASS or (sheet, row) in used_rows:
                continue
            candidates = [
                m for m in members if (sheet, m[1]) not in used_cols
            ]
            if len(candidates) < 1 or len(members) - 1 < MIN_FAMILY - 1:
                continue
            ref, col, formula, rows0, signature = rng.choice(candidates)
            after: str | None = None
            if defect == "plug":
                after = formula + f"-{rng.choice(PLUGS)}"
            elif defect == "off-by-one":
                after = _shift_start(formula)
            elif defect == "bleed":
                sibling_cols = [m[1] for m in members if m[0] != ref]
                toward = (
                    col + 1
                    if any(c > col for c in sibling_cols)
                    else col - 1
                )
                if toward >= 1:
                    after = _widen(formula, col, toward)
            elif defect == "mis-drag":
                for dc in (2, -2):
                    at_col = col + dc
                    if at_col < 1 or not rows0:
                        continue
                    at_row = rng.choice(rows0)
                    stray = f"{get_column_letter(at_col)}{at_row}"
                    if f"{sheet}!{stray}" in book.cells:
                        after = formula + f"+{stray}"
                        break
            if after is None or after == formula:
                continue
            plants.append(
                {
                    "host": host.name,
                    "class": defect,
                    "ref": ref,
                    "sheet": sheet,
                    "before": formula,
                    "after": after,
                }
            )
            used_rows.add((sheet, row))
            used_cols.add((sheet, col))
            made += 1

    #: Truth on disk before the engine ever sees the planted file.
    truth_path.write_text(json.dumps(plants, indent=1))

    shutil.copy(host, out)
    with zipfile.ZipFile(host) as zf:
        targets = sheet_files(zf)
        contents = {name: zf.read(name) for name in zf.namelist()}
    planted_ok: list[dict] = []
    for one in plants:
        target = targets.get(one["sheet"])
        if target is None:
            one["planted"] = False
            continue
        xml = contents[target].decode("utf-8")
        bare_ref = one["ref"].rsplit("!", 1)[-1]
        new_xml = rewrite_formula(xml, bare_ref, one["after"])
        if new_xml is None:
            one["planted"] = False
            continue
        contents[target] = new_xml.encode("utf-8")
        one["planted"] = True
        planted_ok.append(one)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in contents.items():
            zf.writestr(name, data)
    truth_path.write_text(json.dumps(plants, indent=1))
    print(
        f"{host.name}: {len(planted_ok)} planted of {len(plants)} drawn "
        f"({', '.join(one['class'] for one in planted_ok)})"
    )


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--scan":
        scan(Path(sys.argv[2]))
    elif len(sys.argv) == 5:
        plant(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), int(sys.argv[4]))
    else:
        print(__doc__)
        raise SystemExit(2)
