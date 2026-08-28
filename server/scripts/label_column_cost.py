"""What the label-column exclusion costs, per corpus.

The registered instrument (`docs/pierce/label-column-cost.md`).
Rebuilds each sheet's grid exactly as `_read_sheet` does, asks
`_label_column` which column is excluded, and counts what sits in it.

    uv run python -m scripts.label_column_cost <corpus-dir> <out.json> [pattern]

Counts nothing by heuristic that the engine does not already decide:
the label column is whatever `_label_column` returns, so the blind
spot measured here is the engine's actual blind spot.
"""

import json
import sys
import warnings
from decimal import Decimal
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")


def _label_like(value: Decimal) -> bool:
    """Registered in advance: a year, or a small whole number.

    Crude on purpose. Its job is to separate « a year in the label
    column » from « a cash figure in the label column », and it is
    fixed before the draw so it cannot be tuned afterwards.
    """
    if value != value.to_integral_value():
        return False
    whole = int(value)
    if 1900 <= whole <= 2100:
        return True
    return 0 <= whole <= 100


def one(path: Path) -> dict[str, Any]:
    from polar.tieout.workbook import (
        _formula,
        _grid_of,
        _label_column,
        _decimal,
    )
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    row: dict[str, Any] = {"file": path.name}
    totals = {
        "formulas": 0, "formulas_in_label": 0,
        "numerics": 0, "numerics_in_label": 0,
        "content_like_in_label": 0, "label_like_in_label": 0,
        "sheets": 0,
    }
    examples: list[dict[str, Any]] = []
    try:
        for name in formulas.sheetnames:
            sheet = formulas[name]
            if not hasattr(sheet, "max_row"):
                continue
            grid = _grid_of(sheet, values[name])
            if not grid.written and not grid.values:
                continue
            last_row = max((r for r, _ in grid.written), default=0)
            last_column = max((c for _, c in grid.written), default=0)
            if not last_row or not last_column:
                continue
            totals["sheets"] += 1
            label_column = _label_column(grid, last_row, last_column)
            for (r, c), written in grid.written.items():
                is_formula = _formula(written) is not None
                number = _decimal(grid.values.get((r, c)))
                if is_formula:
                    totals["formulas"] += 1
                    if c == label_column:
                        totals["formulas_in_label"] += 1
                        if len(examples) < 12:
                            examples.append({
                                "sheet": name, "row": r, "column": c,
                                "formula": str(written)[:70],
                                "value": str(number) if number is not None else None,
                            })
                elif number is not None:
                    totals["numerics"] += 1
                    if c == label_column:
                        totals["numerics_in_label"] += 1
                        if _label_like(number):
                            totals["label_like_in_label"] += 1
                        else:
                            totals["content_like_in_label"] += 1
                            if len(examples) < 12:
                                examples.append({
                                    "sheet": name, "row": r, "column": c,
                                    "value": str(number), "formula": None,
                                })
    finally:
        formulas.close()
        values.close()
    row.update(totals)
    row["examples"] = examples
    return row


def main() -> None:
    root, out = Path(sys.argv[1]), Path(sys.argv[2])
    pattern = sys.argv[3] if len(sys.argv) > 3 else "*.xls[xmb]"
    results: list[dict[str, Any]] = []
    for path in sorted(root.rglob(pattern)):
        if path.name.startswith("._"):
            continue
        try:
            results.append(one(path))
        except Exception as problem:
            results.append({"file": path.name,
                            "error": f"{type(problem).__name__}: {str(problem)[:120]}"})
        r = results[-1]
        print(
            f"{path.name[:44]:<44} "
            f"F {r.get('formulas_in_label', '-')}/{r.get('formulas', '-')}  "
            f"N {r.get('content_like_in_label', '-')}/{r.get('numerics', '-')}"
            f"{'  ' + r['error'] if 'error' in r else ''}",
            flush=True,
        )
        out.write_text(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
