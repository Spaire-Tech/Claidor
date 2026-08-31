"""Piece 8 — the B6 decision measurement (docs/pierce/b6-decision.md).

    cd server && uv run python -m scripts.b6_decision MODEL_KEY OUT.json

Measures, on a model B5 can mine informatively today, the decision
quantity the plan's fifth amendment names: how many mined rules does
one real authoring decision break? Two parts:

- **M1, static**: for every cell, R(c) = how many committed stable
  rules a defect at that cell could reach (a ceiling by
  reachability over the dependency graph). Reports the rule graph's
  field of view.
- **M2, live**: the PR24 revision study's two verified regression
  classes (formula overwritten with a constant; a constant hardcoded
  into a formula tail — the study's own `0.490096707821704`) planted
  at the best-case site (highest R among rule-term cells) and at an
  R = 0 control site, 40 perturbed runs each under round 4's exact
  typing, every committed rule checked against every kept run. A
  clean control runs first; if any rule fails on the unplanted file
  the harness is defective and the round stops.

Sites are chosen by the registered rules, never by hand. The plant
classes are the study's verbatim; the bias (designed plants, sites
adversarially in B6's favour) is stated in the round document.
"""

import json
import random
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

from polar.tieout.recalc.mine import Rule, sample_with_families
from polar.tieout.recalc.narrow import precedent_cone
from polar.tieout.recalc.uno_calc import UnoCalculator, find_install
from polar.tieout.workbook import read_workbook
from scripts.recalc_behave import perturb
from scripts.recalc_mine import MODELS, families_of, inferred_inputs

#: The committed round-4 stable rule sets — the product's own rules,
#: the ones both stability gates passed on (c6-stability.md).
ROUNDS = {
    "h7-fds": "docs/pierce/logs/dynamo/round4-h7fds.json",
    "h7-fp": "docs/pierce/logs/dynamo/round4-h7fp.json",
}

#: Severn Trent `N1885`'s tail constant, verbatim from
#: revision-defect-results.md.
HARDCODE = 0.490096707821704

RUNS = 40
SEED_PLANTS = 7
SEED_CONTROL = 8
#: A rule is broken when it fails in more than half the kept runs.
BROKEN_SHARE = 0.5


def load_rules(repo: Path, key: str) -> list[Rule]:
    record = json.loads((repo / ROUNDS[key]).read_text())
    return [
        Rule(terms=tuple((ref, int(sign)) for ref, sign in entry["terms"]))
        for entry in record["stable"]
    ]


def rule_reach(cells: dict, rules: list[Rule]) -> dict[str, int]:
    """R(c): how many rules a defect at `c` could reach, as a ceiling."""
    reach: dict[str, int] = dict.fromkeys(cells, 0)
    for rule in rules:
        can_break: set[str] = set()
        for ref in rule.refs:
            can_break.add(ref)
            can_break |= precedent_cone(cells, ref)
        for ref in can_break:
            if ref in reach:
                reach[ref] += 1
    return reach


def distribution(reach: dict[str, int], refs: list[str]) -> dict:
    counts = sorted(reach[ref] for ref in refs)
    if not counts:
        return {"cells": 0}
    touched = sum(1 for n in counts if n > 0)
    return {
        "cells": len(counts),
        "r_zero": len(counts) - touched,
        "r_at_least_one": touched,
        "share_reachable": round(touched / len(counts), 4),
        "median": counts[len(counts) // 2],
        "max": counts[-1],
    }


def plant(source: Path, target: Path, ref: str, kind: str, stored: float) -> str:
    """One regression, the study's classes verbatim. Returns the edit."""
    book = load_workbook(source)
    sheet, at = ref.rsplit("!", 1)
    cell = book[sheet][at]
    formula = cell.value
    if not (isinstance(formula, str) and formula.startswith("=")):
        raise ValueError(f"{ref} does not hold a formula ({formula!r})")
    if kind == "constant-overwrite":
        cell.value = stored
        edit = f"{ref}: {formula!r} -> {stored!r}"
    elif kind == "tail-hardcode":
        cell.value = f"{formula}-{HARDCODE!r}"
        edit = f"{ref}: {formula!r} -> {cell.value!r}"
    else:
        raise ValueError(kind)
    book.save(target)
    return edit


def formula_diff(clean: Path, planted: Path) -> list[str]:
    """Every cell whose content differs — what a raw diff already sees."""
    a, b = load_workbook(clean), load_workbook(planted)
    changed = []
    for name in a.sheetnames:
        ws_a, ws_b = a[name], b[name]
        rows = max(ws_a.max_row, ws_b.max_row)
        columns = max(ws_a.max_column, ws_b.max_column)
        for row in ws_a.iter_rows(min_row=1, max_row=rows, max_col=columns):
            for cell in row:
                if cell.value != ws_b[cell.coordinate].value:
                    changed.append(f"{name}!{cell.coordinate}")
    return changed


def run_condition(
    calculator: UnoCalculator,
    source: Path,
    work: Path,
    tag: str,
    typed: list,
    families: list,
    watched: list[str],
    rules: list[Rule],
    seed: int,
    inject: dict[str, float] | None = None,
) -> dict:
    """RUNS perturbed recalculations of `source`; every rule vs every run.

    `inject` carries the constant a constant-overwrite plant put where
    a formula was — the driver returns formula cells only, and the
    planted cell's value is exactly the planted constant.
    """
    rng = random.Random(seed)
    expected = [ref for ref in watched if ref not in (inject or {})]
    kept = 0
    drops = 0
    failures: dict[int, int] = dict.fromkeys(range(len(rules)), 0)
    for index in range(RUNS):
        copy = work / f"{source.stem}-{tag}-{index}.xlsx"
        perturb(source, copy, sample_with_families(typed, families, rng))
        try:
            values = calculator.recalculate(str(copy)).values
        except Exception:
            drops += 1
            copy.unlink(missing_ok=True)
            continue
        numeric = {
            ref: value
            for ref, value in values.items()
            if isinstance(value, float) and ref in set(expected)
        }
        copy.unlink(missing_ok=True)
        if len(numeric) < len(expected):
            drops += 1
            continue
        numeric.update(inject or {})
        kept += 1
        for position, rule in enumerate(rules):
            if not rule.holds(numeric):
                failures[position] += 1
    broken = [
        position
        for position, count in failures.items()
        if kept and count / kept > BROKEN_SHARE
    ]
    return {
        "runs": RUNS,
        "kept": kept,
        "drops": drops,
        "rules_failed_any": sum(1 for count in failures.values() if count),
        "rules_broken": len(broken),
        "broken": [
            {"terms": list(rules[position].terms), "failed_in": failures[position]}
            for position in broken
        ],
    }


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    key, out = sys.argv[1], Path(sys.argv[2])
    if find_install() is None:
        print("no LibreOffice >= 25.8 here — refusing")
        return 1
    repo = Path(__file__).parent.parent.parent
    path, sheet, _typing = MODELS[key]
    source = Path(path)
    rules = load_rules(repo, key)
    work = out.parent / f"{out.stem}-work"
    work.mkdir(parents=True, exist_ok=True)

    book = read_workbook(str(source))
    cells = book.cells
    watched = sorted(
        ref
        for ref, cell in cells.items()
        if cell.sheet == sheet and cell.formula is not None
    )
    print(f"{source.name}: {len(rules)} committed stable rules, "
          f"{len(watched)} watched formula cells", flush=True)

    # --- M1: the rule graph's field of view -------------------------
    started = time.monotonic()
    reach = rule_reach(cells, rules)
    every = list(cells)
    formulas = [ref for ref, cell in cells.items() if cell.formula is not None]
    constants = [ref for ref, cell in cells.items() if cell.formula is None]
    m1 = {
        "all_cells": distribution(reach, every),
        "formula_cells": distribution(reach, formulas),
        "constant_cells": distribution(reach, constants),
        "seconds": round(time.monotonic() - started, 1),
    }
    for name, stats in m1.items():
        if name != "seconds":
            print(f"  M1 {name}: {stats}", flush=True)

    # --- Site selection, by the registered rules --------------------
    term_cells = sorted({ref for rule in rules for ref in rule.refs})
    best = max(term_cells, key=lambda ref: (reach.get(ref, 0), ref))
    best = min(
        (ref for ref in term_cells if reach.get(ref, 0) == reach.get(best, 0)),
    )
    zero_candidates = [ref for ref in watched if reach.get(ref, 0) == 0]
    control_site = zero_candidates[0] if zero_candidates else None
    print(f"  sites: best-case {best} (R={reach.get(best)}), "
          f"R=0 control {control_site}", flush=True)

    # --- M2: plants, runs, rules ------------------------------------
    typed, _how = inferred_inputs(cells)
    families = families_of(cells, typed, sheet)
    stored = {ref: cells[ref].value for ref in term_cells}

    plants = [("a1", best, "constant-overwrite"), ("b1", best, "tail-hardcode")]
    if control_site is not None:
        plants.append(("c1", control_site, "tail-hardcode"))

    calculator = UnoCalculator(document_timeout=1800)
    calculator.start()
    conditions: dict[str, dict] = {}
    try:
        conditions["control"] = run_condition(
            calculator, source, work, "control", typed, families,
            watched, rules, SEED_CONTROL,
        )
        print(f"  control: {conditions['control']}", flush=True)
        if conditions["control"]["rules_failed_any"]:
            print("  CONTROL FAILED — harness defective, stopping", flush=True)
            out.write_text(json.dumps(
                {"model": source.name, "m1": m1, "control": conditions["control"],
                 "verdict": "control failed — no plant result is valid"}, indent=1))
            return 1
        for tag, site, kind in plants:
            planted_file = work / f"{source.stem}-{tag}.xlsx"
            base = stored[site]
            edit = plant(source, planted_file, site, kind, float(base))
            changed = formula_diff(source, planted_file)
            inject = {site: float(base)} if kind == "constant-overwrite" else None
            result = run_condition(
                calculator, planted_file, work, tag, typed, families,
                watched, rules, SEED_PLANTS, inject=inject,
            )
            result.update({
                "site": site, "class": kind, "edit": edit,
                "reach_ceiling": reach.get(site, 0),
                "diff_changed_cells": changed,
                "diff_names_plant_alone": changed == [site],
            })
            conditions[tag] = result
            print(f"  {tag} [{kind} at {site}]: broke {result['rules_broken']} "
                  f"of {len(rules)} rules (any-run {result['rules_failed_any']}); "
                  f"diff sees {len(changed)} changed cell(s)", flush=True)
    finally:
        calculator.stop()

    out.write_text(json.dumps({
        "model": source.name,
        "rules": len(rules),
        "watched": len(watched),
        "m1": m1,
        "sites": {"best": best, "best_reach": reach.get(best, 0),
                  "zero_control": control_site},
        "conditions": conditions,
    }, indent=1))
    print(f"wrote {out}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
