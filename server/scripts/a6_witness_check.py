"""Criterion 1 — does the BIFF census count what it claims to?

The census counts FORMULA (0x0006) records in the workbook stream.
The claim under test is that this equals the number of formula-bearing
cells, including cells inside a shared-formula group (SHRFMLA, 0x04BC),
which BIFF8 is documented to emit one FORMULA record for.

The second, independent count is LibreOffice's: convert the file and
count formula cells in the .xlsx through openpyxl. Two different
implementations reading the same bytes. Where they agree the census is
trustworthy; where they disagree the disagreement is the finding, and
this round stops rather than judging a route by a witness that cannot
be trusted.

    uv run python -m scripts.a6_witness_check <manifest.json> <out.json>

Pairing comes from the manifest, never from a basename: two files in
the population share a name while differing in content, and a flat
output directory silently overwrote one conversion with the other.
"""

import json
import struct
import sys
import warnings
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")


def census(path: Path) -> dict[str, int]:
    """FORMULA, SHRFMLA and ARRAY record counts, straight from the stream."""
    import olefile

    ole = olefile.OleFileIO(str(path))
    stream = "Workbook" if ole.exists("Workbook") else "Book"
    data = ole.openstream(stream).read()
    ole.close()

    counts = {"formula": 0, "shrfmla": 0, "array": 0, "table": 0}
    pos, end = 0, len(data)
    while pos + 4 <= end:
        code, size = struct.unpack("<HH", data[pos : pos + 4])
        pos += 4
        if pos + size > end:
            break
        if code == 0x0006:
            counts["formula"] += 1
        elif code == 0x04BC:
            counts["shrfmla"] += 1
        elif code == 0x0221:
            counts["array"] += 1
        elif code == 0x0236:
            counts["table"] += 1
        pos += size
    return counts


def converted_formulas(path: Path) -> int:
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False, read_only=True)
    total = 0
    #: A cell whose *text* begins with "=" is not a formula. These
    #: corpora carry human notes like "=tput learn x unyld dice" as
    #: plain text, and counting them by their leading character
    #: inflated the count and made the record census look wrong.
    #: openpyxl's data_type is the discriminator: "f" is a formula,
    #: "s" is a string that merely looks like one.
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    total += 1
    book.close()
    return total


def main() -> None:
    manifest, out = Path(sys.argv[1]), Path(sys.argv[2])
    results: list[dict[str, Any]] = []
    for pair in json.loads(manifest.read_text()):
        original = Path(pair["original"])
        if not original.exists():
            continue
        converted = Path(pair["converted"])
        row: dict[str, Any] = {"file": original.name, "path": str(original)}
        try:
            row["census"] = census(original)
        except Exception as problem:
            row["census_error"] = f"{type(problem).__name__}: {problem}"
        if converted.exists():
            try:
                row["libreoffice"] = converted_formulas(converted)
            except Exception as problem:
                row["libreoffice_error"] = (
                    f"{type(problem).__name__}: {str(problem)[:120]}"
                )
        else:
            row["libreoffice_error"] = "not converted"
        results.append(row)
        print(json.dumps(row)[:220], flush=True)
        out.write_text(json.dumps(results, indent=1))


if __name__ == "__main__":
    main()
