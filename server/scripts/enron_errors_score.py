"""Score the audit against the Enron error corpus
(`docs/pierce/enron-errors.md`).

    PYTHONPATH=. uv run python scripts/enron_errors_score.py CORPUS_DIR OUT.json

CORPUS_DIR holds the unpacked TU Graz archive plus `xlsx/` (the 26
workbooks converted through LibreOffice), `labels.json` (the
properties files' faulty cells, `sheet-index!col!row`) and
`sheetnames.json` (index → name per converted workbook). The unit is
the catalogue error (`enron-errors.xlsx`); an error is covered when
any finding on its workbook touches one of its cells, under the
CUSTODES scorer's touch convention (cell, roster, filled rectangle).
"""

from __future__ import annotations

import json
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from polar.tieout.audit import audit
from polar.tieout.structure import period_axes
from polar.tieout.workbook import read_workbook

A1 = re.compile(r"^(?:(.*)!)?\$?([A-Z]{1,3})\$?(\d+)$")
RECT = re.compile(r"filled across \d+ cells \(([^)]+?) to ([^)]+?)\)")


def _expand(a: str, b: str, default_sheet: str):
    ma, mb = A1.match(a.strip()), A1.match(b.strip())
    if not ma or not mb:
        return
    sheet = ma.group(1) or default_sheet
    c1, c2 = (
        column_index_from_string(ma.group(2)),
        column_index_from_string(mb.group(2)),
    )
    r1, r2 = int(ma.group(3)), int(mb.group(3))
    if (c2 - c1 + 1) * (abs(r2 - r1) + 1) > 5000:
        return
    for col in range(min(c1, c2), max(c1, c2) + 1):
        for row in range(min(r1, r2), max(r1, r2) + 1):
            yield sheet, f"{get_column_letter(col)}{row}"


def finding_cells(finding) -> set[tuple[str, str]]:
    sheet = finding.sheet
    cells: set[tuple[str, str]] = set()
    for token in [finding.ref, *(finding.cells or "").split(", ")]:
        token = token.strip()
        if not token:
            continue
        m = A1.match(token if "!" in token else f"{sheet}!{token}")
        if m:
            cells.add((m.group(1) or sheet, f"{m.group(2)}{m.group(3)}"))
    for m in RECT.finditer(finding.detail or ""):
        cells.update(_expand(m.group(1), m.group(2), sheet))
    return cells


def catalogue(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Error Nr")
    keys = [str(c) for c in rows[head]]
    out = []
    for r in rows[head + 1 :]:
        if not r or not isinstance(r[0], (int, float)):
            continue
        out.append({k: r[i] for i, k in enumerate(keys) if i < len(r)})
    return out


def main(argv: list[str]) -> None:
    corpus, out = Path(argv[1]), Path(argv[2])
    labels = json.loads((corpus / "labels.json").read_text())
    names = json.loads((corpus / "sheetnames.json").read_text())
    errors = catalogue(corpus / "enron-errors.xlsx")

    #: Faulty cells per workbook number, resolved to (sheet name, A1).
    #: The properties files give a sheet *index*, and the converted
    #: workbook's order does not always match it (hidden and chart
    #: sheets): error 27's « EGS EXP » is index 2 in the properties and
    #: the sixth sheet after conversion. The catalogue's sheet name
    #: wins when the workbook has it; the index is the fallback.
    sheet_named: dict[str, str] = {}
    for error in errors:
        number = f"{int(error['Spreadsheet Nr']):02d}"
        name = str(error["Faulty worksheet"])
        if number in names and name in names[number]:
            sheet_named.setdefault(number, name)
    faulty: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for number, entry in labels.items():
        for token in entry["cells"]:
            index, col, row = token.split("!")
            sheet = sheet_named.get(number) or names[number][int(index)]
            faulty[number].add((sheet, f"{col}{row}"))

    started = time.time()
    touched: dict[str, dict[tuple[str, str], set[str]]] = {}
    raised: Counter[str] = Counter()
    per_book: dict[str, int] = {}
    for number in sorted(labels):
        book = read_workbook(str(corpus / "xlsx" / f"{number}.xlsx"))
        result = audit(book, axes=period_axes(book))
        per_book[number] = len(result.findings)
        hits: dict[tuple[str, str], set[str]] = defaultdict(set)
        for finding in result.findings:
            raised[finding.rule] += 1
            for cell in finding_cells(finding):
                if cell in faulty[number]:
                    hits[cell].add(finding.rule)
        touched[number] = hits
    seconds = round(time.time() - started, 1)

    #: Per catalogue error: its cells (from the properties file, by
    #: workbook number and sheet) and whether any is touched.
    scored = []
    not_held: list[int] = []
    for error in errors:
        number = f"{int(error['Spreadsheet Nr']):02d}"
        if number not in labels:
            #: The catalogue names 30 workbooks; the archive packages 26.
            #: An error whose workbook is not held is outside the
            #: denominator, and reported as such.
            not_held.append(int(error["Error Nr"]))
            continue
        sheet = str(error["Faulty worksheet"])
        cells_text = str(error["Faulty cells"])
        cells = {c for c in faulty.get(number, set()) if c[0] == sheet}
        if not cells:  # the catalogue's sheet name differs from the file's; fall back to the workbook
            cells = set(faulty.get(number, set()))
        covered_by = sorted(
            {r for c in cells for r in touched.get(number, {}).get(c, set())}
        )
        scored.append(
            {
                "error": int(error["Error Nr"]),
                "workbook": number,
                "file": str(error["Faulty Spreadsheet"]),
                "sheet": sheet,
                "cells": cells_text,
                "type": str(error.get("Type")),
                "subtype": str(error.get("Subtype")),
                "cell_type": str(error.get("Cell type")),
                "info": str(error.get("Additional information")),
                "covered": bool(covered_by),
                "covered_by": covered_by,
            }
        )
    covered = [e for e in scored if e["covered"]]
    cells_total = sum(len(v) for v in faulty.values())
    cells_covered = sum(len(v) for v in touched.values())

    record = {
        "errors": len(scored),
        "not_held": not_held,
        "covered": len(covered),
        "cells_total": cells_total,
        "cells_covered": cells_covered,
        "by_type": {
            k: [
                sum(1 for e in scored if e[key] == k and e["covered"]),
                sum(1 for e in scored if e[key] == k),
            ]
            for key in ("type", "subtype", "cell_type")
            for k in sorted({e[key] for e in scored})
        },
        "raised_by_rule": dict(raised.most_common()),
        "raised_per_book": per_book,
        "seconds": seconds,
        "scored": scored,
    }
    out.write_text(json.dumps(record, indent=1, ensure_ascii=False))
    print(
        f"COVERAGE {len(covered)}/{len(scored)} errors (not held: {not_held}); cells {cells_covered}/{cells_total}; findings raised {sum(raised.values())} in {seconds}s"
    )
    for e in scored:
        mark = "HIT " if e["covered"] else "miss"
        print(
            f"  {mark} #{e['error']:2d} wb{e['workbook']} {e['sheet'][:22]:22} {e['cells'][:28]:28} {e['type'][:12]:12} {e['subtype'][:10]:10} {e['cell_type'][:7]:7} {e['info'][:26]:26} {','.join(e['covered_by'])}"
        )
    print("raised by rule:", dict(raised.most_common()))


if __name__ == "__main__":
    main(sys.argv)
