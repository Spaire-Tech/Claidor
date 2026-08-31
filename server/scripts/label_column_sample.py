"""Enumerate every label-column formula cell, so 20 can be drawn fairly.

Criterion 2 of `label-column-cost.md` requires hand-reading a sample
of the cells the exclusion hides. The cost instrument keeps only a
dozen examples per file, which is not a sample — this walks the
corpora again and records every one, with the context needed to
judge it: the row's label, the formula, the cached value, and what
the two columns to the right hold on the same row.

    uv run python -m scripts.label_column_sample <corpus-dir> <out.json> [pattern]
"""

import json
import sys
import warnings
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")


def one(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from polar.tieout.workbook import (
        _decimal,
        _formula,
        _grid_of,
        _label_column,
        _shown,
    )

    formulas = load_workbook(path, data_only=False, read_only=True)
    values = load_workbook(path, data_only=True, read_only=True)
    found: list[dict[str, Any]] = []
    try:
        for name in formulas.sheetnames:
            sheet = formulas[name]
            if not hasattr(sheet, "max_row"):
                continue
            grid = _grid_of(sheet, values[name])
            last_row = max((r for r, _ in grid.written), default=0)
            last_column = max((c for _, c in grid.written), default=0)
            if not last_row or not last_column:
                continue
            label_column = _label_column(grid, last_row, last_column)
            for (r, c), written in grid.written.items():
                if c != label_column:
                    continue
                text = _formula(written)
                if text is None:
                    continue
                neighbours = []
                for step in (1, 2):
                    other = grid.written.get((r, c + step))
                    neighbours.append(
                        (_formula(other) or _shown(grid, r, c + step) or "")[:44]
                    )
                found.append(
                    {
                        "file": path.name,
                        "sheet": name,
                        "ref": f"{get_column_letter(c)}{r}",
                        "label_column": get_column_letter(label_column),
                        "formula": text[:120],
                        "value": str(_decimal(grid.values.get((r, c)))),
                        "left_of_it": (_shown(grid, r, c - 1) or "")[:36]
                        if c > 1
                        else "",
                        "right": neighbours,
                    }
                )
    finally:
        formulas.close()
        values.close()
    return found


def main() -> None:
    root, out = Path(sys.argv[1]), Path(sys.argv[2])
    pattern = sys.argv[3] if len(sys.argv) > 3 else "*.xls[xmb]"
    every: list[dict[str, Any]] = []
    for path in sorted(root.rglob(pattern)):
        if path.name.startswith("._"):
            continue
        try:
            rows = one(path)
        except Exception as problem:
            print(f"{path.name}: {type(problem).__name__}", flush=True)
            continue
        every.extend(rows)
        print(f"{path.name[:46]:<46} {len(rows):>6} label-column formulas", flush=True)
        out.write_text(json.dumps(every, indent=1))


if __name__ == "__main__":
    main()
