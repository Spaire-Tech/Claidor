import sys
from openpyxl import load_workbook
from scripts.recalc_units_convention import is_percent_format
for path in sys.argv[1:]:
    book = load_workbook(path, read_only=False, data_only=True)
    print(f"=== {path.split('/')[-1]}")
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    continue
                if not is_percent_format(cell.number_format or ""):
                    continue
                if abs(float(v)) <= 1.5:
                    continue
                label = ""
                for c in range(1, 8):
                    t = sheet.cell(row=cell.row, column=c).value
                    if isinstance(t, str) and t.strip():
                        label = t.strip()[:52]
                        break
                print(f"  {sheet.title}!{cell.coordinate} = {v!r}  fmt={cell.number_format!r}  label={label!r}")
    book.close()
