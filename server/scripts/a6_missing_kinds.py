"""What the converted file actually holds where a number went missing.

Closes criterion 2 of `a6-intake.md`: every disagreement enumerated
and hand-read. A cell counted « missing » is not necessarily absent
— it may have arrived as a type our reader does not elect, which is
the same thing to the engine and a different thing to the truth.
"""

import collections
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

from scripts.a6_fidelity import _converted, _witness_xls, _witness_xlsb

warnings.filterwarnings("ignore")

CORPUS = Path(__file__).parent / "corpus_sft"


def main() -> None:
    converted_dir = Path(sys.argv[1])
    for name in sys.argv[2:]:
        original = CORPUS / name
        witness = (
            _witness_xlsb(original)
            if original.suffix.lower() == ".xlsb"
            else _witness_xls(original)
        )
        converted_path = converted_dir / (original.stem + ".xlsx")
        after = _converted(converted_path)
        missing = {
            k: v for k, v in witness["values"].items() if k not in after["values"]
        }
        print(f"=== {name}: {len(missing)} missing")
        if not missing:
            continue

        wanted = collections.defaultdict(set)
        for sheet, r, c in missing:
            wanted[sheet].add((r, c))

        book = load_workbook(converted_path, data_only=False, read_only=True)
        kinds: collections.Counter = collections.Counter()
        examples: dict = {}
        for sheet_name, cells in wanted.items():
            if sheet_name not in book.sheetnames:
                kinds["sheet absent"] += len(cells)
                continue
            sheet = book[sheet_name]
            rows_wanted = {r for r, _ in cells}
            for index, row in enumerate(sheet.iter_rows()):
                if index not in rows_wanted:
                    continue
                for column, cell in enumerate(row):
                    if (index, column) not in cells:
                        continue
                    value = getattr(cell, "value", None)
                    kind = type(value).__name__
                    kinds[kind] += 1
                    examples.setdefault(
                        kind, f"{sheet_name}!r{index + 1}c{column + 1} = {value!r}"
                    )
        book.close()
        for kind, count in kinds.most_common():
            print(
                f"    {count:>6}  arrives as {kind:<10} e.g. {examples.get(kind, '')}"
            )


if __name__ == "__main__":
    main()
