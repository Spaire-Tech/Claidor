"""A3 label mining — the classifier the registration describes.

Implements `docs/pierce/custodes-mining.md` exactly: the missed truth
cells (truth minus covered, per the committed scorer's own machinery),
each classified by content type and its row/column formula context
into the registered buckets, and the first twelve cells per bucket
dumped with their neighbourhoods for the hand reading.

    uv run python -m scripts.custodes_score   # first, to build the work dir
    uv run python -m scripts.custodes_mine

Prints the bucket table; writes `custodes_work/mining_samples.txt`
with the hand-reading material. Decides nothing: verdicts live in the
mining document, and any adopted pattern still goes through the loop
on our own corpora before it ships.
"""

import re
import sys
import warnings
from collections import Counter, defaultdict

warnings.filterwarnings("ignore")

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from scripts.custodes_score import (  # noqa: E402
    WORK,
    _findings,
    _truth_from_comments,
)

WINDOW = 3

#: Classifier-local shape: refs → offsets, numbers → #. Mirrors the
#: engine's normalization in spirit; decides nothing outside mining.
REF = re.compile(r"(?<![A-Za-z0-9_])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(])")
NUM = re.compile(r"(?<![A-Za-z0-9_\x00])\d+(?:\.\d+)?")


def shape(formula: str, row: int, col: int) -> str:
    holes: list[str] = []

    def stash(m: re.Match[str]) -> str:
        ca, letters, ra, digits = m.groups()
        c = 0
        for ch in letters:
            c = c * 26 + ord(ch) - 64
        cc = f"C{letters}" if ca else f"C[{c - col:+d}]"
        rr = f"R{digits}" if ra else f"R[{int(digits) - row:+d}]"
        holes.append(rr + cc)
        return f"\x00{len(holes) - 1}\x00"

    out = REF.sub(stash, formula)
    out = NUM.sub("#", out)
    for i, hole in enumerate(holes):
        out = out.replace(f"\x00{i}\x00", hole)
    return out


def a1(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def parse_ref(cell: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Z]{1,3})(\d+)", cell)
    if m is None:
        raise ValueError(cell)
    c = 0
    for ch in m.group(1):
        c = c * 26 + ord(ch) - 64
    return c, int(m.group(2))


def main() -> None:
    subjects = sorted(p.stem for p in (WORK / "subjects").iterdir())
    truth = _truth_from_comments(subjects)
    findings = _findings()

    covered: set[tuple[str, str, str]] = set()
    for book, _rule, cells in findings:
        for sheet, cell in cells:
            if (book, sheet, cell) in truth:
                covered.add((book, sheet, cell))
    missed = sorted(truth - covered)
    print(f"missed truth cells: {len(missed)} (expected 1690)")

    books: dict[str, object] = {}

    def sheet_of(book: str, name: str):
        if book not in books:
            path = WORK / "xlsx" / f"{book}.xlsx"
            books[book] = load_workbook(path) if path.exists() else None
        wb = books[book]
        if wb is None:
            return None
        for candidate in wb.sheetnames:
            if candidate.strip() == name.strip():
                return wb[candidate]
        return None

    buckets: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    samples: dict[str, list[str]] = defaultdict(list)

    for book, sheet_name, cell in missed:
        ws = sheet_of(book, sheet_name)
        if ws is None:
            buckets["unresolvable"].append((book, sheet_name, cell))
            continue
        try:
            col, row = parse_ref(cell)
        except ValueError:
            buckets["unresolvable"].append((book, sheet_name, cell))
            continue
        held = ws.cell(row=row, column=col).value
        if held is None:
            buckets["unresolvable"].append((book, sheet_name, cell))
            continue

        def neighbours(dr: int, dc: int) -> list[tuple[int, int, object]]:
            out = []
            for step in range(1, WINDOW + 1):
                for sign in (-1, 1):
                    r, c = row + dr * step * sign, col + dc * step * sign
                    if r < 1 or c < 1:
                        continue
                    got = ws.cell(row=r, column=c).value
                    if got is not None:
                        out.append((r, c, got))
            return out

        row_win = neighbours(0, 1)
        col_win = neighbours(1, 0)

        def family(window: list[tuple[int, int, object]]) -> tuple[int, str | None]:
            shapes = Counter(
                shape(v, r, c)
                for r, c, v in window
                if isinstance(v, str) and v.startswith("=")
            )
            if not shapes:
                return 0, None
            best, n = shapes.most_common(1)[0]
            return (n, best if n >= 2 else None)

        row_n, row_fam = family(row_win)
        col_n, col_fam = family(col_win)
        is_formula = isinstance(held, str) and held.startswith("=")

        if not is_formula:
            if row_fam:
                bucket = "row-family-gap"
            elif col_fam:
                bucket = "column-family-gap"
            elif row_n:
                bucket = "row-context-weak"
            elif col_n:
                bucket = "column-context-weak"
            else:
                bucket = "no-formula-context"
        else:
            own = shape(held, row, col)
            if row_fam and own != row_fam:
                bucket = "differs-from-row-family"
            elif col_fam and own != col_fam:
                bucket = "differs-from-column-family"
            elif row_n or col_n:
                bucket = "formula-matches-context"
            else:
                bucket = "formula-no-context"

        buckets[bucket].append((book, sheet_name, cell))
        if len(samples[bucket]) < 12:
            lines = [f"--- {bucket} · {book} · {sheet_name}!{cell}"]
            lines.append(f"    holds: {held!r}")
            for label, window in (("row", row_win), ("col", col_win)):
                for r, c, v in sorted(window, key=lambda t: (t[0], t[1])):
                    lines.append(f"    {label} {a1(c, r)}: {v!r}")
            samples[bucket].append("\n".join(lines))

    print(f"\n{'bucket':28} {'cells':>6}")
    for bucket, members in sorted(buckets.items(), key=lambda kv: -len(kv[1])):
        print(f"{bucket:28} {len(members):6}")

    out = WORK / "mining_samples.txt"
    out.write_text(
        "\n\n".join(
            f"== {bucket} ({len(buckets[bucket])} cells) ==\n\n"
            + "\n\n".join(samples[bucket])
            for bucket in sorted(samples)
        )
    )
    print(f"\nhand-reading material: {out}")


if __name__ == "__main__":
    sys.exit(main())
