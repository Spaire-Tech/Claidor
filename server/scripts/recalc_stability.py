"""B5's two stability preconditions — the gates C6 waits on.

    cd server && uv run python -m scripts.recalc_stability MODEL OUT.json [RUNS]

The plan's third amendment (27 Aug) makes these **binding before any
C6 diff code exists anywhere**, and cheap enough to run first:

1. **Seed stability** — mine one unmodified model five times under
   five seeds. The rule sets must agree, or « v12 broke a rule » is
   seed noise rather than a finding.
2. **Cosmetic invariance** — insert blank rows, rename a sheet,
   reformat a block; the mined rule sets must be **identical**, or
   the claimed advantage over positional diff is unproven.

The comparison is **by label, never by cell reference**. That is not
a convenience: inserting rows moves every downstream cell, so a
reference-keyed comparison would report total disagreement for a
model that behaves identically — which is precisely the failure of
positional diffing that behavioural mining is supposed to escape. A
rule is identified by the (sign, row label, column label) triples of
its terms.

Cosmetic edits are made **by the engine** (`UnoCalculator.cosmetic`),
so inserted rows carry their formulas; a file library that shifted
cells without rewriting references would produce a broken model and
a meaningless test.
"""

import json
import random
import sys
import time
from dataclasses import replace
from pathlib import Path
from typing import Any

from polar.tieout.recalc.mine import (
    Family,
    Rule,
    cleanse,
    mine_signed_sums,
    sample_with_families,
    stable_rules,
)
from polar.tieout.recalc.uno_calc import UnoCalculator, find_install
from polar.tieout.workbook import read_workbook
from scripts.recalc_mine import MODELS, families_of, inferred_inputs

#: Round 1's seed anchors, paired. **The compared object is the
#: product's rule set, and the product's rule set is the intersection
#: of two independent minings** (`stable_rules` — « a rule that
#: appears in one sampling and not the other was luck »). So one rule
#: set costs one seed *pair*, and the five sets the amendment asks for
#: need five disjoint pairs. The first of each pair is round 1's own
#: seed, so the ten individual minings are still directly comparable
#: to the five-single-seed result of 27 August.
SEED_PAIRS = ((11, 12), (22, 23), (33, 34), (44, 45), (55, 56))


def signature(
    rule: Rule, labels: dict[str, tuple[str, str]]
) -> tuple[tuple[Any, ...], ...]:
    """A rule's identity in the model's own words, not its geometry."""
    return tuple(
        sorted(
            (coefficient, *labels.get(ref, (ref, "")))
            for ref, coefficient in rule.terms
        )
    )


def _renamed_sheet(variant: Path, original: str) -> str:
    """The variant's new sheet name, as the file actually stored it."""
    from openpyxl import load_workbook

    book = load_workbook(variant, read_only=True)
    try:
        names = list(book.sheetnames)
    finally:
        book.close()
    for name in names:
        if name != original and name.startswith(original[:15]):
            return name
    raise RuntimeError(f"no renamed sheet found in {variant.name}: {names}")


def typing_for(
    path: Path, sheet: str, typing: Any, hand: bool
) -> tuple[Any, list[Family]]:
    """The typed inputs and constrained families, the product's way.

    **Inferred is the default and hand is the fallback**, because the
    product types automatically and this round exists to measure the
    product. `--hand` stays reachable so the 27 August result remains
    reproducible from this same file.

    The families are not optional. Round 4 found round 2's draws
    **illegal** — « weight on embedded debt » and « weight on new
    debt » perturbed apart into a capital structure that cannot
    exist — and `sample_with_families` is the fix. Sampling without
    them here would measure a sampler the product no longer uses.
    """
    cells = read_workbook(str(path)).cells
    if hand:
        constants = {
            ref: float(cell.value)
            for ref, cell in cells.items()
            if cell.formula is None and cell.value is not None
        }
        return typing(constants), []
    typed, _how = inferred_inputs(cells)
    return typed, families_of(cells, typed, sheet)


def mine_once(
    calculator: UnoCalculator,
    path: Path,
    sheet: str,
    typing: Any,
    runs: int,
    seed: int,
    typed_override: Any = None,
    families: list[Family] | None = None,
    hand: bool = False,
) -> tuple[list[Rule], dict[str, tuple[str, str]], dict[str, int], Any, list[Family]]:
    book = read_workbook(str(path))
    cells = book.cells
    if typed_override is not None:
        typed, kin = typed_override, (families or [])
    else:
        typed, kin = typing_for(path, sheet, typing, hand)
    watched = sorted(
        ref for ref, cell in cells.items() if cell.sheet == sheet and cell.formula
    )
    labels = {
        ref: ((cells[ref].row_label or ""), (cells[ref].column_label or ""))
        for ref in watched
    }
    rng = random.Random(seed)
    kept: list[dict[str, float]] = []
    drops = {"engine-error": 0, "failed": 0}
    calculator.open(str(path))
    try:
        for _ in range(runs):
            try:
                values = calculator.set_and_recalculate(
                    sample_with_families(typed, kin, rng), sheets=[sheet]
                ).values
            except Exception:
                drops["failed"] += 1
                continue
            numeric = {
                ref: value
                for ref, value in values.items()
                if isinstance(value, float) and ref in set(watched)
            }
            if len(numeric) < len(watched):
                drops["engine-error"] += 1
            else:
                kept.append(numeric)
    finally:
        calculator.close()
    return cleanse(mine_signed_sums(kept, watched)), labels, drops, typed, kin


def product_rule_set(
    calculator: UnoCalculator,
    path: Path,
    sheet: str,
    typing: Any,
    runs: int,
    pair: tuple[int, int],
    typed_override: Any = None,
    families: list[Family] | None = None,
    hand: bool = False,
) -> tuple[list[Rule], list[list[Rule]], dict[str, tuple[str, str]], Any, list[Family]]:
    """One rule set as the product makes one: two minings, intersected.

    Returns the intersection *and* both individual minings, because
    the individual sets are what 27 August compared and keeping them
    costs nothing — the same machine time answers both questions.
    """
    singles: list[list[Rule]] = []
    labels: dict[str, tuple[str, str]] = {}
    typed = typed_override
    kin = families
    for seed in pair:
        rules, labels, drops, typed, kin = mine_once(
            calculator,
            path,
            sheet,
            typing,
            runs,
            seed,
            typed_override=typed,
            families=kin,
            hand=hand,
        )
        print(f"    seed {seed}: {len(rules)} rules, drops {drops}", flush=True)
        singles.append(rules)
    return stable_rules(singles[0], singles[1]), singles, labels, typed, kin


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    argv = [a for a in sys.argv[1:] if a != "--hand"]
    hand = "--hand" in sys.argv
    key, out = argv[0], Path(argv[1])
    runs = int(argv[2]) if len(argv) > 2 else 200
    if find_install() is None:
        print("no LibreOffice >= 25.8 here — refusing")
        return 1
    path_str, sheet, typing = MODELS[key]
    source = Path(path_str)
    work = out.parent / f"{out.stem}-work"
    work.mkdir(parents=True, exist_ok=True)

    calculator = UnoCalculator(document_timeout=1800)
    calculator.start()
    record: dict[str, Any] = {
        "model": source.name,
        "runs": runs,
        "typing": "hand" if hand else "inferred",
    }
    started = time.monotonic()
    try:
        # --- 1. seed stability -------------------------------------
        sets: list[set[Any]] = []
        single_sets: list[set[Any]] = []
        typed = kin = None
        for pair in SEED_PAIRS:
            print(f"  pair {pair}:", flush=True)
            stable, singles, labels, typed, kin = product_rule_set(
                calculator,
                source,
                sheet,
                typing,
                runs,
                pair,
                typed_override=typed,
                families=kin,
                hand=hand,
            )
            sets.append({signature(rule, labels) for rule in stable})
            single_sets += [{signature(r, labels) for r in one} for one in singles]
            print(f"    -> product rule set: {len(stable)}", flush=True)
        common = set.intersection(*sets) if sets else set()
        union = set.union(*sets) if sets else set()
        #: The statistic the amendment's question actually needs: a rule
        #: in some sets and not others is exactly the false « v12 broke
        #: a rule ». Reported beside the mean pairwise Jaccard, which is
        #: what `agreement()` computes, so this round is comparable to
        #: every earlier one.
        pairwise = [
            len(a & b) / len(a | b) if (a | b) else 1.0
            for i, a in enumerate(sets)
            for b in sets[i + 1 :]
        ]
        singles_common = set.intersection(*single_sets) if single_sets else set()
        singles_union = set.union(*single_sets) if single_sets else set()
        record["seed_stability"] = {
            "seed_pairs": [list(p) for p in SEED_PAIRS],
            "counts": [len(s) for s in sets],
            "in_all_five": len(common),
            "in_any": len(union),
            "core_over_union": round(len(common) / len(union), 4) if union else 1.0,
            "mean_pairwise_jaccard": round(sum(pairwise) / len(pairwise), 4)
            if pairwise
            else 1.0,
            "identical": all(s == sets[0] for s in sets),
            #: The ten individual minings, compared the way 27 August
            #: compared its five — kept so the two rounds are the same
            #: question asked of different rule sets.
            "single_minings": {
                "counts": [len(s) for s in single_sets],
                "in_all": len(singles_common),
                "in_any": len(singles_union),
                "identical": all(s == single_sets[0] for s in single_sets),
            },
            #: What flickers, named rather than counted, so a shortfall
            #: is a list of sentences and not a number to argue about.
            "flickering": [
                [list(term) for term in sorted(rule)]
                for rule in sorted(union - common, key=str)[:40]
            ],
        }
        print(
            f"seed stability: {[len(s) for s in sets]} rules per set, "
            f"{len(common)} in all five of {len(union)} in any "
            f"(core/union {record['seed_stability']['core_over_union']}), "
            f"identical={record['seed_stability']['identical']}",
            flush=True,
        )

        # --- 2. cosmetic invariance --------------------------------
        variant = work / f"{source.stem}-cosmetic.xlsx"
        calculator.cosmetic(
            str(source),
            {
                # Rows inserted *above* the modelled block, so every
                # watched cell moves — the whole point of the test.
                "insert_rows": [{"sheet": sheet, "at": 0, "count": 3}],
                "rename": [{"from": sheet, "to": (sheet[:20] + " (renamed)")[:31]}],
            },
            str(variant),
        )
        # Excel caps a sheet name at 31 characters and truncates on
        # save, so the new name is chosen to fit and then **read back
        # from the stored file** rather than assumed — the first
        # attempt addressed a sheet the file did not contain and
        # every run failed (lane log).
        renamed = _renamed_sheet(variant, sheet)
        base_rules, _, base_labels, base_typed, base_kin = product_rule_set(
            calculator,
            source,
            sheet,
            typing,
            runs,
            SEED_PAIRS[0],
            typed_override=typed,
            families=kin,
            hand=hand,
        )

        # The typing is done once on the original and *translated* to
        # the variant's coordinates — three rows down, on the renamed
        # sheet. Re-deriving it from the variant would ask the typing
        # to recognise cells that have moved, which is a different
        # (and later) question than whether the mining is invariant.
        def _move(ref: str) -> str:
            _, _, at = ref.partition("!")
            column = "".join(c for c in at if c.isalpha())
            row = int("".join(c for c in at if c.isdigit())) + 3
            return f"{renamed}!{column}{row}"

        shifted = [replace(item, ref=_move(item.ref)) for item in base_typed]
        #: The families move with the typing. A family is « these inputs
        #: sum to a constant the model never lets them leave »; leaving
        #: its refs on the original sheet would silently drop the
        #: constraint on the variant and draw the illegal capital
        #: structures round 4 exists to prevent — the same class of
        #: harness defect as the un-shifted typing, one level down.
        shifted_kin = [
            replace(f, refs=tuple(_move(r) for r in f.refs)) for f in base_kin
        ]
        var_rules, _, var_labels, _, _ = product_rule_set(
            calculator,
            variant,
            renamed,
            typing,
            runs,
            SEED_PAIRS[0],
            typed_override=shifted,
            families=shifted_kin,
            hand=hand,
        )
        base_set = {signature(r, base_labels) for r in base_rules}
        var_set = {signature(r, var_labels) for r in var_rules}
        #: The reference-keyed comparison, reported **as well**. It is
        #: expected to collapse — every watched cell moved three rows —
        #: and that collapse is the measurement of what C6 would be
        #: worth without label keying, so it is a number rather than an
        #: assertion.
        base_refs = {frozenset(r.terms) for r in base_rules}
        var_refs = {frozenset(r.terms) for r in var_rules}
        by_ref = (
            len(base_refs & var_refs) / len(base_refs | var_refs)
            if (base_refs | var_refs)
            else 1.0
        )
        record["cosmetic_invariance"] = {
            "variant": variant.name,
            "edits": "3 blank rows inserted at the top, sheet renamed",
            "rules_original": len(base_set),
            "rules_variant": len(var_set),
            "identical_by_label": base_set == var_set,
            "only_in_original": len(base_set - var_set),
            "only_in_variant": len(var_set - base_set),
            "jaccard_by_reference": round(by_ref, 4),
            "differences": [
                {"side": side, "rule": [list(t) for t in sorted(rule)]}
                for side, group in (
                    ("original", base_set - var_set),
                    ("variant", var_set - base_set),
                )
                for rule in sorted(group, key=str)[:20]
            ],
        }
        print(
            f"cosmetic invariance: {len(base_set)} vs {len(var_set)} rules, "
            f"identical_by_label={base_set == var_set}, "
            f"by-reference Jaccard {by_ref:.4f}",
            flush=True,
        )
    finally:
        calculator.stop()

    record["seconds"] = round(time.monotonic() - started, 1)
    out.write_text(json.dumps(record, indent=1))
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
