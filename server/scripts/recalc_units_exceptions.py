"""Read the percent-format outliers, one by one, with their labels.

    cd server && uv run python -m scripts.recalc_units_exceptions FILE [FILE ...]

`recalc_units_convention.py` counts how many percent-formatted cells
exceed 1.5. A count is not a finding: on our corpus the 24 outliers
turned out to be three different things, and I published a wrong
characterisation of them from reading only the first four. This
prints every one with its row label so the next person reads them
instead of trusting a number — including mine.
"""

import sys

from openpyxl import load_workbook

from scripts.recalc_units_convention import is_percent_format


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    for path in sys.argv[1:]:
        book = load_workbook(path, read_only=False, data_only=True)
        print(f"=== {path.split('/')[-1]}")
        try:
            for sheet in book.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if not isinstance(value, (int, float)) or isinstance(
                            value, bool
                        ):
                            continue
                        if not is_percent_format(cell.number_format or ""):
                            continue
                        if abs(float(value)) <= 1.5:
                            continue
                        label = ""
                        for column in range(1, 8):
                            text = sheet.cell(row=cell.row, column=column).value
                            if isinstance(text, str) and text.strip():
                                label = text.strip()[:52]
                                break
                        print(
                            f"  {sheet.title}!{cell.coordinate} = {value!r}  "
                            f"fmt={cell.number_format!r}  label={label!r}"
                        )
        finally:
            book.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
