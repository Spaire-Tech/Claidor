"""A6 — is a converted file the same file?

The registered instrument (`docs/pierce/a6-intake.md`). For each
held model it reads the **original** through an independent witness
and the **converted** file through our own reader, and compares what
the engine would see.

    uv run python -m scripts.a6_fidelity <converted-dir> <out.json>

Witnesses, chosen so our own decompiling is never in the loop:

* `.xlsb` — `pyxlsb`, a separate implementation. Its record types
  are used directly: cell records carry values, and **BrtCellBool
  (type 4) is a boolean cell, not a formula** — verified at the
  cells, because pyxlsb's own constant table names 8–11 the formula
  records and a naive read of that table gets this wrong.
* `.xls` — a raw BIFF record census over the workbook stream, so
  neither our `legacy.py` nor xlrd's cell model is trusted. FORMULA
  is 0x0006; BOOLERR (0x0205) is the boolean cell.

Neither witness is asked for formula *text* — only for whether a
cell carries a formula at all, which is what the round's criterion 3
needs and is all either can honestly attest to.
"""

import io
import json
import struct
import sys
import warnings
import zipfile
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

CORPUS = Path(__file__).parent / "corpus_sft"

HELD = [
    "barrhead_model.xlsb",
    "largs_model.xlsb",
    "dalbeattie_model.xlsb",
    "our_lady_st_patricks_model.xlsb",
    "inverclyde_model.xls",
]

#: Relative tolerance for a numeric match, as registered.
TOLERANCE = 1e-9


def _close(a: float, b: float) -> bool:
    if a == b:
        return True
    scale = max(abs(a), abs(b))
    return scale > 0 and abs(a - b) / scale <= TOLERANCE


# --- the witnesses -------------------------------------------------------


def _witness_xlsb(path: Path) -> dict[str, Any]:
    """Values and formula-carrying cells, straight from the records."""
    from pyxlsb import biff12
    from pyxlsb.reader import BIFF12Reader

    #: 8-11 are the formula records; 4 is a *boolean cell*, which
    #: pyxlsb's table names BOOL and which a careless read counts as a
    #: formula. Verified at the cells before this instrument was used.
    FORMULA = {8, 9, 10, 11}
    BOOLEAN = {4}
    #: Type 7 (BrtCellIsst) carries an *index into the shared string
    #: table*, which is an integer and is not a number in the cell.
    #: Counting it as one made 10,397 string cells look like numeric
    #: cells the conversion had lost.
    STRING = {6, 7}

    archive = zipfile.ZipFile(path)
    order: list[str] = []
    book = [n for n in archive.namelist() if n.endswith("workbook.bin")][0]
    for typ, obj in BIFF12Reader(fp=io.BytesIO(archive.read(book))):
        if typ == biff12.SHEET:
            order.append(obj.name)

    values: dict[tuple[str, int, int], float] = {}
    formulas = booleans = 0
    for index, sheet in enumerate(order, start=1):
        part = f"xl/worksheets/sheet{index}.bin"
        if part not in archive.namelist():
            continue
        row = -1
        for typ, obj in BIFF12Reader(fp=io.BytesIO(archive.read(part))):
            if typ == biff12.ROW:
                row = obj.r
                continue
            if typ in FORMULA:
                formulas += 1
            if typ in BOOLEAN:
                booleans += 1
            if typ in STRING or typ in BOOLEAN or typ in FORMULA:
                continue
            value = getattr(obj, "v", None)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                values[(sheet, row, obj.c)] = float(value)
    archive.close()
    return {"sheets": order, "values": values, "formulas": formulas,
            "booleans": booleans}


def _witness_xls(path: Path) -> dict[str, Any]:
    """A raw BIFF census, plus xlrd for the values only."""
    import olefile
    import xlrd

    ole = olefile.OleFileIO(str(path))
    stream = "Workbook" if ole.exists("Workbook") else "Book"
    data = ole.openstream(stream).read()
    ole.close()

    formulas = booleans = 0
    pos, end = 0, len(data)
    while pos + 4 <= end:
        code, size = struct.unpack("<HH", data[pos:pos + 4])
        pos += 4
        if pos + size > end:
            break
        if code == 0x0006:
            formulas += 1
        elif code == 0x0205:
            booleans += 1
        pos += size

    book = xlrd.open_workbook(str(path))
    values: dict[tuple[str, int, int], float] = {}
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_NUMBER:
                    values[(sheet.name, r, c)] = float(cell.value)
    names = [s.name for s in book.sheets()]
    book.release_resources()
    return {"sheets": names, "values": values, "formulas": formulas,
            "booleans": booleans}


# --- the converted side, read as the engine reads it ---------------------


def _converted(path: Path) -> dict[str, Any]:
    import datetime as _dt

    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel

    def _serial(value: Any) -> float | None:
        """A date-formatted cell returns as a datetime, not a number.

        Ignoring those made ~10,900 perfectly converted cells look
        missing. They are compared on the serial the file actually
        stores.
        """
        if isinstance(value, _dt.datetime):
            return float(to_excel(value))
        if isinstance(value, _dt.date):
            return float(to_excel(_dt.datetime(value.year, value.month, value.day)))
        if isinstance(value, _dt.timedelta):
            return value.total_seconds() / 86400.0
        return None

    book = load_workbook(path, data_only=False, read_only=True)
    values: dict[tuple[str, int, int], float] = {}
    formulas: dict[str, int] = {"total": 0, "boolean_shaped": 0, "empty": 0}
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                if isinstance(value, str) and value.startswith("="):
                    formulas["total"] += 1
                    text = value.strip()
                    if text in ("=TRUE()", "=FALSE()"):
                        formulas["boolean_shaped"] += 1
                    elif text == "=":
                        formulas["empty"] += 1
                elif isinstance(value, (int, float)) and not isinstance(value, bool):
                    values[(sheet.title, cell.row - 1, cell.column - 1)] = float(value)
                else:
                    serial = _serial(value)
                    if serial is not None:
                        values[(sheet.title, cell.row - 1, cell.column - 1)] = serial
    names = list(book.sheetnames)
    book.close()
    return {"sheets": names, "values": values, "formulas": formulas}


def one(original: Path, converted: Path) -> dict[str, Any]:
    row: dict[str, Any] = {"file": original.name}
    if not converted.exists():
        row["error"] = "no converted file"
        return row
    try:
        witness = (
            _witness_xlsb(original)
            if original.suffix.lower() == ".xlsb"
            else _witness_xls(original)
        )
    except Exception as problem:
        row["error"] = f"witness failed: {type(problem).__name__}: {problem}"
        return row
    try:
        after = _converted(converted)
    except Exception as problem:
        row["refused"] = f"{type(problem).__name__}: {str(problem)[:200]}"
        row["witness_sheets"] = len(witness["sheets"])
        row["witness_values"] = len(witness["values"])
        row["witness_formulas"] = witness["formulas"]
        row["witness_booleans"] = witness["booleans"]
        return row

    before_values = witness["values"]
    after_values = after["values"]
    matched = changed = missing = 0
    disagreements: list[dict[str, Any]] = []
    for key, value in before_values.items():
        other = after_values.get(key)
        if other is None:
            missing += 1
            if len(disagreements) < 40:
                disagreements.append({"kind": "missing", "at": list(key),
                                      "original": value})
        elif _close(value, other):
            matched += 1
        else:
            changed += 1
            if len(disagreements) < 40:
                disagreements.append({"kind": "changed", "at": list(key),
                                      "original": value, "converted": other})
    added = sum(1 for key in after_values if key not in before_values)

    row.update(
        witness_sheets=len(witness["sheets"]),
        converted_sheets=len(after["sheets"]),
        sheets_lost=sorted(set(witness["sheets"]) - set(after["sheets"]))[:20],
        witness_values=len(before_values),
        converted_values=len(after_values),
        matched=matched,
        changed=changed,
        missing=missing,
        added=added,
        witness_formulas=witness["formulas"],
        witness_booleans=witness["booleans"],
        converted_formulas=after["formulas"],
        disagreements=disagreements,
    )
    return row


def main() -> None:
    converted_dir, out = Path(sys.argv[1]), Path(sys.argv[2])
    results = []
    for name in HELD:
        original = CORPUS / name
        if not original.exists():
            results.append({"file": name, "error": "original not on disk"})
            continue
        converted = converted_dir / (original.stem + ".xlsx")
        row = one(original, converted)
        results.append(row)
        print(json.dumps(row.get("disagreements") and
                         {k: v for k, v in row.items() if k != "disagreements"}
                         or row)[:400], flush=True)
        out.write_text(json.dumps(results, indent=1))


if __name__ == "__main__":
    main()
